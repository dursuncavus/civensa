"""Render an EK-1 report snapshot using the owner's Word styles; no database access."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Length, Twips

from .numbers import format_number

TEMPLATE_PATH = Path(__file__).parent / 'templates_docx' / 'faaliyet_raporu.docx'
BODY = 'Civensa Gövde'
NOTE = 'Civensa Not'
CAPTION = 'Civensa Tablo Başlık'
DRAFT = 'TASLAK — onaylanmamış içerik'
STYLES = [*(f'Civensa Başlık {i}' for i in range(1, 5)), BODY, 'Civensa Madde',
          'Civensa Tablo', CAPTION, NOTE, 'Civensa Kapak Başlık',
          *(f'TOC {i}' for i in range(1, 4))]


def _text(value):
    return '—' if value is None or value == '' else str(value)


def _field(paragraph, instruction, placeholder=''):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


class _ReportContent:
    """Content and number rules shared by the Word and PDF renderers."""

    def blocks(self, blocks):
        for block in blocks:
            if block['type'] == 'paragraph':
                self.paragraph(block['text'])
            elif block['type'] == 'bullets':
                for item in block['items']:
                    self.paragraph(item, 'Civensa Madde')
            elif block['type'] == 'table' and block['rows']:
                self.table(block['rows'][0], block['rows'][1:])

    def resources(self, code):
        for table in self.payload.get('resources', {}).values():
            if table.get('section_code') != code or not table['rows']:
                continue
            self.paragraph(table['title'], CAPTION)
            columns = list(table['columns'])
            if not self.payload.get('department_id'):
                columns = [('department', 'Müdürlük'), *columns]
            self.table([label for _, label in columns], [
                [format_number(row.get(key)) if key in {'count', 'area'} else row.get(key)
                 for key, _ in columns] for row in table['rows']])
            for key, total in table.get('totals', {}).items():
                label = dict(columns).get(key, key)
                self.paragraph(f"{label} toplamı: {format_number(total['value'])} · Boş değer: {total['missing']}", NOTE)

    def financial(self, code):
        # Shared storage labels only; rendering performs no service calls or queries.
        from .activity_financial_service import EXPENSE_FIELDS, REVENUE_FIELDS, ENTRY_FIELDS
        data = self.payload.get('financial', {})
        if code == 'III.A.1':
            for kind, table in data.get('tables', {}).items():
                title = {'expenditure': 'Ekonomik Koda Göre Giderler',
                         'programme': 'Program ve Alt Programa Göre Giderler',
                         'revenue': 'Gelirler'}[kind]
                fields = list(REVENUE_FIELDS if kind == 'revenue' else EXPENSE_FIELDS)
                fields += [('net', 'Net Tahsilat')] if kind == 'revenue' else [('final', 'Hesaplanan Yıl Sonu Ödeneği')]
                fields += [('rate', 'Gerçekleşme %')]
                if kind == 'expenditure':
                    fields += [('pp_budget', 'PP Bütçesi')]
                # Keep the full data readable: at most six measures per table,
                # with the identifying columns repeated on each continuation.
                for start in range(0, len(fields), 6):
                    group = fields[start:start + 6]
                    headers = ['Müdürlük', 'Sınıflandırma', *[label for _, label in group]]
                    rows = [[row.get('department'), row.get('classification', row.get('economic_code')),
                             *[format_number(row.get(key), money=key != 'rate') for key, _ in group]]
                            for row in table['rows']]
                    self.table(headers, rows, title)
                if table['rows']:
                    self.table(['Müdürlük', 'Sınıflandırma', 'Açıklama', 'Birim Açıklaması'], [
                        [row.get('department'), row.get('classification', row.get('economic_code')),
                         row.get('explanation'), row.get('birim_explanation')] for row in table['rows']],
                        title + ' · Açıklamalar')
                self.paragraph('Toplamlar (boş değerler toplamın dışında tutulur)', CAPTION)
                self.table(['Kalem', 'Toplam', 'Boş değer'], [
                    [label, format_number(table['totals'][key]['value'], money=key != 'rate'),
                     table['totals'][key]['missing']] for key, label in fields if key in table.get('totals', {})])
            policy = data.get('policy', {})
            self.paragraph(f"Sapma eşikleri: %{format_number(policy.get('lower_rate'))} – %{format_number(policy.get('upper_rate'))}", NOTE)
        for kind, slot, title in [('assets', 'III.A.2', 'Varlık ve Yükümlülükler'),
                                  ('audit', 'III.A.3', 'Malî Denetim Kayıtları'),
                                  ('aid', 'III.A.4', 'Yardım Yapılan Kuruluşlar'),
                                  ('explanations', 'III.A.1', 'Birim Sapma Açıklamaları')]:
            if code != slot:
                continue
            for suffix, extra in [('', []), ('_directorates', [('department', 'Müdürlük')])]:
                rows = data.get('entries', {}).get(kind + suffix, [])
                if rows:
                    columns = [*extra, *ENTRY_FIELDS[kind]]
                    self.paragraph(title + (' · Müdürlük Kayıtları' if suffix else ''), CAPTION)
                    self.table([label for _, label in columns], [
                        [format_number(row.get(key), money=True) if key in {'amount', 'previous'} else row.get(key)
                         for key, _ in columns] for row in rows])
        if code == 'III.A.2':
            for statement in data.get('statements', []):
                self.paragraph(statement['name'], CAPTION)
                self.paragraph(f"Dönem: {_text(statement['period'])} · Tarih: {_text(statement['statement_date'])}")
                self.paragraph(statement['explanation'])
                self.paragraph(f"Kaynak dosya: {_text(statement['filename'])} · SHA-256: {_text(statement['sha256'])}", NOTE)
                if statement['figures']:
                    self.table(['Özet Kalem', 'Tutar'], [[row['label'], format_number(row['amount'], money=True)]
                                                          for row in statement['figures']])

    def core(self):
        data = self.payload.get('core', {})
        self.paragraph(data.get('message', 'Gösterge sonucu bulunmuyor.'), NOTE)
        rows = []
        for item in data.get('indicators', []):
            targets = '; '.join(format_number(target['value']) if target.get('state') == 'verified'
                                else _text(target.get('raw_text')) for target in item.get('targets', []))
            actual = format_number(item.get('actual')) if item.get('actual_state') == 'verified' else _text(item.get('raw_text'))
            evidence = '; '.join(' · '.join(str(source.get(key, '')) for key in ('title', 'locator', 'source_revision'))
                                 for source in item.get('evidence', []))
            rows.append([item['name'], item.get('unit'), targets, actual,
                         '\n'.join(filter(None, [item.get('data_source'), evidence]))])
        self.table(['Gösterge', 'Ölçü Birimi', 'Hedef', 'Gerçekleşme', 'Kaynak'], rows)
        for item in data.get('indicators', []):
            if item.get('explanation'):
                self.paragraph(f"{item['name']}: {item['explanation']}")
            deviation = item.get('deviation', {})
            if deviation.get('delta') is not None:
                self.paragraph(f"{item['name']} · Hedeften fark: {format_number(deviation['delta'])}")
            if deviation.get('corrective_action'):
                self.paragraph(f"Düzeltici faaliyet: {deviation['corrective_action']}")

    def assurances(self, items):
        for item in items:
            statement = item['rendered']
            if statement.get('error'):
                self.paragraph(statement['error'], NOTE)
                continue
            self.paragraph(f"{item['code']} · {statement['title']}", CAPTION)
            self.paragraph(statement['heading'], CAPTION)
            for text in statement['paragraphs']:
                if text.endswith('(Yer-Tarih)'):
                    text = text[:-11] + f"({statement['variables']['place']} - {statement['date_display']})"
                self.paragraph(text)
            for label, value in zip(statement['signature_labels'], ['________________', statement['name'], statement['title_value']]):
                self.paragraph(f'{label}: {_text(value)}')
            if statement['variables']['signatories']:
                self.table(['Ad-Soyad', 'Ünvan', 'Göreve Başlangıç', 'Görev Bitişi'], [
                    [person.get(key) for key in ('name', 'title', 'start_date', 'end_date')]
                    for person in statement['variables']['signatories']], 'Görev Dönemleri')
            if statement.get('reservations_note'):
                self.paragraph(statement['reservations_note'], NOTE)
                for reservation in statement['variables']['reservations']:
                    self.paragraph(f"{reservation['text']}\nGerekçe: {reservation['justification']}", 'Civensa Madde')
            for footnote in statement.get('footnotes', []):
                self.paragraph(footnote, NOTE)


class _Report(_ReportContent):
    def __init__(self, payload, template_path):
        self.payload = payload
        self.doc = Document(template_path)
        for name in STYLES:
            if name not in self.doc.styles:
                raise ValueError(f'Word şablonunda “{name}” stili eksik. Şablonu bu stille kaydedin.')
        # The template owns styles and page setup; its sample body is replaced.
        for node in list(self.doc._element.body):
            if node.tag != qn('w:sectPr'):
                self.doc._element.body.remove(node)
        section = self.doc.sections[0]
        self.portrait = (min(section.page_width, section.page_height),
                         max(section.page_width, section.page_height))
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = self.portrait
        self.doc.settings.odd_and_even_pages_header_footer = False
        section.different_first_page_header_footer = True
        for part in (section.first_page_header, section.first_page_footer):
            for node in list(part._element):
                part._element.remove(node)
            part.add_paragraph(style=BODY)
        header = section.header.paragraphs[0]
        header.clear()
        header.style = NOTE
        if not payload['frozen']:
            header.add_run(DRAFT)
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.style = NOTE
        if payload['frozen']:
            footer.add_run(f"Onaylı kopya · SHA-256: {payload['sha256'][:16]} · ")
        footer.add_run('Sayfa ')
        _field(footer, 'PAGE', '1')
        update = self.doc.settings.element.find(qn('w:updateFields'))
        if update is not None:
            self.doc.settings.element.remove(update)

    def paragraph(self, text, style=BODY):
        return self.doc.add_paragraph('' if text == '' else _text(text), style)

    def page_break(self):
        self.paragraph('').add_run().add_break(WD_BREAK.PAGE)

    def toc(self, sections):
        self.paragraph("Sayfa numaraları için Word'de içindekilere sağ tıklayıp Alanı Güncelleştir seçin.", NOTE)
        # A complex field can contain a multi-paragraph cached result. fldSimple
        # cannot hold these TOC paragraphs reliably in Microsoft Word.
        paragraph = self.paragraph('')
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        paragraph.add_run()._r.append(begin)
        instruction = OxmlElement('w:instrText')
        instruction.set(qn('xml:space'), 'preserve')
        instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
        paragraph.add_run()._r.append(instruction)
        separate = OxmlElement('w:fldChar')
        separate.set(qn('w:fldCharType'), 'separate')
        paragraph.add_run()._r.append(separate)
        for section in sections:
            level = section['depth'] + 1
            if level <= 3:
                paragraph = self.paragraph(section['title'], f'TOC {level}')
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        paragraph.add_run()._r.append(end)

    def orientation(self, landscape):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.doc.paragraphs[-1].style = BODY
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        width, height = self.portrait
        section.page_width, section.page_height = (height, width) if landscape else (width, height)
        section.different_first_page_header_footer = False
        return section

    def table(self, headers, rows, title=None):
        if not headers:
            return
        wide = len(headers) > 6
        section = self.orientation(True) if wide else self.doc.sections[-1]
        if title:
            self.paragraph(title, CAPTION)
        available = Length(section.page_width - section.left_margin - section.right_margin)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Civensa Tablo'
        table.autofit = False
        # Set both the grid and cell widths; Word otherwise retains the default grid.
        width = Twips(available.twips // len(headers))
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:w'), str(width.twips * len(headers)))
        table._tbl.tblPr.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
        for column in table.columns:
            column.width = width
        repeat = OxmlElement('w:tblHeader')
        repeat.set(qn('w:val'), 'true')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, row in enumerate([headers, *rows]):
            cells = table.rows[0].cells if index == 0 else table.add_row().cells
            for cell, value in zip(cells, row):
                cell.width = width
                cell.text = _text(value)
                for paragraph in cell.paragraphs:
                    paragraph.style = CAPTION if index == 0 else BODY
        if wide:
            self.orientation(False)


def _render_activity_report(report):
    payload = report.payload
    report.paragraph(str(payload['year']), 'Civensa Kapak Başlık')
    report.paragraph('FAALİYET RAPORU', 'Civensa Kapak Başlık')
    report.paragraph(payload.get('organization'), 'Civensa Kapak Başlık')
    if not payload.get('organization'):
        report.paragraph('Kurum adı bu onaylı kopyada saklanmamış. Ayrıntı için Strateji ile iletişime geçin.', NOTE)
    if payload.get('department_id'):
        report.paragraph(payload['department'], 'Civensa Kapak Başlık')
    report.page_break()
    report.paragraph('İçindekiler', 'Civensa Kapak Başlık')
    sections = sorted(payload['sections'], key=lambda row: row['position'])
    report.toc(sections)
    report.page_break()
    for section in sections:
        report.paragraph(section['title'], f"Civensa Başlık {min(section['depth'] + 1, 4)}")
        if not payload['frozen'] and section.get('missing'):
            report.paragraph('Eksikler:\n' + '\n'.join('• ' + message for message in section['missing']), NOTE)
        if section['status'] == 'not_applicable':
            report.paragraph(f"Uygulanamaz: {section['reason']}")
            continue
        report.blocks(section['blocks'])
        report.resources(section['code'])
        report.financial(section['code'])
        if section['code'] == 'III.B.2.i':
            report.core()
        report.assurances(section.get('assurances', []))


def build_activity_report_docx(payload, template_path) -> bytes:
    report = _Report(payload, template_path)
    _render_activity_report(report)
    output = BytesIO()
    report.doc.save(output)
    # OPC parts are deterministic; ZIP's wall-clock member timestamps are not.
    # Normalize the archive so retained releases download identically later.
    stable = BytesIO()
    with ZipFile(output) as source, ZipFile(stable, 'w', compression=ZIP_DEFLATED) as target:
        for part in source.infolist():
            part.date_time = (1980, 1, 1, 0, 0, 0)
            target.writestr(part, source.read(part.filename))
    return stable.getvalue()
