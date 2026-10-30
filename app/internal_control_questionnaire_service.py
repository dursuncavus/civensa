"""CIV-085: uploaded questions, directorate-owned snapshots, CIV-079 workflow."""
import json
from io import BytesIO
from zipfile import BadZipFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from . import internal_control_monitoring_service as m
from .internal_control_models import COMPONENTS
from .internal_control_questionnaire_models import (ICQuestionList as QuestionList,
    ICQuestionAttachment as Attachment, ICQuestionBatch as Batch, ICQuestionEvent as Event)
from .models import Department
from .pp_errors import pp_error

ANSWERS = {'yes': 'Evet', 'no': 'Hayır', 'developing': 'Geliştirilmekte'}
HEADERS = ['Sıra', 'Bileşen', 'Kod', 'Soru']
SCORING_ASSUMPTION = '2022 uygulaması — Strateji teyit edecek'


def template():
    book = Workbook()
    book.active.title = 'Sorular'
    book.active.append(HEADERS)
    for col, width in zip('ABCD', (10, 28, 20, 90)):
        book.active.column_dimensions[col].width = width
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def parse_questions(content):
    if len(content) > 5 * 1024 * 1024:
        m.fail('Satır 1: Soru dosyası 5 MB sınırını aşıyor. En çok 200 soruluk Excel şablonunu yükleyin.', 'file')
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (ValueError, BadZipFile, OSError, KeyError):
        m.fail('Satır 1: Excel dosyası okunamadı. Soru şablonunu .xlsx olarak yükleyin.', 'file')
    errors, result, codes, orders = [], [], set(), set()
    try:
        rows = book.active.iter_rows(values_only=True)
        if list(next(rows, ()))[:4] != HEADERS:
            errors.append('Satır 1: Başlıklar geçersiz. Sıra | Bileşen | Kod | Soru şablonunu kullanın.')
        aliases = {key.casefold(): key for key in COMPONENTS}
        aliases.update({value.casefold(): key for key, value in COMPONENTS.items()})
        for line, row in enumerate(rows, 2):
            raw = list(row[:4]) + [None] * max(0, 4 - len(row))
            if all(v is None or str(v).strip() == '' for v in raw):
                continue
            if len(result) == 200:
                errors.append(f'Satır {line}: 200 soru sınırı aşıldı. Dosyayı en çok 200 soruya indirin.')
                break
            order, component, code, question = [str(v).strip() if v is not None else '' for v in raw]
            # Excel numeric cells containing whole numbers are accepted.
            if not order.isascii() or not order.isdecimal() or not 1 <= int(order) <= 999999:
                errors.append(f'Satır {line}: Sıra geçersiz. Pozitif tam sayı girin.')
                order = 0
            else:
                order = int(order)
            if order in orders:
                errors.append(f'Satır {line}: Sıra {order} yineleniyor. Benzersiz sıra girin.')
            if not code or len(code) > 60 or code.casefold() in codes:
                errors.append(f'Satır {line}: Kod boş, uzun veya yineleniyor ({code}). En çok 60 karakterlik benzersiz kod girin.')
            component = aliases.get(component.casefold())
            if not component:
                errors.append(f'Satır {line}: Bileşen tanınmıyor. KOS/RDS/KFS/BİS/İS veya katalog adını seçin.')
            if not question or len(question) > 10000:
                errors.append(f'Satır {line}: Soru boş veya çok uzun. 1–10.000 karakterlik soru yazın.')
            if any(isinstance(v, str) and v.startswith('=') for v in raw):
                errors.append(f'Satır {line}: Formül içeren alan okunamaz. Alanları düz metin olarak girin.')
            result.append(dict(order=order, component=component, code=code, question=question))
            orders.add(order)
            codes.add(code.casefold())
        if not result:
            errors.append('Satır 2: Soru listesi boş. Şablona en az bir soru ekleyin.')
    finally:
        book.close()
    if errors:
        raise pp_error(422, errors[0], 'file', messages=errors)
    return sorted(result, key=lambda r: r['order'])


def save_list(db, actor, name, content):
    actor = m.lock_actor(db, m.require_access(db, actor, strategy=True), strategy=True)
    name = str(name).strip()
    if not 1 <= len(name) <= 200:
        m.fail('Soru listesi adı boş veya uzun. 1–200 karakterlik ad girin.', 'name')
    questions = parse_questions(content)
    record = QuestionList(name=name, questions_json=m.auth.encode(questions), actor_id=actor.id,
        actor_name=actor.full_name, created_at=m.auth.server_now(db))
    db.add(record)
    db.flush()
    return record


def attachment(db, period_id):
    return db.scalar(select(Attachment).where(Attachment.period_id == period_id).order_by(Attachment.revision.desc()).limit(1))


def questions(db, period_id):
    record = attachment(db, period_id)
    return json.loads(db.get(QuestionList, record.list_id).questions_json) if record else []


def batches(db, period_id):
    record = attachment(db, period_id)
    if not record:
        return []
    return db.scalars(select(Batch).where(Batch.period_id == period_id,
        Batch.department_id.in_(json.loads(record.departments_json))).order_by(Batch.label)).all()


def attach(db, actor, period_id, list_id, department_ids, expected_revision):
    m.require_strategy_director(m.require_access(db, actor, strategy=True))
    period = m.get_period(db, actor, period_id, lock=True, writable=True)
    actor = m.lock_actor(db, actor, strategy=True)
    m.require_strategy_director(actor)
    old = attachment(db, period.id)
    if m.number(expected_revision) != (old.revision if old else 0):
        m.conflict()
    if not db.get(QuestionList, list_id):
        m.fail('Soru listesi bulunamadı. Kaydedilmiş bir liste seçin.', 'list_id')
    selected = sorted({m.number(v, 'department_ids') for v in department_ids if str(v).strip()})
    if not selected:
        m.fail('Katılımcı birim seçilmedi. En az bir müdürlük seçin.', 'department_ids')
    previous = set(json.loads(old.departments_json)) if old else set()
    saved = db.scalar(select(Event.id).join(Batch).where(Batch.period_id == period.id, Event.action == 'save').limit(1))
    if saved and (old.list_id != list_id or not previous.issubset(selected)):
        m.fail('Bu dönemde yanıt kaydedilmiş. Soru listesini ve katılımcıları koruyun; yalnız yeni müdürlük ekleyebilirsiniz.', 'list_id', 409)
    departments = {d.id: d for d in db.scalars(select(Department).where(Department.id.in_(selected)).with_for_update())}
    if any(did not in departments or (did not in previous and not departments[did].is_active) for did in selected):
        m.fail('Katılımcı müdürlük bulunamadı veya etkin değil. Etkin müdürlükleri seçin.', 'department_ids')
    record = Attachment(period_id=period.id, list_id=list_id, revision=(old.revision if old else 0)+1,
        departments_json=m.auth.encode(selected), actor_id=actor.id, actor_name=actor.full_name,
        created_at=m.auth.server_now(db))
    db.add(record)
    db.flush()
    existing = set(db.scalars(select(Batch.department_id).where(Batch.period_id == period.id)))
    for did in selected:
        if did not in existing:
            db.add(Batch(period_id=period.id, department_id=did, label=departments[did].name))
    db.flush()
    return record


def latest_event(db, batch_id):
    return db.scalar(select(Event).where(Event.batch_id == batch_id).order_by(Event.revision.desc()).limit(1))


def state(db, batch_id):
    last = latest_event(db, batch_id)
    return (last.to_status, last.revision) if last else ('draft', 0)


def get_batch(db, actor, batch_id, *, lock=False, expected_revision=None):
    actor = m.require_access(db, actor)
    batch = db.get(Batch, batch_id)
    if not batch:
        m.fail('Soru formu bulunamadı. Dönem ekranından formu seçin.', status=404)
    if actor.role == 'department' and batch.department_id != actor.department_id:
        m.fail('Bu soru formuna erişiminiz yok. Kendi müdürlüğünüzün formunu açın.', status=403)
    period = m.get_period(db, actor, batch.period_id, lock=lock, writable=lock)
    if lock:
        db.scalar(select(Batch).where(Batch.id == batch_id).with_for_update())
        actor = m.lock_actor(db, actor)
        if actor.role == 'department' and batch.department_id != actor.department_id:
            m.fail('Müdürlük yetkiniz değişti. Kendi müdürlüğünüzün formunu açın.', status=403)
    attached = attachment(db, period.id)
    if not attached or batch.department_id not in json.loads(attached.departments_json):
        m.fail('Müdürlük bu dönemin katılımcısı değil. Dönem ekranını kontrol edin.', status=404)
    status, revision = state(db, batch_id)
    if expected_revision is not None and m.number(expected_revision) != revision:
        m.conflict()
    return actor, period, batch, status, revision


def rows(db, batch):
    last = latest_event(db, batch.id)
    answers = {r['code']: r for r in json.loads(last.snapshot_json)} if last else {}
    return [{**q, **answers.get(q['code'], dict(answer=None, explanation=''))} for q in questions(db, batch.period_id)]


def event(db, actor, batch, revision, action, origin, destination, content, reason=''):
    record = Event(batch_id=batch.id, revision=revision+1, action=action, from_status=origin,
        to_status=destination, actor_id=actor.id, actor_name=actor.full_name, reason=reason,
        snapshot_json=m.auth.encode(content), created_at=m.auth.server_now(db))
    db.add(record)
    db.flush()
    return record


def save_answers(db, actor, batch_id, expected_revision, data, *, submit=False, expected_list_id=None):
    actor, period, batch, status, revision = get_batch(db, actor, batch_id, lock=True, expected_revision=expected_revision)
    if expected_list_id is not None and m.number(expected_list_id, 'list_id') != attachment(db, period.id).list_id:
        m.conflict()
    if actor.role != 'department' or actor.department_id != batch.department_id:
        m.fail('Soru formunu yalnız ilgili müdürlük doldurabilir. Kendi müdürlüğünüzün hesabını kullanın.', status=403)
    if m.today(db) > period.deadline:
        m.fail('İzleme giriş süresi doldu. Son giriş tarihini Strateji Geliştirme birimiyle kontrol edin.', status=409)
    if status not in ('draft', 'returned'):
        m.fail('Soru formu incelemede veya onaylı. Strateji biriminden gerekçeli iade ya da yeniden açma isteyin.', status=409)
    content = []
    for q in questions(db, period.id):
        answer = str(data.get('answer_' + q['code'], '')).strip() or None
        explanation = str(data.get('explanation_' + q['code'], '')).strip()
        if answer is not None and answer not in ANSWERS:
            m.fail(f"{q['code']} yanıtı geçersiz. Evet, Hayır veya Geliştirilmekte seçin.", 'answer_' + q['code'])
        if len(explanation) > 1000:
            m.fail(f"{q['code']} açıklaması 1.000 karakteri aşıyor. Açıklamayı kısaltın.", 'explanation_' + q['code'])
        content.append(dict(code=q['code'], answer=answer, explanation=explanation))
    if submit:
        m.require_submit(actor)
        missing = [r['code'] for r in content if r['answer'] is None]
        if missing:
            m.fail('Yanıtı eksik sorular: ' + ', '.join(missing) + '. Her soru için yanıt seçip yeniden gönderin.', 'request')
    record = event(db, actor, batch, revision, 'save', status, status, content)
    if submit:
        record = transition(db, actor, batch.id, 'submit', record.revision)
    return record


def transition(db, actor, batch_id, action, expected_revision, reason=''):
    actor, period, batch, status, revision = get_batch(db, actor, batch_id, lock=True, expected_revision=expected_revision)
    content = [dict(code=r['code'], answer=r['answer'], explanation=r['explanation']) for r in rows(db, batch)]
    if action == 'submit':
        if actor.role != 'department' or actor.department_id != batch.department_id:
            m.fail('Soru formunu yalnız ilgili müdürlük gönderebilir. Sorumlu müdürlüğün hesabını kullanın.', status=403)
        m.require_submit(actor)
        if m.today(db) > period.deadline:
            m.fail('İzleme gönderim süresi doldu. Son giriş tarihini Strateji birimiyle kontrol edin.', status=409)
        if status not in ('draft', 'returned'):
            m.conflict()
        missing = [r['code'] for r in content if r['answer'] is None]
        if missing:
            m.fail('Yanıtı eksik sorular: ' + ', '.join(missing) + '. Her soru için yanıt seçip yeniden gönderin.')
        destination = 'submitted'
    elif action in ('return', 'approve', 'reopen'):
        m.require_strategy_director(actor)
        if status != ('approved' if action == 'reopen' else 'submitted'):
            m.conflict()
        destination = {'return': 'returned', 'approve': 'approved', 'reopen': 'draft'}[action]
    else:
        m.fail('Soru formu işlemi tanınmıyor. Ekrandaki işlem düğmesini kullanın.', 'action')
    reason = str(reason).strip()
    if action in ('return', 'reopen') and not reason:
        m.fail('İade veya yeniden açma gerekçesi eksik. Gerekçe yazıp yeniden deneyin.', 'reason')
    if action == 'approve':
        boundary = db.scalar(select(Event.revision).where(Event.batch_id == batch.id,
            Event.action.in_(['return', 'reopen'])).order_by(Event.revision.desc()).limit(1)) or 0
        if db.scalar(select(Event.id).where(Event.batch_id == batch.id, Event.revision > boundary,
            Event.actor_id == actor.id, Event.action == 'save').limit(1)):
            m.fail(m.SOD_ERROR.format(record='İç Kontrol soru formu'), status=403)
    return event(db, actor, batch, revision, action, status, destination, content, reason)


def results(db, actor, period):
    m.require_access(db, actor, strategy=True)
    selected = batches(db, period.id)
    included, pending, approved, units = [], [], [], []
    for batch in selected:
        last = latest_event(db, batch.id)
        if last and last.to_status == 'approved':
            included.append(batch.label)
            content = json.loads(last.snapshot_json)
            approved.extend(content)
            units.append((batch, content))
        else:
            pending.append(batch.label)
    def tally(label, answers):
        counts = {key: sum(r['answer'] == key for r in answers) for key in ANSWERS}
        return dict(label=label, counts=counts, total=len(answers), rates={k: m.rate_text(v, len(answers)) for k, v in counts.items()})

    def scored(label, answers):
        result = tally(label, answers)
        counts, total = result['counts'], result['total']
        realised = counts['yes'] + counts['developing']
        return dict(result, score=2 * counts['yes'] + counts['developing'] if total else None,
                    max_score=2 * total, realisation_count=realised,
                    realisation_rate=m.rate_text(realised, total))

    qs = questions(db, period.id)
    by_question = [{**tally(q['code'], [r for r in approved if r['code'] == q['code']]), **q} for q in qs]
    by_component = []
    for code, label in COMPONENTS.items():
        codes = {q['code'] for q in qs if q['component'] == code}
        by_component.append(tally(label, [r for r in approved if r['code'] in codes]))
    return dict(questions=by_question, components=by_component, total=scored('Kurum toplamı', approved),
        directorates=[dict(scored(batch.label, content), batch_id=batch.id) for batch, content in units],
        included=included, pending=pending, participant_count=len(selected))
