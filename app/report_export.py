"""Approved department performance report DOCX generation."""
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


KIND_LABELS = {
    "actual": "Gerçekleşme", "target": "Hedef", "budget": "Bütçe",
    "midyear_spend": "Yıl Ortası Harcama", "year_end_estimate": "Yıl Sonu Tahmini",
    "forecast": "Tahmin",
}
FUNDING_LABELS = {"total": "Toplam", "inside": "Bütçe İçi", "outside": "Bütçe Dışı"}


def _text(value):
    return "—" if value is None or str(value).strip() == "" else str(value)


def _number(value, money=False):
    if value is None:
        return "—"
    try:
        number = float(value)
    except (TypeError, ValueError):
        return _text(value)
    formatted = f"{number:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"{formatted} TL" if money else formatted


def _set_cell(cell, value, bold=False):
    cell.text = _text(value)
    for run in cell.paragraphs[0].runs:
        run.bold = bold
        run.font.size = Pt(9)


def _heading(document, title):
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Heading 1"]
    paragraph.add_run(title)
    return paragraph


def build_department_docx(payload: dict) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Pt(42)
    section.bottom_margin = Pt(42)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{payload['period']} PERFORMANS RAPORU")
    run.bold = True
    run.font.size = Pt(18)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(_text(payload.get("department"))).bold = True

    _heading(document, "Performans Göstergeleri")
    indicators = payload.get("indicators") or []
    if not indicators:
        document.add_paragraph("Kayıt bulunmuyor.")
    for index, indicator in enumerate(indicators, 1):
        document.add_heading(f"{index}. {_text(indicator.get('name'))}", level=2)
        meta = document.add_table(rows=3, cols=2)
        meta.style = "Table Grid"
        for row, pair in enumerate([
            ("Program / Alt Program", f"{_text(indicator.get('program'))} / {_text(indicator.get('subprogram'))}"),
            ("Ölçü Birimi", indicator.get("unit")),
            ("Veri Kaynağı", indicator.get("data_source")),
        ]):
            _set_cell(meta.cell(row, 0), pair[0], True); _set_cell(meta.cell(row, 1), pair[1])
        values = sorted(indicator.get("values") or [], key=lambda v: (v.get("year", 0), v.get("kind", "")))
        if values:
            table = document.add_table(rows=1, cols=3); table.style = "Table Grid"
            for col, label in enumerate(["Yıl", "Tür", "Değer"]): _set_cell(table.cell(0, col), label, True)
            for value in values:
                cells = table.add_row().cells
                _set_cell(cells[0], value.get("year")); _set_cell(cells[1], KIND_LABELS.get(value.get("kind"), value.get("kind")))
                _set_cell(cells[2], _number(value.get("value")))
        document.add_paragraph(f"Açıklama: {_text(indicator.get('explanation'))}")
        document.add_paragraph(f"Hesaplama Yöntemi: {_text(indicator.get('calculation_method'))}")

    _heading(document, "Faaliyetler ve Maliyetler")
    activities = payload.get("activities") or []
    if not activities:
        document.add_paragraph("Kayıt bulunmuyor.")
    for index, activity in enumerate(activities, 1):
        document.add_heading(f"{index}. {_text(activity.get('name'))}", level=2)
        document.add_paragraph(f"Program / Alt Program: {_text(activity.get('program'))} / {_text(activity.get('subprogram'))}")
        document.add_paragraph(f"Açıklama: {_text(activity.get('explanation'))}")
        costs = sorted(activity.get("costs") or [], key=lambda v: (v.get("year", 0), v.get("kind", ""), v.get("funding_type", "")))
        table = document.add_table(rows=1, cols=4); table.style = "Table Grid"
        for col, label in enumerate(["Yıl", "Tür", "Finansman", "Tutar"]): _set_cell(table.cell(0, col), label, True)
        if costs:
            for cost in costs:
                cells = table.add_row().cells
                _set_cell(cells[0], cost.get("year")); _set_cell(cells[1], KIND_LABELS.get(cost.get("kind"), cost.get("kind")))
                _set_cell(cells[2], FUNDING_LABELS.get(cost.get("funding_type"), cost.get("funding_type")))
                _set_cell(cells[3], _number(cost.get("amount"), money=True))
        else:
            cells = table.add_row().cells
            _set_cell(cells[0], "—"); _set_cell(cells[1], "—"); _set_cell(cells[2], "—"); _set_cell(cells[3], "—")

    _heading(document, "Personel ve SGK Maliyetleri")
    costs = sorted(payload.get("department_costs") or [], key=lambda v: (v.get("year", 0), v.get("kind", "")))
    table = document.add_table(rows=1, cols=3); table.style = "Table Grid"
    for col, label in enumerate(["Yıl", "Tür", "Tutar"]): _set_cell(table.cell(0, col), label, True)
    if costs:
        for cost in costs:
            cells = table.add_row().cells
            _set_cell(cells[0], cost.get("year")); _set_cell(cells[1], KIND_LABELS.get(cost.get("kind"), cost.get("kind")))
            _set_cell(cells[2], _number(cost.get("amount"), money=True))
    else:
        cells = table.add_row().cells
        _set_cell(cells[0], "—"); _set_cell(cells[1], "—"); _set_cell(cells[2], "—")

    document.add_paragraph()
    footer = document.add_paragraph("Civensa tarafından onaylı verilerden oluşturulmuştur.")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs: run.italic = True; run.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
