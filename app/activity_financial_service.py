"""Financial validation, scoped editing, computed consolidation and XLSX previews."""
import hashlib
import io
import json
from itertools import islice
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from lxml.etree import XMLSyntaxError
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from . import activity_sections_service as sections
from .structure_service import period_name
from .models import (ARBudgetRow, ARFinancialPolicy, ARFinancialEntry, ARFinancialStatement,
    ARFinancialImport, ARDocument, Department, ReportingPeriod, Program, SubProgram, Publication)
from .numbers import parse_number, format_number
from .activity_contract_service import digest

EXPENSE_CODES = {'01':'Personel Giderleri','02':'Sosyal Güvenlik Kurumlarına Devlet Primi Giderleri',
    '03':'Mal ve Hizmet Alım Giderleri','04':'Faiz Giderleri','05':'Cari Transferler',
    '06':'Sermaye Giderleri','07':'Sermaye Transferleri','08':'Borç Verme','09':'Yedek Ödenekler'}
REVENUE_CODES = {'01':'Vergi Gelirleri','03':'Teşebbüs ve Mülkiyet Gelirleri',
    '04':'Alınan Bağış ve Yardımlar ile Özel Gelirler','05':'Diğer Gelirler',
    '06':'Sermaye Gelirleri','08':'Alacaklardan Tahsilat','09':'Red ve İadeler'}
# Level-two codes from the economic classification, never arbitrary text codes.
EXPENSE_LEVEL2 = {'01':range(1,6),'02':range(1,6),'03':range(1,10),'04':range(1,4),
    '05':range(1,10),'06':range(1,10),'07':range(1,3),'08':range(1,3),'09':range(1,10)}
REVENUE_LEVEL2 = {'01':range(1,10),'03':range(1,7),'04':range(1,6),'05':range(1,10),
    '06':range(1,3),'08':range(1,3),'09':range(1,3)}
EXPENSE_FIELDS = [('initial','Başlangıç Ödeneği'),('added','Eklenen Ödenek'),('deducted','Düşülen Ödenek'),
    ('entered_final','Bildirilen Yıl Sonu Ödeneği'),('actual','Harcama (Gerçekleşme)'),
    ('cancelled','İptal Edilen Ödenek'),('carried','Ertesi Yıla Devreden Ödenek')]
REVENUE_FIELDS = [('estimate','Bütçe Tahmini'),('actual','Tahsilat (Gerçekleşme)'),('refunds','Red ve İadeler')]
ENTRY_FIELDS = {
    'assets':[('label','Varlık / yükümlülük'),('amount','Yıl Sonu Tutarı'),('previous','Önceki Yıl Tutarı'),('explanation','Açıklama')],
    'aid':[('name','Kuruluş Adı'),('aid_type','Yardımın Türü'),('amount','Tutar'),('basis','Dayanak'),('explanation','Faaliyet / Sonuç Açıklaması')],
    'audit':[('audit_type','Denetim Türü'),('auditor','Denetçi / Kurum'),('period','Dönem'),('report_date','Rapor Tarihi'),
        ('report_number','Rapor Sayısı'),('findings','Bulgu Özeti'),('measures','Alınan Önlemler'),('status','Durum')],
    'explanations':[('row_key','Bütçe Satırı'),('explanation','Sapma Açıklaması')]}
SLOTS = {'expenditure':'III.A.1','programme':'III.A.1','revenue':'III.A.1','policy':'III.A.1',
    'explanations':'III.A.1','statements':'III.A.2','assets':'III.A.2','audit':'III.A.3','aid':'III.A.4'}
DEFAULT_STATEMENTS = ('Bilanço','Faaliyet Sonuçları Tablosu')
MONEY_FIELDS = {k for k,_ in EXPENSE_FIELDS + REVENUE_FIELDS}
MAX_FILE = 10 * 1024 * 1024


def idare_document(db, doc):
    return db.scalar(select(ARDocument).where(ARDocument.organization_id==doc.organization_id,
        ARDocument.period_id==doc.period_id, ARDocument.scope=='idare'))


def authorize(db, user, doc, kind):
    sections.actor(user)
    if kind not in SLOTS:
        sections.fail('Malî tablo bulunamadı. III-A bölümünden geçerli bir tablo seçin.',404)
    if kind=='policy' and user.role!='strategy':
        sections.fail('Sapma eşiklerini yalnız Strateji belirleyebilir. Strateji ile iletişime geçin.',403)
    allowed = user.role=='strategy' or (not doc.department_id and sections.finance(db,user)) or (
        doc.department_id==user.department_id and kind in {'aid','audit','explanations'})
    if not allowed:
        sections.fail('Bu malî veriyi düzenleyemezsiniz. İdare tutarları için Mali Hizmetler, birim açıklamaları için kendi raporunuzu kullanın.',403)
    if doc.department_id and kind not in {'aid','audit','explanations'}:
        sections.fail('Bütçe tutarları ve temel malî tablolar idare raporunda yönetilir. İdare III-A bölümünü açın.',403)


def start_write(db,user,document_id,kind,revision):
    doc=sections.document(db,user,document_id)
    authorize(db,user,doc,kind)
    sections.writable(db,user,doc,revision)
    return doc


def finish(db,user,doc,kind,before,after,action='financial_save'):
    _,content=sections.section_pair(db,doc,SLOTS[kind])
    content.status='draft'
    content.edited_by=user.id
    from datetime import datetime, UTC
    content.edited_at=datetime.now(UTC)
    doc.revision+=1
    sections.event(db,user,doc,SLOTS[kind],action,before,after)
    db.flush()
    return doc


def text(value,label,required=False):
    if not isinstance(value,str) or len(value)>10000 or (required and not value.strip()):
        sections.fail(f'{label}: Geçerli metin girin; zorunlu alanları doldurun (en fazla 10.000 karakter).')
    return value.strip()


def number(value,label,negative=False):
    try:
        if isinstance(value,bool):raise ValueError()
        result=parse_number(value)
        if result is not None and (abs(result)>=Decimal('1e16') or result!=result.quantize(Decimal('.01')) or (not negative and result<0)):
            raise ValueError()
        return result
    except (ValueError,ArithmeticError):
        sections.fail(f'{label}: Sayı okunamadı veya izin verilmeyen negatif tutar var. En fazla iki ondalıkla 12.500,23 girin; bilinmiyorsa boş bırakın.')


def integer(value,label,optional=False):
    if optional and value in (None,''):return None
    try:
        if isinstance(value,bool) or str(int(value))!=str(value):raise ValueError()
        return int(value)
    except (ValueError,TypeError):sections.fail(f'{label}: Listeden geçerli bir kayıt seçin.')


def code_valid(value,revenue=False):
    levels=REVENUE_LEVEL2 if revenue else EXPENSE_LEVEL2
    parts=value.split('.')
    return (len(parts)==1 and value in levels) or (len(parts)==2 and parts[0] in levels and
        parts[1] in {f'{v:02}' for v in levels[parts[0]]})


def budget_columns(kind):
    ids=[('department_id','Müdürlük No')] if kind!='revenue' else []
    ids+= [('program_id','Program No'),('subprogram_id','Alt Program No')] if kind=='programme' else [('economic_code','Ekonomik Kod')]
    return ids + (REVENUE_FIELDS if kind=='revenue' else EXPENSE_FIELDS) + [('explanation','Sapma Açıklaması')]


def normalize_budget(db,doc,kind,rows,imported=False):
    if kind not in {'expenditure','programme','revenue'} or not isinstance(rows,list) or len(rows)>5000:
        sections.fail('Bütçe tablosu geçersiz. En fazla 5.000 satır kullanın.')
    result=[];errors=[];seen=set();positions={}
    for index,row in enumerate(rows,1):
        try:
            if not isinstance(row,dict) or set(row)-dict(budget_columns(kind)).keys():
                sections.fail('Sütunlar geçersiz. Ekrandaki sütunları veya indirilen şablonu kullanın.')
            item={k:None for k in MONEY_FIELDS}
            item.update(department_id=None,program_id=None,subprogram_id=None,economic_code='')
            if kind!='revenue':
                dep_id=integer(row.get('department_id'),'Müdürlük')
                dep=db.get(Department,dep_id)
                scope=db.scalar(select(ARDocument.id).where(ARDocument.organization_id==doc.organization_id,
                    ARDocument.period_id==doc.period_id,ARDocument.department_id==dep_id))
                if not dep or dep.organization_id!=doc.organization_id or not scope:
                    sections.fail('Bilinmeyen müdürlük. Bu yıl kapsamındaki müdürlük numarasını kullanın.')
                item['department_id']=dep_id
            if kind=='programme':
                pid=integer(row.get('program_id'),'Program')
                sid=integer(row.get('subprogram_id'),'Alt program',True)
                program=db.get(Program,pid);sub=db.get(SubProgram,sid) if sid else None
                if not program or program.period_id!=doc.period_id or (sid and (not sub or sub.program_id!=pid or sub.period_id!=doc.period_id)):
                    sections.fail('Program / alt program bilinmiyor. Rapor yılına ait eşleşen programı seçin.')
                item.update(program_id=pid,subprogram_id=sid)
            else:
                code=str(row.get('economic_code','')).strip()
                if not code_valid(code,kind=='revenue'):
                    sections.fail('Bilinmeyen ekonomik kod. Listeden düzey 1 veya 01.01 biçiminde düzey 2 kod seçin.')
                item['economic_code']=code
            for field,label in (REVENUE_FIELDS if kind=='revenue' else EXPENSE_FIELDS):
                item[field]=number(row.get(field),label)
            item['explanation']=text(row.get('explanation',''),'Sapma açıklaması')
            item['row_key']=f"{item['department_id'] or 0}:{item['economic_code']}:{item['program_id'] or 0}:{item['subprogram_id'] or 0}"
            if imported and kind!='revenue':
                computed=calculated(item,kind)['final']
                if computed is not None and computed<0:sections.fail(f'Yıl Sonu Ödeneği {format_number(computed)} negatif. Başlangıç, eklenen ve düşülen tutarları düzeltin.')
                if computed is not None and item['entered_final'] is not None and item['entered_final']!=computed:
                    sections.fail(f"Yıl Sonu Ödeneği tutarsız: hesaplanan {format_number(computed)}, bildirilen {format_number(item['entered_final'])}, fark {format_number(item['entered_final']-computed)}. Dosyadaki ödenekleri düzeltin.")
            if item['row_key'] in seen:sections.fail('Yinelenen satır. Aynı müdürlük / kod veya program satırını tek satırda birleştirin.')
            seen.add(item['row_key']);positions[item['row_key']]=index;result.append(item)
        except HTTPException as exc:
            errors.append({'row':index,'message':f'{index}. satır: {exc.detail}'})
    # Reject overlapping roll-up levels without counting a parent twice.
    dimensions={(item['department_id'],item['economic_code'],item['program_id'],item['subprogram_id']) for item in result}
    for item in result:
        index=positions[item['row_key']]
        code=item['economic_code'];dep=item['department_id']
        if '.' in code and (dep,code.split('.')[0],None,None) in dimensions:
            errors.append({'row':index,'message':f"{code}: Üst kod {code.split('.')[0]} da girilmiş. Çift toplamı önlemek için yalnız bir sınıflandırma düzeyi kullanın."})
        if kind=='programme' and item['subprogram_id'] is not None and (dep,'',item['program_id'],None) in dimensions:
            errors.append({'row':index,'message':'Program toplamı ve alt program satırları birlikte girilmiş. Yalnız bir düzeyi kullanın.'})
    return result,errors


def budget_raw(row):
    return {k:getattr(row,k) for k in ['row_key','department_id','economic_code','program_id','subprogram_id','explanation',*sorted(MONEY_FIELDS)]}


def stored_rows(db,doc,kind):
    return db.scalars(select(ARBudgetRow).where(ARBudgetRow.document_id==doc.id,ARBudgetRow.kind==kind).order_by(ARBudgetRow.row_key)).all()


def save_budget(db,user,document_id,kind,rows,expected_revision):
    doc=start_write(db,user,document_id,kind,expected_revision)
    normalized,errors=normalize_budget(db,doc,kind,rows)
    if errors:sections.fail('\n'.join(e['message'] for e in errors))
    existing=stored_rows(db,doc,kind);before=[budget_raw(r) for r in existing]
    by_key={r.row_key:r for r in existing}
    for item in normalized:
        record=by_key.pop(item['row_key'],None)
        if record:
            for key,value in item.items():setattr(record,key,value)
            # Corrections keep import provenance, so a later import still previews replacement.
        else:db.add(ARBudgetRow(document_id=doc.id,kind=kind,**item))
    for record in by_key.values():db.delete(record)
    return finish(db,user,doc,kind,before,normalized)


def entries(db,doc,kind):
    return [json.loads(r.payload_json,parse_float=Decimal) for r in db.scalars(select(ARFinancialEntry).where(
        ARFinancialEntry.document_id==doc.id,ARFinancialEntry.kind==kind).order_by(ARFinancialEntry.position))]


def save_entries(db,user,document_id,kind,rows,expected_revision):
    doc=start_write(db,user,document_id,kind,expected_revision)
    if kind not in ENTRY_FIELDS or not isinstance(rows,list) or len(rows)>2000:
        sections.fail('Geçersiz malî kayıt tablosu. En fazla 2.000 satır kullanın.')
    normalized=[];seen=set()
    source=idare_document(db,doc)
    own_keys={r.row_key for r in db.scalars(select(ARBudgetRow).where(ARBudgetRow.document_id==source.id,
        ARBudgetRow.department_id==doc.department_id))} if doc.department_id else set()
    for i,row in enumerate(rows,1):
        if not isinstance(row,dict) or set(row)-dict(ENTRY_FIELDS[kind]).keys():sections.fail(f'{i}. satır: Geçerli sütunları kullanın.')
        item={}
        for key,label in ENTRY_FIELDS[kind]:
            item[key]=number(row.get(key),f'{i}. satır, {label}',negative=kind=='assets') if key in {'amount','previous'} else text(row.get(key,''),f'{i}. satır, {label}',required=True)
        if kind=='audit':
            if item['audit_type'] not in {'iç denetim','dış denetim (Sayıştay)'}:sections.fail(f'{i}. satır: İç denetim veya dış denetim (Sayıştay) seçin.')
            validate_date(item['report_date'])
        if kind=='explanations':
            if item['row_key'] not in own_keys:sections.fail(f'{i}. satır: Kendi müdürlüğünüze ait bütçe satırını seçin.',403)
            if item['row_key'] in seen:sections.fail(f'{i}. satır: Sapma açıklaması yinelenmiş. Her bütçe satırını bir kez kullanın.')
            seen.add(item['row_key'])
        normalized.append(item)
    before=entries(db,doc,kind)
    for r in db.scalars(select(ARFinancialEntry).where(ARFinancialEntry.document_id==doc.id,ARFinancialEntry.kind==kind)):db.delete(r)
    for position,item in enumerate(normalized):db.add(ARFinancialEntry(document_id=doc.id,kind=kind,position=position,payload_json=sections.numeric_canonical(item)))
    return finish(db,user,doc,kind,before,normalized)


def save_policy(db,user,document_id,lower_rate,upper_rate,expected_revision):
    doc=start_write(db,user,document_id,'policy',expected_revision)
    lower=number(lower_rate,'Alt eşik');upper=number(upper_rate,'Üst eşik')
    if any(v is not None and v>=Decimal('1e10') for v in (lower,upper)):sections.fail('Sapma eşiği çok büyük. 10 milyardan küçük bir oran girin.')
    if lower is not None and upper is not None and lower>upper:sections.fail('Alt eşik üst eşikten büyük. Eşik sırasını düzeltin.')
    policy=db.scalar(select(ARFinancialPolicy).where(ARFinancialPolicy.document_id==doc.id))
    before={'lower_rate':policy.lower_rate,'upper_rate':policy.upper_rate} if policy else None
    if not policy:policy=ARFinancialPolicy(document_id=doc.id);db.add(policy)
    policy.lower_rate=lower;policy.upper_rate=upper
    return finish(db,user,doc,'policy',before,{'lower_rate':lower,'upper_rate':upper})


def validate_date(value):
    try:
        if date.fromisoformat(value).isoformat()!=value:raise ValueError()
    except (ValueError,TypeError):sections.fail('Rapor tarihi geçersiz. Yıl-ay-gün biçiminde geçerli bir tarih seçin.')


def validate_xlsx(data):
    if len(data)>MAX_FILE:sections.fail('Dosya 10 MB sınırını aşıyor. Daha küçük bir dosya yükleyin.')
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist())>50*1024*1024:raise ValueError()
        return load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
    except (BadZipFile,ValueError,KeyError,TypeError,OSError,ParseError,XMLSyntaxError,InvalidFileException):
        sections.fail('XLSX dosyası okunamadı veya çok büyük. İndirilen şablonu kullanarak yeniden yükleyin.')


def seed_statements(db,doc):
    if not doc.department_id:
        for name in DEFAULT_STATEMENTS:db.add(ARFinancialStatement(document_id=doc.id,name=name))


def statement_data(row):
    return {'id':row.id,'name':row.name,'period':row.period,'statement_date':row.statement_date,
        'explanation':row.explanation,'figures':json.loads(row.figures_json,parse_float=Decimal),
        'filename':row.filename,'sha256':row.sha256}


def save_statement(db,user,document_id,data,expected_revision,file_data=None,filename=''):
    doc=start_write(db,user,document_id,'statements',expected_revision)
    name=text(data.get('name',''),'Tablo adı',True)
    if len(name)>200:sections.fail('Tablo adı çok uzun. En fazla 200 karakter kullanın.')
    row=db.scalar(select(ARFinancialStatement).where(ARFinancialStatement.document_id==doc.id,ARFinancialStatement.name==name))
    before=statement_data(row) if row else None
    if not row:row=ARFinancialStatement(document_id=doc.id,name=name,period='',statement_date='',explanation='',figures_json='[]',filename='',sha256='');db.add(row)
    row.period=text(data.get('period',''),'Dönem')
    if len(row.period)>200:sections.fail('Dönem çok uzun. En fazla 200 karakter kullanın.')
    row.statement_date=text(data.get('statement_date',''),'Tablo tarihi')
    if row.statement_date:validate_date(row.statement_date)
    row.explanation=text(data.get('explanation',''),'Açıklamalar')
    figures=data.get('figures',[])
    if not isinstance(figures,list) or len(figures)>100:sections.fail('Özet rakamlar geçersiz. En fazla 100 etiket ve tutar girin.')
    if any(not isinstance(f,dict) or set(f)-{'label','amount'} for f in figures):sections.fail('Özet rakam alanları geçersiz. Etiket ve tutar kullanın.')
    row.figures_json=sections.numeric_canonical([{'label':text(f.get('label',''),'Özet rakam etiketi',True),
        'amount':number(f.get('amount'),'Özet tutar',negative=True)} for f in figures])
    if file_data is not None:
        ext=Path(filename).suffix.lower()
        if len(file_data)>MAX_FILE or not file_data:sections.fail('Dosya boş veya 10 MB sınırını aşıyor. PDF veya XLSX yükleyin.')
        if ext=='.xlsx':validate_xlsx(file_data).close()
        elif ext!='.pdf' or not file_data.startswith(b'%PDF-'):sections.fail('Dosya biçimi geçersiz. PDF veya XLSX yükleyin.')
        row.filename=Path(filename).name[:200];row.file_data=file_data;row.sha256=hashlib.sha256(file_data).hexdigest()
    return finish(db,user,doc,'statements',before,statement_data(row))


def calculated(item,kind):
    data=dict(item)
    if kind=='revenue':
        data['net']=item['actual']-item['refunds'] if item['actual'] is not None and item['refunds'] is not None else None
        numerator=data['net'];denominator=item['estimate']
    else:
        data['final']=item['initial']+item['added']-item['deducted'] if all(item[k] is not None for k in ('initial','added','deducted')) else None
        numerator=item['actual'];denominator=data['final']
    data['rate']=(numerator/denominator*100).quantize(Decimal('.01')) if numerator is not None and denominator is not None and denominator>0 else None
    return data


def totals(rows,kind,missing_sources=0):
    fields=[k for k,_ in (REVENUE_FIELDS if kind=='revenue' else EXPENSE_FIELDS) if k!='entered_final']
    fields+=['net'] if kind=='revenue' else ['final']
    result={}
    for key in fields:
        values=[r[key] for r in rows];known=[v for v in values if v is not None]
        result[key]={'value':sum(known,Decimal(0)) if known else None,'missing':sum(v is None for v in values)+missing_sources}
    numerator='net' if kind=='revenue' else 'actual'
    denominator='estimate' if kind=='revenue' else 'final'
    a,b=result[numerator],result[denominator]
    result['rate']={'value':(a['value']/b['value']*100).quantize(Decimal('.01')) if not a['missing'] and not b['missing'] and a['value'] is not None and b['value'] is not None and b['value']>0 else None,'missing':max(a['missing'],b['missing'])}
    return result


def pp_references(db,doc):
    year=db.get(ReportingPeriod,doc.period_id).year
    pub=db.scalar(select(Publication).where(Publication.period_id==doc.period_id).order_by(Publication.version.desc()).limit(1))
    if not pub:return {},None
    if digest(pub.snapshot_json)!=pub.sha256:sections.fail('PP yayınının doğrulama kodu uyuşmuyor. Strateji ile yayını kontrol edin.',409)
    payload=json.loads(pub.snapshot_json)
    categories={'personnel':'01','sgk':'02','goods_services':'03','interest':'04','current_transfers':'05',
        'capital':'06','capital_transfers':'07','lending':'08','reserve':'09'}
    result={}
    for department in payload.get('departments',[]):
        for row in department.get('economics',[]):
            code=categories.get(row.get('category'))
            if not code or row.get('year')!=year or row.get('value_kind')!='budget' or row.get('funding_type')!='inside':continue
            key=(department['department_id'],code)
            if row.get('amount') is None or row.get('state','verified')!='verified':result[key]=None
            elif key not in result or result[key] is not None:result[key]=result.get(key,Decimal(0))+Decimal(str(row['amount']))
    return result,pub.id


def budget_label(db,item,period_id):
    dep=db.get(Department,item.get('department_id')) if item.get('department_id') else None
    parts=[period_name(db,period_id,dep) if dep else 'İdare']
    if item.get('economic_code'):parts.append(item['economic_code'])
    else:
        program=db.get(Program,item.get('program_id'))
        sub=db.get(SubProgram,item.get('subprogram_id')) if item.get('subprogram_id') else None
        parts.extend([program.name if program else 'Program',sub.name if sub else 'Programın tamamı'])
    return ' / '.join(parts)


def has_slot_data(financial,code):
    if code=='III.A.1':return any(t['rows'] for t in financial['tables'].values())
    if code=='III.A.2':return bool(financial['entries']['assets'] or any(s['sha256'] or s['explanation'] or s['period'] or s['statement_date'] or s['figures'] for s in financial['statements']))
    if code=='III.A.3':return bool(financial['entries']['audit'])
    if code=='III.A.4':return bool(financial['entries']['aid'])
    return False


def financial_data(db,doc):
    source=idare_document(db,doc)
    policy=db.scalar(select(ARFinancialPolicy).where(ARFinancialPolicy.document_id==source.id))
    result={'source_document_id':source.id,'source_revision':source.revision,'policy':{'lower_rate':policy.lower_rate if policy else None,'upper_rate':policy.upper_rate if policy else None},
        'tables':{},'checks':[],'missing':{k:[] for k in ('III.A.1','III.A.2','III.A.3','III.A.4')},'statements':[],'entries':{},'imports':[]}
    refs,pub_id=pp_references(db,doc)
    units=db.scalars(select(ARDocument).where(ARDocument.organization_id==doc.organization_id,ARDocument.period_id==doc.period_id,ARDocument.department_id.is_not(None))).all()
    overrides={}
    for unit in units:
        for entry in entries(db,unit,'explanations'):overrides[entry['row_key']]=entry['explanation']
    for kind in ('expenditure','programme','revenue'):
        if doc.department_id and kind=='revenue':continue
        rows=[]
        for record in stored_rows(db,source,kind):
            if doc.department_id and record.department_id!=doc.department_id:continue
            item=calculated(budget_raw(record),kind)
            item.update(id=record.id,import_id=record.import_id,
                department=period_name(db,doc.period_id,db.get(Department,record.department_id)) if record.department_id else 'İdare',
                pp_budget=refs.get((record.department_id,record.economic_code)),pp_publication_id=pub_id,
                birim_explanation=overrides.get(record.row_key,''))
            rows.append(item)
            label=budget_label(db,item,doc.period_id)
            rate=item['rate'];p=result['policy']
            if rate is not None and ((p['lower_rate'] is not None and rate<p['lower_rate']) or (p['upper_rate'] is not None and rate>p['upper_rate'])) and not (item['explanation'] or item['birim_explanation']):
                result['missing']['III.A.1'].append(f"Sapma açıklaması gerekli: {label}, gerçekleşme %{format_number(rate)}. Satırın sapma açıklamasını doldurun.")
            required=[k for k,_ in (REVENUE_FIELDS if kind=='revenue' else EXPENSE_FIELDS) if k!='entered_final']
            blanks=[label for key,label in (REVENUE_FIELDS if kind=='revenue' else EXPENSE_FIELDS) if key in required and item[key] is None]
            if blanks:result['missing']['III.A.1'].append(f"{label}: {', '.join(blanks)} boş; bilinmiyorsa boş kalır, sıfırsa 0 girin.")
            if kind!='revenue' and item['entered_final'] is not None and item['final'] is not None and item['entered_final']!=item['final']:
                result['checks'].append(f"{label}: Yıl Sonu Ödeneği tutarsız; hesaplanan {format_number(item['final'])}, bildirilen {format_number(item['entered_final'])}, fark {format_number(item['entered_final']-item['final'])}. Başlangıç, eklenen ve düşülen ödeneği kontrol edin.")
            if kind!='revenue' and item['final'] is not None and item['final']<0:result['checks'].append(f"{label}: Yıl sonu ödeneği {format_number(item['final'])} negatif. Düşülen ödeneği kontrol edin.")
        groups={}
        for row in rows:groups.setdefault(row['department_id'],[]).append(row)
        expected=([doc.department_id] if doc.department_id else [u.department_id for u in units]) if kind!='revenue' else []
        missing_deps=[dep for dep in expected if dep not in groups]
        for dep in missing_deps:
            name=period_name(db,doc.period_id,db.get(Department,dep))
            result['missing']['III.A.1'].append(f"{name}: {'Ekonomik kod' if kind=='expenditure' else 'Program / alt program'} gider satırları eksik. Mali Hizmetler ilgili müdürlüğün tutarlarını girmeli; gerçek tutar sıfırsa 0 kullanılmalı.")
            groups[dep]=[]
        result['tables'][kind]={'rows':rows,'totals':totals(rows,kind,len(missing_deps)),'by_department':[
            {'department_id':dep,'department':group[0]['department'] if group else period_name(db,doc.period_id,db.get(Department,dep)),
             'totals':totals(group,kind,int(not group))} for dep,group in groups.items()]}
    for kind in ('expenditure','programme'):
        if not result['tables'][kind]['rows']:result['missing']['III.A.1'].append(f"{'Ekonomik kod' if kind=='expenditure' else 'Program / alt program'} gider tablosu boş. Mali Hizmetler verileri geldiğinde giriş veya XLSX içe aktarma yapın.")
    if not doc.department_id and not result['tables']['revenue']['rows']:result['missing']['III.A.1'].append('Gelir tablosu boş. Gelir ekonomik kodlarına göre tahmin ve tahsilatı girin.')
    exp=result['tables']['expenditure']['rows'];prog=result['tables']['programme']['rows']
    dep_ids=sorted({r['department_id'] for r in exp+prog})
    for dep_id in dep_ids+([None] if not doc.department_id else []):
        er=[r for r in exp if dep_id is None or r['department_id']==dep_id];pr=[r for r in prog if dep_id is None or r['department_id']==dep_id]
        et=totals(er,'expenditure');pt=totals(pr,'programme')
        label='İdare' if dep_id is None else period_name(db,doc.period_id,db.get(Department,dep_id))
        for field in et:
            if field=='rate':continue
            a,b=et[field],pt[field]
            if a['missing'] or b['missing'] or a['value'] is None or b['value'] is None:
                result['checks'].append(f"{label}, {dict(EXPENSE_FIELDS).get(field,'Yıl Sonu Ödeneği')}: ekonomik kod toplamı {format_number(a['value'])}, program toplamı {format_number(b['value'])}. Boş değerler nedeniyle karşılaştırma tamamlanamadı; eksik tutarları girin.")
            elif a['value']!=b['value']:
                result['checks'].append(f"{label}, {dict(EXPENSE_FIELDS).get(field,'Yıl Sonu Ödeneği')}: ekonomik kod toplamı {format_number(a['value'])}, program toplamı {format_number(b['value'])}, fark {format_number(a['value']-b['value'])}. İki sınıflandırmanın satırlarını kontrol edin.")
    result['missing']['III.A.1']+=result['checks']
    for kind in ('assets','aid','audit','explanations'):
        result['entries'][kind]=entries(db,doc,kind)
        if not doc.department_id and kind in {'aid','audit'}:
            result['entries'][kind+'_directorates']=[{'department':period_name(db,doc.period_id,db.get(Department,u.department_id)),**item} for u in units for item in entries(db,u,kind)]
    if not doc.department_id:
        result['statements']=[statement_data(s) for s in db.scalars(select(ARFinancialStatement).where(ARFinancialStatement.document_id==doc.id).order_by(ARFinancialStatement.id))]
        names={s['name'] for s in result['statements']}
        for name in DEFAULT_STATEMENTS:
            if name not in names:result['missing']['III.A.2'].append(f'{name}: Dosya, dönem / tarih ve Açıklamalar girin.')
        for s in result['statements']:
            missing=[label for key,label in [('sha256','dosya'),('explanation','Açıklamalar')] if not s[key]]
            if not (s['period'] or s['statement_date']):missing.append('dönem / tarih')
            for figure in s['figures']:
                if figure['amount'] is None:missing.append(f"{figure['label']} özet tutarı")
            if missing:result['missing']['III.A.2'].append(f"{s['name']}: {', '.join(missing)} eksik. Tablo kaydını tamamlayın.")
        if not result['entries']['assets']:result['missing']['III.A.2'].append('Varlık ve yükümlülük özetini girin veya bölümde gerekçeli Uygulanamaz seçin.')
        for row in result['entries']['assets']:
            if row['amount'] is None or row['previous'] is None:result['missing']['III.A.2'].append(f"{row['label']}: Yıl sonu ve önceki yıl tutarını tamamlayın; sıfırsa 0 girin.")
        result['imports']=[{'id':r.id,'kind':r.kind,'created_at':r.created_at.isoformat(),'sha256':r.source_sha256,
            'applied':bool(db.scalar(select(ARBudgetRow.id).where(ARBudgetRow.import_id==r.id).limit(1))) or bool(db.scalar(select(sections.ARContentEvent.id).where(sections.ARContentEvent.document_id==doc.id,sections.ARContentEvent.action=='financial_import',sections.ARContentEvent.after_json.contains(f'"import_id":{r.id},'))))}
            for r in db.scalars(select(ARFinancialImport).where(ARFinancialImport.document_id==doc.id).order_by(ARFinancialImport.id.desc()))]
    for row in result['entries']['aid']:
        if row['amount'] is None:result['missing']['III.A.4'].append(f"{row['name']}: Yardım tutarı boş. Bilinen tutarı veya sıfırsa 0 girin.")
    if not result['entries']['aid']:result['missing']['III.A.4'].append('Yardım yapılan kuruluşları girin; yardım yoksa Uygulanamaz seçip gerekçe yazın.')
    audits={r['audit_type'] for r in result['entries']['audit']}
    for audit_type in ('iç denetim','dış denetim (Sayıştay)'):
        if audit_type not in audits:result['missing']['III.A.3'].append(f'{audit_type}: Denetim kaydı girin veya bölümün uygulanamaz gerekçesini yazın.')
    return result


def template(kind):
    if kind not in {'expenditure','programme','revenue'}:sections.fail('Şablon bulunamadı. Gider, program veya gelir seçin.',404)
    book=Workbook();sheet=book.active;sheet.title='Malî Bilgiler'
    sheet.append([label for _,label in budget_columns(kind)])
    sheet.freeze_panes='A2'
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width=28
        if col[0].value=='Ekonomik Kod':
            sheet.column_dimensions[col[0].column_letter].number_format='@'
    output=io.BytesIO();book.save(output);return output.getvalue()


def preview_import(db,user,document_id,kind,data,expected_revision):
    doc=start_write(db,user,document_id,kind,expected_revision)
    book=validate_xlsx(data)
    sheet=book.active
    if (sheet.max_row or 0)>5001 or (sheet.max_column or 0)>30:book.close();sections.fail('XLSX çok fazla satır veya sütun içeriyor. En fazla 5.000 veri satırı kullanın.')
    try:
        rows=list(islice(sheet.iter_rows(values_only=True),5002))
    except (BadZipFile,ValueError,KeyError,TypeError,OSError,ParseError,XMLSyntaxError,InvalidFileException):
        sections.fail('XLSX satırları okunamadı. Güncel şablonu indirerek dosyayı yeniden oluşturun.')
    finally:
        book.close()
    if len(rows)>5001:sections.fail('XLSX çok fazla satır içeriyor. En fazla 5.000 veri satırı kullanın.')
    columns=budget_columns(kind)
    if not rows or list(rows[0])!=[label for _,label in columns]:sections.fail('XLSX başlıkları şablonla uyuşmuyor. Güncel Türkçe şablonu indirip sütunları koruyun.')
    incoming=[];line_numbers=[]
    for line,values in enumerate(rows[1:],2):
        if all(v is None for v in values):continue
        incoming.append({key:(value if value is not None else '') for (key,_),value in zip(columns,values)})
        line_numbers.append(line)
    normalized,errors=normalize_budget(db,doc,kind,incoming,imported=True)
    for error in errors:
        error['sheet_row']=line_numbers[error['row']-1]
        prefix=f"{error['row']}. satır: "
        message=error['message'].removeprefix(prefix)
        error['message']=f"XLSX {error['sheet_row']}. satır: "+message
    old=stored_rows(db,doc,kind)
    manual={r.row_key for r in old if r.import_id is None}
    for item in normalized:
        if item['row_key'] in manual:errors.append({'row':None,'message':f"{budget_label(db,item,doc.period_id)}: Elle girilmiş satırla çakışıyor. Elle girilmiş kaydı korumak için dosyadaki satırı kaldırın."})
    imported={r.row_key:budget_raw(r) for r in old if r.import_id is not None}
    new={r['row_key']:r for r in normalized}
    diff=[{'row_key':key,'label':budget_label(db,new.get(key) or imported[key],doc.period_id),'action':'removed' if key not in new else 'added' if key not in imported else 'changed',
        'before':imported.get(key),'after':new.get(key)} for key in sorted(imported.keys()|new.keys()) if imported.get(key)!=new.get(key)]
    preview={'rows':normalized,'errors':errors,'differences':diff}
    record=ARFinancialImport(document_id=doc.id,kind=kind,revision=doc.revision,source_sha256=hashlib.sha256(data).hexdigest(),
        preview_json=sections.numeric_canonical(preview),actor_id=user.id)
    db.add(record);db.flush()
    return {'id':record.id,'kind':kind,**preview}


def import_preview(db,user,document_id,import_id):
    doc=sections.document(db,user,document_id)
    authorize(db,user,doc,'expenditure')
    record=db.get(ARFinancialImport,import_id)
    if not record or record.document_id!=doc.id:sections.fail('İçe aktarma önizlemesi bulunamadı. Dosyayı yeniden yükleyin.',404)
    return record,{'id':record.id,'kind':record.kind,**json.loads(record.preview_json,parse_float=Decimal)}


def apply_import(db,user,document_id,import_id,expected_revision):
    record,preview=import_preview(db,user,document_id,import_id)
    doc=start_write(db,user,document_id,record.kind,expected_revision)
    if record.revision!=doc.revision:sections.fail('Önizlemeden sonra rapor değişti. Farkları tekrar görmek için dosyayı yeniden yükleyin.',409)
    if preview['errors']:sections.fail('Önizlemede hatalar var. Satır hatalarını düzeltip dosyayı yeniden yükleyin.')
    old=stored_rows(db,doc,record.kind)
    before=[budget_raw(r) for r in old if r.import_id is not None]
    if any(r.import_id is None and r.row_key in {i['row_key'] for i in preview['rows']} for r in old):sections.fail('Elle girilen kayıtla çakışma var. Dosyayı düzeltip yeniden önizleyin.',409)
    for row in old:
        if row.import_id is not None:db.delete(row)
    db.flush()
    for item in preview['rows']:db.add(ARBudgetRow(document_id=doc.id,kind=record.kind,import_id=record.id,**item))
    return finish(db,user,doc,record.kind,before,{'import_id':record.id,'rows':preview['rows']},'financial_import')
