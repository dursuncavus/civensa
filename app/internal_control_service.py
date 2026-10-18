"""Excel dry-run and append-only plan catalogue. Callers own commit/rollback."""
from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import PurePath
import re
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import exists, or_, select, update

from .authorization_service import encode, require_internal_control_access as require_access
from .file_security import validate_office_zip
from .import_service import department_key
from .models import Department, Organization
from .internal_control_models import (COMPONENTS, STANDARD_COMPONENTS, ICPlan, ICStandard,
    ICCondition, ICAction, ICActionParty, ICUnitAlias, ICImport, ICImportAudit)

HEADERS = ['Bileşen Kodu', 'Standart Kodu', 'Genel Şart Kodu', 'Genel Şart Metni',
           'Mevcut Durum', 'Eylem Kod No', 'Öngörülen Eylem', 'Sorumlu Birim(ler)',
           'İşbirliği Yapılacak Birim(ler)', 'Çıktı / Sonuç', 'Başlama Tarihi',
           'Tamamlanma Tarihi', 'Açıklama']
TITLE_HEADER = 'Standart Başlığı'
INFO_FIELDS = ['Plan Başlığı', 'Başlangıç Yılı', 'Bitiş Yılı', 'Sürüm', 'Onay Tarihi']
MAX_ROWS = 20000
MAX_UPLOAD = 15 * 1024 * 1024
ALL_DEPARTMENTS = department_key('Tüm Müdürlükler')
BODIES = {department_key(v): v for v in ['Üst Yönetim', 'İç Kontrol İzleme ve Yönlendirme Kurulu', 'Çalışma Grubu']}


def normalize_code(value):
    value = re.sub(r'[\s.]+', '.', str(value).strip().upper().replace('İ', 'I')).strip('.')
    match = re.fullmatch(r'(KOS|RDS|KFS|BIS|IS)\.?([0-9]+(?:\.[0-9]+)*)', value)
    if not match:
        return value
    prefix = {'BIS': 'BİS', 'IS': 'İS'}.get(match[1], match[1])
    return prefix + '.'.join(part.lstrip('0') or '0' for part in match[2].split('.'))


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def split_units(value):
    # A numbered list need not contain whitespace between consecutive names.
    return [name.strip() for name in re.split(r'(?:\d+\s*[-–)]+\s*|[\n;]+)', value) if name.strip()]


def parse_date(raw):
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def issue(level, sheet, row, column, value, message, todo):
    return dict(level=level, sheet=sheet, row=row, column=column, value=value,
                message=message, todo=todo)


def export_book(book):
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def literal_row(sheet, values):
    # Source text in templates/reports must never become executable formulas.
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Plan'
    literal_row(sheet, HEADERS + [TITLE_HEADER])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:N1'
    for col in range(1, 15):
        sheet.column_dimensions[get_column_letter(col)].width = 28
    info = book.create_sheet('Bilgi')
    literal_row(info, ['Alan', 'Değer'])
    for field in INFO_FIELDS:
        literal_row(info, [field, ''])
    instructions = book.create_sheet('Yönerge')
    for line in [
        'Her eylemi Plan sayfasında ayrı satıra yazın. İlk 13 sütunun başlıklarını değiştirmeyin.',
        'Eylemi olmayan genel şart için eylem sütunlarını boş bırakın; Mevcut Durum alanını doldurun.',
        'Kod örnekleri: KOS1, KOS 1.1, KOS 1.1.1. Boşluk ve nokta farkları eşleştirmede dikkate alınmaz.',
        'Birimleri yeni satır, noktalı virgül veya 1-Birim2-Birim şeklinde ayırın.',
        'Tüm Müdürlükler tek ortak sorumluluk olarak kaydedilir; Üst Yönetim gibi kurul adları kullanılabilir.',
        'Tarih: gg.aa.yyyy, yıl, dönem açıklaması veya Sürekli. Özgün tarih metni korunur.',
        'Bilgi sayfasında başlık, yıllar ve sürüm zorunludur. Onay Tarihi isteğe bağlıdır (gg.aa.yyyy).',
        'Standart Başlığı isteğe bağlıdır; başlığı kaynağınızdan aynen aktarın.',
        'Önce doğrulayın, eşleşmeyen birimleri eşleştirin, tekrar doğrulayın ve taslak planı kaydedin.',
        'Yeni sürüm içe aktarmak eski planı silmez. Etkinleştirme önceki etkin planı arşivler.',
    ]:
        literal_row(instructions, [line])
    instructions.column_dimensions['A'].width = 130
    return export_book(book)


def parse_workbook(data, filename):
    source = {'metadata': {}, 'rows': [], 'issues': []}
    def block(sheet, row, column, value, message, todo):
        source['issues'].append(issue('blocking', sheet, row, column, value, message, todo))
    try:
        if PurePath(filename).suffix.lower() != '.xlsx':
            raise ValueError('Dosya uzantısı .xlsx olmalıdır.')
        validate_office_zip(data, '.xlsx')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (HTTPException, ValueError, BadZipFile, KeyError, OSError, ParseError, SyntaxError, EOFError, TypeError, InvalidFileException) as exc:
        block('Dosya', 0, '', filename, f'Excel dosyası okunamadı: {getattr(exc, "detail", str(exc))}',
              'Civensa şablonunu indirip geçerli .xlsx olarak yükleyin.')
        return source
    try:
        for name in ('Plan', 'Bilgi'):
            if name not in book.sheetnames:
                block(name, 1, '', '', f'{name} sayfası bulunamadı.', 'Şablondaki sayfa adını geri yükleyin.')
        if source['issues']:
            return source
        sheet = book['Plan']
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            block('Plan', 0, '', str(sheet.max_row), 'Plan satır sınırını aşıyor.', f'En fazla {MAX_ROWS} veri satırı yükleyin.')
            return source
        rows = sheet.iter_rows(values_only=True, max_col=100)
        headers = [cell_text(v) for v in next(rows, ())]
        for col, expected in enumerate(HEADERS):
            if headers.count(expected) != 1:
                block('Plan', 1, get_column_letter(col + 1), expected, f'{expected} başlığı eksik veya değiştirilmiş.',
                      'Şablondaki başlığı aynen kullanın; yinelenen başlıkları kaldırın.')
        if source['issues']:
            return source
        indexes = [headers.index(h) for h in HEADERS]
        title_idx = headers.index(TITLE_HEADER) if TITLE_HEADER in headers else None
        for num, values in enumerate(rows, 2):
            if num > MAX_ROWS + 1:
                block('Plan', num, '', '', 'Plan satır sınırını aşıyor.', f'En fazla {MAX_ROWS} veri satırı yükleyin.')
                break
            if not any(v is not None for v in values):
                continue
            cells = [cell_text(values[idx]) if idx < len(values) else '' for idx in indexes]
            title = cell_text(values[title_idx]) if title_idx is not None and title_idx < len(values) else ''
            source['rows'].append({'row': num, 'cells': cells, 'standard_title': title})
        for num, values in enumerate(book['Bilgi'].iter_rows(values_only=True, max_col=2), 1):
            if num > 100:
                block('Bilgi', num, '', '', 'Bilgi sayfası çok uzun.', 'Şablondaki beş bilgi alanını kullanın.')
                break
            key = cell_text(values[0]) if values else ''
            if key in INFO_FIELDS:
                if key in source['metadata']:
                    block('Bilgi', num, 'A', key, 'Bilgi alanı yinelenmiş.', 'Her bilgi alanını bir kez yazın.')
                source['metadata'][key] = cell_text(values[1]) if len(values) > 1 else ''
    except (ValueError, BadZipFile, KeyError, OSError, ParseError, SyntaxError, EOFError, TypeError) as exc:
        block('Dosya', 0, '', filename, 'Excel sayfaları okunamadı; dosya bozuk veya geçersiz.',
              'Civensa şablonunu doldurup geçerli .xlsx olarak yeniden yükleyin.')
    finally:
        book.close()
    return source


def unit_resolver(db):
    """Build the matching catalogue once per dry-run, including ambiguous names."""
    departments = defaultdict(list)
    for department in db.scalars(select(Department)):
        departments[department_key(department.name)].append(department.id)
    aliases = {a.raw_name: a for a in db.scalars(select(ICUnitAlias))}

    def resolve(raw):
        key = department_key(raw)
        if key == ALL_DEPARTMENTS:
            return dict(department_id=None, body_name=None, all_departments=True)
        if key in aliases:
            alias = aliases[key]
            return dict(department_id=alias.department_id, body_name=alias.body_name, all_departments=False)
        matches = departments[key]
        if len(matches) == 1:
            return dict(department_id=matches[0], body_name=None, all_departments=False)
        if key in BODIES:
            return dict(department_id=None, body_name=BODIES[key], all_departments=False)
        return None
    return resolve


def validate_source(db, source):
    issues = list(source['issues'])
    metadata = source['metadata']
    def add(level, row, col, value, message, todo, sheet='Plan'):
        issues.append(issue(level, sheet, row, col, value, message, todo))
    for idx, field in enumerate(INFO_FIELDS[:4], 2):
        if not metadata.get(field):
            add('blocking', idx, 'B', '', f'{field} boş.', 'Bilgi sayfasındaki zorunlu alanı doldurun.', 'Bilgi')
    years = []
    for idx, field in enumerate(INFO_FIELDS[1:3], 3):
        raw = metadata.get(field, '')
        try:
            year = int(raw)
            if not 1900 <= year <= 2200:
                raise ValueError()
            years.append(year)
        except ValueError:
            add('blocking', idx, 'B', raw, f'{field} geçersiz.', '1900–2200 arasında tam yıl yazın.', 'Bilgi')
    if len(years) == 2 and years[0] > years[1]:
        add('blocking', 4, 'B', str(years[1]), 'Bitiş yılı başlangıçtan önce.', 'Yıl aralığını düzeltin.', 'Bilgi')
    if metadata.get('Onay Tarihi') and not parse_date(metadata['Onay Tarihi']):
        add('blocking', 6, 'B', metadata['Onay Tarihi'], 'Onay tarihi kesin tarih olmalı.', 'gg.aa.yyyy olarak yazın veya boş bırakın.', 'Bilgi')
    conditions, actions, unmatched = {}, [], {}
    codes = set()
    resolve_unit = unit_resolver(db)
    for entry in source['rows']:
        row, cells = entry['row'], entry['cells']
        component, standard_raw, cond_raw, condition_text, state, action_raw, action_text, owners, collaborators, output, start, due, note = cells
        standard, condition, action = map(normalize_code, (standard_raw, cond_raw, action_raw))
        if standard not in STANDARD_COMPONENTS:
            add('blocking', row, 'B', standard_raw, 'Standart kodu tanınmıyor.', '18 resmî standart kodundan birini seçin.')
        elif {'BIS': 'BİS', 'IS': 'İS'}.get(component.upper().replace('İ', 'I'), component.upper()) != STANDARD_COMPONENTS[standard]:
            add('blocking', row, 'A', component, 'Bileşen ve standart kodu uyuşmuyor.', f'{STANDARD_COMPONENTS[standard]} bileşenini kullanın.')
        if len(condition) > 60 or not re.fullmatch(re.escape(standard) + r'\.\d+', condition):
            add('blocking', row, 'C', cond_raw, 'Genel şart kodu standartla uyuşmuyor.', f'{standard}.1 biçiminde doğru kodu yazın.')
        if not condition_text:
            add('blocking', row, 'D', '', 'Genel şart metni boş.', 'Kaynak plandaki genel şart metnini yazın.')
        if condition in conditions:
            first = conditions[condition]
            if (first['text'], first['current_state_text'], first['standard_title']) != (condition_text, state, entry['standard_title']):
                add('warning', row, 'D/E/N', condition_text, 'Aynı genel şartın metni, mevcut durumu veya başlığı farklı; ilk satır korunacak.',
                    f'{first["sort_order"]}. satırla karşılaştırıp farklılığı düzeltin veya ilk satırı kabul ederek kaydedin.')
        else:
            conditions[condition] = dict(standard_code=standard, standard_title=entry['standard_title'],
                condition_code=condition, condition_code_raw=cond_raw, text=condition_text,
                current_state_text=state, sort_order=row)
        # Empty action columns represent a condition that needs no new action.
        if not action_raw and not any(cells[6:13]):
            continue
        if not action_raw or len(action) > 60 or not re.fullmatch(r'(?:KOS|RDS|KFS|BİS|İS)\d+\.\d+\.\d+', action):
            add('blocking', row, 'F', action_raw, 'Eylem kodu eksik veya geçersiz.', 'KOS 1.1.1 gibi tam eylem kodunu yazın.')
        if action in codes:
            add('blocking', row, 'F', action_raw, 'Eylem kodu yinelenmiş.', 'Her eylem kodunu plan içinde yalnızca bir kez kullanın.')
        codes.add(action)
        if not action_text:
            add('blocking', row, 'G', '', 'Eylem kodu var ancak eylem metni boş.', 'Öngörülen Eylem alanını doldurun.')
        if action and action.rsplit('.', 1)[0] != condition:
            add('warning', row, 'F', action_raw, 'Eylem kodunun öneki genel şart koduyla uyuşmuyor.', 'Kaynak kodu kontrol edin; özgün kod korunacaktır.')
        if not split_units(owners):
            add('blocking', row, 'H', owners, 'Eylemin sorumlu birimi yok.', 'En az bir sorumlu birim veya kurul yazın.')
        parties = []
        for role, raw, col in [('responsible', owners, 'H'), ('cooperating', collaborators, 'I')]:
            for order, name in enumerate(split_units(raw)):
                target = resolve_unit(name)
                if target is None:
                    key = department_key(name)
                    if key not in unmatched:
                        unmatched[key] = name
                        add('blocking', row, col, name, 'Birim adı eşleştirilemedi.', 'Eşleştirme bölümünden müdürlük veya müdürlük dışı kurul seçip tekrar doğrulayın.')
                else:
                    parties.append(dict(role=role, raw_name=name, sort_order=order, **target))
        for raw, col in [(start, 'K'), (due, 'L')]:
            if raw and department_key(raw) != department_key('Sürekli') and not re.fullmatch(r'\d{4}', raw) and not parse_date(raw):
                add('warning', row, col, raw, 'Tarih kesin bir güne çevrilemedi; özgün metin korunacak.', 'Gerekiyorsa gg.aa.yyyy yazın; dönem metnini korumak için uyarıyı kabul edin.')
        actions.append(dict(condition_code=condition, action_code=action, action_code_raw=action_raw,
            text=action_text, output_text=output, start_raw=start, due_raw=due, due_date=parse_date(due),
            is_continuous=department_key(due) == department_key('Sürekli'), note=note, sort_order=row, parties=parties))
    if not source['rows'] and not source['issues']:
        add('blocking', 2, 'C', '', 'Plan sayfasında veri yok.', 'En az bir genel şart satırı doldurun.')
    return dict(issues=issues, unmatched=unmatched, conditions=list(conditions.values()), actions=actions,
                counts={'conditions': len(conditions), 'actions': len(actions)},
                can_commit=not any(i['level'] == 'blocking' for i in issues))


def report_payload(report):
    return {k: report[k] for k in ('issues', 'unmatched', 'counts', 'can_commit')}


def audit(db, actor, batch, result, report, plan_id=None):
    db.add(ICImportAudit(import_id=batch.id, plan_id=plan_id, actor_id=actor.id,
        source_sha256=batch.source_sha256, result=result, report_json=encode(report_payload(report))))
    db.flush()


def create_import(db, actor, data, filename):
    actor = require_access(db, actor, strategy=True)
    source = parse_workbook(data, filename)
    batch = ICImport(source_filename=PurePath(filename).name, source_sha256=sha256(data).hexdigest(),
                     source_json=encode(source), created_by=actor.id)
    db.add(batch)
    db.flush()
    report = dry_run(db, actor, batch)
    return batch, report


def dry_run(db, actor, batch):
    actor = require_access(db, actor, strategy=True)
    report = validate_source(db, json.loads(batch.source_json))
    audit(db, actor, batch, 'dry-run', report)
    return report


def map_unit(db, actor, batch, raw_name, department_id=None, body_name=None):
    actor = require_access(db, actor, strategy=True)
    report = validate_source(db, json.loads(batch.source_json))
    key = department_key(raw_name)
    if key not in report['unmatched']:
        raise HTTPException(409, 'Birim artık eşleştirme beklemiyor. Doğrulama listesini yenileyin.')
    if bool(department_id) == bool(body_name and body_name.strip()):
        raise HTTPException(422, 'Birim eşleştirmesi için bir müdürlük veya müdürlük dışı kurul adı seçin.')
    if department_id and not db.get(Department, department_id):
        raise HTTPException(422, 'Seçilen müdürlük bulunamadı. Müdürlüğü listeden seçin.')
    db.add(ICUnitAlias(raw_name=key, department_id=department_id,
                      body_name=body_name.strip() if body_name else None, created_by=actor.id))
    db.flush()


def ensure_replaceable(db, previous):
    """CIV-049 hook: reject replacement here once monitoring reports exist.

    All imports already append new versions; action IDs remain stable forever.
    Future monitoring rows must reference ICAction.id and retain their plan_id.
    """
    if previous.status == 'archived':
        raise HTTPException(409, 'Arşivlenmiş plan değiştirilemez. Yeni bir plan veya etkin planın yeni sürümünü yükleyin.')


def commit_import(db, actor, batch_id, *, supersedes_id=None):
    actor = require_access(db, actor, strategy=True)
    with db.begin_nested():
        batch = db.scalar(select(ICImport).where(ICImport.id == batch_id).with_for_update())
        if not batch:
            raise HTTPException(404, 'Yükleme bulunamadı. Dosyayı yeniden yükleyin.')
        if db.scalar(select(ICPlan.id).where(ICPlan.import_id == batch.id)):
            raise HTTPException(409, 'Bu yükleme zaten kaydedildi. Plan listesini açın.')
        previous = db.scalar(select(ICPlan).where(ICPlan.id == supersedes_id).with_for_update().execution_options(populate_existing=True)) if supersedes_id else None
        if supersedes_id and not previous:
            raise HTTPException(404, 'Önceki plan bulunamadı. Plan listesini yenileyin.')
        if previous:
            ensure_replaceable(db, previous)
        source = json.loads(batch.source_json)
        report = validate_source(db, source)  # Never trust a client-supplied dry-run result.
        if not report['can_commit']:
            raise HTTPException(422, 'Plan kaydedilemedi. Engelleyici doğrulama maddelerini düzeltip tekrar doğrulayın.')
        metadata = source['metadata']
        plan = ICPlan(title=metadata['Plan Başlığı'], start_year=int(metadata['Başlangıç Yılı']),
            end_year=int(metadata['Bitiş Yılı']), version_label=metadata['Sürüm'],
            approval_date=parse_date(metadata.get('Onay Tarihi', '')), status='draft',
            source_filename=batch.source_filename, source_sha256=batch.source_sha256,
            import_id=batch.id, supersedes_id=supersedes_id, created_by=actor.id)
        db.add(plan)
        db.flush()
        for code, component in STANDARD_COMPONENTS.items():
            if not db.get(ICStandard, code):
                db.add(ICStandard(code=code, component_code=component))
        db.flush()
        condition_ids = {}
        for values in report['conditions']:
            condition = ICCondition(plan_id=plan.id, **values)
            db.add(condition)
            standard = db.get(ICStandard, values['standard_code'])
            if not standard.title and values['standard_title']:
                standard.title = values['standard_title']
            db.flush()
            condition_ids[condition.condition_code] = condition.id
        for values in report['actions']:
            values = dict(values)
            parties = values.pop('parties')
            condition_id = condition_ids[values.pop('condition_code')]
            action = ICAction(plan_id=plan.id, condition_id=condition_id, **values)
            db.add(action)
            db.flush()
            db.add_all([ICActionParty(action_id=action.id, **p) for p in parties])
        db.flush()  # Freeze only after all action parties have been inserted.
        audit(db, actor, batch, 'commit', report, plan.id)
        db.flush()
    return plan


def activate_plan(db, actor, plan_id):
    actor = require_access(db, actor, strategy=True)
    # Serialize with every activation, including the first (no active row yet).
    db.scalars(select(Organization).order_by(Organization.id).with_for_update()).all()
    plan = db.scalar(select(ICPlan).where(ICPlan.id == plan_id).with_for_update().execution_options(populate_existing=True))
    if not plan or plan.status != 'draft':
        raise HTTPException(409, 'Yalnızca taslak plan etkinleştirilebilir. Plan listesini yenileyin.')
    db.execute(update(ICPlan).where(ICPlan.status == 'active').values(status='archived'))
    db.execute(update(ICPlan).where(ICPlan.id == plan.id).values(status='active'))
    batch = db.get(ICImport, plan.import_id)
    audit(db, actor, batch, 'activation', validate_source(db, json.loads(batch.source_json)), plan.id)
    db.flush()
    return plan


def directorate_actions(db, actor, role):
    actor = require_access(db, actor)
    if actor.role != 'department':
        raise HTTPException(403, 'Bu liste müdürlük kullanıcılarına açıktır. Plan kataloğunu açın.')
    return db.scalars(select(ICAction).join(ICPlan).where(ICPlan.status == 'active', exists().where(
        ICActionParty.action_id == ICAction.id, ICActionParty.role == role,
        or_(ICActionParty.department_id == actor.department_id, ICActionParty.all_departments.is_(True))))
        .order_by(ICAction.sort_order)).all()


def validation_bytes(report):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Doğrulama'
    literal_row(sheet, ['Düzey', 'Sayfa', 'Satır', 'Sütun', 'Değer', 'Mesaj', 'Yapılacak İşlem'])
    for i in report['issues']:
        literal_row(sheet, ['Engelleyici' if i['level'] == 'blocking' else 'Uyarı',
                           i['sheet'], i['row'], i['column'], i['value'], i['message'], i['todo']])
    sheet.freeze_panes = 'A2'
    summary = book.create_sheet('Özet')
    literal_row(summary, ['Kontrol', 'Sonuç'])
    literal_row(summary, ['Kaydedilebilir', 'Evet' if report['can_commit'] else 'Hayır'])
    literal_row(summary, ['Genel şart sayısı', report['counts']['conditions']])
    literal_row(summary, ['Eylem sayısı', report['counts']['actions']])
    literal_row(summary, ['Engelleyici madde', sum(i['level'] == 'blocking' for i in report['issues'])])
    literal_row(summary, ['Uyarı', sum(i['level'] == 'warning' for i in report['issues'])])
    return export_book(book)
