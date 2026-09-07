import io, re
from pathlib import Path
from docx import Document
from openpyxl import load_workbook

CANONICAL_FIELDS = ["Bütçe Yılı", "Program Adı", "İdare Birim Adı", "Alt Program Adı", "Gerekçe ve Açıklamalar", "Alt Program Hedefi", "Performans Göstergeleri", "Ölçü Birimi", "Planlanan", "Yıl Sonu Gerçekleşme Tahmini", "Hedef", "Tahmin", "Göstergeye İlişkin Açıklama", "Hesaplama Yöntemi", "Verinin Kaynağı", "Sorumlu İdare", "Alt Program Kapsamında Yürütülecek Faaliyet Maliyetleri", "Faaliyetler", "Bütçe İçi", "Bütçe Dışı", "Personel ve SGK Giderleri", "Genel Toplam", "Faaliyetlere İlişkin Açıklamalar"]

def _norm(s: str) -> str:
    s=(s or "").lower().replace("ı","i").replace("ş","s").replace("ğ","g").replace("ü","u").replace("ö","o").replace("ç","c")
    s=re.sub(r"\s+"," ",s)
    return re.sub(r"[^a-z0-9% ]+","",s).strip()

def _tr_title(s: str) -> str:
    trans=str.maketrans({"I":"ı","İ":"i","Ş":"ş","Ğ":"ğ","Ü":"ü","Ö":"ö","Ç":"ç"})
    low=s.translate(trans).lower()
    return " ".join(w[:1].upper()+w[1:] if w else w for w in low.split())

def _collect_docx(data: bytes):
    doc=Document(io.BytesIO(data)); items=[]
    for p in doc.paragraphs:
        t=p.text.strip()
        if t: items.append(t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t=" ".join(x.strip() for x in cell.text.splitlines() if x.strip())
                if t: items.append(t)
    return items

def _collect_xlsx(data: bytes):
    wb=load_workbook(io.BytesIO(data), read_only=True, data_only=True); items=[]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    t=str(cell.value).strip()
                    if t: items.append(t)
    return items

def analyze_template(filename: str, data: bytes):
    suffix=Path(filename).suffix.lower()
    if suffix==".docx": items=_collect_docx(data)
    elif suffix in {".xlsx",".xlsm"}: items=_collect_xlsx(data)
    else: raise ValueError("Yalnızca DOCX ve XLSX şablonları destekleniyor.")
    normalized=[_norm(x) for x in items]
    present=[]; missing=[]
    aliases={"Göstergeye İlişkin Açıklama":["gostergeye iliskin aciklama","gostergelere iliskin aciklama"],"Alt Program Hedefi":["alt program hedefi","alt program hedefleri"],"Performans Göstergeleri":["performans gostergeleri","performans gostergesi"],"Faaliyetlere İlişkin Açıklamalar":["faaliyetlere iliskin aciklamalar","faaliyetlere iliskin aciklama"]}
    for field in CANONICAL_FIELDS:
        candidates=aliases.get(field,[_norm(field)])
        ok=any(any(c in text for c in candidates) for text in normalized)
        (present if ok else missing).append(field)
    canon_norm={_norm(x) for x in CANONICAL_FIELDS}; extras=[]
    for raw,n in zip(items,normalized):
        if 2<=len(raw)<=90 and n not in canon_norm and (raw.isupper() or raw.endswith(":")):
            if raw not in extras: extras.append(raw)
        if len(extras)>=25: break
    departments=[]
    for raw in items:
        upper=raw.upper().strip()
        if "MÜDÜRLÜĞÜ" in upper and "PERFORMANS ESASLI" in upper:
            name=re.sub(r"\s+20\d{2}\s+YILI.*$","",upper).strip()
            if name and name not in departments: departments.append(_tr_title(name))
    return {"present":present,"missing":missing,"extra":extras,"departments":departments,"item_count":len(items),"file_type":suffix.lstrip('.')}
