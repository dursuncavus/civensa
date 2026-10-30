"""CIV-085 export, independent of existing monitoring export functions."""
from openpyxl import Workbook
from . import internal_control_monitoring_service as m
from . import internal_control_questionnaire_service as q
from .internal_control_service import literal_row, export_book


def questionnaire_export(db, actor, period_id, *, draft=False):
    m.require_access(db, actor, strategy=True)
    period = m.get_period(db, actor, period_id, lock=True)
    results = q.results(db, actor, period)
    if not draft and (not results['participant_count'] or results['pending']):
        m.fail('Nihai soru formu sonuçları hazır değil. Katılımcıların onayını tamamlayın veya TASLAK indirin.', status=409)
    book = Workbook()
    book.remove(book.active)
    headers = ['Evet', 'Hayır', 'Geliştirilmekte']
    for name, rows in [('Bileşen özeti', results['components'] + [results['total']]),
                       ('Soru özeti', results['questions'])]:
        sheet = book.create_sheet(name)
        literal_row(sheet, ['TASLAK' if draft else 'Onaylı veriler', m.period_label(period)])
        literal_row(sheet, ['Onaylı müdürlük', len(results['included']), 'Katılımcı', results['participant_count']])
        literal_row(sheet, ['Henüz onaylanmayanlar', ', '.join(results['pending'])])
        question_sheet = name == 'Soru özeti'
        literal_row(sheet, (['Kod', 'Bileşen', 'Soru'] if question_sheet else ['Bileşen']) + headers)
        for row in rows:
            literal_row(sheet, ([row['code'], q.COMPONENTS[row['component']], row['question']] if question_sheet else [row['label']]) +
                [row['rates'][key] for key in q.ANSWERS])
    sheet = book['Bileşen özeti']
    literal_row(sheet, [])
    literal_row(sheet, ['Müdürlük puanları'])
    literal_row(sheet, [q.SCORING_ASSUMPTION])
    literal_row(sheet, ['Müdürlük', *headers, 'Toplam puan', 'Azami puan', 'Gerçekleşme oranı'])
    for row in results['directorates'] + [results['total']]:
        literal_row(sheet, [row['label'], *[row['counts'][key] for key in q.ANSWERS],
            row['score'], row['max_score'], row['realisation_rate']])
    sheet = book.create_sheet('Müdürlük yanıtları')
    literal_row(sheet, ['TASLAK' if draft else 'Onaylı veriler', m.period_label(period)])
    literal_row(sheet, ['Müdürlük', 'Durum', 'Kod', 'Soru', 'Yanıt', 'Açıklama'])
    for batch in q.batches(db, period.id):
        status, _ = q.state(db, batch.id)
        for row in q.rows(db, batch):
            literal_row(sheet, [batch.label, m.STATES[status], row['code'], row['question'],
                q.ANSWERS.get(row['answer'], ''), row['explanation']])
    for sheet in book:
        sheet.freeze_panes = 'B5'
        sheet.column_dimensions['A'].width = 34
        for col in 'BCDEFG':
            sheet.column_dimensions[col].width = 30
    return export_book(book)
