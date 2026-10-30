"""Word/Excel monitoring exports from immutable entry versions and approval snapshots."""
import json
from collections import Counter
from io import BytesIO
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
from openpyxl import Workbook
from sqlalchemy import select
from .internal_control_models import ICPlan, ICCondition, ICAction, COMPONENTS, STANDARD_COMPONENTS
from .numbers import format_number
from .internal_control_service import require_access, literal_row, export_book
from . import internal_control_monitoring_service as m
from .internal_control_monitoring_models import ICMonitoringPeriod as Period, ICMonitoringBatch as Batch, ICMonitoringEvent as Event, ICMonitoringEntry as Entry, ICMonitoringTarget as Target

HEADERS = ['Bileşen', 'Standart', 'Genel şart', 'Eylem kodu', 'Eylem', 'Sorumlu birim',
           'Durum', 'Gerçekleşme oranı', 'Açıklama', 'Tamamlanmama nedeni', 'Kanıt']


def report_row(db, batch, target, entry):
    action = db.get(ICAction, target.action_id)
    condition = db.get(ICCondition, action.condition_id)
    return [COMPONENTS[STANDARD_COMPONENTS[condition.standard_code]],
            condition.standard_code, condition.condition_code_raw, action.action_code_raw,
            action.text, batch.label, m.STATUSES[entry.status] if entry else 'Giriş eksik',
            entry.percentage if entry else None, entry.explanation if entry else '',
            entry.reason if entry else '', entry.evidence if entry else '']


def period_data(db, actor, period_id, *, draft=False):
    require_access(db, actor, strategy=True)
    period = m.get_period(db, actor, period_id, lock=True)
    batches = db.scalars(select(Batch).where(Batch.period_id == period.id).order_by(Batch.id)).all()
    complete = bool(batches) and all(m.state(db, b.id)[0] == 'approved' for b in batches)
    if not complete and not draft:
        m.fail('Nihai izleme raporu hazır değil. Tüm sorumlu birimlerin onayını tamamlayın veya TASLAK indirin.', status=409)
    result = []
    for batch in batches:
        if m.state(db, batch.id)[0] == 'approved':
            approval = m.latest_event(db, batch.id)
            for item in json.loads(approval.snapshot_json):
                result.append(report_row(db, batch, db.get(Target, item['target_id']), db.get(Entry, item['entry_id'])))
        else:
            for row in m.rows(db, batch):
                result.append(report_row(db, batch, row['target'], row['entry']))
    return period, db.get(ICPlan, period.plan_id), sort_rows(result), draft


def sort_rows(rows):
    order = {c: i for i, c in enumerate(STANDARD_COMPONENTS)}
    return sorted(rows, key=lambda r: (order[r[1]], r[2], r[3], r[5]))


def summary(rows):
    counts = Counter((r[0], r[6]) for r in rows)
    return [[name, status, counts[name, status]] for name in COMPONENTS.values()
            for status in [*[label for code, label in m.STATUSES.items()
                            if code not in ('partial', 'not_approved') or any(r[6] == label for r in rows)], 'Giriş eksik']]


def xlsx_bytes(title, subtitle, rows, *, draft=False):
    book = Workbook()
    cover = book.active
    cover.title = 'Özet'
    literal_row(cover, ['TASLAK' if draft else 'Onaylı veriler', title])
    literal_row(cover, ['Dönem', subtitle])
    literal_row(cover, ['Bileşen', 'Durum', 'Sayı'])
    for row in summary(rows):
        literal_row(cover, row)
        cover.cell(cover.max_row, 3).number_format = '#,##0'
    sheet = book.create_sheet('İzleme')
    literal_row(sheet, HEADERS)
    for row in rows:
        literal_row(sheet, row)
        cell = sheet.cell(sheet.max_row, 8)
        cell.value = row[7] / 100 if row[7] is not None else None
        cell.number_format = '0%'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in ('A', 'C', 'D', 'F', 'G', 'I', 'J', 'K'):
        sheet.column_dimensions[column].width = 28
    sheet.column_dimensions['E'].width = 55
    from openpyxl.styles import Alignment, Font
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    return export_book(book)


def docx_bytes(title, subtitle, rows, *, draft=False):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(.45)
    doc.styles['Normal'].font.size = Pt(9)
    if draft:
        section.header.paragraphs[0].text = 'TASLAK — Onay süreci tamamlanmamış olabilir'
    doc.add_heading('TASLAK · İç Kontrol İzleme Raporu' if draft else 'İç Kontrol İzleme Raporu', 0)
    doc.add_paragraph(title)
    doc.add_paragraph(subtitle)
    doc.add_paragraph('Sayılar sorumlu birim–eylem girişlerini gösterir. Tüm müdürlüklerin eylemleri her müdürlük için ayrı sayılır.')
    doc.add_page_break()
    doc.add_heading('Bileşen ve durum özeti', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, ['Bileşen', 'Durum', 'Sayı']):
        cell.text = value
    for row in summary(rows):
        for cell, value in zip(table.add_row().cells, row):
            cell.text = format_number(value) if isinstance(value, int) else str(value)
    for component in COMPONENTS.values():
        selected = [r for r in rows if r[0] == component]
        if not selected:
            continue
        doc.add_heading(component, level=1)
        table = doc.add_table(rows=1, cols=10)
        table.style = 'Table Grid'
        for cell, value in zip(table.rows[0].cells, HEADERS[1:]):
            cell.text = value
        for row in selected:
            values = row[1:]
            values[6] = '' if row[7] is None else f'%{row[7]}'
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value) if value is not None else ''
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def period_export(db, actor, period_id, format, *, draft=False):
    period, plan, rows, draft = period_data(db, actor, period_id, draft=draft)
    if format not in ('docx', 'xlsx'):
        m.fail('Rapor dosya biçimi geçersiz. Word veya Excel bağlantısını seçin.', 'format')
    render = docx_bytes if format == 'docx' else xlsx_bytes
    return render(plan.title + ' · ' + plan.version_label, m.period_label(period), rows, draft=draft)


def annual_export(db, actor, plan_id, year):
    require_access(db, actor, strategy=True)
    year = m.number(year, 'year')
    plan = db.get(ICPlan, plan_id)
    if not plan or not plan.start_year <= year <= plan.end_year:
        m.fail('Yıllık rapor yılı plan kapsamında değil. Plan yılları içinden yıl seçin.', 'year')
    # Freeze period streams consistently with decisions and period reports.
    periods = db.scalars(select(Period).where(Period.plan_id == plan_id, Period.year == year)
                         .order_by(Period.half).with_for_update()).all()
    latest = {}
    for period in periods:
        for batch in db.scalars(select(Batch).where(Batch.period_id == period.id).order_by(Batch.id)):
            approval = db.scalar(select(Event).where(Event.batch_id == batch.id, Event.action == 'approve')
                                 .order_by(Event.revision.desc()).limit(1))
            if not approval:
                continue
            for item in json.loads(approval.snapshot_json):
                target, entry = db.get(Target, item['target_id']), db.get(Entry, item['entry_id'])
                latest[batch.scope_key, target.action_id] = report_row(db, batch, target, entry)
    return xlsx_bytes(plan.title, f'{year} · 15 Mayıs değerlendirmesi için son onaylı I/II dönem verileri', sort_rows(list(latest.values())))


OFFICIAL_HEADERS = ['Standart Kod No', 'Kamu İç Kontrol Standardı ve Genel Şartı',
    'Mevcut Durum', 'Eylem Kod No', 'Öngörülen Eylem veya Eylemler',
    'Sorumlu Birim veya Çalışma grubu üyeleri', 'İşbirliği Yapılacak Birim',
    'Çıktı/Sonuç', 'Başlangıç Tarihi', 'Bitiş Tarihi', 'Açıklama']
COMPONENT_HEADERS = ['Bileşenler', 'Eylem Sayısı', 'Gerçekleşen', 'Gerçekleşmeyen',
                     'Devam Eden', 'Tamamlanan Toplam', 'Gerçekleşme Oranı (%)']
OFFICIAL_PHRASES = {'completed': '(TAMAMLANDI)', 'ongoing': 'Eylem Devam Etmektedir.',
    'partial': '(KISMEN TAMAMLANDI (DEVAM EDİYOR))', 'delayed': 'Eylem Devam Etmektedir.',
    'not_started': '(GERÇEKLEŞMEDİ)', 'not_approved': 'Eylem Uygun Görülmemiştir.'}


def official_data(db, actor, period_id, *, draft=False):
    from .internal_control_models import ICActionParty
    from .models import Department
    require_access(db, actor, strategy=True)
    period = m.get_period(db, actor, period_id, lock=True)
    batches = db.scalars(select(Batch).where(Batch.period_id == period.id)).all()
    if not draft and (not batches or any(m.state(db, b.id)[0] != 'approved' for b in batches)):
        m.fail('Nihai izleme raporu hazır değil. Tüm sorumlu birimlerin onayını tamamlayın veya TASLAK indirin.', status=409)
    records = m.action_records(db, period, draft=draft)
    metrics = m.official_metrics(db, actor, period, records=records)
    grouped = {}
    for record in records:
        grouped.setdefault(record['action'].condition_id, []).append(record)
    standard_order = {s: i for i, s in enumerate(STANDARD_COMPONENTS)}
    conditions = sorted(db.scalars(select(ICCondition).where(ICCondition.plan_id == period.plan_id)).all(),
                        key=lambda c: (standard_order[c.standard_code], c.sort_order, c.id))
    parties = {}
    for party in db.scalars(select(ICActionParty).join(ICAction).where(ICAction.plan_id == period.plan_id)
                            .order_by(ICActionParty.sort_order, ICActionParty.id)):
        name = 'Tüm Müdürlükler' if party.all_departments else db.get(Department, party.department_id).name if party.department_id else party.body_name
        parties.setdefault((party.action_id, party.role), []).append(name)
    rows = []
    for condition in conditions:
        prefix = [condition.condition_code_raw,
                  '\n'.join(v for v in (condition.standard_title, condition.text) if v), condition.current_state_text]
        if not grouped.get(condition.id):
            rows.append(prefix + [''] * 8)
        for record in grouped.get(condition.id, []):
            action = record['action']
            entries = [u['entry'] for u in record['units']]
            status = m.action_status(entries)
            if status == 'ongoing' and any(e and e.status == 'partial' for e in entries):
                status = 'partial'
            if record['due'] is False and not any(u['approved_entry'] for u in record['units']):
                explanation = ['Eylem süresi bu dönemde dolmamaktadır.']
            else:
                explanation = [OFFICIAL_PHRASES[status]]
            for unit in record['units']:
                entry = unit['entry']
                if entry and entry.explanation:
                    explanation.append(unit['batch'].label + ': ' + entry.explanation)
            rows.append(prefix + [action.action_code_raw, action.text,
                '\n'.join(parties.get((action.id, 'responsible'), [])),
                '\n'.join(parties.get((action.id, 'cooperating'), [])),
                action.output_text, action.start_raw, action.due_raw, '\n'.join(explanation)])
    return period, db.get(ICPlan, period.plan_id), rows, metrics


def official_export(db, actor, period_id, format, *, draft=False):
    if format not in ('docx', 'xlsx'):
        m.fail('Rapor dosya biçimi geçersiz. Word veya Excel bağlantısını seçin.', 'format')
    narrative = None
    if format == 'docx':
        from .internal_control_narrative_export import narrative_for_word, insert_narrative
        require_access(db, actor, strategy=True)
        narrative_period = m.get_period(db, actor, period_id, lock=True)
        narrative = narrative_for_word(db, narrative_period, draft=draft)
    period, plan, rows, metrics = official_data(db, actor, period_id, draft=draft)
    title, subtitle = plan.title + ' · ' + plan.version_label, m.period_label(period)
    cover = [f"Gerçekleşme Oranı: {metrics['period_rate']}",
             f"{'İlk 6 Aylık' if period.half == 1 else 'İkinci 6 Aylık'} Başarı Yüzdesi: {metrics['plan_rate']}",
             f"Plan toplamı: {format_number(metrics['plan_total'])} eylem",
             'Sayılar eylem bazındadır; yalnız onaylı sorumlu birim girişleri oranlara katılır.']
    if metrics['warning']:
        cover.append(metrics['warning'])
    if format == 'xlsx':
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        sheet = book.active
        sheet.title = 'Özet'
        literal_row(sheet, ['TASLAK' if draft else 'Onaylı veriler', title])
        literal_row(sheet, ['Dönem', subtitle])
        for line in cover:
            literal_row(sheet, [line])
        literal_row(sheet, COMPONENT_HEADERS)
        for row in metrics['components']:
            literal_row(sheet, row)
        sheet.column_dimensions['A'].width = 70
        for col in 'BCDEFG':
            sheet.column_dimensions[col].width = 25
        sheet = book.create_sheet('Resmî İzleme')
        literal_row(sheet, OFFICIAL_HEADERS)
        for row in rows:
            literal_row(sheet, row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for col in 'ABCDEFGHIJK':
            sheet.column_dimensions[col].width = 32 if col not in 'BEK' else 55
        for sheet in book:
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        return export_book(book)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(.45)
    doc.styles['Normal'].font.size = Pt(8)
    if draft:
        section.header.paragraphs[0].text = 'TASLAK — Onay süreci tamamlanmamış olabilir'
    doc.add_heading(('TASLAK · ' if draft else '') + 'İç Kontrol İzleme Raporu · Resmî format', 0)
    for line in [title, subtitle, *cover]:
        doc.add_paragraph(line)
    if narrative:
        insert_narrative(doc, db, period, narrative)
    doc.add_page_break()
    for headers, contents in [(COMPONENT_HEADERS, metrics['components']), (OFFICIAL_HEADERS, rows)]:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for row in contents:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = format_number(value) if isinstance(value, int) else str(value)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
