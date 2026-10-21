"""Both publication formats consume the same frozen, ordered content blocks."""
import base64
import io
import json
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape
from .numbers import format_number
from .program_service import cost_columns, economic_columns, ECONOMIC_CATEGORIES, PROGRAM_GROUPS

LABELS = {'actual':'Gerçekleşme','planned':'Planlanan','year_end_estimate':'Yıl Sonu Tahmini',
          'target':'Hedef','forecast':'Tahmin','budget':'Bütçe','midyear_spend':'Haziran Sonu Harcama'}


def cell(row, field, money=False):
    if not row: return '—'
    if row.get('state', 'verified') != 'verified':
        return f"{row.get('raw_text', '')} [Doğrulanmalı]"
    return format_number(row.get(field), money=money)


def blocks(payload):
    year = payload['period']
    out = [('heading', f'{year} Yılı Performans Programı')]
    if payload.get('draft'): out.append(('text','TASLAK — eksik ve doğrulanmamış veriler içerebilir.'))
    for s in payload.get('sections', []):
        if json.loads(s['content_json']).get('placement','before')=='after':continue
        out.append(('heading', s['title']))
        content = json.loads(s['content_json'])
        if s['kind']=='image': out.append(('image',content))
        elif s['kind']=='table': out.append(('table', content['rows']))
        else: out.append(('text', content.get('text','')))
    cols = [(year-2,'actual'),(year-1,'planned'),(year-1,'year_end_estimate'),(year,'target'),(year+1,'forecast'),(year+2,'forecast')]
    costcols=cost_columns(year)
    for d in payload.get('departments', []):
        if d.get('scope') and d['scope']['status']=='excluded': continue
        out.extend([('page',None), ('heading',d['department'])])
        if d.get('scope') and d['scope']['status']=='zero':
            out.append(('text',d['scope']['reason']));continue
        for card in d['cards']:
            out.append(('heading', card.get('source_title') or 'Performans Göstergeleri'))
            for key,label in [('administration','İdare / Birim Adı'),('program','Program'),('subprogram','Alt Program'),('rationale','Gerekçe ve Açıklamalar'),('goal','Alt Programın Hedefi'),('responsible','Sorumlu İdare'),('service_area','Hizmet Alanı'),('strategic_link','Stratejik Plan / PH Bağlantısı')]:
                if card.get(key): out.append(('text',f'{label}: {card[key]}'))
            inds=sorted([i for i in d['indicators'] if i.get('card_id')==card['id']],key=lambda i:(i.get('position',0),i['id']))
            table=[['Kod / Performans Göstergesi','Ölçü Birimi']+[f'{y}\n{LABELS[k]}' for y,k in cols]]
            for i in inds:
                vals={(v['year'],v['value_kind']):v for v in i['values']}
                table.append([' '.join(filter(None,[i.get('code'),i['name']])),i['unit']]+[cell(vals.get(c),'value') for c in cols])
            out.append(('table',table))
            for key,label in [('explanation','Göstergeye İlişkin Açıklama'),('calculation_method','Hesaplama Yöntemi'),('data_source','Verinin Kaynağı'),('footnote','Dipnot')]:
                if card.get(key):out.append(('text',f'{label}: {card[key]}'))
            for i in inds:
                for key,label in [('explanation','Açıklama'),('calculation_method','Hesaplama Yöntemi'),('data_source','Verinin Kaynağı')]:
                    if i.get(key) and i[key]!=card.get(key):out.append(('text',f'{i["name"]} — {label}: {i[key]}'))
        unassigned=[i for i in d['indicators'] if not i.get('card_id')]
        for i in unassigned:out.append(('text',f'Kart ataması bekliyor: {i["name"]}'))
        out.append(('heading','Faaliyetler ve Maliyetler'))
        # Once per department, even when several cards/programs share this table.
        table=[['Faaliyet / Finansman']+[f'{y}\n{LABELS[k]}' for y,k in costcols]]
        activities=sorted(d['activities'],key=lambda a:(a.get('position',0),a['id']))
        for a in activities:
            vals={(v['year'],v['value_kind'],v['funding_type']):v for v in a['costs']}
            for fund,label in [('inside','Bütçe İçi'),('outside','Bütçe Dışı'),('total','Toplam')]:
                table.append([f'{a.get("code", "")} {a["name"]}\n{label} (TL)']+[cell(vals.get((y,k,fund)),'amount',True) for y,k in costcols])
        personnel={(v['year'],v['value_kind']):v for v in d['department_costs'] if v['category']=='personnel_sgk'}
        reconciled = {(r['year'], r['value_kind']): r for r in payload.get('cost_reconciliation', [])
                      if r['level'] == 'department_total' and r['department_id'] == d['department_id']}
        activity_totals=[];general=[]
        for y,k in costcols:
            vals=[next((v for v in a['costs'] if (v['year'],v['value_kind'],v['funding_type'])==(y,k,'total')),None) for a in activities]
            complete=bool(vals) and all(v and v.get('amount') is not None and v.get('state','verified')=='verified' for v in vals)
            total=sum((Decimal(v['amount']) for v in vals),Decimal(0)) if complete else None
            if (y,k) in reconciled:
                current = reconciled[y,k]['amount']
                total = Decimal(current) if current is not None else None
            p=personnel.get((y,k));pval=Decimal(p['amount']) if p and p.get('amount') is not None and p.get('state','verified')=='verified' else None
            activity_totals.append(format_number(total,True));general.append(format_number(total+pval,True) if total is not None and pval is not None else 'Eksik veri')
        table.extend([['Faaliyetler Toplamı (TL)']+activity_totals,['Personel ve SGK Giderleri (TL)']+[cell(personnel.get(c),'amount',True) for c in costcols],['Genel Toplam (TL)']+general])
        out.append(('table',table))
        for a in activities:
            out.append(('text',f'{a["name"]} — {a.get("program", "")} / {a.get("subprogram", "")}\n{a.get("explanation", "")}'))
        out.append(('heading','Ekonomik Sınıflandırma'))
        table=[['Ekonomik Ad / Program Grubu / Kaynak']+[f'{y}\n{LABELS[k]}' for y,k in economic_columns(year)]]
        groups=sorted({(v['category'],v['program_group'],v['funding_type']) for v in d['economics']})
        for cat,group,fund in groups:
            vals={(v['year'],v['value_kind']):v for v in d['economics'] if (v['category'],v['program_group'],v['funding_type'])==(cat,group,fund)}
            table.append([f'{ECONOMIC_CATEGORIES[cat]} / {PROGRAM_GROUPS[group]} / {"Bütçe İçi" if fund=="inside" else "Bütçe Dışı"}']+[cell(vals.get(c),'amount',True) for c in economic_columns(year)])
        out.append(('table',table))
    if payload.get('economics'):
        out.extend([('page',None),('heading','Ekonomik Sınıflandırmaya Göre Ödenekler')])
        table=[['Ekonomik Ad','Program Grubu','Finansman']+[str(y) for y in range(year,year+3)]]
        for cat,group,fund in sorted({(v['category'],v['program_group'],v['funding_type']) for v in payload['economics']}):
            vals={v['year']:v['amount'] for v in payload['economics'] if (v['category'],v['program_group'],v['funding_type'])==(cat,group,fund)}
            table.append([ECONOMIC_CATEGORIES[cat],PROGRAM_GROUPS[group],{'inside':'Bütçe İçi','outside':'Bütçe Dışı'}[fund]]+[format_number(vals.get(y),True) for y in range(year,year+3)])
        out.append(('table',table))
    if payload.get('program_costs'):
        out.append(('heading','Program, Alt Program ve Faaliyet Maliyetleri'))
        table=[['Program / Alt Program / Faaliyet','Finansman']+[str(y) for y in range(year,year+3)]]
        for aid in sorted({v['activity_id'] for v in payload['program_costs']}):
            items=[v for v in payload['program_costs'] if v['activity_id']==aid];a=items[0]
            for fund in ['inside','outside','total']:
                vals={v['year']:v['amount'] for v in items if v['funding_type']==fund}
                table.append([f'{a["program"]} / {a["subprogram"]} / {a["name"]}',{'inside':'Bütçe İçi','outside':'Bütçe Dışı','total':'Toplam'}[fund]]+[format_number(vals.get(y),True) for y in range(year,year+3)])
        out.append(('table',table))
    if payload.get('cost_reconciliation'):
        out.append(('heading', 'Maliyet Ara Toplamları ve Kaynak Mutabakatı'))
        table = [['Program / Alt Program / Müdürlük', 'Yıl / Değer Türü', 'Hesaplanan (TL)',
                  'Kaynak (TL)', 'Fark (TL)', 'Durum / Yapılacak İşlem']]
        for row in payload['cost_reconciliation']:
            table.append([row['label'], f"{row['year']} {LABELS[row['value_kind']]}",
                format_number(row['amount'], True) if row['complete'] else 'Eksik ara toplam',
                format_number(row['source_amount'], True) if row['source_amount'] is not None else 'kaynak toplamı yok',
                format_number(row['difference'], True),
                '\n'.join([row['status']] + row['errors'] + row['warnings'])])
        out.append(('table', table))
    if payload.get('catalogue'):
        out.append(('heading','Faaliyetlerden Sorumlu Harcama Birimleri'))
        names={d['department_id']:d['department'] for d in payload['departments']}
        table=[['Program','Alt Program','Faaliyet','Sorumlu Harcama Birimleri']]
        for a in payload['catalogue']:
            table.append([f'{a["program_code"]} {a["program_name"]}',f'{a["subprogram_code"]} {a["subprogram_name"]}',f'{a["code"]} {a["name"]}','\n'.join(names.get(d,str(d)) for d in a['departments'])])
        out.append(('table',table))
    included=[d for d in payload.get('departments',[]) if not d.get('scope') or d['scope']['status']!='excluded']
    if any(d['economics'] for d in included):
        out.append(('heading','Müdürlüklere Göre Bütçe Dağılımı'))
        table=[['Müdürlük']+[str(y) for y in range(year,year+3)]];chart_values=[]
        for d in included:
            vals=[]
            for y in range(year,year+3):
                values=[Decimal(v['amount']) for v in d['economics'] if v['year']==y and v['value_kind']==('budget' if y==year else 'forecast')]
                total=sum(values,Decimal(0)) if values or (d.get('scope') and d['scope']['status']=='zero') else None
                vals.append(total)
            table.append([d['department']]+[format_number(v,True) for v in vals])
            if vals[0] is not None:chart_values.append((d['department'],vals[0]))
        out.append(('table',table))
        if chart_values:out.extend(chart_blocks(chart_values,f'{year} Bütçesinin Müdürlüklere Dağılımı'))
        economic_values=[]
        for cat,label in ECONOMIC_CATEGORIES.items():
            vals=[Decimal(v['amount']) for d in included for v in d['economics'] if v['category']==cat and v['year']==year and v['value_kind']=='budget']
            if vals:economic_values.append((label,sum(vals,Decimal(0))))
        if economic_values:out.extend(chart_blocks(economic_values,f'{year} Bütçesinin Ekonomik Dağılımı'))
    for s in payload.get('sections',[]):
        c=json.loads(s['content_json'])
        if c.get('placement')!='after':continue
        out.append(('heading',s['title']))
        out.append(('image',c) if s['kind']=='image' else ('table',c['rows']) if s['kind']=='table' else ('text',c.get('text','')))
    return out


def chart_blocks(values,title):
    # Values come exclusively from the same snapshot as the tables. Floating-point
    # conversion is limited to pixel geometry; labels and percentages use Decimal.
    from PIL import Image,ImageDraw,ImageFont
    font=ImageFont.truetype(str(Path(__file__).parent/'fonts'/'DejaVuSans.ttf'),20)
    total=sum((v for _,v in values),Decimal(0));maximum=max((v for _,v in values),default=Decimal(0))
    result=[]
    for start in range(0,len(values),12):
        part=sorted(values,key=lambda x:x[1],reverse=True)[start:start+12]
        image=Image.new('RGB',(1800,70+len(part)*70),'white');draw=ImageDraw.Draw(image)
        for index,(label,amount) in enumerate(part):
            y=20+index*70
            label=label if len(label)<65 else label[:62]+'…'
            draw.text((10,y),label,font=font,fill='#152738')
            length=int(amount/maximum*650) if maximum else 0
            draw.rectangle((770,y,770+length,y+27),fill='#385f7e')
            share=amount/total*100 if total else Decimal(0)
            draw.text((1440,y),f'{format_number(amount,True)} TL / %{format_number(share.quantize(Decimal(".01")))}',font=font,fill='#152738')
        buffer=io.BytesIO();image.save(buffer,format='PNG')
        result.extend([('heading',title+(f' ({start//12+1})' if len(values)>12 else '')),('image',{'data':base64.b64encode(buffer.getvalue()).decode(),'caption':''})])
    return result


def export_book(payload, format):
    content=blocks(payload)
    if format=='docx':return _docx(content),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format=='pdf':return _pdf(content),'application/pdf'
    raise ValueError('Desteklenmeyen çıktı biçimi')


def _docx(content):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(29.7);section.page_height=Cm(21)
    section.left_margin=section.right_margin=Cm(1.5)
    section.top_margin=section.bottom_margin=Cm(1.4)
    doc.styles['Normal'].font.name='DejaVu Sans';doc.styles['Normal'].font.size=Pt(9)
    for name in ['Title','Heading 1','Heading 2']:
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        doc.styles[name].font.name='DejaVu Sans'
        for border in list(doc.styles[name]._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for index,(kind,value) in enumerate(content):
        if kind=='page':doc.add_page_break()
        elif kind=='heading':
            paragraph=doc.add_heading(value,level=0 if index==0 else 1)
            for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
        elif kind=='text':doc.add_paragraph(value)
        elif kind=='image':
            from PIL import Image as PILImage
            raw=base64.b64decode(value['data'])
            with PILImage.open(io.BytesIO(raw)) as im: scale=min(24/im.width,14.5/im.height); w=im.width*scale; h=im.height*scale
            doc.add_picture(io.BytesIO(raw),width=Cm(w),height=Cm(h))
            if value.get('caption'):doc.add_paragraph(value['caption'])
        elif kind=='table' and value:
            n=len(value[0]);table=doc.add_table(rows=0,cols=n);table.style='Table Grid';table.autofit=False
            widths=([26.7*.31]+[26.7*.69/(n-1)]*(n-1)) if n>=6 else [26.7/n]*n
            for col,width in zip(table.columns,widths):col.width=Cm(width)
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                border=OxmlElement('w:'+side);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            table._tbl.tblPr.append(borders)
            for index,row in enumerate(value):
                cells=table.add_row().cells
                if index==0:
                    marker=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(marker)
                for target,text,width in zip(cells,row,widths):
                    target.width=Cm(width)
                    if index==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EFF4');target._tc.get_or_add_tcPr().append(shade)
                    target.text=str(text)
                    for p in target.paragraphs:
                        for run in p.runs:run.font.size=Pt(7.5);run.bold=index==0
            doc.add_paragraph()
    footer=section.footer.paragraphs[0]
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    output=io.BytesIO();doc.save(output);return output.getvalue()


def _pdf(content):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, LongTable, TableStyle, Spacer, PageBreak, Image
    fontdir=Path(__file__).parent/'fonts'
    for name,file in [('Civensa','DejaVuSans.ttf'),('CivensaBold','DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
    normal=ParagraphStyle('normal',fontName='Civensa',fontSize=9,leading=13,spaceAfter=7)
    small=ParagraphStyle('small',fontName='Civensa',fontSize=7,leading=10,wordWrap='CJK')
    heading=ParagraphStyle('heading',fontName='CivensaBold',fontSize=14,leading=19,spaceAfter=12,keepWithNext=True)
    def para(text,style):return Paragraph(escape(str(text)).replace('\n','<br/>'),style)
    output=io.BytesIO();width,height=landscape(A4)
    doc=SimpleDocTemplate(output,pagesize=(width,height),leftMargin=36,rightMargin=36,topMargin=36,bottomMargin=36)
    story=[]
    for kind,value in content:
        if kind=='page':story.append(PageBreak())
        elif kind=='heading':story.append(para(value,heading))
        elif kind=='text':
            for line in value.split('\n'):story.append(para(line or ' ',normal))
        elif kind=='image':
            im=Image(io.BytesIO(base64.b64decode(value['data'])));scale=min((width-72)/im.imageWidth,(height-130)/im.imageHeight)
            im.drawWidth=im.imageWidth*scale;im.drawHeight=im.imageHeight*scale;story.append(im)
            if value.get('caption'):story.append(para(value['caption'],normal))
        elif kind=='table' and value:
            n=len(value[0]);available=width-72
            widths=([available*.31]+[available*.69/(n-1)]*(n-1)) if n>=6 else [available/n]*n
            table=LongTable([[para(v,small) for v in row] for row in value],colWidths=widths,repeatRows=1,splitByRow=1,splitInRow=1)
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.3,colors.HexColor('#bcc6d0')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e9eff4')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
            story.extend([table,Spacer(1,10)])
    def footer(canvas,doc):
        canvas.setFont('Civensa',8);canvas.drawRightString(width-36,20,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return output.getvalue()
