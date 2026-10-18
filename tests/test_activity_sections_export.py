"""CIV-046: actual report payloads, access checks and editable Word styles."""
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from sqlalchemy import select, update

from app.ar_export import build_activity_report_docx, TEMPLATE_PATH, STYLES, DRAFT
from app import activity_sections_service as service
from app.models import ARSectionContent, Organization
from test_activity_sections import section_client, login, ready_and_freeze, save, resource

pytestmark = pytest.mark.postgres


def export(client, doc):
    response = client.get(f'/activity-sections/reports/{doc.id}/export.docx')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument')
    assert response.headers['cache-control'] == 'private, no-store'
    return response, Document(BytesIO(response.content))


def text(document):
    return '\n'.join([p.text for p in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def test_definition_order_cover_styles_draft_notes_and_downloads(section_client, db, sample):
    client, docs, _, _ = section_client
    doc = docs['birim:' + str(sample['department'].id)]
    payload = service.report_data(db, sample['strategy'], doc.id)
    response, word = export(client, doc)
    assert '2027-faaliyet-raporu-test-mudurlugu-taslak.docx' in response.headers['content-disposition']
    assert [p.text for p in word.paragraphs if p.style.name.startswith('Civensa Başlık ')] == [s['title'] for s in payload['sections']]
    assert [p.style.name for p in word.paragraphs if p.style.name.startswith('Civensa Başlık ')] == [f"Civensa Başlık {s['depth'] + 1}" for s in payload['sections']]
    assert 'Test Belediyesi' in text(word) and 'Test Müdürlüğü' in text(word)
    assert all(DRAFT in s.header.paragraphs[0].text for s in word.sections)
    assert any('Bölüm içeriğini gözden geçirip' in p.text and p.style.name == 'Civensa Not' for p in word.paragraphs)
    assert set(STYLES) <= {s.name for s in Document(TEMPLATE_PATH).styles}
    assert all(p.style.name in STYLES for p in word.paragraphs)
    assert all(t.style.name == 'Civensa Tablo' for t in word.tables)
    assert word.settings.element.find(qn('w:updateFields')) is None
    assert 'TOC' in word.element.xml and '\\o "1-3"' in word.element.xml
    assert "Sayfa numaraları için Word'de içindekilere sağ tıklayıp Alanı Güncelleştir seçin." in text(word)
    for path in ['/activity-sections', f'/activity-sections/reports/{doc.id}']:
        html = client.get(path).text
        assert 'Word olarak indir' in html and f'/reports/{doc.id}/export.docx' in html
    _, administration = export(client, docs['idare'])
    assert '2027' in text(administration)
    assert [p.text for p in administration.paragraphs if p.style.name.startswith('Civensa Başlık ')] == [s['title'] for s in service.report_data(db, sample['strategy'], docs['idare'].id)['sections']]


def test_access_own_other_idare_and_strategy(section_client, sample):
    client, docs, other, _ = section_client
    own = docs['birim:' + str(sample['department'].id)]
    login(client, 'test@example.invalid')
    export(client, own)
    for doc in [docs['idare'], docs['birim:' + str(other.id)]]:
        response = client.get(f'/activity-sections/reports/{doc.id}/export.docx')
        assert response.status_code == 403 and 'rapor' in response.json()['detail']
        assert f'/reports/{doc.id}/export.docx' not in client.get('/activity-sections').text
    login(client, 'strategy@example.invalid')
    for doc in docs.values():
        export(client, doc)


def test_blocks_resources_null_zero_not_applicable_landscape(section_client, db, sample):
    client, docs, _, _ = section_client
    doc = docs['birim:' + str(sample['department'].id)]
    token = login(client, 'strategy@example.invalid')
    assert save(client, token, doc, 'I.A', blocks=[{'type': 'paragraph', 'text': 'Örnek açıklama'},
        {'type': 'bullets', 'items': ['Birinci madde']},
        {'type': 'table', 'rows': [[f'Sütun {i}' for i in range(10)], ['Uzun değer' * 15] * 10]}]).status_code == 200
    assert save(client, token, doc, 'I.B', status='not_applicable', reason='Bu yıl yetki devri yok.', blocks=[]).status_code == 200
    assert resource(client, token, doc, 'hardware', [{'category':'Bilgisayar','count':0}, {'category':'Yazıcı','count':None}]).status_code == 200
    _, word = export(client, doc)
    rendered = text(word)
    assert 'Uygulanamaz: Bu yıl yetki devri yok.' in rendered
    assert 'Örnek açıklama' in rendered
    assert any(p.text == 'Birinci madde' and p.style.name == 'Civensa Madde' for p in word.paragraphs)
    assert any(s.orientation == WD_ORIENT.LANDSCAPE for s in word.sections)
    assert word.sections[-1].orientation == WD_ORIENT.PORTRAIT
    hardware = next(t for t in word.tables if t.cell(0, 0).text == 'Türü')
    assert [r.cells[1].text for r in hardware.rows[1:]] == ['0', '—']
    max_width = max(s.page_width - s.left_margin - s.right_margin for s in word.sections)
    for table in word.tables:
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        assert sum(c.width for c in table.columns) <= max_width
        assert table.autofit is False


def test_financial_all_slots_money_and_indicator_table(section_client, db, sample):
    client, docs, _, _ = section_client
    payload = service.report_data(db, sample['strategy'], docs['idare'].id)
    row = dict(department='Örnek Müdürlüğü', classification='01 · Personel', economic_code='01', initial=Decimal('12500.23'),
               added=Decimal('0'), deducted=None, entered_final=None, actual=Decimal('0'), cancelled=None, carried=None,
               final=None, rate=None, pp_budget=None, explanation='Merkez açıklaması', birim_explanation='Birim açıklaması')
    payload['financial']['tables']['expenditure']['rows'] = [row]
    payload['financial']['tables']['revenue']['rows'] = [dict(row, estimate=Decimal('12500.23'), refunds=None, net=None)]
    payload['financial']['entries'] = {'assets':[dict(label='Bina', amount=Decimal('12500.23'), previous=None, explanation='Varlık notu')],
        'audit':[dict(audit_type='İç denetim', auditor='Örnek kurum', period='2027', report_date='2027-12-31', report_number='1', findings='Bulgu', measures='Tedbir', status='Açık')],
        'aid':[dict(name='Örnek kuruluş', aid_type='Nakdî', amount=0, basis='Karar', explanation='Sonuç')],
        'explanations':[dict(row_key='01', explanation='Sapma gerekçesi')],
        'audit_directorates':[dict(department='Örnek birim', audit_type='Dış denetim', findings='Birim bulgusu')]}
    payload['financial']['statements'][0].update(period='2027', statement_date='2027-12-31', explanation='Bilanço açıklaması', filename='ornek.pdf', sha256='a' * 64,
        figures=[dict(label='Aktif toplamı', amount=Decimal('12500.23'))])
    payload['core']['indicators'] = [dict(name='Hizmet sayısı', unit='Adet', targets=[dict(value='12500.23', state='verified')], actual=0,
        actual_state='verified', data_source='Kayıt defteri', evidence=[dict(title='Sayım', locator='Sayfa 1', source_revision='2027')])]
    word = Document(BytesIO(build_activity_report_docx(payload, TEMPLATE_PATH)))
    rendered = text(word)
    for value in ['12.500,23', '0,00', '—', 'Merkez açıklaması', 'Birim açıklaması', 'Birim bulgusu',
                  'Varlık notu', 'Tedbir', 'Örnek kuruluş', 'Bilanço açıklaması', 'Aktif toplamı', 'Kayıt defteri', 'Sayım']:
        assert value in rendered
    indicator = next(t for t in word.tables if t.cell(0, 0).text == 'Gösterge')
    assert [c.text for c in indicator.rows[1].cells][:4] == ['Hizmet sayısı', 'Adet', '12.500,23', '0']


def test_freeze_export_preserves_body_org_resources_and_assurances(section_client, db, sample, monkeypatch):
    client, docs, _, _ = section_client
    doc = docs['birim:' + str(sample['department'].id)]
    token = login(client, 'strategy@example.invalid')
    assert resource(client, token, doc, 'hardware', [{'category':'Donmuş bilgisayar','count':5}]).status_code == 200
    ready_and_freeze(client, token, doc)
    # Disposable fixture schemas deliberately have no migration guards: simulate live
    # source changes and prove the HTTP export still reads only the frozen snapshot.
    db.execute(update(ARSectionContent).where(ARSectionContent.document_id == doc.id).values(blocks_json='[{"type":"paragraph","text":"Sonradan değişen canlı içerik"}]'))
    db.get(Organization, doc.organization_id).name = 'Değişen canlı kurum'
    db.commit()
    from app import assurance_service
    def unexpected_live_resolution(*args):
        raise AssertionError('Frozen assurance text must come from the snapshot')
    monkeypatch.setattr(assurance_service, 'view', unexpected_live_resolution)
    response, word = export(client, doc)
    assert '-onayli.docx' in response.headers['content-disposition']
    assert all(DRAFT not in s.header.paragraphs[0].text for s in word.sections)
    assert all('Onaylı kopya · SHA-256:' in s.footer.paragraphs[0].text for s in word.sections)
    assert 'PAGE' in word.sections[0].footer._element.xml
    rendered = text(word)
    assert 'Yıllık değerlendirme' in rendered and 'Donmuş bilgisayar' in rendered
    assert 'Test Belediyesi' in rendered and 'Değişen canlı kurum' not in rendered
    assert 'Sonradan değişen canlı içerik' not in rendered
    assert 'HARCAMA YETKİLİSİNİN İÇ KONTROL GÜVENCE BEYANI' in rendered
    assert 'İmza: ________________' in rendered


def test_template_restyling_survives_export(section_client, db, sample, tmp_path):
    _, docs, _, _ = section_client
    payload = service.report_data(db, sample['strategy'], docs['idare'].id)
    template = Document(TEMPLATE_PATH)
    template.styles['Civensa Başlık 1'].font.color.rgb = RGBColor.from_string('FF0000')
    template.styles['Civensa Gövde'].font.size = Pt(13)
    path = tmp_path / 'restyled.docx'
    template.save(path)
    word = Document(BytesIO(build_activity_report_docx(payload, path)))
    assert word.styles['Civensa Başlık 1'].font.color.rgb == RGBColor.from_string('FF0000')
    assert word.styles['Civensa Gövde'].font.size == Pt(13)
    assert not any(r.rPr is not None for p in word.paragraphs for r in p._p.r_lst)
    with ZipFile(BytesIO(build_activity_report_docx(payload, path))) as archive:
        assert archive.testzip() is None


def test_core_pinned_targets_latest_actual_and_frozen_snapshot(section_client, db, sample):
    from app import activity_contract_service as contract
    from app.models import Publication, ReportSection, Submission, DepartmentScope, ReportIndicator
    from test_activity_contract import actual_payload
    client, docs, other, finance = section_client
    token = login(client, 'strategy@example.invalid')
    db.scalar(select(Submission)).status = 'approved'
    db.add(ReportSection(period_id=sample['period'].id, title='Genel Bilgiler', kind='text', content_json='{}'))
    for dep in (other, finance):
        scope = db.scalar(select(DepartmentScope).where(DepartmentScope.department_id == dep.id))
        scope.status = 'excluded'
        scope.reason = 'Örnek yayın kapsamı dışında'
    db.commit()
    response = client.post('/report-builder/publish', data={'csrf_token':token, 'period_id':sample['period'].id})
    assert response.status_code == 303, response.text
    publication = db.scalar(select(Publication))
    revision = contract.create_report_revision(db, sample['strategy'], publication.id, sample['department'].organization_id)
    db.commit()
    indicator = db.scalar(select(ReportIndicator).where(ReportIndicator.report_revision_id == revision.id))
    response = client.post(f'/api/activity-reports/indicators/{indicator.id}/actuals',
        headers={'X-CSRF-Token':token}, json=actual_payload())
    assert response.status_code == 200, response.text
    doc = docs['birim:' + str(sample['department'].id)]
    before = service.report_data(db, sample['strategy'], doc.id)['core']['indicators'][0]
    assert before['name'] == 'Hizmet sayısı' and Decimal(before['targets'][0]['value']) == Decimal('1.75')
    assert before['actual'] == Decimal('2.750001') and before['evidence'][0]['title'] == 'Kayıt defteri'
    ready_and_freeze(client, token, doc)
    response = client.post(f'/api/activity-reports/indicators/{indicator.id}/actuals',
        headers={'X-CSRF-Token':token}, json=actual_payload(expected_revision=1, value='99'))
    assert response.status_code == 200, response.text
    _, word = export(client, doc)
    table = next(t for t in word.tables if t.cell(0, 0).text == 'Gösterge')
    assert [c.text for c in table.rows[1].cells][:4] == ['Hizmet sayısı', 'Adet', '1,75', '2,750001']
    assert 'Kayıt defteri' in table.rows[1].cells[4].text


def test_template_error_is_actionable_turkish(section_client, monkeypatch, tmp_path):
    from app import ar_export
    client, docs, _, _ = section_client
    monkeypatch.setattr(ar_export, 'TEMPLATE_PATH', tmp_path / 'missing.docx')
    response = client.get(f"/activity-sections/reports/{docs['idare'].id}/export.docx")
    assert response.status_code == 503
    assert 'Word şablonu açılamadı' in response.json()['detail']
    assert 'Sistem sorumlusundan' in response.json()['detail']


def assert_current_turkish_word_document(document):
    settings = document.settings.element
    assert settings.find(qn('w:updateFields')) is None
    modes = [node for node in settings.iter(qn('w:compatSetting'))
             if node.get(qn('w:name')) == 'compatibilityMode']
    assert len(modes) == 1 and modes[0].get(qn('w:val')) == '15'
    defaults = document.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    assert defaults.find(qn('w:lang')).get(qn('w:val')) == 'tr-TR'
    for key in ('eastAsia', 'bidi'):
        original = Document().styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr')).find(qn('w:lang'))
        assert defaults.find(qn('w:lang')).get(qn('w:' + key)) == original.get(qn('w:' + key))
    fonts = defaults.find(qn('w:rFonts'))
    assert fonts.get(qn('w:ascii')) == fonts.get(qn('w:hAnsi')) == 'Arial'
    assert fonts.get(qn('w:asciiTheme')) is None and fonts.get(qn('w:hAnsiTheme')) is None
    for style in document.styles:
        if style.name.startswith('Civensa '):
            assert style.element.rPr.find(qn('w:lang')).get(qn('w:val')) == 'tr-TR'
    assert document.styles['Civensa Gövde'].font.name == 'Arial'
    assert document.styles['Civensa Gövde'].font.size == Pt(10)
    for level in range(1, 5):
        style = document.styles[f'Civensa Başlık {level}']
        assert style.element.pPr.find(qn('w:outlineLvl')).get(qn('w:val')) == str(level - 1)
    for level in range(1, 4):
        style = document.styles[f'TOC {level}']
        assert abs(style.paragraph_format.left_indent.twips - (level - 1) * document.styles['TOC 2'].paragraph_format.left_indent.twips) <= 1
    first = document.sections[0]
    assert first.different_first_page_header_footer
    assert first._sectPr.find(qn('w:titlePg')) is not None
    assert first.first_page_header._element.xpath('.//w:t | .//w:fldChar | .//w:fldSimple') == []
    assert first.first_page_footer._element.xpath('.//w:t | .//w:fldChar | .//w:fldSimple') == []


@pytest.mark.parametrize('frozen', [False, True])
def test_word_mac_settings_cached_toc_and_cover(section_client, db, sample, frozen):
    _, docs, _, _ = section_client
    payload = service.report_data(db, sample['strategy'], docs['idare'].id)
    payload.update(frozen=frozen, sha256='a' * 64)
    word = Document(BytesIO(build_activity_report_docx(payload, TEMPLATE_PATH)))
    assert_current_turkish_word_document(Document(TEMPLATE_PATH))
    assert_current_turkish_word_document(word)
    instruction = ''.join(word.element.xpath('.//w:instrText/text()'))
    assert 'TOC' in instruction and '\\o "1-3"' in instruction and '\\u' in instruction
    entries = []
    within_result = False
    finished = False
    for paragraph in word.paragraphs:
        markers = paragraph._p.xpath('.//w:fldChar/@w:fldCharType')
        if 'separate' in markers:
            within_result = True
        if paragraph.style.name in {'TOC 1', 'TOC 2', 'TOC 3'}:
            assert within_result and not finished
            assert paragraph._p.xpath('.//w:tab | .//w:fldSimple') == []
            entries.append((paragraph.style.name, paragraph.text))
        if 'end' in markers:
            finished = True
            within_result = False
    assert finished
    assert entries == [(f"TOC {s['depth'] + 1}", s['title']) for s in payload['sections'] if s['depth'] < 3]
    assert 'Sayfa numaraları için Word\'de içindekilere sağ tıklayıp Alanı Güncelleştir seçin.' in text(word)
    assert all(not section.different_first_page_header_footer for section in list(word.sections)[1:])
    assert (DRAFT in word.sections[0].header.paragraphs[0].text) is (not frozen)


def test_checked_in_demo_has_current_word_settings_and_cached_toc():
    from pathlib import Path
    word = Document(Path(__file__).resolve().parents[1] / 'docs/civ046/ornek-taslak.docx')
    assert_current_turkish_word_document(word)
    headings = [(int(p.style.name.rsplit(' ', 1)[1]), p.text) for p in word.paragraphs
                if p.style.name.startswith('Civensa Başlık ')]
    entries = [(int(p.style.name[-1]), p.text) for p in word.paragraphs if p.style.name in {'TOC 1', 'TOC 2', 'TOC 3'}]
    assert entries == [(level, title) for level, title in headings if level <= 3]
    assert 'Hiçbir içindekiler tablosu öğesine rastlanmadı' not in text(word)
