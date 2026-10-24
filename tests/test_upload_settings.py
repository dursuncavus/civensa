"""Upload defaults and environment overrides, without application/database mutation."""
import subprocess
import os
import sys

import pytest
from fastapi import HTTPException
from app import config
from app.activity_evidence_service import validate_file


@pytest.mark.parametrize('data', [None, b''])
def test_empty_evidence_is_none_safe_and_actionable(data):
    with pytest.raises(HTTPException) as caught:
        validate_file(data, 'kanit.pdf')
    assert caught.value.status_code == 422
    assert caught.value.detail == {'message': 'Kanıt dosyası boş. PDF veya XLSX dosyasını yeniden seçip yükleyin.', 'field': 'source_file'}


def test_upload_setting_defaults_and_overrides(monkeypatch):
    for key in ('CIVENSA_ATTACHMENT_MAX_MB', 'CIVENSA_PP_IMAGE_MAX_MB', 'CIVENSA_PP_IMAGE_MAX_PIXELS'):
        monkeypatch.delenv(key, raising=False)
    assert (config.attachment_max_bytes(), config.pp_image_max_bytes(), config.pp_image_max_pixels()) == (10*1024**2, 5*1024**2, 16000000)
    monkeypatch.setenv('CIVENSA_ATTACHMENT_MAX_MB', '2')
    monkeypatch.setenv('CIVENSA_PP_IMAGE_MAX_MB', '3')
    monkeypatch.setenv('CIVENSA_PP_IMAGE_MAX_PIXELS', '4000000')
    assert (config.attachment_max_bytes(), config.pp_image_max_bytes(), config.pp_image_max_pixels()) == (2*1024**2, 3*1024**2, 4000000)
    # Import-time attachment constants must also use settings in both consumers.
    result = subprocess.run([sys.executable, '-c',
        'from app.activity_evidence_service import MAX_FILE as evidence; '
        'from app.activity_financial_service import MAX_FILE as financial; '
        'assert evidence == financial == 2*1024**2'], env=os.environ.copy(), capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize('value', ['0', '-1', '1.5', 'bozuk', ''])
@pytest.mark.parametrize('setting', ['CIVENSA_ATTACHMENT_MAX_MB', 'CIVENSA_PP_IMAGE_MAX_MB', 'CIVENSA_PP_IMAGE_MAX_PIXELS'])
def test_invalid_upload_setting_is_actionable(monkeypatch, setting, value):
    monkeypatch.setenv(setting, value)
    with pytest.raises(RuntimeError, match='pozitif bir tam sayı.*düzeltip uygulamayı yeniden başlatın'):
        config.positive_setting(setting, 5)


@pytest.mark.parametrize('pixels,image_kept', [('4', True), ('3', False)])
def test_word_import_uses_configured_image_pixel_limit(monkeypatch, pixels, image_kept):
    from io import BytesIO
    from docx import Document
    from PIL import Image
    from app.import_service import parse_pp_import
    monkeypatch.setenv('CIVENSA_PP_IMAGE_MAX_PIXELS', pixels)
    document = Document()
    document.add_heading('Başkan Sunuşu', level=1)
    document.add_paragraph('Korunacak sunuş metni')
    picture = BytesIO()
    Image.new('RGB', (2, 2), 'white').save(picture, format='PNG')
    picture.seek(0)
    document.add_picture(picture)
    output = BytesIO()
    document.save(output)
    preview = parse_pp_import(output.getvalue())
    assert any(section['kind'] == 'image' for section in preview['general_sections']) == image_kept
    if not image_kept:
        assert any('PNG/JPEG' in warning and 'yeniden ekleyin' in warning and '0,000003 megapiksel' in warning
                   for warning in preview['general_section_warnings'])
