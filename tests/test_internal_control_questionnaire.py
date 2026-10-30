"""CIV-085 PostgreSQL acceptance. Every question is explicitly synthetic."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
from uuid import uuid4
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import internal_control_monitoring_service as m
from app import internal_control_monitoring_export as legacy_export
from app import internal_control_questionnaire_service as q
from app.internal_control_questionnaire_export import questionnaire_export
from app.models import Department, User
from test_internal_control_monitoring import monitoring, fill

pytestmark = pytest.mark.postgres
ROOT = Path(__file__).resolve().parents[1]


def question_file(rows=None, count=6):
    book = load_workbook(BytesIO(q.template()))
    rows = rows if rows is not None else [[i, list(q.COMPONENTS)[(i-1) % 5], f'T{i}', f'Test sorusu {i}'] for i in range(1, count+1)]
    for row in rows:
        book.active.append(row)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture
def questionnaire(db, monitoring):
    d = monitoring
    third = Department(organization_id=d['department'].organization_id, name='Test Üçüncü Müdürlüğü')
    db.add(third)
    db.flush()
    third_user = User(email='third@civ085.invalid', full_name='Üçüncü birim', role='department',
        department_id=third.id, password_hash='unused', can_submit=True)
    db.add(third_user)
    catalogue = q.save_list(db, d['strategy'], 'Sentetik soru listesi', question_file([
        [1, 'KOS', 'T1', 'Test sorusu 1'], [2, 'Kontrol Ortamı', 'T2', 'Test sorusu 2'],
        [3, 'RDS', 'T3', 'Test sorusu 3'], [4, 'Risk Değerlendirme', 'T4', 'Test sorusu 4']]))
    q.attach(db, d['strategy'], d['period'].id, catalogue.id, [d['department'].id, d['collaborator'].department_id, third.id], 0)
    db.commit()
    selected = q.batches(db, d['period'].id)
    by_dep = {b.department_id: b for b in selected}
    return {**d, 'catalogue': catalogue, 'question_batches': selected, 'question_owned': by_dep[d['department'].id],
        'question_other': by_dep[d['collaborator'].department_id], 'question_third': by_dep[third.id], 'third_user': third_user}


def answer_data(*answers):
    return {f'answer_T{i}': a for i, a in enumerate(answers, 1)}


def answer_and_approve(db, d, batch, actor, answers):
    q.save_answers(db, actor, batch.id, q.state(db, batch.id)[1], answer_data(*answers), submit=True)
    q.transition(db, d['strategy'], batch.id, 'approve', q.state(db, batch.id)[1])
    db.commit()


def test_six_question_upload_immutable_versions(db, monitoring):
    d = monitoring
    content = question_file()
    assert len(q.parse_questions(content)) == 6
    first = q.save_list(db, d['strategy'], 'Sentetik altı soru', content)
    second = q.save_list(db, d['staff'], 'Sentetik altı soru', content)
    db.commit()
    assert first.id != second.id and first.actor_id == d['strategy'].id
    assert first.created_at and json.loads(first.questions_json)[0]['question'] == 'Test sorusu 1'
    with pytest.raises(HTTPException):
        q.save_list(db, d['user'], 'Yetkisiz', content)
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 2


@pytest.mark.parametrize('change,message,line', [
    ('empty', 'Soru boş', 'Satır 3'), ('code', 'Kod', 'Satır 3'),
    ('order', 'Sıra', 'Satır 3'), ('component', 'Bileşen', 'Satır 3'),
    ('too_many', '200 soru', 'Satır 202'), ('bad_order', 'Sıra geçersiz', 'Satır 3'),
    ('formula', 'Formül', 'Satır 3'), ('blank_code', 'Kod', 'Satır 3'),
])
def test_upload_row_errors_atomic(db, monitoring, change, message, line):
    rows = [[i, 'KOS', f'T{i}', f'Test sorusu {i}'] for i in range(1, 202 if change == 'too_many' else 7)]
    if change == 'empty': rows[1][3] = ''
    if change == 'code': rows[1][2] = 't1'
    if change == 'order': rows[1][0] = 1
    if change == 'component': rows[1][1] = 'Bilinmeyen'
    if change == 'bad_order': rows[1][0] = -1
    if change == 'formula': rows[1][3] = '=1+1'
    if change == 'blank_code': rows[1][2] = ''
    with pytest.raises(HTTPException) as exc:
        q.save_list(db, monitoring['strategy'], 'Sentetik hatalı', question_file(rows))
    messages = ' '.join(e['message'] for e in exc.value.pp_errors)
    assert message in messages and line in messages
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 0


@pytest.mark.parametrize('content', [b'', b'not an excel file'])
def test_unreadable_upload(db, monitoring, content):
    with pytest.raises(HTTPException) as exc:
        q.save_list(db, monitoring['strategy'], 'Sentetik', content)
    assert 'Satır 1' in exc.value.detail and 'yeniden' not in str(content)
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 0


def test_attach_permissions_open_freeze_and_add(db, questionnaire):
    d = questionnaire
    period_id, list_id = d['period'].id, d['catalogue'].id
    ids = [b.department_id for b in d['question_batches']]
    for actor in (d['user'], d['staff']):
        with pytest.raises(HTTPException) as exc:
            q.attach(db, actor, period_id, list_id, ids, 1)
        assert exc.value.status_code == 403
    # Before the first answer both catalogue and selected participants can change.
    replacement = q.save_list(db, d['strategy'], 'Sentetik yeni sürüm', question_file(count=4))
    q.attach(db, d['strategy'], period_id, replacement.id, ids[:2], 1)
    q.attach(db, d['strategy'], period_id, list_id, ids, 2)
    db.commit()
    q.save_answers(db, d['user'], d['question_owned'].id, 0, answer_data('yes'))
    db.commit()
    for new_list, participants in ((replacement.id, ids), (list_id, ids[:1])):
        with pytest.raises(HTTPException) as exc:
            q.attach(db, d['strategy'], period_id, new_list, participants, 3)
        assert exc.value.status_code == 409
    addition = Department(organization_id=d['department'].organization_id, name='Test Sonradan Müdürlüğü')
    db.add(addition)
    db.flush()
    q.attach(db, d['strategy'], period_id, list_id, ids + [addition.id], 3)
    db.commit()
    assert len(q.batches(db, period_id)) == 4
    assert q.state(db, d['question_owned'].id) == ('draft', 1)
    with pytest.raises(HTTPException):
        q.attach(db, d['strategy'], period_id, list_id, ids, 3)
    m.close_period(db, d['strategy'], period_id, 1, 'Test kapatma')
    db.commit()
    with pytest.raises(HTTPException) as exc:
        q.attach(db, d['strategy'], period_id, list_id, ids + [addition.id], 4)
    assert exc.value.status_code == 409


def test_attachment_invalid_fields_and_revision(db, questionnaire):
    d = questionnaire
    for list_id, departments, revision in [(999999, [d['department'].id], 1),
        (d['catalogue'].id, [], 1), (d['catalogue'].id, [999999], 1),
        (d['catalogue'].id, [d['department'].id], 0)]:
        with pytest.raises(HTTPException):
            q.attach(db, d['strategy'], d['period'].id, list_id, departments, revision)
    assert q.attachment(db, d['period'].id).revision == 1


def test_partial_ownership_submit_missing_deadline_stale(db, questionnaire, monkeypatch):
    d = questionnaire
    b = d['question_owned']
    for actor, target in ((d['strategy'], b), (d['staff'], b), (d['user'], d['question_other'])):
        with pytest.raises(HTTPException) as exc:
            q.save_answers(db, actor, target.id, 0, answer_data('yes'))
        assert exc.value.status_code == 403
    with pytest.raises(HTTPException): q.get_batch(db, d['user'], d['question_other'].id)
    q.save_answers(db, d['user'], b.id, 0, {**answer_data('yes'), 'explanation_T2': 'Test açıklaması'})
    db.commit()
    assert [r['answer'] for r in q.rows(db, b)] == ['yes', None, None, None]
    with pytest.raises(HTTPException) as exc:
        q.transition(db, d['user'], b.id, 'submit', 1)
    assert all(code in exc.value.detail for code in ('T2', 'T3', 'T4'))
    assert 'T1' not in exc.value.detail
    with pytest.raises(HTTPException) as exc:
        q.save_answers(db, d['user'], b.id, 0, answer_data('no'))
    assert exc.value.status_code == 409
    monkeypatch.setattr(m, 'today', lambda db: d['period'].deadline + timedelta(days=1))
    for operation in (lambda: q.save_answers(db, d['user'], b.id, 1, answer_data('no')),
                      lambda: q.transition(db, d['user'], b.id, 'submit', 1)):
        with pytest.raises(HTTPException) as exc: operation()
        assert exc.value.status_code == 409 and 'süresi doldu' in exc.value.detail


@pytest.mark.parametrize('data', [answer_data('invalid'), {'explanation_T1': 'a' * 1001}])
def test_answer_validation_is_atomic(db, questionnaire, data):
    d = questionnaire
    with pytest.raises(HTTPException) as exc:
        q.save_answers(db, d['user'], d['question_owned'].id, 0, data)
    assert 'T1' in exc.value.detail
    assert q.state(db, d['question_owned'].id) == ('draft', 0)


def test_workflow_reasons_roles_history_and_submit_authority(db, questionnaire):
    d = questionnaire
    b = d['question_owned']
    d['user'].can_submit = False
    db.commit()
    with pytest.raises(HTTPException):
        q.save_answers(db, d['user'], b.id, 0, answer_data('yes', 'no', 'developing', 'yes'), submit=True)
    assert q.state(db, b.id) == ('draft', 0)
    d['user'].can_submit = True
    db.commit()
    q.save_answers(db, d['user'], b.id, 0, answer_data('yes', 'no', 'developing', 'yes'), submit=True)
    db.commit()
    assert q.state(db, b.id) == ('submitted', 2)
    for action in ('return', 'approve', 'reopen'):
        for actor in (d['user'], d['staff']):
            with pytest.raises(HTTPException) as exc:
                q.transition(db, actor, b.id, action, 2, 'Test gerekçesi')
            assert exc.value.status_code == 403
    with pytest.raises(HTTPException): q.transition(db, d['strategy'], b.id, 'return', 2, ' ')
    q.transition(db, d['strategy'], b.id, 'return', 2, 'Test iade gerekçesi')
    db.commit()
    q.save_answers(db, d['user'], b.id, 3, answer_data('no', 'no', 'developing', 'yes'), submit=True)
    q.transition(db, d['strategy'], b.id, 'approve', 5)
    db.commit()
    approval = q.latest_event(db, b.id)
    with pytest.raises(HTTPException): q.transition(db, d['strategy'], b.id, 'reopen', 6, '')
    q.transition(db, d['strategy'], b.id, 'reopen', 6, 'Test yeniden açma')
    db.commit()
    assert q.state(db, b.id) == ('draft', 7)
    assert db.get(q.Event, approval.id).to_status == 'approved'
    assert q.results(db, d['strategy'], d['period'])['included'] == []
    m.close_period(db, d['strategy'], d['period'].id, 1, 'Test kapatma')
    db.commit()
    with pytest.raises(HTTPException): q.save_answers(db, d['user'], b.id, 7, answer_data('yes'))


def test_separation_of_duties_survives_role_change(db, questionnaire):
    d = questionnaire
    b = d['question_owned']
    q.save_answers(db, d['user'], b.id, 0, answer_data('yes', 'yes', 'yes', 'yes'), submit=True)
    db.commit()
    d['user'].department_id = d['third_user'].department_id
    d['user'].role = 'strategy'
    d['user'].position = 'director'
    db.commit()
    with pytest.raises(HTTPException) as exc:
        q.transition(db, d['user'], b.id, 'approve', 2)
    assert exc.value.status_code == 403 and 'Görevler ayrılığı' in exc.value.detail
    q.transition(db, d['strategy'], b.id, 'approve', 2)
    db.commit()
    assert q.state(db, b.id)[0] == 'approved'


def test_hand_computed_three_directorates_four_questions(db, questionnaire):
    d = questionnaire
    # T1 Y/Y, T2 N/D, T3 D/N, T4 Y/N. Third (all Y) stays submitted.
    answer_and_approve(db, d, d['question_owned'], d['user'], ['yes','no','developing','yes'])
    answer_and_approve(db, d, d['question_other'], d['collaborator'], ['yes','developing','no','no'])
    q.save_answers(db, d['third_user'], d['question_third'].id, 0, answer_data('yes','yes','yes','yes'), submit=True)
    db.commit()
    result = q.results(db, d['strategy'], d['period'])
    assert len(result['included']) == 2 and result['pending'] == [d['question_third'].label]
    assert [r['counts'] for r in result['questions']] == [
        {'yes':2,'no':0,'developing':0}, {'yes':0,'no':1,'developing':1},
        {'yes':0,'no':1,'developing':1}, {'yes':1,'no':1,'developing':0}]
    assert result['components'][0]['counts'] == {'yes':2,'no':1,'developing':1}
    assert result['components'][1]['counts'] == {'yes':1,'no':2,'developing':1}
    assert result['total']['counts'] == {'yes':3,'no':3,'developing':2}
    assert result['total']['rates'] == {'yes':'37,50% (3/8)', 'no':'37,50% (3/8)', 'developing':'25,00% (2/8)'}
    assert result['questions'][0]['rates']['yes'] == '100,00% (2/2)'
    assert result['components'][2]['rates']['yes'] == '— (0/0)'
    assert m.rate_text(2, 3) == '66,67% (2/3)'
    q.transition(db, d['strategy'], d['question_third'].id, 'approve', 2)
    db.commit()
    assert q.results(db, d['strategy'], d['period'])['questions'][3]['rates']['yes'] == '66,67% (2/3)'
    q.transition(db, d['strategy'], d['question_other'].id, 'reopen', 3, 'Test yeniden açma')
    db.commit()
    assert q.results(db, d['strategy'], d['period'])['total']['total'] == 8


def test_zero_denominator_and_draft_export_formula_safety(db, questionnaire):
    d = questionnaire
    result = q.results(db, d['strategy'], d['period'])
    assert result['total']['rates'] == {key: '— (0/0)' for key in q.ANSWERS}
    with pytest.raises(HTTPException) as exc:
        questionnaire_export(db, d['strategy'], d['period'].id)
    assert exc.value.status_code == 409
    q.save_answers(db, d['user'], d['question_owned'].id, 0, {'answer_T1':'no', 'explanation_T1':'=1+1'})
    db.commit()
    book = load_workbook(BytesIO(questionnaire_export(db, d['strategy'], d['period'].id, draft=True)))
    assert book.sheetnames == ['Bileşen özeti', 'Soru özeti', 'Müdürlük yanıtları']
    assert all(sheet['A1'].value == 'TASLAK' for sheet in book)
    assert book.worksheets[0]['B5'].value == '— (0/0)'
    answer_rows = list(book.worksheets[2].iter_rows(min_row=3))
    literal = next(r[5] for r in answer_rows if r[5].value == '=1+1')
    assert literal.data_type == 's'
    assert any(r[4].value is None for r in answer_rows)


def test_final_export_only_all_current_approved(db, questionnaire):
    d = questionnaire
    for batch, actor in [(d['question_owned'], d['user']), (d['question_other'], d['collaborator']), (d['question_third'], d['third_user'])]:
        answer_and_approve(db, d, batch, actor, ['yes','no','developing','yes'])
    book = load_workbook(BytesIO(questionnaire_export(db, d['strategy'], d['period'].id)))
    assert all(s['A1'].value == 'Onaylı veriler' for s in book)
    assert book.worksheets[0]['B10'].value == '50,00% (6/12)'


@pytest.mark.parametrize('table', ['ic_question_lists', 'ic_question_attachments', 'ic_question_batches', 'ic_question_events'])
def test_raw_history_update_delete_truncate_refused(db, questionnaire, table):
    d = questionnaire
    q.save_answers(db, d['user'], d['question_owned'].id, 0, answer_data('yes'))
    db.commit()
    schema = db.get_bind().get_execution_options()['schema_translate_map'][None]
    table = f'"{schema}".{table}'
    for sql in (f'UPDATE {table} SET id=id', f'DELETE FROM {table}', f'TRUNCATE {table} CASCADE'):
        with pytest.raises(DBAPIError) as exc:
            with db.begin_nested(): db.execute(text(sql))
        assert 'geçmişi değiştirilemez' in str(exc.value)


def test_raw_event_permissions_and_snapshot_guard(db, questionnaire):
    d = questionnaire
    b = d['question_owned']
    valid = [dict(code=f'T{i}', answer='yes', explanation='') for i in range(1,5)]
    for actor_id, snapshot, action in [(d['strategy'].id, valid, 'save'),
        (d['user'].id, valid[:1], 'save'),
        (d['user'].id, [{**r, 'answer':'unknown'} for r in valid], 'save'),
        (d['user'].id, valid, 'approve')]:
        with pytest.raises(DBAPIError):
            with db.begin_nested():
                db.execute(q.Event.__table__.insert().values(batch_id=b.id, revision=1, action=action,
                    from_status='draft', to_status='draft' if action=='save' else 'approved', actor_id=actor_id,
                    actor_name='Test aktör', reason='', snapshot_json=json.dumps(snapshot)))
    assert q.state(db, b.id) == ('draft', 0)


def test_http_only_own_form_and_retains_incomplete_submit(db, questionnaire, client):
    from app.main import app, require_user
    d = questionnaire
    b = d['question_owned']
    assert client.get(f'/internal-control/questionnaires/batches/{d["question_other"].id}').status_code == 403
    assert client.get('/internal-control/questionnaires').status_code == 403
    response = client.get(f'/internal-control/our-actions?period_id={d["period"].id}')
    assert 'Test sorusu 1' in response.text and d['question_third'].label not in response.text
    token = response.context['request'].session.get('csrf_token', '')
    response = client.post(f'/internal-control/questionnaires/batches/{b.id}/answers', data={
        'csrf_token':token, 'list_id':d['catalogue'].id, 'expected_revision':'0', 'operation':'submit', 'answer_T1':'yes', 'explanation_T1':'Korunan test açıklaması'})
    assert response.status_code == 422 and 'T2, T3, T4' in response.text
    assert 'Korunan test açıklaması' in response.text
    assert 'name="answer_T1" value="yes" checked' in response.text
    assert q.state(db, b.id) == ('draft', 0)
    app.dependency_overrides[require_user] = lambda: d['strategy']
    assert client.get(f'/internal-control/questionnaires/batches/{b.id}').status_code == 200
    response = client.post(f'/internal-control/questionnaires/batches/{b.id}/answers', data={
        'csrf_token':token, 'list_id':d['catalogue'].id, 'expected_revision':'0', 'answer_T1':'yes'})
    assert response.status_code == 403


def test_preview_does_not_save_and_signed_confirmation(db, monitoring, client):
    from app.main import app, require_user
    d = monitoring
    app.dependency_overrides[require_user] = lambda: d['strategy']
    response = client.get('/internal-control/questionnaires')
    token = response.context['request'].session.get('csrf_token', '')
    response = client.post('/internal-control/questionnaires/preview', data={'csrf_token':token, 'name':'Sentetik yükleme'},
        files={'file':('test.xlsx', question_file(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200 and 'Test sorusu 6' in response.text
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 0
    preview = response.context['preview_token']
    assert client.post('/internal-control/questionnaires/save', data={'csrf_token':token, 'preview_token':preview+'x'}).status_code == 422
    client.post('/internal-control/questionnaires/save', data={'csrf_token':token, 'preview_token':preview})
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 1
    invalid = question_file([[1,'KOS','T1','']])
    response = client.post('/internal-control/questionnaires/preview', data={'csrf_token':token, 'name':'Hatalı sentetik'}, files={'file':('test.xlsx',invalid)})
    assert response.status_code == 422 and 'Satır 2' in response.text
    assert db.scalar(select(func.count()).select_from(q.QuestionList)) == 1


def test_existing_exports_bytes_unchanged_by_questionnaires(db, questionnaire, monkeypatch):
    # Freeze only file creation clocks; compare every byte, including core.xml.
    import zipfile
    import time
    import openpyxl.packaging.core
    import openpyxl.writer.excel
    class FixedDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2026, 10, 30, 12, 0, 0, tzinfo=tz)
    fixed = SimpleNamespace(datetime=FixedDateTime, timezone=timezone)
    monkeypatch.setattr(openpyxl.packaging.core, 'datetime', fixed)
    monkeypatch.setattr(openpyxl.writer.excel, 'datetime', fixed)
    monkeypatch.setattr(zipfile.time, 'localtime', lambda *args: time.struct_time((2026,10,30,12,0,0,4,303,0)))
    d = questionnaire
    def outputs():
        return [render(db, d['strategy'], d['period'].id, kind, draft=True)
            for render in (legacy_export.period_export, legacy_export.official_export)
            for kind in ('xlsx','docx')] + [legacy_export.annual_export(db, d['strategy'], d['plan'].id, '2027')]
    before = outputs()
    answer_and_approve(db, d, d['question_owned'], d['user'], ['yes','no','yes','no'])
    assert before == outputs()


def test_migration_postgresql_roundtrip_refusal_and_single_head():
    base = create_engine(os.environ['CIVENSA_TEST_POSTGRES_URL'])
    schema = 'civ085_migration_' + uuid4().hex
    with base.begin() as conn: conn.execute(CreateSchema(schema))
    url = make_url(os.environ['CIVENSA_TEST_POSTGRES_URL']).update_query_dict({'options':'-csearch_path='+schema}).render_as_string(hide_password=False)
    engine = create_engine(url)
    env = {**os.environ, 'DATABASE_URL':url}
    def migrate(direction, target):
        return subprocess.run([sys.executable,'-m','alembic',direction,target], cwd=ROOT, env=env, capture_output=True, text=True)
    try:
        for direction, revision in [('upgrade','c086a1b2c3d4'), ('upgrade','head'), ('downgrade','c086a1b2c3d4'), ('upgrade','head')]:
            result = migrate(direction,revision)
            assert result.returncode == 0, result.stderr
        heads = subprocess.run([sys.executable,'-m','alembic','heads'], cwd=ROOT, env=env, capture_output=True,text=True)
        assert heads.stdout.strip() == 'c085a1b2c3d4 (head)'
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO organizations(id,name) VALUES (1,'Sentetik')"))
            conn.execute(text("INSERT INTO departments(id,organization_id,name) VALUES (1,1,'Test Müdürlüğü')"))
            conn.execute(text("INSERT INTO users(id,email,full_name,role,password_hash,is_active,department_id,position) VALUES (1,'test@civ085.invalid','Strateji','strategy','unused',true,1,'director')"))
        with Session(engine) as session:
            q.save_list(session,session.get(User,1),'Sentetik göç listesi',question_file())
            session.commit()
            before = session.execute(text('SELECT * FROM ic_question_lists')).all()
            session.rollback()
        result = migrate('downgrade','c086a1b2c3d4')
        assert result.returncode != 0 and 'Soru formu verisi var' in result.stderr
        with engine.connect() as conn:
            assert before == conn.execute(text('SELECT * FROM ic_question_lists')).all()
            assert conn.execute(text('SELECT version_num FROM alembic_version')).scalar() == 'c085a1b2c3d4'
    finally:
        engine.dispose()
        with base.begin() as conn: conn.execute(DropSchema(schema,cascade=True))
        base.dispose()


def test_stale_question_list_form_refused(db, questionnaire):
    d = questionnaire
    replacement = q.save_list(db, d['strategy'], 'Sentetik değişen liste', question_file(count=4))
    q.attach(db, d['strategy'], d['period'].id, replacement.id,
        [b.department_id for b in d['question_batches']], 1)
    db.commit()
    with pytest.raises(HTTPException) as exc:
        q.save_answers(db, d['user'], d['question_owned'].id, 0,
            answer_data('yes'), expected_list_id=d['catalogue'].id)
    assert exc.value.status_code == 409
    assert q.state(db, d['question_owned'].id) == ('draft',0)


def test_concurrent_answers_accept_one_revision(db, questionnaire):
    d = questionnaire
    actor_id, batch_id = d['user'].id, d['question_owned'].id
    db.commit()
    def attempt(answer):
        with Session(db.get_bind(), expire_on_commit=False) as session:
            try:
                q.save_answers(session, session.get(User, actor_id), batch_id, 0, answer_data(answer))
                session.commit()
                return 200
            except HTTPException as exc:
                session.rollback()
                return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as executor:
        assert sorted(executor.map(attempt,['yes','no'])) == [200,409]
    db.expire_all()
    assert q.state(db, batch_id) == ('draft',1)


def test_database_freezes_list_and_participation_after_answer(db, questionnaire):
    d = questionnaire
    q.save_answers(db, d['user'], d['question_owned'].id, 0, answer_data('yes'))
    replacement = q.save_list(db, d['strategy'], 'Sentetik başka liste', question_file())
    db.commit()
    ids = [b.department_id for b in d['question_batches']]
    for list_id, selected, actor_id in [(replacement.id,ids,d['strategy'].id),
        (d['catalogue'].id,ids[:1],d['strategy'].id), (d['catalogue'].id,ids,d['staff'].id)]:
        with pytest.raises(DBAPIError):
            with db.begin_nested():
                db.execute(q.Attachment.__table__.insert().values(period_id=d['period'].id,list_id=list_id,
                    revision=2,departments_json=json.dumps(selected),actor_id=actor_id,actor_name='Test aktör'))
    assert q.attachment(db,d['period'].id).revision == 1
