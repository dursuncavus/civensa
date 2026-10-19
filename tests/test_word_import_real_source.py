"""CIV-010: real 2027 source parity, coverage, and PostgreSQL persistence."""
import hashlib
import io
import json
from collections import Counter
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import func, select

from app.data_importer import parse_performance_docx
from app.import_service import apply_import_batch
from app.import_service import consolidate_source_departments
from app.models import Activity, ImportBatch, Indicator, EconomicValue, ReportCard, Department
from app.program_service import validate_department, activity_source_disagreements
from test_word_import import import_client, import_context, word_document

pytestmark = pytest.mark.postgres
FIXTURES = Path(__file__).parent / 'fixtures'
ALLOWED_BASELINE_DIFFERENCES = ({
    'department': 'Strateji Geliştirme Müdürlüğü',
    'block_index': 0,
    'field': 'rationale',
    'baseline': '',
    'actual': (
        "5018 ve 5393 sayılı kanunların yasal zorunlulukları doğrultusunda; belediyemizin "
        "mali ve idari yönetim kapasitesini artırmak, kaynakları verimli kullanmak ve hesap "
        "verebilirliği sağlamak amaçlanmaktadır .Bu doğrultuda, 2027-2029 döneminde stratejik "
        "hedeflere ulaşma oranını %90'ın üzerine çıkarmak ve iç kontrol sistemine tam uyum "
        "sağlamak hedeflenmiştir. Belirlenen amaçlara ulaşmak için; dijital takip sistemi "
        "(ERP) üzerinden üçer aylık performans ve risk analizleri yapılacak, süreçler "
        "güncellenecek, harcamacı birimlere düzenli eğitimler verilecek ve bütçe-performans "
        "uyumu tam kapasiteyle işletilecektir. Yürütülecek bu çalışmalar, belediyemizin "
        "mevcut öz kaynakları, personeli ve ilgili hizmet alımı bütçe kalemleri bünyesinde "
        "koordine edilecektir."
    ),
    'reason': (
        'Tablo 134 gerekçesi etiket tabanlı okuyucuyla artık okunuyor; 960d62c konumsal '
        'okuyucusu bu alanı boş bırakıyordu.'
    ),
},)


def source_snapshot(parsed):
    """Retain every source field whose loss or drift would change imported programme data."""
    return [{
        'name': department['name'],
        'blocks': [{
            **{field: block[field] for field in (
                'budget_year', 'program', 'administration', 'subprogram', 'rationale', 'goal',
            )},
            'indicators': [{
                'name': indicator['name'],
                'values': [{
                    field: value.get(field) for field in ('year', 'kind', 'raw_text')
                } for value in indicator['values']],
            } for indicator in block['indicators']],
        } for block in department['blocks']],
        'activities': [{
            'name': activity['name'],
            'costs': [{
                field: cost.get(field)
                for field in ('year', 'kind', 'funding_type', 'raw_text')
            } for cost in activity['costs']],
        } for activity in department['activities']],
        'personnel': [{
            field: value.get(field) for field in ('year', 'kind', 'raw_text')
        } for value in department['personnel']],
    } for department in parsed['departments']]


def apply_documented_baseline_allowlist(actual, expected):
    """Verify and mask only reviewed field improvements before exact snapshot comparison."""
    adjusted = deepcopy(actual)
    for allowed in ALLOWED_BASELINE_DIFFERENCES:
        actual_departments = [
            department for department in adjusted
            if department['name'] == allowed['department']
        ]
        expected_departments = [
            department for department in expected
            if department['name'] == allowed['department']
        ]
        assert len(actual_departments) == len(expected_departments) == 1, allowed['reason']
        actual_block = actual_departments[0]['blocks'][allowed['block_index']]
        expected_block = expected_departments[0]['blocks'][allowed['block_index']]
        assert expected_block[allowed['field']] == allowed['baseline'], allowed['reason']
        assert actual_block[allowed['field']] == allowed['actual'], allowed['reason']
        actual_block[allowed['field']] = allowed['baseline']
    return adjusted


@pytest.fixture(scope='module')
def real_source():
    return (FIXTURES / 'pp2027_source_no_media.docx').read_bytes()


@pytest.fixture(scope='module')
def real_parsed(real_source):
    return parse_performance_docx(real_source)


def test_real_source_matches_960d62c_baseline(real_source, real_parsed):
    baseline = json.loads((FIXTURES / 'pp2027_baseline.json').read_text())
    assert hashlib.sha256(real_source).hexdigest() == baseline['fixture_sha256']
    summary = dict(real_parsed['summary'])
    # CIV-005 deliberately includes the one nested table, unlike 960d62c.
    assert summary.pop('table_count') == baseline['summary']['table_count'] + 1 == 162
    expected = {k: v for k, v in baseline['summary'].items() if k != 'table_count'}
    # CIV-042: two valid ₺-prefixed totals no longer mark ten economic
    # category/funding cells for review. The frozen source fixture stays intact.
    assert expected['needs_review_count'] == 440
    expected['needs_review_count'] = 430
    assert summary == expected
    assert real_parsed['errors'] == real_parsed['warnings'] == []
    economic_snapshot = [{'department': d['name'], 'tables': d.get('economic_tables', [])}
                         for d in real_parsed['departments']]
    digest = hashlib.sha256(json.dumps(economic_snapshot, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    # Approved semantic change only: valid currency-prefix amounts and their
    # propagated review flags. Raw book cells and all other fields remain exact.
    assert digest == '46c76bc9427a3fd491215a429860d9f5a2e544863c1464564ef4a42f7b5f16a9'
    currency_totals = [v for d in real_parsed['departments'] for t in d.get('economic_tables', [])
                       for values in t['totals'].values() for v in values if v['raw_text'].startswith('₺')]
    assert len(currency_totals) == 2
    assert all((v['amount'], v['state'], v['raw_text']) == ('10424938.95', 'verified', '₺10.424.938,95')
               for v in currency_totals)
    headers = {'Performans Göstergesi', 'Performans Göstergeleri', 'Ölçü Birimi',
               '2025 Yılı Gerçekleşme', '2027 Hedef', '2028 Tahmin', '2029 Tahmini'}
    assert not any(i['name'] in headers for d in real_parsed['departments'] for b in d['blocks'] for i in b['indicators'])
    actual = source_snapshot(real_parsed)
    assert apply_documented_baseline_allowlist(actual, baseline['departments']) == baseline['departments']
    human_resources = next(
        department for department in actual
        if department['name'] == 'İnsan Kaynakları ve Eğitim Müdürlüğü'
    )
    assert human_resources['blocks'][1]['subprogram'] == 'Üst Yönetim, İdari ve Mali Hizmetler'
    coverage = real_parsed['coverage_items']
    assert len(coverage) == 32
    assert [item['table_number'] for item in coverage[:28]] == list(range(1, 29))
    assert all(f"Tablo {item['table_number']}: {item['first_cell_text']}" in item['message']
               for item in coverage)


def test_real_source_preview_is_applicable(import_client, import_context, real_source, real_parsed):
    response = import_client.post(
        f'/api/periods/{import_context.period.id}/data-import-preview',
        files={'file': ('pp2027.docx', real_source)})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['summary']['unresolved_department_count'] == 0
    assert result['can_apply'] is True
    assert result['conflicts'] == []
    assert result['errors'] == []
    assert result['coverage_items'] == real_parsed['coverage_items']
    assert not any('Tablo ' in warning for warning in result['warnings'])
    detail = import_client.get(f"/api/import-batches/{result['batch_id']}").json()
    assert detail['can_apply'] is True
    assert detail['conflicts'] == []
    assert detail['coverage_items'] == result['coverage_items']
    assert detail['errors'] == []
    duplicates = [r for r in result['review_items'] if r.get('cause', '').startswith('duplicate_')]
    assert Counter(r['cause'] for r in duplicates) == {'duplicate_programme_block': 11, 'duplicate_indicator': 8}
    assert detail['review_items'] == result['review_items']
    assert result['summary']['needs_review_count'] == len(result['review_items']) == 498


def test_real_source_duplicate_evidence_and_economic_causes(real_parsed):
    plan = consolidate_source_departments(real_parsed)
    social = [r for r in plan['review_items'] if r.get('cause', '').startswith('duplicate_')
              and r['department']=='Sosyal Hizmetler Müdürlüğü']
    assert len(social) == 6
    assert social[0]['cause']=='duplicate_programme_block' and social[0]['table_numbers']==[100, 101]
    for index, item in enumerate(social[1:], 1):
        assert f'SH36/PG0{index}.' in item['record']
        assert item['table_numbers'] == ([100, 101, 102] if index < 5 else [100, 101])
    values = [v for d in real_parsed['departments'] for t in d.get('economic_tables', []) for v in t['values']]
    reasons = [c for v in values for c in v.get('review_causes', [])]
    # Counts overlap: one value/column can have several independent causes.
    assert Counter(c['cause'] for c in reasons) == {
        'inside_total_mismatch': 180, 'source_cell_unverified': 130,
        'program_group_unallocated': 140, 'personnel_total_mismatch': 8}
    assert Counter(cause for cause, message in {(c['cause'], c['message']) for c in reasons}) == {
        'inside_total_mismatch': 22, 'source_cell_unverified': 19,
        'program_group_unallocated': 15, 'personnel_total_mismatch': 2}
    assert Counter(cause for v in values for cause in {c['cause'] for c in v.get('review_causes', [])}) == {
        'inside_total_mismatch': 180, 'source_cell_unverified': 90,
        'program_group_unallocated': 140, 'personnel_total_mismatch': 8}
    assert all('Tablo ' in c['message'] and '202' in c['message'] for c in reasons)
    assert not any('Ekonomik satırlar, kaynak toplamları' in v.get('issue', '') for v in values)


def test_real_source_apply_persists_all_records(db, import_context, real_parsed):
    assert db.bind.dialect.name == 'postgresql'
    assert db.scalar(select(func.count()).select_from(Indicator)) == 0
    assert db.scalar(select(func.count()).select_from(Activity)) == 0
    batch = ImportBatch(period_id=import_context.period.id, filename='pp2027.docx',
                        status='preview', payload_json=json.dumps(real_parsed),
                        created_by=import_context.user.id)
    db.add(batch)
    db.commit()
    result = apply_import_batch(db, batch)
    db.commit()
    assert result['indicators'] == real_parsed['summary']['indicator_count'] == 263
    assert result['activities'] == real_parsed['summary']['activity_count'] == 288
    assert db.scalar(select(func.count()).select_from(Indicator)) == 263
    assert db.scalar(select(func.count()).select_from(Activity)) == 288
    assert batch.status == 'applied'
    assert result['cards'] == db.scalar(select(func.count()).select_from(ReportCard)) == 64
    assert result['economic_values'] == db.scalar(select(func.count()).select_from(EconomicValue)) == 1290
    assert all(row.source_card_id and json.loads(row.source_context_json)['card_ids']
               for row in db.scalars(select(EconomicValue)))
    sample = db.scalar(select(Department).where(Department.name == 'Özel Kalem Müdürlüğü'))
    errors, warnings = validate_department(db, import_context.period.id, sample.id)
    relevant = [e for e in errors if any(w in e.lower() for w in ['kart', 'doğrulan', 'ekonomik'])]
    assert not relevant, '\n'.join(relevant)
    assert len(errors) == 5 and all('kataloğu bağlantısı eksik' in e for e in errors), errors
    # CIV-011's source totals and CIV-026's economic import coexist on D-23.
    disagreements = Counter()
    for activity in db.scalars(select(Activity)):
        source_disagreements = activity_source_disagreements(activity)
        for key in source_disagreements:
            disagreements[key] += 1
        for value in activity.cost_values:
            if value.source_amount is not None:
                assert value.amount == value.source_amount
                if f'{value.year}:{value.value_kind}' in source_disagreements:
                    if value.funding_type == 'total':
                        assert value.state == 'needs_review'
                        assert value.source_raw_text == value.raw_text
    assert disagreements == {'2027:budget': 13, '2028:forecast': 14, '2029:forecast': 18}



@pytest.mark.parametrize('year,valid', [('2027/1', True), ('2027 / 2027', True),
                                       ('', False), ('2027/2028', False)])
def test_budget_year_extracts_one_distinct_year(year, valid):
    parsed = parse_performance_docx(word_document(year=year))
    assert bool(parsed['errors']) is not valid
    if valid:
        assert parsed['summary']['source_years'] == [2027]
    else:
        assert any('Tablo 1' in error and 'Bütçe Yılı' in error for error in parsed['errors'])


def test_punctuated_labels_and_section_title_before_plural_header():
    doc = Document(io.BytesIO(word_document(variant='plural')))
    table = doc.tables[0]
    table.cell(0, 0).text = 'BÜTÇE / YILI :'
    table.cell(1, 0).text = 'PROGRAM-ADI.'
    table.cell(2, 0).text = 'İdare/Birim   Adı:'
    table.cell(3, 0).text = 'Altprogram Adı:'
    title = table.add_row()
    title.cells[0].merge(title.cells[-1]).text = 'PERFORMANS GÖSTERGELERİ'
    table.rows[6]._tr.addprevious(title._tr)
    stream = io.BytesIO()
    doc.save(stream)
    parsed = parse_performance_docx(stream.getvalue())
    assert parsed['errors'] == []
    assert parsed['summary']['indicator_count'] == 1
    assert parsed['departments'][0]['blocks'][0]['program'] == 'Çevre Programı'
    assert parsed['departments'][0]['blocks'][0]['subprogram'] == 'Yeşil Alanlar'


@pytest.mark.parametrize(('row_index', 'field_name'), [
    (1, 'Program Adı'),
    (2, 'İdare/Birim Adı'),
    (3, 'Alt Program Adı'),
])
def test_missing_required_indicator_field_is_blocking(row_index, field_name):
    doc = Document(io.BytesIO(word_document()))
    for cell in doc.tables[0].rows[row_index].cells[2:]:
        cell.text = ''
    stream = io.BytesIO()
    doc.save(stream)

    parsed = parse_performance_docx(stream.getvalue())

    assert parsed['departments'] == []
    assert any(
        f'Tablo 1: {field_name} alanı eksik veya boş.' in error
        for error in parsed['errors']
    )


@pytest.mark.parametrize('label', ['Bütçe Yılı:', 'Program Adı', 'Performans Göstergesi', 'Faaliyetler'])
def test_malformed_data_table_remains_blocking(label):
    doc = Document(io.BytesIO(word_document()))
    doc.add_table(rows=1, cols=1).cell(0, 0).text = label
    stream = io.BytesIO()
    doc.save(stream)
    parsed = parse_performance_docx(stream.getvalue())
    assert len(parsed['errors']) == 1
    assert 'Tablo 3' in parsed['errors'][0]
    assert parsed['coverage_items'] == []
