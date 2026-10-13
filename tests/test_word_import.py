"""Synthetic Ministry Word imports: current guarantees and VERİ-01 regressions.

Every database scenario uses the disposable, per-test PostgreSQL schema from
conftest.py. No municipal document or application database is used.
"""
import io
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session

import app.import_service as service
from app.data_importer import parse_performance_docx
from app.db import get_db
from app.main import app, require_strategy
from app.models import (
    Activity, ActivityCostValue, Department, DepartmentCostValue, ImportBatch,
    Indicator, IndicatorValue, Organization, Program, ReportCard, ReportingPeriod,
    Submission, SubProgram, User,
)
from app.template_parser import analyze_template

pytestmark = pytest.mark.postgres


def word_document(*, year='2027', department='Park Bahçeler Müdürlüğü',
                  values=('1.234,50', '12', '0'), variant=None):
    """Build the expected label/value rows, indicator header and cost triples."""
    doc = Document()
    table = doc.add_table(rows=0, cols=5)
    labels = [
        ('Bütçe Yılı', year), ('Program Adı', 'Çevre Programı'),
        ('İdare Birim Adı', 'Örnek Belediyesi / ' + department),
        ('Alt Program Adı', 'Yeşil Alanlar'),
        ('Gerekçe ve Açıklamalar', 'Yeşil alan ihtiyacı'),
        ('Alt Programın Hedefi', 'Parkları artırmak'),
    ]
    for label, value in labels:
        cells = table.add_row().cells
        cells[0].merge(cells[1]).text = label
        cells[2].merge(cells[4]).text = value
    headers = ('Performans Göstergesi', 'Ölçü Birimi', '2026 Gerçekleşme',
               '2027 Hedef', '2028 Tahmin')
    for cell, text in zip(table.add_row().cells, headers):
        cell.text = text
    if variant == 'two_row_header':
        for col in range(2, 5):
            table.cell(6, col).text = headers[col].split()[0]
        for cell, text in zip(table.add_row().cells, ('', '', 'Gerçekleşme', 'Hedef', 'Tahmin')):
            cell.text = text
    data_row = table.add_row()
    for cell, text in zip(data_row.cells, ('Park sayısı', 'Adet', *values)):
        cell.text = text
    for label, text in [
        ('Göstergeye İlişkin Açıklama', 'Kayıtlı parklar'),
        ('Hesaplama Yöntemi', 'Parkları saymak'),
        ('Verinin Kaynağı', 'Park kayıtları'), ('Sorumlu İdare', department),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].merge(cells[4]).text = text
    if variant in {'tracked_change', 'content_control'}:
        # Move the visible numeric run into genuine OOXML wrappers. Cell.text
        # currently ignores both w:ins and w:sdtContent descendants.
        paragraph = data_row.cells[3].paragraphs[0]._p
        run = paragraph[0]
        paragraph.remove(run)
        if variant == 'tracked_change':
            wrapper = OxmlElement('w:ins')
            from docx.oxml.ns import qn
            wrapper.set(qn('w:id'), '1')
            wrapper.set(qn('w:author'), 'Sentetik Yazar')
            wrapper.set(qn('w:date'), '2026-10-10T00:00:00Z')
            wrapper.append(run)
        else:
            wrapper = OxmlElement('w:sdt')
            wrapper.append(OxmlElement('w:sdtPr'))
            content = OxmlElement('w:sdtContent')
            content.append(run)
            wrapper.append(content)
        paragraph.append(wrapper)
    if variant == 'plural':
        table.cell(6, 0).text = 'Performans Göstergeleri'
    if variant == 'short':
        # Eight rows still contain all identification/header/data rows.
        for row in list(table.rows)[8:]:
            table._tbl.remove(row._tr)
    if variant == 'missing_year_row':
        table._tbl.remove(table.rows[0]._tr)
    if variant == 'nested':
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0)._tc.append(table._tbl)
    costs = doc.add_table(rows=0, cols=4)
    for texts in [
        ('Faaliyetler', '2027 Bütçe', '2028 Tahmin', '2029 Tahmin'),
        ('Park bakımı', '1.234,50', '0', ''),
        ('Bütçe İçi', '1.200,00', '0', ''),
        ('Bütçe Dışı', '34,50', '0', ''),
        ('Personel ve SGK Giderleri', '100,00', '0', ''),
        ('Faaliyetlere İlişkin Açıklamalar', '', '', ''),
        ('Faaliyet Adı: Park bakımı', '', '', ''),
        ('Parkların bakımı yapılır.', '', '', ''),
    ]:
        for cell, text in zip(costs.add_row().cells, texts):
            cell.text = text
    if variant == 'unrecognised':
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Tanınmayan tablo'
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def import_context(db):
    org = Organization(name='Örnek Belediyesi')
    period = ReportingPeriod(year=2027, template_status='verified')
    user = User(email='import@example.invalid', full_name='Strateji Kullanıcısı',
                role='strategy', password_hash='unused')
    db.add_all([org, period, user])
    db.commit()
    return SimpleNamespace(org=org, period=period, user=user)


@pytest.fixture
def import_client(db, import_context):
    def database():
        yield db
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[require_strategy] = lambda: import_context.user
    # Surface D7 as the actual HTTP response rather than raising in TestClient.
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            yield client
    finally:
        app.dependency_overrides.clear()


def preview(client, context, **options):
    response = client.post(
        f'/api/periods/{context.period.id}/data-import-preview',
        files={'file': ('sentetik.docx', word_document(**options),
                        'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
    )
    assert response.status_code == 200, response.text
    return response.json()


def apply(client, batch_id):
    response = client.post(f'/api/import-batches/{batch_id}/apply')
    assert response.status_code == 200, response.text
    return response.json()


def parsed_values(**options):
    parsed = parse_performance_docx(word_document(**options))
    return parsed['departments'][0]['blocks'][0]['indicators'][0]['values']


def test_preview_keeps_payload_and_does_not_write_department_data(db, import_client, import_context):
    result = preview(import_client, import_context)
    assert result['status'] == 'preview'
    assert result['can_apply'] is True
    assert result['summary']['source_years'] == [2027]
    assert result['summary']['indicator_count'] == 1
    assert result['summary']['activity_count'] == 1
    assert result['departments'][0]['name'] == 'Park ve Bahçeler Müdürlüğü'
    batch = db.get(ImportBatch, result['batch_id'])
    assert json.loads(batch.payload_json) == parse_performance_docx(word_document())
    assert db.scalar(select(func.count()).select_from(Department)) == 0
    assert db.scalar(select(func.count()).select_from(Indicator)) == 0
    detail = import_client.get(f'/api/import-batches/{batch.id}').json()
    assert detail['can_apply'] is True
    assert detail['can_revert'] is False
    template = analyze_template('sentetik.docx', word_document())
    assert 'Performans Göstergeleri' in template['present']
    assert 'Bütçe Yılı' in template['present']


def test_department_alias_matches_existing_department(db, import_client, import_context):
    dept = Department(organization_id=import_context.org.id, name='Park ve Bahçeler Müdürlüğü')
    db.add(dept)
    db.commit()
    result = preview(import_client, import_context)
    assert result['mappings'][0]['action'] == 'match'
    assert result['mappings'][0]['department_id'] == dept.id
    assert result['mappings'][0]['target_name'] == 'Park ve Bahçeler Müdürlüğü'
    applied = apply(import_client, result['batch_id'])
    assert applied['created_departments'] == 0
    assert db.scalar(select(func.count()).select_from(Department)) == 1


def test_apply_persists_indicators_costs_and_raw_text(db, import_client, import_context):
    result = apply(import_client, preview(import_client, import_context)['batch_id'])
    assert (result['departments'], result['created_departments'], result['indicators'],
            result['activities'], result['personnel_values']) == (1, 1, 1, 1, 2)
    ind = db.scalar(select(Indicator))
    assert (ind.name, ind.unit, ind.explanation, ind.calculation_method, ind.data_source) == (
        'Park sayısı', 'Adet', 'Kayıtlı parklar', 'Parkları saymak', 'Park kayıtları')
    values = {v.year: v for v in db.scalars(select(IndicatorValue)).all()}
    assert values[2026].value == Decimal('1234.50')
    assert values[2026].raw_text == '1.234,50'
    assert values[2027].value_kind == 'target'
    assert values[2028].value == 0
    activity = db.scalar(select(Activity))
    assert activity.name == 'Park bakımı'
    assert activity.explanation == 'Parkların bakımı yapılır.'
    assert activity.subprogram_id == ind.subprogram_id
    cost = db.scalar(select(ActivityCostValue).where(
        ActivityCostValue.year == 2027, ActivityCostValue.funding_type == 'total'))
    assert (cost.amount, cost.raw_text, cost.state) == (Decimal('1234.50'), '1.234,50', 'verified')
    personnel = db.scalar(select(DepartmentCostValue).where(DepartmentCostValue.year == 2027))
    assert (personnel.amount, personnel.raw_text) == (Decimal('100'), '100,00')
    assert db.scalar(select(ReportCard)).goal == 'Parkları artırmak'
    assert db.scalar(select(Submission)).status == 'draft'


def test_ambiguous_number_needs_review_and_blank_is_not_zero(db, import_client, import_context):
    options = {'values': ('1,234', '', '0')}
    values = parsed_values(**options)
    assert [(v['value'], v['raw_text'], v['state']) for v in values] == [
        (None, '1,234', 'needs_review'), (None, '', 'missing'), ('0', '0', 'verified')]
    apply(import_client, preview(import_client, import_context, **options)['batch_id'])
    persisted = {v.year: v for v in db.scalars(select(IndicatorValue)).all()}
    assert 2027 not in persisted
    assert (persisted[2026].value, persisted[2026].raw_text, persisted[2026].state) == (
        None, '1,234', 'needs_review')
    assert (persisted[2028].value, persisted[2028].raw_text, persisted[2028].state) == (
        Decimal('0'), '0', 'verified')


def test_wrong_year_blocks_preview_and_apply(db, import_client, import_context):
    result = preview(import_client, import_context, year='2026')
    assert result['can_apply'] is False
    assert result['warnings'][0] == 'Dosya bütçe yılı [2026]; seçilen Civensa dönemi 2027. Uygulama engellendi.'
    response = import_client.post(f"/api/import-batches/{result['batch_id']}/apply")
    assert response.status_code == 409
    assert 'Yanlış döneme veri yazılması engellendi.' in response.json()['detail']
    assert db.get(ImportBatch, result['batch_id']).status == 'preview'
    assert db.scalar(select(func.count()).select_from(Department)) == 0


@pytest.mark.parametrize('status', ['submitted', 'approved'])
def test_reviewed_packages_cannot_be_overwritten(db, import_client, import_context, status):
    dept = Department(organization_id=import_context.org.id, name='Park ve Bahçeler Müdürlüğü')
    db.add(dept)
    db.flush()
    submission = Submission(department_id=dept.id, period_id=import_context.period.id,
                            status=status, return_note='İnceleme kaydı')
    db.add(submission)
    db.commit()
    batch_id = preview(import_client, import_context)['batch_id']
    response = import_client.post(f'/api/import-batches/{batch_id}/apply')
    assert response.status_code == 409
    assert response.json()['detail'] == 'İncelemede/onaylı müdürlükler üzerine veri yazılamaz: Park ve Bahçeler Müdürlüğü'
    db.refresh(submission)
    assert (submission.status, submission.return_note) == (status, 'İnceleme kaydı')
    assert db.scalar(select(func.count()).select_from(Indicator)) == 0


def test_rollback_refuses_changed_data(db, import_client, import_context):
    batch_id = preview(import_client, import_context)['batch_id']
    apply(import_client, batch_id)
    assert import_client.get(f'/api/import-batches/{batch_id}').json()['can_revert'] is True
    ind = db.scalar(select(Indicator))
    ind.name = 'Sonradan değişen gösterge'
    db.commit()
    detail = import_client.get(f'/api/import-batches/{batch_id}').json()
    assert detail['can_revert'] is False
    assert detail['revert_reason'] == 'İçe aktarma sonrasında veriler değiştiği için otomatik geri alma güvenli değil.'
    response = import_client.post(f'/api/import-batches/{batch_id}/revert')
    assert response.status_code == 409
    db.refresh(ind)
    assert ind.name == 'Sonradan değişen gösterge'
    assert db.get(ImportBatch, batch_id).status == 'applied'


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D1: İzlenen değişiklik ve içerik denetimindeki değerler kaybolmamalı.')
@pytest.mark.parametrize('variant', ['tracked_change', 'content_control'])
def test_d1_wrapped_visible_value_is_preserved(variant):
    value = parsed_values(variant=variant)[1]
    assert (value['value'], value['raw_text'], value['state']) == ('12', '12', 'verified')


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D2: Çoğul başlık, kısa ve iç içe tablolar sessizce atlanmamalı.')
@pytest.mark.parametrize('variant', ['plural', 'short', 'nested'])
def test_d2_supported_table_variants_are_recognised(variant):
    parsed = parse_performance_docx(word_document(variant=variant))
    assert parsed['summary']['indicator_count'] == 1
    assert parsed['departments'][0]['blocks'][0]['indicators'][0]['name'] == 'Park sayısı'


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D3: İki satırlı başlık hedef ve gerçekleşme ayrımını korumalı.')
def test_d3_two_row_header_retains_value_kinds():
    values = parsed_values(variant='two_row_header')
    assert [(v['year'], v['kind'], v['value']) for v in values] == [
        (2026, 'actual', '1234.50'), (2027, 'target', '12'), (2028, 'forecast', '0')]


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D4: Yılı eksik veya tanınan tablosu olmayan önizleme uygulanabilir olmamalı.')
@pytest.mark.parametrize('options', [
    {'variant': 'missing_year_row'}, {'year': ''}, {'variant': 'unrecognised'},
], ids=['missing-year-row', 'empty-year', 'no-recognised-table'])
def test_d4_incomplete_preview_is_not_applicable(import_client, import_context, options):
    result = preview(import_client, import_context, **options)
    detail = import_client.get(f"/api/import-batches/{result['batch_id']}").json()
    assert (result['can_apply'], detail['can_apply']) == (False, False)


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D5: Aynı önizlemenin eşzamanlı uygulanması mükerrer müdürlük oluşturmamalı.')
def test_d5_concurrent_apply_creates_department_once(db, import_client, import_context, monkeypatch):
    batch_id = preview(import_client, import_context)['batch_id']
    # Avoid unrelated Program/SubProgram creation races. Hold both genuine plans
    # after reading departments, before either transaction inserts a department.
    program = Program(period_id=import_context.period.id, name='Çevre Programı')
    db.add(program)
    db.flush()
    db.add(SubProgram(period_id=import_context.period.id, program_id=program.id, name='Yeşil Alanlar'))
    db.commit()
    bind = db.get_bind()
    rendezvous = threading.Barrier(2, timeout=10)
    original = service.build_department_plan

    def synchronized_plan(session, parsed):
        plan = original(session, parsed)
        rendezvous.wait()
        return plan

    monkeypatch.setattr(service, 'build_department_plan', synchronized_plan)

    def worker():
        with Session(bind, expire_on_commit=False) as session:
            # Bound SQL waits and propagate infrastructure failures (not xfail).
            from sqlalchemy import text
            session.execute(text("SET LOCAL lock_timeout = '10s'"))
            session.execute(text("SET LOCAL statement_timeout = '15s'"))
            batch = session.get(ImportBatch, batch_id)
            try:
                result = service.apply_import_batch(session, batch)
                session.commit()
                return result
            except ValueError as error:
                session.rollback()
                if str(error) != 'Bu içe aktarma paketi artık önizleme durumunda değil.':
                    raise
                return None  # An explicit stale-preview rejection is valid.

    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(worker) for _ in range(2)]
        results = [future.result(timeout=25) for future in futures]
    db.expire_all()
    count = db.scalar(select(func.count()).select_from(Department).where(
        Department.name == 'Park ve Bahçeler Müdürlüğü'))
    assert count == 1, f'Aynı müdürlük {count} kez oluşturuldu: {results}'


@pytest.mark.xfail(strict=True, raises=AssertionError,
                   reason='VERİ-01 D6: İade edilen müdürlüğün inceleme notu içe aktarmada korunmalı.')
def test_d6_return_note_survives_import(db, import_client, import_context):
    dept = Department(organization_id=import_context.org.id, name='Park ve Bahçeler Müdürlüğü')
    db.add(dept)
    db.flush()
    sub = Submission(department_id=dept.id, period_id=import_context.period.id,
                     status='returned', return_note='Hedef değerini düzeltin.')
    db.add(sub)
    db.commit()
    apply(import_client, preview(import_client, import_context)['batch_id'])
    db.refresh(sub)
    assert sub.return_note == 'Hedef değerini düzeltin.'


# D7 did not reproduce for either new departments or legacy data restoration.
# Keep the successful case as characterization, not a fabricated xfail.
def test_safe_rollback_restores_legacy_indicator(db, import_client, import_context):
    # Existing legacy indicator data has no ReportCard (accepted by apply).
    dept = Department(organization_id=import_context.org.id, name='Park ve Bahçeler Müdürlüğü')
    program = Program(period_id=import_context.period.id, name='Eski program')
    db.add_all([dept, program])
    db.flush()
    sp = SubProgram(period_id=import_context.period.id, program_id=program.id, name='Eski alt program')
    db.add(sp)
    db.flush()
    ind = Indicator(period_id=import_context.period.id, department_id=dept.id,
                    subprogram_id=sp.id, name='Eski gösterge', unit='Adet')
    db.add(ind)
    db.flush()
    db.add(IndicatorValue(indicator_id=ind.id, year=2027, value_kind='target',
                          value=Decimal('5'), raw_text='5', state='verified'))
    db.commit()
    batch_id = preview(import_client, import_context)['batch_id']
    apply(import_client, batch_id)
    detail = import_client.get(f'/api/import-batches/{batch_id}').json()
    assert detail['can_revert'] is True
    response = import_client.post(f'/api/import-batches/{batch_id}/revert')
    assert response.status_code == 200, response.text
    assert response.json()['status'] == 'reverted'
    db.expire_all()
    assert db.scalar(select(Indicator)).name == 'Eski gösterge'


def test_safe_rollback_removes_new_department(db, import_client, import_context):
    batch_id = preview(import_client, import_context)['batch_id']
    apply(import_client, batch_id)
    assert import_client.get(f'/api/import-batches/{batch_id}').json()['can_revert'] is True
    response = import_client.post(f'/api/import-batches/{batch_id}/revert')
    assert response.status_code == 200, response.text
    assert response.json()['status'] == 'reverted'
    db.expire_all()
    for model in (Department, Indicator, IndicatorValue, Activity, ActivityCostValue,
                  DepartmentCostValue, ReportCard, Submission, SubProgram, Program):
        assert db.scalar(select(func.count()).select_from(model)) == 0
    assert db.get(ImportBatch, batch_id).status == 'reverted'
