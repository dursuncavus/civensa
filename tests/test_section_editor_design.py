"""012: frozen PostgreSQL controls, audit safety and disclosure keyboard behavior."""
from datetime import datetime, UTC
from collections import Counter
from html.parser import HTMLParser
import json
import re
from pathlib import Path

import pytest
from sqlalchemy import select
from app import activity_sections_service as sections
from app.activity_contract_service import digest
from app.models import ARContentEvent, ARUnitWorkflowEvent
from test_activity_sections import section_client, login
from test_sidebar_ui import browser

pytestmark = pytest.mark.postgres
CONTRACT = Path(__file__).parent / 'fixtures' / 'section_editor_012_contract.json'
CODES = ['I.A', 'I.C.1', 'III.A.1', 'III.B.1', 'III.B.2.i', 'EKLER']


class Controls(HTMLParser):
    """Hidden fields/options/disclosures included; no field values or CSRF secrets."""
    def __init__(self, html):
        super().__init__()
        self.controls = []
        self.bindings = []
        self.form_target = None
        self.in_main = False
        self.export = False
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'main': self.in_main = True
        if tag == 'form': self.form_target = [attrs.get('action',''),attrs.get('method','get')]
        if tag == 'details' and 'data-section-export' in attrs: self.export = True
        if not self.in_main or tag not in {'a','button','form','input','select','textarea','summary','option'}:
            return
        if tag == 'summary' and self.export: return  # Only newly added presentation control.
        control = [tag, sorted([k,v] for k,v in attrs.items()
            if k not in {'value','class','style','tabindex'} )]
        self.controls.append(control)
        if tag in {'input','select','textarea','option','button'}:
            self.bindings.append(control + [self.form_target])

    def handle_endtag(self, tag):
        if tag == 'main': self.in_main = False
        if tag == 'form': self.form_target = None
        if tag == 'details': self.export = False


def configure(db, sample, doc, state, actor):
    sample['strategy'].position = 'staff'
    db.flush()
    user = sample['user']
    user.position = 'director' if actor == 'director' else 'staff'
    user.can_submit = actor == 'delegated'
    if state != 'draft-empty':
        sections.save_section(db, user, doc.id, 'I.A',
            [{'type':'paragraph','text':'Sentetik hizmet değerlendirmesi.'}], 'draft', '', doc.revision)
    if state in {'submitted','returned','approved'}:
        payload = sections.numeric_canonical(sections.report_data(db, user, doc.id, use_freeze=False, include_typeset=True))
        for index, (action, before, after) in enumerate(
            [('submit','draft','submitted')] +
            ([('return','submitted','returned')] if state == 'returned' else
             [('approve','submitted','approved')] if state == 'approved' else []), 1):
            db.add(ARUnitWorkflowEvent(document_id=doc.id, revision=doc.revision+index,
                action=action, from_status=before, to_status=after,
                reason='Sentetik iade notu: hizmet kaynağını açıklayın.' if action == 'return' else '',
                actor_id=sample['strategy'].id, actor_name='Sentetik Strateji Müdürü',
                payload_json=payload if action != 'return' else None,
                sha256=digest(payload) if action != 'return' else None,
                content_event_cursor=0, created_at=datetime(2026,10,25,9,tzinfo=UTC)))
    db.flush()
    if actor == 'strategy': sample['strategy'].position = 'director'
    db.commit()


@pytest.mark.parametrize('actor', ['staff','delegated','director','strategy'])
@pytest.mark.parametrize('state', ['draft-empty','draft-content','returned','submitted','approved'])
def test_012_preserves_all_fields_and_action_targets(section_client, db, sample, actor, state):
    client, docs, _, _ = section_client
    doc = docs['birim:'+str(sample['department'].id)]
    configure(db, sample, doc, state, actor)
    login(client, sample['strategy'].email if actor == 'strategy' else sample['user'].email)
    frozen = json.loads(CONTRACT.read_text())
    template = (CONTRACT.parents[2] / 'app/templates/activity_sections_editor.html').read_text()
    assert not Counter(frozen['template_blocks']) - Counter(re.findall(r'{%.*?%}',template))
    assert not Counter(frozen['template_targets']) - Counter(re.findall('(?<=\\s)(?:(?:id|name|href|action|method|role)="[^"]*"|data-[\\w-]+(?:="[^"]*")?)',template))
    for code in CODES:
        response = client.get(f'/activity-sections/reports/{doc.id}?code={code}')
        assert response.status_code == 200
        key = f'{actor}/{state}/{code}'
        parsed = Controls(response.text)
        actual = parsed.controls
        expected = frozen['contracts'][frozen['cases'][key]]
        assert Counter(json.dumps(c, sort_keys=True) for c in actual) == Counter(json.dumps(c, sort_keys=True) for c in expected), key
        # In the frozen output all field/button controls belong to the preceding
        # form record. Also verify actual nesting, so moving an unchanged field
        # into another action (or outside its form) cannot pass a multiset check.
        target, bindings = None, []
        for tag, attrs in expected:
            if tag == 'form':
                form = dict(attrs)
                target = [form.get('action',''),form.get('method','get')]
            elif tag in {'input','select','textarea','option','button'}:
                bindings.append([tag,attrs,target])
        assert parsed.bindings == bindings, key


@pytest.mark.parametrize('width', [1440,390])
def test_012_export_audit_focus_and_safe_summary(browser, section_client, db, sample, width):
    from playwright.sync_api import expect
    client, docs, _, _ = section_client
    doc = docs['birim:'+str(sample['department'].id)]
    configure(db, sample, doc, 'draft-content', 'director')
    login(client, sample['user'].email)
    # Untrusted content is rendered as text; malformed old payloads degrade safely.
    db.add(ARContentEvent(document_id=doc.id, section_code='I.A', action='section_save',
        before_json='invalid', after_json='<script>window.bad=true</script>', actor_id=sample['user'].id))
    db.add(ARContentEvent(document_id=doc.id, section_code='I.A', action='section_save',
        before_json='{}', after_json='{"reason":""}', actor_id=sample['user'].id))
    db.commit()
    with browser.new_context(viewport={'width':width,'height':900}) as context:
        errors = []
        def serve(route):
            assert route.request.method == 'GET'
            response = client.get(route.request.url)
            headers = dict(response.headers); headers.pop('set-cookie', None)
            route.fulfill(status=response.status_code, headers=headers, body=response.content)
        context.route('http://testserver/**', serve)
        page = context.new_page()
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(f'http://testserver/activity-sections/reports/{doc.id}?code=I.A')
        status_columns = {}
        for link in page.locator('.ar-tree a').all():
            box, status = link.bounding_box(), link.locator('small').bounding_box()
            assert box['width'] <= 44 * 16
            assert status['x'] + status['width'] <= box['x'] + box['width'] + 1
            status_columns.setdefault(round(box['x']), []).append(round(status['x']))
        assert all(max(edges) - min(edges) <= 1 for edges in status_columns.values())
        trigger = page.locator('[data-section-export] > summary')
        trigger.focus(); page.keyboard.press('Enter')
        expect(page.locator('[data-section-export]')).to_have_attribute('open','')
        page.keyboard.press('Tab')
        expect(page.get_by_role('link',name='Word olarak indir',exact=True)).to_be_focused()
        page.keyboard.press('Escape'); expect(trigger).to_be_focused()
        expect(page.locator('[data-section-audit]').first.locator('[data-section-change]')).to_have_text(
            'Kayıt ayrıntılarını Teknik ayrıntı bölümünde inceleyebilirsiniz.')
        audit = page.locator('[data-section-audit]').last
        expect(audit.locator('[data-section-change]')).to_contain_text('Durum: Boş → Taslak')
        expect(audit.locator('[data-section-change]')).to_contain_text('İçerik: 1 paragraf eklendi')
        expect(audit.locator('pre').first).not_to_be_visible()
        detail = audit.locator('summary')
        detail.focus(); page.keyboard.press('Enter')
        expect(audit.locator('pre').first).to_be_visible()
        page.keyboard.press('Escape'); expect(detail).to_be_focused()
        assert page.locator('[data-section-audit] script').count() == 0
        assert page.evaluate('window.bad === undefined')
        field = page.locator('[name=text]'); field.focus()
        style = field.evaluate('el => {const s=getComputedStyle(el);return [s.outlineColor,s.outlineWidth,s.boxShadow,s.borderColor]}')
        assert style == ['rgb(30, 107, 101)','2px','none','rgb(84, 91, 100)']
        assert page.evaluate('document.documentElement.scrollWidth === innerWidth')
        assert not errors


def test_012_audit_names_batch_lookup_and_missing_user_fallback(section_client, db, sample):
    from types import SimpleNamespace
    from sqlalchemy import event
    client, docs, _, _ = section_client
    doc = docs['birim:'+str(sample['department'].id)]
    configure(db, sample, doc, 'draft-content', 'director')
    sample['user'].full_name = 'Sentetik <Deniz> Yılmaz'
    db.add(ARContentEvent(document_id=doc.id, section_code='I.A', action='section_save',
        before_json='{}', after_json='{}', actor_id=sample['strategy'].id))
    db.commit()
    login(client, sample['user'].email)
    queries = []
    def record(conn, cursor, statement, parameters, context, executemany):
        if re.search(r'users\.id IN\s*\(', statement): queries.append(statement)
    event.listen(db.bind, 'before_cursor_execute', record)
    try:
        response = client.get(f'/activity-sections/reports/{doc.id}?code=I.A')
    finally:
        event.remove(db.bind, 'before_cursor_execute', record)
    assert response.status_code == 200
    assert len(queries) == 1
    assert response.context['event_actor_names'] == {
        sample['user'].id: sample['user'].full_name,
        sample['strategy'].id: sample['strategy'].full_name}
    assert 'Sentetik &lt;Deniz&gt; Yılmaz' in response.text
    assert f'· Kullanıcı {sample["user"].id} ·' not in response.text
    # A dangling historical actor cannot be inserted with the current FK. Render
    # the same route context with a historical event whose user is unavailable.
    missing = SimpleNamespace(actor_id=999999, created_at=datetime(2026,10,25,tzinfo=UTC),
        action='section_save', before_json='{}', after_json='{}')
    html = response.template.render({**response.context, 'events':[missing]})
    assert '· Kullanıcı 999999 ·' in html


def test_012_reopen_shows_only_actor_name(section_client, sample):
    from test_activity_sections import ready_and_freeze
    from test_ar_release_chain import reopen
    client, docs, _, _ = section_client
    token = login(client, sample['strategy'].email)
    doc = docs['idare']
    ready_and_freeze(client, token, doc)
    assert reopen(client, token, doc).status_code == 200
    response = client.get(f'/activity-sections/reports/{doc.id}?code=I.A')
    assert response.status_code == 200
    assert f'Gerekçeyle açıldı · {sample["strategy"].full_name} ·' in response.text
    assert f'(Kullanıcı {sample["strategy"].id})' not in response.text


def test_012_sentence_case_preserves_excel_and_export_headers(section_client, db, sample):
    from io import BytesIO
    from docx import Document
    from openpyxl import load_workbook
    from app import activity_financial_service as finance
    from app.ar_export import build_activity_report_docx, TEMPLATE_PATH
    client, docs, _, _ = section_client
    doc = docs['idare']
    login(client, sample['strategy'].email)
    frozen = json.loads((CONTRACT.parent/'civ101_financial_headers.json').read_text())
    actual = dict(expense=finance.EXPENSE_FIELDS, revenue=finance.REVENUE_FIELDS,
        entries=finance.ENTRY_FIELDS,
        budget={k:finance.budget_columns(k) for k in frozen['budget']})
    assert json.loads(json.dumps(actual)) == frozen
    # Actual XLSX download headers, independently frozen from the base branch.
    for kind, columns in frozen['budget'].items():
        response = client.get(f'/api/activity-sections/reports/{doc.id}/financial/template/{kind}')
        assert response.status_code == 200
        sheet = load_workbook(BytesIO(response.content)).active
        assert list(next(sheet.values)) == [label for _,label in columns]
    payload = sections.report_data(db, sample['strategy'], doc.id)
    for kind, columns in frozen['entries'].items():
        payload['financial']['entries'][kind] = [{key:None if key in {'amount','previous'} else 'Sentetik' for key,_ in columns}]
    word = Document(BytesIO(build_activity_report_docx(payload, TEMPLATE_PATH)))
    headers = [[cell.text for cell in table.rows[0].cells] for table in word.tables]
    assert ['Müdürlük','Sınıflandırma',*[label for _,label in frozen['expense'][:6]]] in headers
    assert ['Müdürlük','Sınıflandırma',*[label for _,label in frozen['revenue']],
        'Net Tahsilat','Gerçekleşme %'] in headers
    for kind in ('assets','audit','aid','explanations'):
        assert [label for _,label in frozen['entries'][kind]] in headers
    expected = json.loads((CONTRACT.parents[2]/'docs/design/rollout/civ-101/header-changes.json').read_text())
    seen = set()
    for code in ('I.C.1','I.C.3','I.C.4','III.A','III.A.1','III.A.2','III.A.3','III.A.4'):
        response = client.get(f'/activity-sections/reports/{doc.id}?code={code}')
        assert response.status_code == 200
        labels = re.findall(r'<th>(.*?)</th>', response.text)
        assert not set(labels) & expected.keys()
        seen.update(labels)
        # Classification names and accessible form labels remain verbatim.
        if code == 'III.A.1':
            assert '06 · Sermaye Giderleri' in response.text
            assert '01 · Vergi Gelirleri' in response.text
            assert '09 · Red ve İadeler' in response.text
            assert 'aria-label="Başlangıç Ödeneği 1"' in response.text
        if code == 'III.A.3':
            assert '<th>Alınan önlemler</th>' in response.text
    own = docs['birim:'+str(sample['department'].id)]
    response = client.get(f'/activity-sections/reports/{own.id}?code=III.A.1')
    assert response.status_code == 200
    seen.update(re.findall(r'<th>(.*?)</th>', response.text))
    assert set(expected.values()) <= seen
