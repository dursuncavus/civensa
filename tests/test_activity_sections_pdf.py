"""CIV-075: product PDF downloads, shared AR content and frozen snapshots."""
import hashlib
import json
import os
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re

import pytest
from docx import Document
from pypdf import PdfReader
from sqlalchemy import select, update

from app import activity_sections_service as service
from app.ar_export import DRAFT, TEMPLATE_PATH, build_activity_report_docx
from app.ar_pdf_export import build_activity_report_pdf
from app.models import ARSectionContent, ARSectionDefinition, Organization
from test_activity_sections import section_client, login, ready_and_freeze, resource, save

pytestmark = pytest.mark.postgres


def compact(text):
    return re.sub(r'\s+', '', text)


def pdf_text(reader):
    return '\n'.join(page.extract_text() for page in reader.pages)


def download(client, doc):
    response = client.get(f'/activity-sections/reports/{doc.id}/export.pdf')
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/pdf'
    assert response.headers['cache-control'] == 'private, no-store'
    assert response.headers['x-content-type-options'] == 'nosniff'
    assert response.content.startswith(b'%PDF-')
    return response, PdfReader(BytesIO(response.content))


def artifact(tmp_path, name, content):
    # Opt-in regeneration writes only these synthetic fixture reports to the repo.
    folder = Path(os.environ.get('CIVENSA_PDF_ARTIFACTS', tmp_path))
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(content)


@pytest.mark.parametrize('scope', ['birim', 'idare'])
def test_pdf_cover_turkish_section_order_draft_and_buttons(section_client, db, sample, tmp_path, scope):
    client, docs, _, _ = section_client
    doc = docs['idare' if scope == 'idare' else f"birim:{sample['department'].id}"]
    contents = db.scalars(select(ARSectionContent).where(ARSectionContent.document_id == doc.id)).all()
    for content in contents:
        definition = db.get(ARSectionDefinition, content.definition_id)
        content.blocks_json = json.dumps([{'type': 'paragraph',
            'text': f"Bölüm sırası {definition.position}: ç ğ ı İ ö ş ü <metin> & veri"}], ensure_ascii=False)
    db.commit()
    payload = service.report_data(db, sample['strategy'], doc.id)
    word = Document(BytesIO(build_activity_report_docx(payload, TEMPLATE_PATH)))
    response, reader = download(client, doc)
    rendered = compact(pdf_text(reader))
    assert f"2027-faaliyet-raporu-{'idare' if scope == 'idare' else 'test-mudurlugu'}-taslak.pdf" in response.headers['content-disposition']
    cover = reader.pages[0].extract_text()
    assert '2027' in cover and 'FAALİYET RAPORU' in cover and 'Test Belediyesi' in cover
    assert ('Test Müdürlüğü' in cover) is (scope == 'birim')
    assert DRAFT not in cover
    assert all(DRAFT in page.extract_text() for page in reader.pages[1:])
    assert compact('Bölüm içeriğini gözden geçirip') in rendered
    assert compact('ç ğ ı İ ö ş ü <metin> & veri') in rendered
    headings = [p.text for p in word.paragraphs if p.style.name.startswith('Civensa Başlık ')]
    assert headings == [s['title'] for s in payload['sections']]
    # The body has a unique marker in each section, so TOC text cannot satisfy order.
    cursor = rendered.index(compact(f"Bölüm sırası {payload['sections'][0]['position']}:"))
    for section, heading in zip(payload['sections'][1:], headings[1:]):
        cursor = rendered.index(compact(heading), cursor)
        if section['code'] != 'EKLER':
            cursor = rendered.index(compact(f"Bölüm sırası {section['position']}:"), cursor)
    for path in ['/activity-sections', f'/activity-sections/reports/{doc.id}']:
        html = client.get(path).text
        assert f'/reports/{doc.id}/export.pdf' in html and 'PDF olarak indir' in html
        assert f'/reports/{doc.id}/export.docx' in html and 'Word olarak indir' in html
    artifact(tmp_path, f'{scope}-taslak.pdf', response.content)


@pytest.mark.parametrize('email', ['test@example.invalid', 'park@example.invalid',
                                  'finance@example.invalid', 'strategy@example.invalid'])
def test_pdf_access_matches_docx(section_client, sample, email):
    client, docs, _, _ = section_client
    login(client, email)
    for doc in docs.values():
        pdf = client.get(f'/activity-sections/reports/{doc.id}/export.pdf')
        word = client.get(f'/activity-sections/reports/{doc.id}/export.docx')
        assert pdf.status_code == word.status_code
        if word.status_code == 403:
            assert pdf.json() == word.json()
            for path in ['/activity-sections', f'/activity-sections/reports/{doc.id}']:
                assert f'/reports/{doc.id}/export.pdf' not in client.get(path).text
        else:
            assert pdf.status_code == 200
    assert client.get('/activity-sections/reports/999999/export.pdf').status_code == client.get(
        '/activity-sections/reports/999999/export.docx').status_code == 404
    client.cookies.clear()
    assert client.get(f"/activity-sections/reports/{docs['idare'].id}/export.pdf").status_code == client.get(
        f"/activity-sections/reports/{docs['idare'].id}/export.docx").status_code == 401


def test_pdf_null_zero_not_applicable_and_landscape_blocks(section_client, sample):
    client, docs, _, _ = section_client
    doc = docs[f"birim:{sample['department'].id}"]
    token = login(client, 'strategy@example.invalid')
    assert resource(client, token, doc, 'hardware', [
        {'category': 'Sıfır bilgisayar', 'count': 0}, {'category': 'Boş yazıcı', 'count': None}]).status_code == 200
    assert save(client, token, doc, 'I.B', status='not_applicable', reason='Bu yıl yetki devri yok.', blocks=[]).status_code == 200
    assert save(client, token, doc, 'I.A', blocks=[{'type': 'bullets', 'items': ['Birinci madde']},
        {'type': 'table', 'rows': [[f'Sütun {i}' for i in range(10)], ['Uzun değer' * 15] * 10]}]).status_code == 200
    _, reader = download(client, doc)
    rendered = compact(pdf_text(reader))
    for value in ['Uygulanamaz: Bu yıl yetki devri yok.', 'Birinci madde', 'Sıfır bilgisayar 0', 'Boş yazıcı —', 'Boş değer: 1']:
        assert compact(value) in rendered
    assert any(page.mediabox.width > page.mediabox.height for page in reader.pages)
    assert reader.pages[-1].mediabox.width < reader.pages[-1].mediabox.height


def test_pdf_button_download_in_browser(section_client, sample, tmp_path):
    from playwright.sync_api import sync_playwright
    client, docs, _, _ = section_client
    doc = docs[f"birim:{sample['department'].id}"]
    with sync_playwright() as playwright:
        options = {'headless': True}
        if os.getenv('CIVENSA_CHROMIUM_PATH'):
            options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
        with playwright.chromium.launch(**options) as browser:
            page = browser.new_page()
            def serve(route):
                response = client.get(route.request.url.replace('http://testserver', ''))
                headers = dict(response.headers)
                headers.pop('set-cookie', None)
                route.fulfill(status=response.status_code, headers=headers, body=response.content)
            page.route('http://testserver/**', serve)
            page.goto(f'http://testserver/activity-sections/reports/{doc.id}')
            page.locator('[data-section-export] > summary').click()
            with page.expect_download() as pending:
                page.get_by_role('link', name='PDF olarak indir', exact=True).click()
            download = pending.value
            assert download.suggested_filename == '2027-faaliyet-raporu-test-mudurlugu-taslak.pdf'
            destination = tmp_path / download.suggested_filename
            download.save_as(destination)
            assert 'Test Müdürlüğü' in pdf_text(PdfReader(destination))


@pytest.mark.parametrize('scope', ['birim', 'idare'])
def test_pdf_frozen_bytes_and_snapshot_stable(section_client, db, sample, monkeypatch, tmp_path, scope):
    client, docs, _, _ = section_client
    doc = docs['idare' if scope == 'idare' else f"birim:{sample['department'].id}"]
    token = login(client, 'strategy@example.invalid')
    own = docs[f"birim:{sample['department'].id}"]
    assert resource(client, token, own, 'hardware', [{'category': 'Donmuş bilgisayar', 'count': 5}]).status_code == 200
    frozen = ready_and_freeze(client, token, doc)
    before, _ = download(client, doc)
    db.execute(update(ARSectionContent).where(ARSectionContent.document_id == doc.id).values(
        blocks_json='[{"type":"paragraph","text":"Sonradan değişen canlı içerik"}]'))
    db.get(Organization, doc.organization_id).name = 'Değişen canlı kurum'
    db.commit()
    from app import assurance_service
    def unexpected_live_resolution(*args):
        raise AssertionError('Frozen assurances must come from snapshot')
    monkeypatch.setattr(assurance_service, 'view', unexpected_live_resolution)
    after, reader = download(client, doc)
    assert before.content == after.content
    assert hashlib.sha256(before.content).digest() == hashlib.sha256(download(client, doc)[0].content).digest()
    assert '-onayli.pdf' in after.headers['content-disposition']
    rendered = compact(pdf_text(reader))
    for value in ['Yıllık değerlendirme', 'Test Belediyesi', 'Donmuş bilgisayar',
                  'İmza: ________________', f"Onaylı kopya · SHA-256: {frozen['sha256'][:16]}"]:
        assert compact(value) in rendered
    for value in [DRAFT, 'Eksikler:', 'Değişen canlı kurum', 'Sonradan değişen canlı içerik']:
        assert compact(value) not in rendered
    artifact(tmp_path, f'{scope}-onayli.pdf', after.content)


def test_pdf_financial_slots_money_core_and_landscape(section_client, db, sample, tmp_path):
    _, docs, _, _ = section_client
    payload = service.report_data(db, sample['strategy'], docs['idare'].id)
    row = dict(department='Örnek Müdürlüğü', classification='01 · Personel', initial=Decimal('12500.23'),
               added=Decimal('0'), deducted=None, actual=Decimal('0'), final=None, rate=None,
               explanation='Merkez açıklaması', birim_explanation='Birim açıklaması')
    payload['financial']['tables']['expenditure']['rows'] = [row]
    payload['financial']['tables']['programme']['rows'] = [dict(row, classification='Alt program')]
    payload['financial']['tables']['revenue']['rows'] = [dict(row, estimate=Decimal('12500.23'), refunds=None, net=None)]
    payload['financial']['entries'] = {
        'assets': [dict(label='Bina', amount=Decimal('12500.23'), previous=None, explanation='Varlık notu')],
        'audit': [dict(audit_type='İç denetim', findings='Bulgu', measures='Tedbir')],
        'aid': [dict(name='Örnek kuruluş', aid_type='Nakdî', amount=0, basis='Karar', explanation='Sonuç')],
        'explanations': [dict(row_key='01', explanation='Sapma gerekçesi')],
        'audit_directorates': [dict(department='Örnek birim', audit_type='Dış denetim', findings='Birim bulgusu')]}
    payload['financial']['statements'][0].update(period='2027', statement_date='2027-12-31',
        explanation='Bilanço açıklaması', filename='ornek.pdf', sha256='a' * 64,
        figures=[dict(label='Aktif toplamı', amount=Decimal('12500.23'))])
    payload['core']['indicators'] = [dict(name='Hizmet sayısı', unit='Adet',
        targets=[dict(value='12500.23', state='verified')], actual=0, actual_state='verified',
        data_source='Kayıt defteri', evidence=[dict(title='Sayım', locator='Sayfa 1', source_revision='2027')],
        explanation='Gösterge açıklaması', deviation=dict(delta=0, corrective_action='Düzeltici çalışma'))]
    content = build_activity_report_pdf(payload)
    reader = PdfReader(BytesIO(content))
    word = Document(BytesIO(build_activity_report_docx(payload, TEMPLATE_PATH)))
    word_text = compact('\n'.join([p.text for p in word.paragraphs] + [
        cell.text for table in word.tables for row in table.rows for cell in row.cells]))
    rendered = compact(pdf_text(reader))
    for value in ['12.500,23', '0,00', '—', 'Merkez açıklaması', 'Birim açıklaması', 'Birim bulgusu',
                  'Varlık notu', 'Tedbir', 'Örnek kuruluş', 'Bilanço açıklaması', 'Aktif toplamı',
                  'Kayıt defteri', 'Sayım', 'Hizmet sayısı', 'Gösterge açıklaması', 'Düzeltici çalışma',
                  'Alt program', 'Net Tahsilat', 'PP Bütçesi', 'Boş değer']:
        assert compact(value) in rendered and compact(value) in word_text
    financial_page = next(page for page in reader.pages if
        compact('Başlangıç Ödeneği') in compact(page.extract_text()) and
        '12.500,23' in page.extract_text())
    assert financial_page.mediabox.width > financial_page.mediabox.height
    assert compact('12.500,23') in compact(financial_page.extract_text())
    artifact(tmp_path, 'idare-mali-tablolar-taslak.pdf', content)
