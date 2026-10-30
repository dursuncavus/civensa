# Frozen official_export from a5c40f6, before CIV-160.
# Executed with the unchanged export helpers on the same synthetic fixture.
def official_export(db, actor, period_id, format, *, draft=False):
    if format not in ('docx', 'xlsx'):
        m.fail('Rapor dosya biçimi geçersiz. Word veya Excel bağlantısını seçin.', 'format')
    period, plan, rows, metrics = official_data(db, actor, period_id, draft=draft)
    title, subtitle = plan.title + ' · ' + plan.version_label, m.period_label(period)
    cover = [f"Gerçekleşme Oranı: {metrics['period_rate']}",
             f"{'İlk 6 Aylık' if period.half == 1 else 'İkinci 6 Aylık'} Başarı Yüzdesi: {metrics['plan_rate']}",
             f"Plan toplamı: {format_number(metrics['plan_total'])} eylem",
             'Sayılar eylem bazındadır; yalnız onaylı sorumlu birim girişleri oranlara katılır.']
    if metrics['warning']:
        cover.append(metrics['warning'])
    if format == 'xlsx':
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        sheet = book.active
        sheet.title = 'Özet'
        literal_row(sheet, ['TASLAK' if draft else 'Onaylı veriler', title])
        literal_row(sheet, ['Dönem', subtitle])
        for line in cover:
            literal_row(sheet, [line])
        literal_row(sheet, COMPONENT_HEADERS)
        for row in metrics['components']:
            literal_row(sheet, row)
        sheet.column_dimensions['A'].width = 70
        for col in 'BCDEFG':
            sheet.column_dimensions[col].width = 25
        sheet = book.create_sheet('Resmî İzleme')
        literal_row(sheet, OFFICIAL_HEADERS)
        for row in rows:
            literal_row(sheet, row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for col in 'ABCDEFGHIJK':
            sheet.column_dimensions[col].width = 32 if col not in 'BEK' else 55
        for sheet in book:
            for row in sheet:
                for cell in row:
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
        return export_book(book)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.left_margin = section.right_margin = Inches(.45)
    doc.styles['Normal'].font.size = Pt(8)
    if draft:
        section.header.paragraphs[0].text = 'TASLAK — Onay süreci tamamlanmamış olabilir'
    doc.add_heading(('TASLAK · ' if draft else '') + 'İç Kontrol İzleme Raporu · Resmî format', 0)
    for line in [title, subtitle, *cover]:
        doc.add_paragraph(line)
    doc.add_page_break()
    for headers, contents in [(COMPONENT_HEADERS, metrics['components']), (OFFICIAL_HEADERS, rows)]:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for row in contents:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = format_number(value) if isinstance(value, int) else str(value)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
