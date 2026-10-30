"""CIV-160 PostgreSQL acceptance; all narrative text is explicitly synthetic."""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
from types import ModuleType, SimpleNamespace
from uuid import uuid4

from docx import Document
from fastapi import HTTPException
import pytest
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import internal_control_narrative_service as n
from app import internal_control_monitoring_service as m
from app import internal_control_monitoring_export as export
from app.models import Department, User
from test_internal_control_monitoring import monitoring, approve
from test_ic_our_actions_design import ic_client
from test_internal_control import workbook, action_row
from test_internal_control_questionnaire import question_file, answer_data
from app import internal_control_questionnaire_service as q

pytestmark = pytest.mark.postgres
ROOT = Path(__file__).resolve().parents[1]


def synthetic(db, period, **changes):
    result = {key: 'Test ' + label.lower() + ' metni …\nTest ikinci satır' for key, label, _ in n.fields(db, period)}
    result['band_label'] = 'Test puan aralığı etiketi'
    return {**result, **changes}


def save(db, d, actor=None, **changes):
    return n.save(db, actor or d['staff'], d['period'].id, n.state(db, d['period'].id)[1], synthetic(db, d['period'], **changes))


def approved(db, d, actor=None):
    record = save(db, d, actor)
    n.transition(db, d['strategy'], d['period'].id, 'submit', record.revision)
    result = n.transition(db, d['strategy'], d['period'].id, 'approve', record.revision+1)
    db.commit()
    return result


def path(d):
    return f'/internal-control/monitoring/periods/{d["period"].id}/narrative'


def strategy_client(client, d):
    from app.main import app, require_user
    app.dependency_overrides[require_user] = lambda: d['strategy']
    response = client.get(path(d))
    assert response.status_code == 200
    return response.context['request'].session.get('csrf_token', '')


@pytest.mark.parametrize('writer', ['staff', 'strategy'])
def test_every_save_is_full_revision_trimmed_and_keeps_newlines(db, monitoring, writer):
    d = monitoring
    first = save(db, d, d[writer], mission='  Test misyon metni …\nTest satır 2  ')
    db.commit()
    before = first.snapshot_json
    second = save(db, d, d[writer], vision='Test yeni vizyon metni …')
    db.commit()
    assert first.revision == 1 and second.revision == 2 and first.snapshot_json == before
    assert json.loads(first.snapshot_json)['mission'] == 'Test misyon metni …\nTest satır 2'
    assert json.loads(second.snapshot_json) == synthetic(db, d['period'], vision='Test yeni vizyon metni …')
    assert first.actor_id == d[writer].id


@pytest.mark.parametrize('key,limit,label', [
    ('mission',2000,'Misyon'), ('vision',2000,'Vizyon'), ('component_KOS',8000,'Kontrol Ortamı'),
    ('strengths',8000,'Güçlü yönler'), ('improvements',8000,'Gelişim alanları'),
    ('recommendations',8000,'Öneriler'), ('band_label',100,'Puan aralığı'), ('band_explanation',2000,'Puan aralığı açıklaması')])
def test_over_length_http_retains_all_values_and_saves_nothing(db, monitoring, client, key, limit, label):
    d = monitoring
    token = strategy_client(client, d)
    values = synthetic(db, d['period'], **{key: 'Ş'*(limit+1)})
    response = client.post(path(d)+'/save', data={**values, 'csrf_token':token, 'expected_revision':0})
    assert response.status_code == 422 and label in response.text and m.format_number(limit) in response.text
    assert response.context['request'].state.pp_attempt_rendered
    assert values[key] in response.text
    assert 'Test ikinci satır' in response.text
    assert n.state(db, d['period'].id) == ('draft', 0)


@pytest.mark.parametrize('route', ['', '/save', '/copy', '/transition'])
def test_directorate_403_on_every_narrative_route(db, monitoring, client, route):
    d = monitoring
    response = client.get(f'/internal-control/our-actions?period_id={d["period"].id}')
    token = response.context['request'].session.get('csrf_token', '')
    response = client.post(path(d)+route, data={'csrf_token':token,'expected_revision':0}) if route else client.get(path(d))
    assert response.status_code == 403
    assert client.get('/internal-control/monitoring/narrative-editor.js').status_code == 403
    assert n.latest(db, d['period'].id) is None


def test_submit_lists_every_empty_required_field_keeps_values(db, monitoring, client):
    d = monitoring
    token = strategy_client(client, d)
    response = client.post(path(d)+'/save', data={'csrf_token':token,'expected_revision':0,
        'operation':'submit','band_explanation':'Test korunan puan açıklaması'})
    assert response.status_code == 422
    for key, label, _ in n.fields(db, d['period']):
        if key != 'band_explanation': assert label in response.text
    assert 'Test korunan puan açıklaması' in response.text
    assert n.latest(db, d['period'].id) is None
    record = save(db, d, band_explanation='')
    n.transition(db, d['staff'], d['period'].id, 'submit', record.revision)
    db.commit()
    assert n.state(db, d['period'].id)[0] == 'submitted'


def test_stale_and_closed_period_refuse_all_mutations(db, monitoring):
    d = monitoring
    save(db, d)
    db.commit()
    for operation in [lambda: n.save(db,d['staff'],d['period'].id,0,{}),
                      lambda: n.transition(db,d['strategy'],d['period'].id,'submit',0)]:
        with pytest.raises(HTTPException) as exc: operation()
        assert exc.value.status_code == 409 and 'yenileyip' in str(exc.value.detail)
    m.close_period(db,d['strategy'],d['period'].id,1,'Test dönem kapatma')
    db.commit()
    for operation in [lambda: n.save(db,d['staff'],d['period'].id,1,{}),
                      lambda: n.transition(db,d['strategy'],d['period'].id,'submit',1),
                      lambda: n.transition(db,d['strategy'],d['period'].id,'approve',1),
                      lambda: n.copy_previous(db,d['staff'],d['period'].id,1)]:
        with pytest.raises(HTTPException) as exc: operation()
        assert exc.value.status_code == 409 and 'kapalı' in str(exc.value.detail)
    assert n.state(db,d['period'].id)[1] == 1


@pytest.fixture
def second_director(db, monitoring):
    department = Department(organization_id=monitoring['department'].organization_id,
        name='Test İkinci Strateji Birimi')
    db.add(department)
    db.flush()
    director = User(email='second-director@civ160.invalid', full_name='Test İkinci Strateji Müdürü',
        role='strategy', position='director', department_id=department.id, password_hash='unused')
    db.add(director)
    db.commit()
    return director


def test_return_and_unchanged_resubmit_do_not_reset_authorship(db, monitoring, second_director):
    d = monitoring
    period_id = d['period'].id
    first = save(db, d, d['strategy'])
    n.transition(db, d['strategy'], period_id, 'submit', first.revision)
    n.transition(db, second_director, period_id, 'return', 2, 'Test iade gerekçesi')
    resubmitted = n.transition(db, d['strategy'], period_id, 'submit', 3)
    db.commit()
    assert resubmitted.snapshot_json == first.snapshot_json
    with pytest.raises(HTTPException) as exc:
        n.transition(db, d['strategy'], period_id, 'approve', 4)
    assert exc.value.status_code == 403
    assert exc.value.detail == m.SOD_ERROR.format(record='İç Kontrol rapor anlatı bölümleri')
    assert n.state(db, period_id) == ('submitted', 4)
    approval = n.transition(db, second_director, period_id, 'approve', 4)
    db.commit()
    assert approval.actor_id == second_director.id
    assert approval.snapshot_json == first.snapshot_json


def test_last_approval_starts_new_authorship_window(db, monitoring, second_director):
    d = monitoring
    period_id = d['period'].id
    first = save(db, d, d['strategy'])
    n.transition(db, second_director, period_id, 'submit', first.revision)
    prior_approval = n.transition(db, second_director, period_id, 'approve', 2)
    n.transition(db, second_director, period_id, 'reopen', 3, 'Test yeni değerlendirme')
    edited = save(db, d, mission='Test personel tarafından düzeltilmiş misyon metni …')
    n.transition(db, d['staff'], period_id, 'submit', edited.revision)
    db.commit()
    with pytest.raises(HTTPException) as exc:
        n.transition(db, d['staff'], period_id, 'approve', 6)
    assert exc.value.status_code == 403
    approval = n.transition(db, d['strategy'], period_id, 'approve', 6)
    db.commit()
    assert first.revision < prior_approval.revision < edited.revision < approval.revision
    assert approval.actor_id == d['strategy'].id
    assert approval.snapshot_json == edited.snapshot_json != first.snapshot_json


def test_workflow_reasons_permissions_sod_and_frozen_history(db, monitoring, second_director):
    d = monitoring
    record = save(db,d,d['strategy'])
    n.transition(db,d['staff'],d['period'].id,'submit',record.revision)
    db.commit()
    with pytest.raises(HTTPException) as exc: n.transition(db,d['strategy'],d['period'].id,'approve',2)
    assert exc.value.status_code == 403 and 'Görevler ayrılığı' in str(exc.value.detail)
    for action,reason in [('return','Test iade'),('approve','')]:
        with pytest.raises(HTTPException) as exc: n.transition(db,d['staff'],d['period'].id,action,2,reason)
        assert exc.value.status_code == 403
    with pytest.raises(HTTPException) as exc: n.transition(db,d['strategy'],d['period'].id,'return',2)
    assert exc.value.status_code == 422
    n.transition(db,d['strategy'],d['period'].id,'return',2,'Test eksikleri tamamlayın')
    save(db,d)
    n.transition(db,d['strategy'],d['period'].id,'submit',4)
    with pytest.raises(HTTPException) as exc:
        n.transition(db,d['strategy'],d['period'].id,'approve',5)
    assert exc.value.status_code == 403
    first_approval = n.transition(db,second_director,d['period'].id,'approve',5)
    db.commit()
    with pytest.raises(HTTPException): save(db,d)
    with pytest.raises(HTTPException) as exc: n.transition(db,d['staff'],d['period'].id,'reopen',6,'Test gerekçe')
    assert exc.value.status_code == 403
    with pytest.raises(HTTPException): n.transition(db,d['strategy'],d['period'].id,'reopen',6)
    reopened = n.transition(db,d['strategy'],d['period'].id,'reopen',6,'Test yeni inceleme')
    db.commit()
    assert reopened.to_status == 'draft' and reopened.snapshot_json == first_approval.snapshot_json
    assert n.latest(db,d['period'].id,approved=True).id == first_approval.id
    assert db.get(n.Event,first_approval.id).action == 'approve'


def test_another_director_can_approve_while_submit_is_not_authorship(db, monitoring):
    d = monitoring
    first = save(db,d,d['strategy'])
    # User administration permits one active director per department.
    other_department = Department(organization_id=d['department'].organization_id,name='Test Diğer Strateji Birimi')
    db.add(other_department); db.flush()
    director = User(email='director@civ160.invalid',full_name='Test Diğer Strateji Müdürü',role='strategy',
        position='director',department_id=other_department.id,password_hash='unused')
    db.add(director); db.commit()
    n.transition(db,director,d['period'].id,'submit',first.revision)
    n.transition(db,director,d['period'].id,'approve',2)
    db.commit()
    assert n.latest(db,d['period'].id).actor_id == director.id


def test_approved_revision_rejects_orm_and_sql_update_delete_truncate(db, monitoring):
    d = monitoring
    record = approved(db,d)
    with pytest.raises(Exception,match='değiştirilemez'):
        with db.begin_nested():
            record.snapshot_json='{}'; db.flush()
    with pytest.raises(Exception,match='değiştirilemez'):
        with db.begin_nested(): db.delete(record); db.flush()
    schema = db.get_bind().get_execution_options()['schema_translate_map'][None]
    table = f'"{schema}".ic_narrative_events'
    for sql in [f'UPDATE {table} SET snapshot_json=\'{{}}\'',f'DELETE FROM {table}',f'TRUNCATE {table} CASCADE']:
        with pytest.raises(DBAPIError,match='geçmişi değiştirilemez'):
            with db.begin_nested(): db.execute(text(sql))
    assert n.latest(db,d['period'].id).snapshot_json == record.snapshot_json


def test_raw_insert_enforces_permissions_snapshot_revision_and_sod(db, monitoring):
    d = monitoring
    valid = synthetic(db,d['period'])
    def insert(actor, revision, action, origin, destination, snapshot, reason=''):
        db.execute(n.Event.__table__.insert().values(period_id=d['period'].id,revision=revision,
            action=action,from_status=origin,to_status=destination,actor_id=actor.id,actor_name=actor.full_name,
            reason=reason,snapshot_json=json.dumps(snapshot)))
    for actor,snapshot in [(d['user'],valid),(d['staff'],{}),(d['staff'],{**valid,'mission':'x'*2001}),
                            (d['staff'],{**valid,'vision':None}),(d['staff'],{**valid,'mission':' test '})]:
        with pytest.raises(DBAPIError):
            with db.begin_nested(): insert(actor,1,'save','draft','draft',snapshot)
    save(db,d,d['strategy']); db.commit()
    with pytest.raises(DBAPIError):
        with db.begin_nested(): insert(d['staff'],1,'save','draft','draft',valid)
    n.transition(db,d['strategy'],d['period'].id,'submit',1); db.commit()
    for actor,snapshot in [(d['staff'],valid),(d['strategy'],valid),(d['strategy'],{**valid,'mission':'Test karar değişikliği'})]:
        with pytest.raises(DBAPIError):
            with db.begin_nested(): insert(actor,3,'approve','submitted','approved',snapshot)
    assert n.state(db,d['period'].id) == ('submitted',2)


def test_copy_previous_latest_approved_same_plan_and_clicking_author(db, monitoring):
    d = monitoring
    first = approved(db,d)
    n.transition(db,d['strategy'],d['period'].id,'reopen',first.revision,'Test ikinci onay hazırlığı')
    updated = save(db,d,mission='Test son onaylı misyon metni …')
    n.transition(db,d['staff'],d['period'].id,'submit',updated.revision)
    source = n.transition(db,d['strategy'],d['period'].id,'approve',updated.revision+1)
    db.commit()
    assert source.snapshot_json != first.snapshot_json
    n.transition(db,d['strategy'],d['period'].id,'reopen',source.revision,'Test yeni taslak')
    save(db,d,mission='Test kaynak sonraki taslak — kopyalanmamalı')
    m.close_period(db,d['strategy'],d['period'].id,1,'Test önceki dönem sonu'); db.commit()
    next_period = m.open_period(db,d['strategy'],d['plan'].id,2027,2,m.today(db)+timedelta(days=1000)); db.commit()
    record = n.copy_previous(db,d['strategy'],next_period.id,0); db.commit()
    assert record.snapshot_json == source.snapshot_json and record.actor_id == d['strategy'].id
    assert record.action == 'save' and record.to_status == 'draft'
    with pytest.raises(HTTPException): n.copy_previous(db,d['strategy'],next_period.id,1)
    n.transition(db,d['staff'],next_period.id,'submit',1)
    with pytest.raises(HTTPException) as exc: n.transition(db,d['strategy'],next_period.id,'approve',2)
    assert exc.value.status_code == 403


def test_copy_no_source_uses_nearest_period_only_and_rejects_other_plan(db, monitoring):
    d = monitoring
    with pytest.raises(HTTPException) as exc: n.copy_previous(db,d['staff'],d['period'].id,0)
    assert 'Önceki dönemde onaylı' in str(exc.value.detail)
    assert n.latest(db,d['period'].id) is None
    approved(db,d)
    m.close_period(db,d['strategy'],d['period'].id,1,'Test kapanış'); db.commit()
    second=m.open_period(db,d['strategy'],d['plan'].id,2027,2,m.today(db)); db.commit()
    m.close_period(db,d['strategy'],second.id,1,'Test anlatısız ara dönem'); db.commit()
    third=m.open_period(db,d['strategy'],d['plan'].id,2028,1,m.today(db)); db.commit()
    with pytest.raises(HTTPException): n.copy_previous(db,d['staff'],third.id,0)
    assert n.latest(db,third.id) is None
    # Make the other plan's nearest earlier period approved: an unscoped lookup
    # would now copy it, so the refusal below proves plan isolation.
    approved(db,{**d,'period':third})
    # A separate imported plan never inherits the prior plan's approval.
    from app import internal_control_service as ic
    batch,_=ic.create_import(db,d['strategy'],workbook([action_row()]),'civ160-other-test.xlsx')
    plan=ic.commit_import(db,d['strategy'],batch.id); db.commit()
    ic.activate_plan(db,d['strategy'],plan.id); db.commit()
    other=m.open_period(db,d['strategy'],plan.id,2028,2,m.today(db)); db.commit()
    with pytest.raises(HTTPException): n.copy_previous(db,d['staff'],other.id,0)
    assert n.latest(db,other.id) is None


def test_directorate_latest_approved_mission_vision_only_and_escaping(db, monitoring, ic_client):
    from test_ic_our_actions_design import render_contract
    client=ic_client
    d=monitoring
    url=f'/internal-control/our-actions?period_id={d["period"].id}'
    assert 'Misyon ve Vizyon' not in client.get(url).text
    saved=save(db,d,mission='Test misyon <script>alert(1)</script>',strengths='Test gizli güçlü yönler')
    db.commit()
    assert 'Misyon ve Vizyon' not in client.get(url).text
    n.transition(db,d['staff'],d['period'].id,'submit',saved.revision)
    n.transition(db,d['strategy'],d['period'].id,'approve',2); db.commit()
    response=client.get(url)
    assert 'Misyon ve Vizyon' in response.text and 'Test misyon &lt;script&gt;' in response.text
    assert 'Test gizli güçlü yönler' not in response.text and '<script>alert(1)</script>' not in response.text
    assert set(response.context['narrative_mission_vision']) == {'mission','vision'}
    render_contract(response)
    n.transition(db,d['strategy'],d['period'].id,'reopen',3,'Test yeni taslak')
    save(db,d,mission='Test gizli taslak misyon'); db.commit()
    response=client.get(url)
    assert 'Test gizli taslak misyon' not in response.text and 'Test misyon &lt;script&gt;' in response.text
    render_contract(response)
    n.transition(db,d['staff'],d['period'].id,'submit',5)
    n.transition(db,d['strategy'],d['period'].id,'approve',6); db.commit()
    assert 'Test gizli taslak misyon' in client.get(url).text


def test_questionnaire_reference_98_of_140_and_no_band_suggestion(db, monitoring, client):
    d=monitoring
    strategy_client(client,d)
    response=client.get(path(d))
    assert response.context['reference']['total']['score'] is None
    assert 'Toplam puan / azami: <strong>—</strong>' in response.text
    catalogue=q.save_list(db,d['staff'],'Test 70 soru',question_file(count=70))
    q.attach(db,d['strategy'],d['period'].id,catalogue.id,[d['department'].id],0); db.commit()
    batch=q.batches(db,d['period'].id)[0]
    q.save_answers(db,d['user'],batch.id,0,answer_data(*(['yes']*47+['no']*19+['developing']*4)),submit=True)
    q.transition(db,d['strategy'],batch.id,'approve',2); db.commit()
    response=client.get(path(d))
    assert '98 / 140' in response.text and '72,86% (51/70)' in response.text and 'Onaylı müdürlük: 1' in response.text
    assert response.context['values']['band_label'] == ''
    assert 'sistem bant hesaplamaz' in response.text
    # The directorate block still appears with questionnaire-backed our-actions.
    approved(db,d)
    from app.main import app,require_user
    app.dependency_overrides[require_user]=lambda:d['user']
    response=client.get(f'/internal-control/our-actions?period_id={d["period"].id}')
    assert 'Misyon ve Vizyon' in response.text and 'Test misyon' in response.text


def baseline_export():
    source=(ROOT/'tests/artifacts/civ160/official_export_before_civ160.py').read_text()
    module=ModuleType('app._civ160_baseline_export')
    module.__dict__.update(export.__dict__)
    exec(compile(source,'CIV-160 base official export','exec'),module.__dict__)
    return module


@pytest.mark.parametrize('draft',[False,True])
def test_no_narrative_word_and_all_excel_legacy_exports_are_byte_identical(db, monitoring, monkeypatch, draft):
    d=monitoring
    for batch in d['batches']: approve(db,d,batch)
    old=baseline_export()
    # ZIP timestamps and workbook modified metadata vary with wall time; pin
    # these serialization clocks and compare actual bytes, not extracted XML.
    from datetime import datetime, timezone
    import openpyxl.packaging.core as core
    import openpyxl.writer.excel as writer
    import zipfile
    class FixedDatetime(datetime):
        @classmethod
        def now(cls,tz=None): return cls(2026,10,30,12,0,0,tzinfo=tz)
    clock=SimpleNamespace(datetime=FixedDatetime, timezone=timezone)
    monkeypatch.setattr(writer,'datetime',clock)
    monkeypatch.setattr(core,'datetime',clock)
    monkeypatch.setattr(zipfile,'time',SimpleNamespace(time=lambda:1793361600, localtime=lambda *a:(2026,10,30,12,0,0,4,303,0)))
    def outputs(module):
        return [module.official_export(db,d['strategy'],d['period'].id,'xlsx',draft=draft),
                module.period_export(db,d['strategy'],d['period'].id,'docx',draft=draft),
                module.period_export(db,d['strategy'],d['period'].id,'xlsx',draft=draft),
                module.annual_export(db,d['strategy'],d['plan'].id,2027)]
    assert export.official_export(db,d['strategy'],d['period'].id,'docx',draft=draft) == old.official_export(db,d['strategy'],d['period'].id,'docx',draft=draft)
    before=outputs(old)
    save(db,d);db.commit()
    assert outputs(export) == before


def test_word_order_components_lines_draft_and_final_refusal(db, monitoring, client):
    d=monitoring
    token=strategy_client(client,d)
    record=save(db,d,mission='Test misyon ilk satır\r\nTest misyon ikinci satır');db.commit()
    url=f'/internal-control/monitoring/periods/{d["period"].id}/official-report.docx'
    response=client.get(url)
    assert response.status_code==409
    assert 'Rapor anlatı bölümleri onaylanmadı.' in response.text
    # The narrative-specific refusal also applies while action batches are pending.
    for batch in d['batches']: approve(db,d,batch)
    assert 'Rapor anlatı bölümleri onaylanmadı.' in client.get(url).text
    draft=Document(BytesIO(export.official_export(db,d['strategy'],d['period'].id,'docx',draft=True)))
    headings=[p.text for p in draft.paragraphs if p.style.name.startswith('Heading')]
    assert all(h.endswith('(taslak metin)') for h in headings)
    n.transition(db,d['staff'],d['period'].id,'submit',record.revision)
    n.transition(db,d['strategy'],d['period'].id,'approve',2);db.commit()
    doc=Document(BytesIO(export.official_export(db,d['strategy'],d['period'].id,'docx')))
    paragraphs=[p.text for p in doc.paragraphs]
    expected=['Misyon','Vizyon','Bileşen değerlendirmeleri',*[label for _,label in n.components(db,d['period'])],
              'Puan aralığı','Güçlü yönler','Gelişim alanları','Öneriler']
    assert [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')] == expected
    assert paragraphs.index('Test misyon ikinci satır')==paragraphs.index('Test misyon ilk satır')+1
    assert '(taslak metin)' not in '\n'.join(paragraphs)
    body=list(doc.element.body)
    narrative_end=next(i for i,e in enumerate(body) if 'Test öneriler' in ''.join(e.itertext()))
    page_break=next(i for i,e in enumerate(body) if e.xpath('.//w:br[@w:type="page"]'))
    assert narrative_end<page_break<next(i for i,e in enumerate(body) if e.tag.endswith('}tbl'))
    n.transition(db,d['strategy'],d['period'].id,'reopen',3,'Test tekrar inceleme');db.commit()
    assert client.get(url).status_code==409
    assert client.get(url.replace('.docx','.xlsx')).status_code==200


def test_concurrent_first_save_serializes_period(db, monitoring):
    d=monitoring
    engine=db.get_bind();db.commit()
    data=synthetic(db,d['period']);pid=d['period'].id;uid=d['staff'].id;db.rollback()
    def attempt():
        with Session(engine) as session:
            try:
                n.save(session,session.get(User,uid),pid,0,data);session.commit();return 200
            except HTTPException as exc:
                session.rollback();return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(lambda _:attempt(),range(2)))==[200,409]


def test_migration_postgres_roundtrip_refuses_data_single_head(db, monitoring):
    from alembic.config import Config
    from alembic.script import ScriptDirectory
    d=monitoring
    script=ScriptDirectory.from_config(Config(str(ROOT/'alembic.ini')))
    assert script.get_heads()==['c160a1b2c3d4']
    assert script.get_revision('c160a1b2c3d4').down_revision=='c085a1b2c3d4'
    # Exercise the actual migration on a disposable, independently migrated schema.
    base=create_engine(os.environ['CIVENSA_TEST_POSTGRES_URL'])
    schema='civ160_migration_'+uuid4().hex
    with base.begin() as conn:conn.execute(CreateSchema(schema))
    url=make_url(os.environ['CIVENSA_TEST_POSTGRES_URL']).update_query_dict({'options':'-csearch_path='+schema}).render_as_string(hide_password=False)
    engine=create_engine(url)
    env={**os.environ,'DATABASE_URL':url}
    def migrate(direction,target):
        return subprocess.run([sys.executable,'-m','alembic',direction,target],cwd=ROOT,env=env,capture_output=True,text=True)
    try:
        for direction,target in [('upgrade','c085a1b2c3d4'),('upgrade','head'),('downgrade','c085a1b2c3d4'),('upgrade','head')]:
            result=migrate(direction,target);assert result.returncode==0,result.stderr
        # Copy the synthetic fixture's catalogue and period to the migrated schema
        # by inserting a valid empty draft via a Session using current services.
        from app import internal_control_service as ic
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO organizations(id,name) VALUES (1,'Test sentetik')"))
            conn.execute(text("INSERT INTO departments(id,organization_id,name) VALUES (1,1,'Test birim')"))
            conn.execute(text("INSERT INTO users(id,email,full_name,role,password_hash,is_active,department_id,position) VALUES (1,'test@civ160.invalid','Test Strateji','strategy','unused',true,1,'director')"))
        with Session(engine) as session:
            actor=session.get(User,1)
            imp,_=ic.create_import(session,actor,workbook([action_row(**{'7':'Test birim'})]),'civ160-test.xlsx')
            plan=ic.commit_import(session,actor,imp.id);session.commit()
            ic.activate_plan(session,actor,plan.id);session.commit()
            period=m.open_period(session,actor,plan.id,2027,1,m.today(session));session.commit()
            n.save(session,actor,period.id,0,{});session.commit()
            before=session.execute(select(n.Event.snapshot_json)).scalar_one()
        result=migrate('downgrade','c085a1b2c3d4')
        assert result.returncode!=0 and 'Rapor anlatı verisi var' in result.stderr
        with engine.connect() as conn:
            assert conn.execute(text('SELECT version_num FROM alembic_version')).scalar_one()=='c160a1b2c3d4'
            assert conn.execute(text('SELECT snapshot_json FROM ic_narrative_events')).scalar_one()==before
    finally:
        engine.dispose()
        with base.begin() as conn:conn.execute(DropSchema(schema,cascade=True))
        base.dispose()


def test_plan_components_use_catalogue_order_and_word_subheadings(db, sample):
    from app import internal_control_service as ic
    from app.internal_control_models import COMPONENTS
    # Reverse import order deliberately: editor/report must use catalogue order.
    standard={'KOS':'KOS1','RDS':'RDS5','KFS':'KFS7','BİS':'BİS13','İS':'İS17'}
    rows=[action_row(**{'0':code,'1':standard[code],'2':standard[code]+'.1',
        '5':standard[code]+'.1.1','3':'Test genel şart','6':'Test eylem','13':'Test standart'})
        for code in reversed(COMPONENTS)]
    imp,report=ic.create_import(db,sample['strategy'],workbook(rows),'civ160-components-test.xlsx')
    assert report['can_commit'],report['issues']
    plan=ic.commit_import(db,sample['strategy'],imp.id);db.commit()
    ic.activate_plan(db,sample['strategy'],plan.id);db.commit()
    period=m.open_period(db,sample['strategy'],plan.id,2027,1,m.today(db));db.commit()
    assert n.components(db,period)==list(COMPONENTS.items())
    values=synthetic(db,period)
    values['component_RDS']=''
    with pytest.raises(HTTPException) as exc:
        n.save(db,sample['strategy'],period.id,0,values,submit=True)
    assert 'Risk Değerlendirme' in str(exc.value.detail)
    n.save(db,sample['strategy'],period.id,0,synthetic(db,period));db.commit()
    doc=Document(BytesIO(export.official_export(db,sample['strategy'],period.id,'docx',draft=True)))
    assert [p.text for p in doc.paragraphs if p.style.name=='Heading 2']==[label+' (taslak metin)' for label in COMPONENTS.values()]
