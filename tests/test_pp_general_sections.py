"""CIV-071 general sections: disposable PostgreSQL regression coverage."""
import json
import threading
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.db import get_db
from app.main import app, create_period_record
from app.import_service import parse_pp_import
from app.models import (Activity, AuditLog, EconomicValue, ImportBatch, Organization,
                        ReportSection, ReportingPeriod)
from app.pp_manifest import completion_manifest, SECTION_TYPES, TABLE_TYPES, table_type
from app.structure_service import copy_general_sections, copy_period_work
from test_word_import import import_context, import_client, word_document, assert_waiting_for_institution_lock

pytestmark = pytest.mark.postgres

def test_period_creation_copies_general_sections(db, sample):
    original = ReportSection(period_id=sample['period'].id, title='Başkan Sunuşu',
                             position=2, kind='text', content_json=json.dumps({
                                 'text': 'Kentimizin geleceği.', 'placement': 'after'}, ensure_ascii=False))
    db.add(original)
    db.commit()
    target = create_period_record(db, sample['strategy'], {'year': '2028', 'base_period_id': original.period_id})
    copied = db.scalar(select(ReportSection).where(ReportSection.period_id == target.id))
    assert copied is not None
    assert (copied.title, copied.position, copied.kind, copied.content_json) == (
        original.title, original.position, original.kind, original.content_json)



def section_document(*, financial=True):
    doc = Document(BytesIO(word_document())) if financial else Document()
    # Put institutional content before the financial tables.
    heading = doc.add_heading('A- Yetki, Görev ve Sorumluluklar', level=1)
    body = doc.add_paragraph('5393 sayılı Kanun kapsamında hizmet sunulur.\nİkinci satır korunur.')
    org = doc.add_heading('B- Teşkilat Yapısı', level=1)
    org_body = doc.add_paragraph('Strateji ve harcama birimleri birlikte çalışır.')
    for index, item in enumerate([heading, body, org, org_body]):
        doc.element.body.insert(index, item._p)
    doc.add_heading('Performans Bilgileri', level=1)
    doc.add_heading('E- Diğer Hususlar', level=2)
    doc.add_paragraph('Dönem sonu açıklaması.')
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


def scan(client, period_id, data=None):
    response = client.post(f'/api/periods/{period_id}/data-import-preview',
        files={'file': ('genel-bolumler.docx', data or section_document())})
    assert response.status_code == 200, response.text
    return response.json()


def records(db, period_id):
    return db.scalars(select(ReportSection).where(ReportSection.period_id == period_id)
                      .order_by(ReportSection.position, ReportSection.id)).all()


def test_period_copy_preserves_edited_target_reports_conflicts_and_repeats_without_writes(db, sample):
    base = sample['period']
    source = ReportSection(period_id=base.id, title='Başkan Sunuşu', kind='text', position=4,
                           content_json='{"text":"Kaynak","placement":"before"}')
    table1 = ReportSection(period_id=base.id, title='Stratejik Amaçlar', kind='table', position=5,
                          content_json='{"rows":[["Başlık"],["Birinci amaç"]],"placement":"before"}')
    table2 = ReportSection(period_id=base.id, title='Stratejik Amaçlar', kind='table', position=6,
                          content_json='{"rows":[["Başlık"],["İkinci amaç"]],"placement":"before"}')
    target = ReportingPeriod(year=2028)
    db.add_all([source, table1, table2, target]); db.flush()
    edited = ReportSection(period_id=target.id, title=source.title, kind='text', position=1,
                           content_json='{"text":"Hedefte elle düzenlendi","placement":"before"}')
    db.add(edited); db.commit()
    result = copy_general_sections(db, sample['strategy'], base, target); db.commit()
    assert result['copied'] == 2 and result['kept'] == 1
    assert result['conflicts'][0]['target_ids'] == [edited.id]
    assert 'korundu' in result['conflicts'][0]['message']
    assert edited.content_json == '{"text":"Hedefte elle düzenlendi","placement":"before"}'
    before = [(r.id, r.content_json) for r in records(db, target.id)]
    audits = db.scalar(select(func.count()).select_from(AuditLog))
    second = copy_general_sections(db, sample['strategy'], base, target); db.commit()
    assert second['copied'] == 0 and second['conflicts'] == result['conflicts']
    assert [(r.id, r.content_json) for r in records(db, target.id)] == before
    assert db.scalar(select(func.count()).select_from(AuditLog)) == audits
    audit = db.scalar(select(AuditLog).where(AuditLog.action == 'section_period_copy'))
    assert audit.user_id == sample['strategy'].id
    # An edit that renames a copied section must still retain its source link.
    copied = next(r for r in records(db, target.id) if r.title == 'Stratejik Amaçlar')
    copied.title = 'Elle değişen başlık'; db.commit()
    result = copy_general_sections(db, sample['strategy'], base, target); db.commit()
    assert result['copied'] == 0 and len(records(db, target.id)) == 3
    assert any(copied.id in c['target_ids'] for c in result['conflicts'])


def test_whole_period_copy_second_run_changes_nothing(db, sample):
    target = create_period_record(db, sample['strategy'], {'year': '2028', 'base_period_id': sample['period'].id})
    from app.models import SubProgram
    new_sub = db.scalar(select(SubProgram).where(SubProgram.period_id == target.id))
    from app.db import Base
    def counts():
        return {name: db.scalar(select(func.count()).select_from(table)) for name, table in Base.metadata.tables.items()}
    before = counts()
    assert copy_period_work(db, sample['strategy'], sample['period'], target, {sample['subprogram'].id: new_sub.id}) == []
    db.commit()
    assert counts() == before


@pytest.mark.parametrize('financial', [False, True])
def test_word_general_sections_preview_apply_revert(db, import_client, import_context, financial):
    pid = import_context.period.id
    preview = scan(import_client, pid, section_document(financial=financial))
    assert preview['can_apply'] and len(preview['general_sections']) == 3
    assert records(db, pid) == []
    assert '5393' in preview['general_sections'][0]['content']['text']
    assert 'İkinci satır' in preview['general_sections'][0]['content']['text']
    assert preview['general_sections'][-1]['title'] == 'Performans Bilgileri — Diğer Hususlar'
    assert not db.scalar(select(Activity.id))
    result = import_client.post(f'/api/import-batches/{preview["batch_id"]}/apply')
    assert result.status_code == 200, result.text
    assert result.json()['sections_created'] == 3
    actual = records(db, pid)
    assert len(actual) == 3
    assert json.loads(actual[-1].content_json)['placement'] == ('after' if financial else 'before')
    # Another manual section is independent of the imported records and survives revert.
    manual = ReportSection(period_id=pid, title='Özel Not', content_json='{"text":"Kalacak"}')
    db.add(manual); db.commit()
    detail = import_client.get(f'/api/import-batches/{preview["batch_id"]}').json()
    assert detail['can_revert']
    response = import_client.post(f'/api/import-batches/{preview["batch_id"]}/revert')
    assert response.status_code == 200, response.text
    assert response.json()['sections_removed'] == 3
    assert [r.id for r in records(db, pid)] == [manual.id]
    assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'data_import_apply')).all()) == 1
    assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'data_import_revert')).all()) == 1


def test_import_keeps_manual_conflict_and_revert_refuses_edited_import(db, import_client, import_context):
    pid = import_context.period.id
    edited = ReportSection(period_id=pid, title='Yetki, Görev ve Sorumluluklar',
                           content_json='{"text":"Elle düzenlenen hedef"}')
    db.add(edited); db.commit()
    preview = scan(import_client, pid)
    assert preview['general_sections'][0]['action'] == 'conflict'
    assert preview['general_sections'][0]['target_ids'] == [edited.id]
    applied = import_client.post(f'/api/import-batches/{preview["batch_id"]}/apply').json()
    assert applied['sections_created'] == 2 and len(applied['section_conflicts']) == 1
    assert edited.content_json == '{"text":"Elle düzenlenen hedef"}'
    created = next(r for r in records(db, pid) if r.id != edited.id)
    created.content_json = '{"text":"İçe aktarma sonrasında düzeltme"}'; db.commit()
    result = import_client.post(f'/api/import-batches/{preview["batch_id"]}/revert')
    assert result.status_code == 409
    assert 'veriler değiştiği' in result.json()['detail']
    assert len(records(db, pid)) == 3


def test_equivalent_section_batches_serialize_and_same_batch_applies_once(db, import_client, import_context):
    from queue import Queue
    previews = [scan(import_client, import_context.period.id, section_document(financial=False)) for _ in range(2)]
    bind = db.get_bind(); pids = Queue()
    def database():
        with Session(bind, expire_on_commit=False) as session:
            session.execute(text("SET LOCAL lock_timeout = '10s'"))
            session.execute(text("SET LOCAL statement_timeout = '15s'"))
            pids.put(session.scalar(text('SELECT pg_backend_pid()')))
            yield session
    previous = app.dependency_overrides[get_db]
    app.dependency_overrides[get_db] = database
    db.scalar(select(Organization).where(Organization.id == import_context.org.id).with_for_update())
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(import_client.post, f'/api/import-batches/{p["batch_id"]}/apply') for p in previews]
            try:
                assert_waiting_for_institution_lock(db, [pids.get(timeout=5) for _ in futures])
            finally:
                db.commit()
            responses = [f.result(timeout=25) for f in futures]
        assert [r.status_code for r in responses] == [200, 200]
        assert sorted(r.json()['sections_created'] for r in responses) == [0, 3]
    finally:
        app.dependency_overrides[get_db] = previous
    assert len(records(db, import_context.period.id)) == 3
    assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'data_import_apply')).all()) == 2
    # Two callers with the same cached preview cannot apply it twice.
    fresh = scan(import_client, import_context.period.id, section_document(financial=False))
    barrier = threading.Barrier(2, timeout=10)
    def same_database():
        with Session(bind, expire_on_commit=False) as session:
            session.execute(text("SET LOCAL lock_timeout = '10s'"))
            session.execute(text("SET LOCAL statement_timeout = '15s'"))
            assert session.get(ImportBatch, fresh['batch_id']).status == 'preview'
            barrier.wait()
            yield session
    app.dependency_overrides[get_db] = same_database
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(import_client.post, f'/api/import-batches/{fresh["batch_id"]}/apply') for _ in range(2)]
            responses = [f.result(timeout=25) for f in futures]
        assert sorted(r.status_code for r in responses) == [200, 409]
    finally:
        app.dependency_overrides[get_db] = previous
    db.expire_all()
    assert len(records(db, import_context.period.id)) == 3
    assert len(db.scalars(select(AuditLog).where(AuditLog.action == 'data_import_apply', AuditLog.entity_id == fresh['batch_id'])).all()) == 1


def test_manifest_full_partial_empty_and_live_edits(db, sample):
    pid = sample['period'].id
    for _, title, _ in SECTION_TYPES:
        db.add(ReportSection(period_id=pid, title=title, kind='text', content_json='{"text":"Kaynak içerik"}'))
    for _, title, _ in TABLE_TYPES:
        db.add(ReportSection(period_id=pid, title=title, kind='table', content_json='{"rows":[["Başlık"],["0"]]}'))
    db.commit()
    full = completion_manifest(db, pid)
    assert all(r['status'] == 'Tamam' and r['count'] > 0 for r in full['rows'])
    assert len({r['key'] for r in full['rows']}) == len(full['rows'])
    empty = ReportingPeriod(year=2029); db.add(empty); db.commit()
    result = completion_manifest(db, empty.id)
    assert all(r['status'] == 'Eksik' and r['count'] == 0 for r in result['rows'])
    partial = ReportingPeriod(year=2030); db.add(partial); db.flush()
    db.add_all([
        ReportSection(period_id=partial.id, title='Başkan Sunuşu', content_json='{"text":"Yeni yıl sunuşu"}'),
        ReportSection(period_id=partial.id, title='Teşkilat Yapısı', content_json='{"text":"   "}'),
        ReportSection(period_id=partial.id, title='Donanım Tablosu', kind='table', content_json='{"rows":[["Başlık"]]}'),
    ]); db.commit()
    result = {r['key']: r for r in completion_manifest(db, partial.id)['rows']}
    assert (result['foreword']['status'], result['foreword']['count']) == ('Tamam', 1)
    assert result['organization']['status'] == result['hardware']['status'] == 'Boş'
    assert result['authority']['status'] == result['costs']['status'] == 'Eksik'
    # No persisted status: a normal editor change immediately changes the list.
    row = db.scalar(select(ReportSection).where(ReportSection.period_id == partial.id, ReportSection.title == 'Teşkilat Yapısı'))
    row.content_json = '{"text":"Yeni teşkilat"}'; db.commit()
    assert next(r for r in completion_manifest(db, partial.id)['rows'] if r['key'] == 'organization')['status'] == 'Tamam'
    # Zero is meaningful financial data; null is not.
    economics = db.scalars(select(EconomicValue).where(EconomicValue.period_id == pid)).all()
    for v in economics: v.amount = None
    db.commit()
    assert next(r for r in completion_manifest(db, pid)['rows'] if r['key'] == 'economics')['status'] == 'Boş'
    economics[0].amount = 0; db.commit()
    assert next(r for r in completion_manifest(db, pid)['rows'] if r['key'] == 'economics')['status'] == 'Tamam'


def test_official_source_table_families_and_narrative_are_covered():
    source = Path(__file__).parent / 'fixtures/pp2027_source_no_media.docx'
    doc = Document(source)
    for number in [*range(1, 29), 159, 160, 161]:
        table = doc.tables[number - 1]
        heading = ' '.join(cell.text for row in table.rows[:3] for cell in row.cells)
        assert table_type(heading), f'Unlisted official table {number}: {heading[:150]}'
    parsed = parse_pp_import(source.read_bytes())
    sections = parsed['general_sections']
    assert any('5393' in s['content'].get('text', '') for s in sections)
    assert any('Kentsel hizmetleri' in s['content'].get('text', '') for s in sections)
    assert any(s['title'] == 'Program, Alt Program ve Faaliyet Maliyetleri' for s in sections)
    assert len([s for s in sections if s['kind'] == 'table']) >= 30
    assert all(len(s['title']) <= 300 for s in sections)


def test_copy_pairs_existing_duplicate_table_families_without_adding_rows(db, sample):
    base = sample['period']; target = ReportingPeriod(year=2028)
    db.add(target); db.flush()
    for pid in (base.id, target.id):
        for index in range(2):
            db.add(ReportSection(period_id=pid, title='Stratejik Amaçlar', kind='table', position=index,
                content_json=json.dumps({'rows': [['Başlık'], [f'Amaç {index}']]})))
    db.commit()
    result = copy_general_sections(db, sample['strategy'], base, target); db.commit()
    assert result['copied'] == 0 and result['conflicts'] == []
    assert len(records(db, target.id)) == 2


def test_visible_word_sections_keep_tracked_text_tables_and_images(db, import_client, import_context):
    from PIL import Image
    doc = Document()
    doc.add_heading('Başkan Sunuşu', level=1)
    visible = doc.add_paragraph('Görünür metin')
    wrapper = OxmlElement('w:ins'); doc.element.body.remove(visible._p)
    wrapper.append(visible._p); doc.element.body.insert(1, wrapper)
    removed = OxmlElement('w:del')
    paragraph = OxmlElement('w:p'); run = OxmlElement('w:r'); value = OxmlElement('w:t')
    value.text = 'Silinmiş kaynak metni'; run.append(value); paragraph.append(run); removed.append(paragraph)
    doc.element.body.insert(2, removed)
    png = BytesIO(); Image.new('RGB', (2, 2), 'white').save(png, format='PNG')
    doc.add_picture(png)
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ['Donanım Tablosu','Adet']): cell.text = value
    for cell, value in zip(table.rows[1].cells, ['Bilgisayar','0']): cell.text = value
    stream = BytesIO(); doc.save(stream)
    preview = scan(import_client, import_context.period.id, stream.getvalue())
    assert preview['can_apply']
    assert {s['kind'] for s in preview['general_sections']} == {'text', 'image', 'table'}
    assert 'Silinmiş' not in str(preview['general_sections'])
    assert any(s['content'].get('text') == 'Görünür metin' for s in preview['general_sections'])
    result = import_client.post(f'/api/import-batches/{preview["batch_id"]}/apply')
    assert result.status_code == 200, result.text
    target = create_period_record(db, import_context.user, {'year':'2028', 'base_period_id':import_context.period.id})
    original = records(db, import_context.period.id); copied = records(db, target.id)
    assert [(s.kind,s.title,s.position,s.content_json) for s in copied] == [(s.kind,s.title,s.position,s.content_json) for s in original]


def test_parent_section_intro_without_subheadings_is_imported(db, import_client, import_context):
    source = Document()
    source.add_heading('Genel Bilgiler', level=1)
    source.add_paragraph('Ana bölümün alt başlığı olmayan genel açıklaması.')
    source.add_heading('Performans Bilgileri', level=1)
    source.add_paragraph('Performans bölümünün giriş açıklaması.')
    buffer = BytesIO(); source.save(buffer)
    preview = scan(import_client, import_context.period.id, buffer.getvalue())
    assert preview['can_apply'] and len(preview['general_sections']) == 2
    assert [s['content']['text'] for s in preview['general_sections']] == [
        'Ana bölümün alt başlığı olmayan genel açıklaması.', 'Performans bölümünün giriş açıklaması.']
    assert records(db, import_context.period.id) == []
    response = import_client.post(f'/api/import-batches/{preview["batch_id"]}/apply')
    assert response.status_code == 200, response.text
    manifest = {r['key']: r for r in completion_manifest(db, import_context.period.id)['rows']}
    assert manifest['general']['status'] == manifest['performance']['status'] == 'Tamam'
    assert import_client.post(f'/api/import-batches/{preview["batch_id"]}/revert').status_code == 200
    assert records(db, import_context.period.id) == []
