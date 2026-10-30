"""D-336 scoring uses synthetic answers; official workbook is validation input only."""
from collections import Counter
from io import BytesIO

from openpyxl import load_workbook
import pytest

from app import internal_control_questionnaire_service as q
from app.internal_control_questionnaire_export import questionnaire_export
from test_internal_control_questionnaire import (
    ROOT, questionnaire, question_file, answer_data, answer_and_approve,
)
from test_internal_control_monitoring import monitoring

pytestmark = pytest.mark.postgres


def test_real_seventy_question_file_validates_without_errors():
    # Read the lead's file only here. Never reproduce official wording in fixtures.
    questions = q.parse_questions((ROOT / 'docs/internal-control/ic-kontrol-soru-formu-70.xlsx').read_bytes())
    assert len(questions) == 70
    assert Counter(row['component'] for row in questions) == {
        'KOS': 24, 'RDS': 16, 'KFS': 12, 'BİS': 11, 'İS': 7,
    }


@pytest.fixture
def seventy_questions(db, questionnaire):
    d = questionnaire
    catalogue = q.save_list(db, d['strategy'], 'Sentetik 70 soru', question_file(count=70))
    q.attach(db, d['strategy'], d['period'].id, catalogue.id,
        [b.department_id for b in d['question_batches']], 1)
    db.commit()
    return {**d, 'catalogue': catalogue}


def test_score_47_yes_19_no_4_developing_excludes_unapproved_and_reopened(db, seventy_questions):
    d = seventy_questions
    answer_and_approve(db, d, d['question_owned'], d['user'], ['yes']*47 + ['no']*19 + ['developing']*4)
    q.save_answers(db, d['collaborator'], d['question_other'].id, 0, answer_data(*(['yes']*70)), submit=True)
    db.commit()
    result = q.results(db, d['strategy'], d['period'])
    assert len(result['directorates']) == 1
    unit = result['directorates'][0]
    assert unit['label'] == d['question_owned'].label
    assert unit['score'] == result['total']['score'] == 98
    assert unit['max_score'] == result['total']['max_score'] == 140
    assert unit['realisation_count'] == result['total']['realisation_count'] == 51
    assert unit['realisation_rate'] == result['total']['realisation_rate'] == '72,86% (51/70)'
    assert result['total']['total'] == 70
    assert set(result['pending']) == {d['question_other'].label, d['question_third'].label}

    # Institution totals sum the approved directorate-question pairs.
    q.transition(db, d['strategy'], d['question_other'].id, 'approve', 2)
    db.commit()
    result = q.results(db, d['strategy'], d['period'])
    assert len(result['directorates']) == 2
    assert result['total']['score'] == 238
    assert result['total']['max_score'] == 280
    assert result['total']['realisation_rate'] == '86,43% (121/140)'
    q.transition(db, d['strategy'], d['question_owned'].id, 'reopen', 3, 'Sentetik yeniden değerlendirme')
    db.commit()
    result = q.results(db, d['strategy'], d['period'])
    assert [r['label'] for r in result['directorates']] == [d['question_other'].label]
    assert result['total']['score'] == result['total']['max_score'] == 140
    assert result['total']['realisation_rate'] == '100,00% (70/70)'


def test_empty_approved_set_is_distinct_from_approved_zero_score(db, seventy_questions):
    d = seventy_questions
    result = q.results(db, d['strategy'], d['period'])
    assert result['directorates'] == []
    assert result['total']['score'] is None
    assert result['total']['max_score'] == 0
    assert result['total']['realisation_rate'] == '— (0/0)'
    answer_and_approve(db, d, d['question_owned'], d['user'], ['no']*70)
    result = q.results(db, d['strategy'], d['period'])
    assert result['directorates'][0]['score'] == result['total']['score'] == 0
    assert result['total']['max_score'] == 140
    assert result['total']['realisation_rate'] == '0,00% (0/70)'


@pytest.mark.parametrize('draft', [True, False])
def test_scoring_export_and_overview_only_include_current_approved_units(db, seventy_questions, client, draft):
    from app.main import app, require_user
    d = seventy_questions
    if not draft:
        q.attach(db, d['strategy'], d['period'].id, d['catalogue'].id,
            [d['question_owned'].department_id], 2)
        db.commit()
    answer_and_approve(db, d, d['question_owned'], d['user'], ['yes']*47 + ['no']*19 + ['developing']*4)
    if draft:
        q.save_answers(db, d['collaborator'], d['question_other'].id, 0, answer_data(*(['yes']*70)), submit=True)
        db.commit()
    book = load_workbook(BytesIO(questionnaire_export(db, d['strategy'], d['period'].id, draft=draft)))
    assert book.sheetnames == ['Bileşen özeti', 'Soru özeti', 'Müdürlük yanıtları']
    assert all(sheet['A1'].value == ('TASLAK' if draft else 'Onaylı veriler') for sheet in book)
    rows = list(book['Bileşen özeti'].values)
    header = ('Müdürlük', 'Evet', 'Hayır', 'Geliştirilmekte', 'Toplam puan', 'Azami puan', 'Gerçekleşme oranı')
    start = rows.index(header) + 1
    assert rows[start:] == [
        (d['question_owned'].label, 47, 19, 4, 98, 140, '72,86% (51/70)'),
        ('Kurum toplamı', 47, 19, 4, 98, 140, '72,86% (51/70)'),
    ]
    assert any(row[0] == '2022 uygulaması — Strateji teyit edecek' for row in rows)
    app.dependency_overrides[require_user] = lambda: d['strategy']
    response = client.get(f'/internal-control/monitoring?period_id={d["period"].id}')
    assert response.status_code == 200
    assert 'Müdürlük puanları' in response.text
    assert '2022 uygulaması — Strateji teyit edecek' in response.text
    assert '72,86% (51/70)' in response.text
    result = response.context['question_results']
    assert result['total']['score'] == 98 and len(result['directorates']) == 1
