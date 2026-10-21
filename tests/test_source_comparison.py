"""CIV-073: lossless Word assertions and period-scoped, read-only comparison."""
import io
import json
from decimal import Decimal

import pytest
from docx import Document
from sqlalchemy import select

from app.data_importer import parse_performance_docx
from app.models import Activity, ImportBatch
from test_word_import import word_document, import_context, import_client

pytestmark = pytest.mark.postgres


def totals_document():
    doc = Document(io.BytesIO(word_document()))
    table = doc.tables[1]
    anchor = table.rows[5]._tr
    for label, amount in [('Toplam', '1.234,50'), ('Genel Toplam', '1.334,50'), ('Genel Bütçe', '1.334,50')]:
        cells = table.add_row().cells
        for cell, value in zip(cells, (label, amount, '0', '')):
            cell.text = value
        anchor.addprevious(cells[0]._tc.getparent())
    result = io.BytesIO(); doc.save(result)
    return result.getvalue()


def test_word_totals_are_assertions_with_location_without_extra_activities():
    parsed = parse_performance_docx(totals_document())
    dept = parsed['departments'][0]
    assert len(dept['activities']) == 1
    baseline = parse_performance_docx(word_document())['departments'][0]['activities'][0]['costs']
    fields = ('year', 'kind', 'funding_type', 'amount', 'raw_text', 'state')
    assert [{k: v[k] for k in fields} for v in dept['activities'][0]['costs']] == [{k: v[k] for k in fields} for v in baseline]
    totals = dept['total_assertions']
    assert len(totals) == 9
    assert {r['label'] for r in totals} == {'Toplam', 'Genel Toplam', 'Genel Bütçe'}
    first = next(r for r in totals if r['label'] == 'Toplam' and r['year'] == 2027)
    assert (first['amount'], first['table_number'], first['row_number'], first['column_number']) == ('1234.50', 2, 6, 2)


def imported(db, context, *, data=None, filename='PP karşılaştırma.docx', catalogue=False):
    from app.import_service import parse_pp_import, apply_import_batch
    from app.serialization import json_dumps
    parsed = parse_pp_import(data or totals_document())
    batch = ImportBatch(period_id=context.period.id, filename=filename,
        payload_json=json_dumps(parsed, ensure_ascii=False))
    db.add(batch); db.flush()
    apply_import_batch(db, batch)
    db.commit()
    binding = json.loads(batch.snapshot_json)['source_comparison']['departments'][0]
    if catalogue:
        from app.models import OfficialActivity, ActivityResponsibility, ActivityCatalogueLink
        activity = db.get(Activity, binding['activity_ids'][0])
        official = OfficialActivity(period_id=context.period.id, program_code='1', program_name='Çevre Programı',
            subprogram_code='1.1', subprogram_name='Yeşil Alanlar', code='1.1.1', name=activity.name)
        db.add(official); db.flush()
        db.add_all([ActivityResponsibility(official_activity_id=official.id, department_id=activity.department_id),
                    ActivityCatalogueLink(official_activity_id=official.id, activity_id=activity.id)])
        db.commit()
    return batch, binding


def comparison(db, context, **kwargs):
    from app.source_comparison_service import source_comparison
    return source_comparison(db, context.period.id, **kwargs)


def test_apply_stores_totals_and_compares_shared_current_subtotal(db, import_context):
    from app.models import ActivityCostValue
    batch, binding = imported(db, import_context, catalogue=True)
    totals = binding['source']['total_assertions']
    assert len(totals) == 9 and all(v['column_number'] for v in totals)
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'total' and r['year'] == 2027]
    assert len(rows) == 3
    assert [r['source_amount'] for r in rows] == ['1.234,50', '1.334,50', '1.334,50']
    assert all(r['status'] == 'Eşleşiyor' for r in rows)
    assert all((r['document'], r['table'], r['column'], r['page']) == (batch.filename, 2, 2, 'Bilinmiyor') for r in rows)
    assert all(len(r['document_sha256']) == 64 for r in rows)
    # Both activity financing rows move; a source assertion is never recalculated.
    for value in db.scalars(select(ActivityCostValue).where(ActivityCostValue.activity_id == binding['activity_ids'][0], ActivityCostValue.year == 2027)):
        if value.funding_type in {'total', 'inside'}:
            value.amount += Decimal('10')
    db.commit()
    changed = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'total' and r['year'] == 2027]
    assert all(r['status'] == 'Fark var' and r['difference'] == '10,00' for r in changed)


@pytest.mark.parametrize('case,expected', [('match', 'Eşleşiyor'), ('changed', 'Fark var'),
    ('missing_store', 'Kayıtta yok'), ('null_store', 'Kayıtta yok'), ('missing_source', 'Kaynakta yok')])
def test_statuses_and_missing_values(db, import_context, case, expected):
    from app.models import IndicatorValue
    _, binding = imported(db, import_context)
    iid = binding['indicator_ids'][0][0]
    current = db.scalar(select(IndicatorValue).where(IndicatorValue.indicator_id == iid, IndicatorValue.year == 2027))
    if case == 'changed': current.value = Decimal('99')
    if case == 'missing_store': db.delete(current)
    if case == 'null_store': current.value = None
    if case == 'missing_source':
        db.add(IndicatorValue(indicator_id=iid, year=2030, value_kind='forecast', value=Decimal('0')))
    db.commit()
    year = 2030 if case == 'missing_source' else 2027
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'indicator' and r['year'] == year)
    assert row['status'] == expected
    if case == 'missing_source': assert row['source_amount'] == '—' and row['stored_amount'] == '0'


def test_recorded_decision_and_stale_decision(db, import_context):
    from app.models import ActivityCostValue
    from app.program_service import save_activity_costs
    _, binding = imported(db, import_context)
    activity = db.get(Activity, binding['activity_ids'][0])
    save_activity_costs(db, activity, 2027, {'cost_2027_budget_inside':'1.210,00',
        'cost_2027_budget_outside':'34,50', 'cost_2027_budget_use_computed':True,
        'cost_2027_budget_reason':'Finansman dağılımı kontrol edildi.'}, actor=import_context.user)
    db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027 and r['funding'] == 'Toplam')
    assert row['status'] == 'Kararla değişti' and 'Finansman dağılımı kontrol edildi.' in row['decision']
    assert import_context.user.full_name in row['decision']
    total = db.scalar(select(ActivityCostValue).where(ActivityCostValue.activity_id == activity.id,
        ActivityCostValue.year == 2027, ActivityCostValue.funding_type == 'total'))
    total.amount += Decimal('1'); db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027 and r['funding'] == 'Toplam')
    assert row['status'] == 'Fark var'


def test_import_conflict_keeps_both_locations_and_reason(db, import_context):
    from app.import_service import decide_import_conflict, apply_import_batch, parse_pp_import, build_department_plan
    from app.serialization import json_dumps
    from test_source_conflicts import duplicated_personnel
    parsed = parse_pp_import(duplicated_personnel())
    batch = ImportBatch(period_id=import_context.period.id, filename='İki kaynak.docx', payload_json=json_dumps(parsed))
    db.add(batch); db.flush()
    conflict = build_department_plan(db, parsed)['conflicts'][0]
    decide_import_conflict(db, batch, conflict['id'], '20,00', 'İkinci kaynak tutarı kontrol edildi.', import_context.user)
    apply_import_batch(db, batch); db.commit()
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'personnel' and r['year'] == 2027]
    assert {r['table'] for r in rows} == {2, 3}
    assert {r['status'] for r in rows} == {'Eşleşiyor', 'Kararla değişti'}
    assert all(r['column'] == 2 and 'İkinci kaynak tutarı kontrol edildi.' in r['decision'] for r in rows)


def test_revert_removes_only_its_active_total_assertions(db, import_context):
    from app.import_service import revert_import_batch
    first, _ = imported(db, import_context, filename='Birinci.docx')
    data = totals_document()
    doc = Document(io.BytesIO(data)); doc.tables[0].cell(2, 2).text = 'Örnek Belediyesi / Fen İşleri Müdürlüğü'
    stream = io.BytesIO(); doc.save(stream)
    second, binding = imported(db, import_context, data=stream.getvalue(), filename='İkinci.docx')
    second_snapshot = second.snapshot_json
    before = comparison(db, import_context)['rows']
    assert len([r for r in before if r['model'] == 'total']) == 18
    revert_import_batch(db, first); db.commit()
    rows = comparison(db, import_context)['rows']
    assert len([r for r in rows if r['model'] == 'total']) == 9
    assert all(r['batch_id'] == second.id for r in rows)
    assert second.snapshot_json == second_snapshot
    assert db.get(Activity, binding['activity_ids'][0]) is not None


def test_old_batch_unknown_location_and_total_notice(db, import_context):
    batch, _ = imported(db, import_context)
    snapshot = json.loads(batch.snapshot_json); snapshot.pop('source_comparison')
    payload = json.loads(batch.payload_json); payload.pop('document_sha256'); payload.pop('source_assertion_version')
    for source in payload['departments']:
        source.pop('total_assertions', None)
        for activity in source['activities']:
            for value in activity['costs']:
                value.pop('column_number', None)
    batch.snapshot_json = json.dumps(snapshot); batch.payload_json = json.dumps(payload)
    db.commit()
    result = comparison(db, import_context)
    assert any('Kaynak toplamı bu içe aktarmada kaydedilmedi' in message for message in result['notices'])
    row = next(r for r in result['rows'] if r['model'] == 'activity')
    assert row['column'] == row['page'] == row['document_sha256'] == 'Bilinmiyor'
    assert row['status'] == 'Eşleşiyor'


def test_department_scope_period_isolation_csv_safety_and_read_only(db, import_context, import_client):
    import csv
    from app.main import app, require_user
    from app.models import DepartmentScope, ReportingPeriod, User, AuditLog
    batch, binding = imported(db, import_context, filename='=TOPLA(1;2).docx')
    other, _ = imported(db, import_context, data=word_document(department='Fen İşleri Müdürlüğü'), filename='Gizli.docx')
    user = User(email='own@example.invalid', full_name='Müdürlük Kullanıcısı', role='department', password_hash='unused', department_id=binding['department_id'])
    closed = ReportingPeriod(year=2028)
    db.add_all([user, closed]); db.flush()
    db.add(DepartmentScope(period_id=closed.id, department_id=user.department_id, status='excluded'))
    db.commit()
    app.dependency_overrides[require_user] = lambda: user
    path = f'/source-comparison?period_id={batch.period_id}'
    audit_count = len(db.scalars(select(AuditLog)).all())
    before = batch.snapshot_json
    page = import_client.get(path)
    assert page.status_code == 200 and 'Gizli.docx' not in page.text
    assert 'Fen İşleri' not in page.text and 'PP Kaynak Karşılaştırması' in page.text
    response = import_client.get(f'/source-comparison.csv?period_id={batch.period_id}')
    assert response.status_code == 200 and 'Gizli.docx' not in response.text
    exported = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig')), delimiter=';'))
    expected = comparison(db, import_context, department_id=user.department_id)['rows']
    assert len(exported) == len(expected) + 1
    assert all(row[5].startswith("'=TOPLA") for row in exported[1:])
    assert [r[16] for r in exported[1:]] == [r['status'] for r in expected]
    assert import_client.get(path + '&department_id=9999').status_code == 403
    assert import_client.get(f'/source-comparison.csv?period_id={closed.id}').status_code == 403
    assert import_client.get('/source-comparison').status_code == 422
    assert batch.snapshot_json == before and len(db.scalars(select(AuditLog)).all()) == audit_count
    app.dependency_overrides[require_user] = lambda: import_context.user
    assert 'Gizli.docx' in import_client.get(path).text
    assert import_client.get(f'/source-comparison?period_id={closed.id}').status_code == 200
    assert not comparison(db, type('Context', (), {'period': closed})())['rows']


def test_economic_cells_and_totals_keep_exact_sources(db, import_context):
    from app.models import EconomicValue
    from test_pp_import_submit_ready import source_document
    _, binding = imported(db, import_context, data=source_document())
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'economic' and r['batch_id']]
    assert len(rows) == 30  # Three category cells and three totals for each of five columns.
    category = next(r for r in rows if r['record'].endswith('Mal ve Hizmet Alım Giderleri') and r['year'] == 2027)
    assert (category['table'], category['row'], category['column'], category['source_amount'], category['status']) == (3, 4, 4, '100,00', 'Eşleşiyor')
    assert all(r['status'] == 'Eşleşiyor' for r in rows)
    current = db.scalar(select(EconomicValue).where(EconomicValue.department_id == binding['department_id'],
        EconomicValue.year == 2027, EconomicValue.category == 'goods_services', EconomicValue.funding_type == 'inside'))
    current.amount += Decimal('1'); db.commit()
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'economic' and r['year'] == 2027]
    changed = [r for r in rows if r['status'] == 'Fark var']
    assert len(changed) == 3  # Category, inside total and grand total all expose the same edit.


def test_institutional_and_unmapped_cells_are_visible_without_double_counting(db, import_context):
    from app.models import ReportSection
    doc = Document(io.BytesIO(totals_document()))
    table = doc.add_table(rows=2, cols=2)
    for cell, value in zip(table.rows[0].cells, ['Donanım Tablosu', 'Adet']): cell.text = value
    for cell, value in zip(table.rows[1].cells, ['Bilgisayar', '23']): cell.text = value
    unmapped = doc.add_table(rows=2, cols=2)
    for cell, value in zip(unmapped.rows[0].cells, ['Sınıflandırılmamış çizelge', 'Tutar']): cell.text = value
    for cell, value in zip(unmapped.rows[1].cells, ['Kaynak satırı', '456,70']): cell.text = value
    stream = io.BytesIO(); doc.save(stream)
    batch, _ = imported(db, import_context, data=stream.getvalue())
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] in {'section', 'unmapped'}]
    assert len(rows) == 2
    assert {(r['raw_text'], r['status']) for r in rows} == {('23', 'Eşleşiyor'), ('456,70', 'Kayıtta yok')}
    section = db.scalar(select(ReportSection).where(ReportSection.period_id == batch.period_id))
    content = json.loads(section.content_json); content['rows'][1][1] = '25'; section.content_json = json.dumps(content)
    db.commit()
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'section']
    assert len(rows) == 1 and rows[0]['status'] == 'Fark var' and rows[0]['difference'] == '2'


def test_known_page_and_rename_preserve_identity_missing_owner_does_not_match_name(db, import_context):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(io.BytesIO(totals_document()))
    caption = OxmlElement('w:tblCaption'); caption.set(qn('w:val'), '{"page": "44"}')
    doc.tables[1]._tbl.tblPr.append(caption)
    stream = io.BytesIO(); doc.save(stream)
    batch, binding = imported(db, import_context, data=stream.getvalue())
    activity = db.get(Activity, binding['activity_ids'][0]); activity.name = 'Güncel faaliyet adı'; db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027)
    assert row['page'] == '44' and row['status'] == 'Eşleşiyor'
    assert row['record'] == 'Faaliyet: Park bakımı'  # Original source identity survives a rename.
    db.delete(activity); db.flush()
    replacement = Activity(period_id=batch.period_id, department_id=binding['department_id'], name='Park bakımı')
    db.add(replacement); db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027)
    assert row['status'] == 'Kayıtta yok'


def test_blank_and_zero_never_match(db, import_context):
    from app.models import ActivityCostValue
    _, binding = imported(db, import_context)
    aid = binding['activity_ids'][0]
    db.add(ActivityCostValue(activity_id=aid, year=2029, value_kind='forecast', funding_type='total', amount=Decimal('0')))
    db.commit()
    result = comparison(db, import_context)['rows']
    zero = next(r for r in result if r['model'] == 'activity' and r['year'] == 2028 and r['funding'] == 'Toplam')
    blank = next(r for r in result if r['model'] == 'activity' and r['year'] == 2029 and r['funding'] == 'Toplam')
    assert zero['status'] == 'Eşleşiyor' and zero['source_amount'] == zero['stored_amount'] == '0,00'
    assert blank['status'] == 'Kaynakta yok' and blank['source_amount'] == '—' and blank['stored_amount'] == '0,00'


def test_total_traces_show_current_additions_and_valid_source_decision(db, import_context):
    from app.models import ActivityCostValue, ActivityCatalogueLink
    from app.program_service import save_activity_costs
    _, binding = imported(db, import_context, catalogue=True)
    activity = db.get(Activity, binding['activity_ids'][0])
    save_activity_costs(db, activity, 2027, {'cost_2027_budget_inside':'1.210,00',
        'cost_2027_budget_outside':'34,50', 'cost_2027_budget_use_computed':True,
        'cost_2027_budget_reason':'Toplam kaynak kararı doğrulandı.'}, actor=import_context.user)
    db.commit()
    totals = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'total' and r['year'] == 2027]
    assert all(r['status'] == 'Kararla değişti' and 'Toplam kaynak kararı doğrulandı.' in r['decision'] for r in totals)
    second = Activity(period_id=activity.period_id, department_id=activity.department_id,
        card_id=activity.card_id, subprogram_id=activity.subprogram_id, name='Yeni güncel faaliyet')
    db.add(second); db.flush()
    official_id = db.scalar(select(ActivityCatalogueLink.official_activity_id).where(ActivityCatalogueLink.activity_id == activity.id))
    db.add(ActivityCatalogueLink(activity_id=second.id, official_activity_id=official_id))
    for year, kind in [(2027, 'budget'), (2028, 'forecast'), (2029, 'forecast')]:
        for funding, amount in [('inside', '5'), ('outside', '0'), ('total', '5')]:
            db.add(ActivityCostValue(activity_id=second.id, year=year, value_kind=kind, funding_type=funding, amount=Decimal(amount)))
    db.commit()
    result = comparison(db, import_context)['rows']
    totals = [r for r in result if r['model'] == 'total' and r['year'] == 2027]
    assert all(r['status'] == 'Fark var' and r['difference'] == '15,00' for r in totals)
    assert all(r['status'] == 'Kaynakta yok' for r in result if r['record'] == 'Faaliyet: Yeni güncel faaliyet')


def test_historical_total_uses_same_calculator_and_unknown_scope_is_not_guessed(db, import_context):
    from test_pp_import_submit_ready import source_document
    doc = Document(io.BytesIO(source_document()))
    costs = doc.tables[1]
    for col in (4, 5, 6):
        for row, amount in [(1, '40'), (2, '40'), (3, '0')]: costs.cell(row, col).text = amount
    anchor = costs.rows[5]._tr
    for label in ('Toplam', 'Toplam özel dağılım'):
        cells = costs.add_row().cells
        for col, cell in enumerate(cells): cell.text = label if col == 0 else '40' if col in (4, 5, 6) else '100'
        anchor.addprevious(cells[0]._tc.getparent())
    stream = io.BytesIO(); doc.save(stream)
    imported(db, import_context, data=stream.getvalue(), catalogue=True)
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'total' and r['year'] == 2026]
    assert len(rows) == 6
    assert all(r['status'] == 'Eşleşiyor' and r['stored_amount'] == '40,00' for r in rows if r['record'] == 'Kaynak toplamı: Toplam')
    assert all(r['status'] == 'Fark var' and r['stored_amount'] == '—' and 'Bilinmiyor' in r['note'] for r in rows if r['record'] != 'Kaynak toplamı: Toplam')


def test_historical_preview_decisions_without_columns_still_apply_and_compare(db, import_context):
    from app.import_service import apply_import_batch, build_department_plan
    from app.serialization import json_dumps
    from test_source_conflicts import duplicated_personnel
    parsed = parse_performance_docx(duplicated_personnel())
    parsed.pop('source_assertion_version'); parsed.pop('document_sha256')
    for dept in parsed['departments']:
        for value in dept['personnel']:
            value.pop('column_number', None)
        for activity in dept['activities']:
            for value in activity['costs']:
                value.pop('column_number', None)
    conflict = build_department_plan(db, parsed)['conflicts'][0]
    candidates = conflict['candidates']
    assert all('column_number' not in location for c in candidates for location in c['locations'])
    decision = {'conflict_id':conflict['id'], 'actor':import_context.user.id, 'time':'2026-10-01T12:00:00Z',
        'department':conflict['department'], 'year':2027, 'kind':'budget', 'funding_type':'total',
        'chosen_amount':candidates[0]['amount'], 'candidates':candidates, 'reason':'Eski önizleme kaynak seçimi doğrulandı.'}
    batch = ImportBatch(period_id=import_context.period.id, filename='Eski önizleme.docx',
        payload_json=json_dumps(parsed), decisions_json=json_dumps([decision]))
    db.add(batch); db.flush(); apply_import_batch(db, batch); db.commit()
    rows = [r for r in comparison(db, import_context)['rows'] if r['model'] == 'personnel' and r['year'] == 2027]
    assert {r['status'] for r in rows} == {'Eşleşiyor', 'Kararla değişti'}
    assert all(r['column'] == 'Bilinmiyor' and 'Eski önizleme kaynak seçimi doğrulandı.' in r['decision'] for r in rows)


def test_current_record_without_any_values_stays_visible(db, import_context):
    _, binding = imported(db, import_context)
    activity = Activity(period_id=import_context.period.id, department_id=binding['department_id'], name='Henüz tutarı girilmemiş faaliyet')
    db.add(activity); db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['record'] == 'Faaliyet: Henüz tutarı girilmemiş faaliyet')
    assert row['status'] == 'Kaynakta yok' and row['stored_amount'] == '—'
    assert 'henüz girilmemiş' in row['note']


def test_existing_budget_decision_explains_only_its_recorded_amount(db, import_context):
    from app.models import ActivityCostValue, BudgetOperation, BudgetOperationLine
    _, binding = imported(db, import_context)
    current = db.scalar(select(ActivityCostValue).where(ActivityCostValue.activity_id == binding['activity_ids'][0],
        ActivityCostValue.year == 2027, ActivityCostValue.funding_type == 'total'))
    original = current.amount
    operation = BudgetOperation(organization_id=import_context.org.id, period_id=import_context.period.id,
        actor_id=import_context.user.id, year=2027, value_kind='budget', phase='pre_publication',
        reference_type='mayor_instruction', reference_date='2026-10-21', reference_number='K-73',
        reason='Kayıtlı bütçe aktarma gerekçesi.', request_json='{}', preview_json='{}')
    db.add(operation); db.flush()
    db.add(BudgetOperationLine(operation_id=operation.id, department_id=binding['department_id'],
        value_table='activity_cost_values', value_id=current.id, label='Park bakımı', funding_type='total',
        before=original, after=original+Decimal('10'), before_json='{}', after_json='{}'))
    current.amount += Decimal('10'); db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027 and r['funding'] == 'Toplam')
    assert row['status'] == 'Kararla değişti' and 'K-73' in row['decision'] and operation.reason in row['decision']
    current.amount += Decimal('1'); db.commit()
    row = next(r for r in comparison(db, import_context)['rows'] if r['model'] == 'activity' and r['year'] == 2027 and r['funding'] == 'Toplam')
    assert row['status'] == 'Fark var' and 'güncel tutarı açıklamıyor' in row['note']


def test_preserved_existing_import_also_compares_every_source_without_changing_values(db, import_context):
    from app.import_service import parse_pp_import, apply_import_batch, _serialize_department_state
    from app.serialization import json_dumps
    _, binding = imported(db, import_context)
    before = _serialize_department_state(db, binding['department_id'], import_context.period.id)
    doc = Document(io.BytesIO(totals_document())); doc.tables[1].cell(1, 1).text = '2.000,00'
    stream = io.BytesIO(); doc.save(stream)
    batch = ImportBatch(period_id=import_context.period.id, filename='Korunan kayıt kaynağı.docx',
        payload_json=json_dumps(parse_pp_import(stream.getvalue())))
    db.add(batch); db.flush()
    result = apply_import_batch(db, batch, preserve_existing=True); db.commit()
    assert result['skipped_existing_departments'] == 1 and result['activities'] == 0
    assert _serialize_department_state(db, binding['department_id'], import_context.period.id) == before
    rows = [r for r in comparison(db, import_context)['rows'] if r['batch_id'] == batch.id]
    assert len(rows) == 24
    changed = next(r for r in rows if r['model'] == 'activity' and r['year'] == 2027 and r['funding'] == 'Toplam')
    assert (changed['source_amount'], changed['stored_amount'], changed['status']) == ('2.000,00', '1.234,50', 'Fark var')
    assert 'Mevcut kayıtlar korunarak' in changed['note']


def test_preserved_economic_values_match_business_scope_when_document_tables_move(db, import_context):
    from app.import_service import parse_pp_import, apply_import_batch
    from app.serialization import json_dumps
    from test_pp_import_submit_ready import source_document
    imported(db, import_context, data=source_document())
    doc = Document(io.BytesIO(source_document()))
    note = doc.add_table(rows=1, cols=1); note.cell(0, 0).text = 'Bağımsız not çizelgesi'
    doc.tables[0]._tbl.addprevious(note._tbl)
    stream = io.BytesIO(); doc.save(stream)
    batch = ImportBatch(period_id=import_context.period.id, filename='Yeni tablo sırası.docx',
        payload_json=json_dumps(parse_pp_import(stream.getvalue())))
    db.add(batch); db.flush(); apply_import_batch(db, batch, preserve_existing=True); db.commit()
    rows = [r for r in comparison(db, import_context)['rows'] if r['batch_id'] == batch.id and r['model'] == 'economic']
    assert len(rows) == 30 and all(r['status'] == 'Eşleşiyor' and r['table'] == 4 for r in rows)
