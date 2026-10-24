"""CIV-084: official reports use synthetic plans and disposable PostgreSQL only."""
from datetime import date, timedelta
from io import BytesIO
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from uuid import uuid4

from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import internal_control_service as ic
from app import internal_control_monitoring_service as m
from app import internal_control_monitoring_export as export
from app.internal_control_models import ICAction
from app.internal_control_monitoring_models import ICMonitoringBatch as Batch, ICMonitoringEntry as Entry
from app.models import User
from test_internal_control import workbook, action_row
from test_internal_control_monitoring import monitoring, values, approve, fill

pytestmark = pytest.mark.postgres
ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'tests/artifacts/civ084'


@pytest.fixture
def official_plan(db, sample, request):
    source = []
    statuses = {}
    index = 0
    for standard, count in [('KOS1', 11), ('RDS5', 8), ('KFS7', 9), ('BİS13', 2)]:
        for n in range(1, count + 1):
            code = f'{standard}.{n}.1'
            source.append(action_row(**{'0': standard.rstrip('0123456789'), '1': standard,
                '2': f'{standard}.{n}', '5': code, '6': f'Sentetik eylem {index + 1}',
                '10': '01.01.2027', '11': '30.06.2027' if index < 25 else '31.12.2027'}))
            statuses[code] = 'completed' if index < 5 or 11 <= index < 15 else 'ongoing' if index < 25 else 'not_started'
            index += 1
    if getattr(request, 'param', '') == 'warning':
        source[-1][11] = 'Plan sonunda'
    assurance = action_row(**{'2': 'KOS1.99', '3': 'Makul güvence şartı',
        '4': 'Mevcut durumda makul güvence sağlanmıştır. Yeni bir eylem belirlemeye ihtiyaç yoktur'})
    assurance[5:13] = [''] * 8
    source.insert(0, assurance)
    batch, report = ic.create_import(db, sample['strategy'], workbook(source), 'civ084-sentetik.xlsx')
    assert report['can_commit'], report['issues']
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    ic.activate_plan(db, sample['strategy'], plan.id)
    db.commit()
    period = m.open_period(db, sample['strategy'], plan.id, 2027, 1, m.today(db) + timedelta(days=1000))
    db.commit()
    owned = db.scalar(select(Batch).where(Batch.period_id == period.id))
    for row in m.rows(db, owned):
        status = statuses[row['action'].action_code_raw]
        m.save_entry(db, sample['user'], owned.id, row['target'].id, m.state(db, owned.id)[1],
                     values(status=status, percentage='100' if status == 'completed' else '25' if status == 'ongoing' else ''))
    m.transition(db, sample['user'], owned.id, 'submit', m.state(db, owned.id)[1])
    db.commit()
    m.transition(db, sample['strategy'], owned.id, 'approve', m.state(db, owned.id)[1])
    db.commit()
    return {**sample, 'plan': plan, 'period': period, 'owned': owned}


def test_synthetic_rates_official_columns_and_assurance(db, official_plan):
    d = official_plan
    metrics = m.official_metrics(db, d['strategy'], d['period'])
    assert (metrics['plan_total'], metrics['due_total'], metrics['completed']) == (30, 25, 9)
    assert metrics['period_rate'] == '36,00% (9/25)'
    assert metrics['plan_rate'] == '30,00% (9/30)'
    assert metrics['components'][0] == ['Kontrol Ortamı', 11, 5, 0, 6, 5, '16,67% (5/30)']
    assert metrics['components'][-1][1:6] == [0, 0, 0, 0, 0]
    assert [sum(row[column] for row in metrics['components']) for column in range(1, 6)] == [30, 9, 0, 16, 9]
    book = load_workbook(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'xlsx')))
    rows = list(book['Resmî İzleme'].values)
    assert list(rows[0]) == export.OFFICIAL_HEADERS and len(rows) == 32
    assert rows[1][3] is None and rows[1][2].startswith('Mevcut durumda makul güvence')
    assert rows[2][3] == 'KOS1.1.1' and rows[2][-1].startswith('(TAMAMLANDI)')
    assert any(row == tuple(export.COMPONENT_HEADERS) for row in book['Özet'].values)
    assert any('36,00% (9/25)' in str(row) for row in book['Özet'].values)
    doc = Document(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'docx')))
    assert [c.text for c in doc.tables[0].rows[0].cells] == export.COMPONENT_HEADERS
    assert [c.text for c in doc.tables[1].rows[0].cells] == export.OFFICIAL_HEADERS
    assert doc.tables[1].rows[1].cells[3].text == ''
    assert '30,00% (9/30)' in '\n'.join(p.text for p in doc.paragraphs)

    # Reopening removes approval while keeping the period's draft entries.
    m.transition(db, d['strategy'], d['owned'].id, 'reopen', m.state(db, d['owned'].id)[1], 'Dönem dışı eylemlerin kontrolü')
    db.commit()
    not_due_codes = {r['action'].action_code_raw for r in m.action_records(db, d['period']) if r['due'] is False}
    assert len(not_due_codes) == 5
    phrase = 'Eylem süresi bu dönemde dolmamaktadır.'
    draft_book = load_workbook(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'xlsx', draft=True)))
    not_due_rows = [row for row in list(draft_book['Resmî İzleme'].values)[1:] if row[3] in not_due_codes]
    assert len(not_due_rows) == 5 and all(row[-1].startswith(phrase) for row in not_due_rows)
    draft_doc = Document(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'docx', draft=True)))
    not_due_rows = [row for row in draft_doc.tables[1].rows[1:] if row.cells[3].text in not_due_codes]
    assert len(not_due_rows) == 5 and all(row.cells[-1].text.startswith(phrase) for row in not_due_rows)


@pytest.mark.parametrize('half,due,start,continuous,expected', [
    (1, date(2027, 6, 30), '', False, True), (1, date(2027, 7, 1), '', False, False),
    (2, date(2027, 12, 31), '', False, True), (2, date(2028, 1, 1), '', False, False),
    (1, date(2026, 1, 1), '', False, True), (1, None, '', False, None),
    (1, None, '2027', True, True), (1, None, '2028', True, False),
    (1, None, '30.06.2027', True, True), (1, None, '01.07.2027', True, False),
    (1, None, '31/12/2026', True, True), (1, None, 'belirsiz', True, None),
])
def test_due_in_period_boundaries(half, due, start, continuous, expected):
    assert m.due_in_period(SimpleNamespace(due_date=due, start_raw=start, is_continuous=continuous),
                           SimpleNamespace(year=2027, half=half)) is expected


@pytest.mark.parametrize('statuses,expected', [(['completed', 'ongoing'], 'ongoing'),
    (['completed', 'completed'], 'completed'), (['not_approved', 'not_approved'], 'not_approved'),
    (['not_approved', 'not_started'], 'not_started'), (['partial', None], 'ongoing'),
    (['delayed', 'not_approved'], 'ongoing'), (['completed', None], 'ongoing'), ([], 'not_started')])
def test_action_status_rules(statuses, expected):
    assert m.action_status([SimpleNamespace(status=s) if s else None for s in statuses]) == expected


def test_two_units_and_reopened_approval_not_counted(db, monitoring):
    d = monitoring
    approve(db, d, d['owned'], status='completed', percentage='100')
    approve(db, d, d['coop'])
    records = m.action_records(db, d['period'])
    shared = next(r for r in records if r['action'].action_code_raw == 'KOS1.2.1')
    assert len(shared['units']) == 2 and shared['status'] == 'ongoing'
    m.transition(db, d['strategy'], d['owned'].id, 'reopen', m.state(db, d['owned'].id)[1], 'Kanıt kontrolü')
    db.commit()
    assert all(r['status'] != 'completed' for r in m.action_records(db, d['period']))


@pytest.mark.parametrize('status,percentage,explanation,valid', [
    ('partial', '', 'Açıklama', False), ('partial', '0', 'Açıklama', False),
    ('partial', '100', 'Açıklama', False), ('partial', '1', 'Açıklama', True),
    ('partial', '99', 'Açıklama', True), ('not_approved', '', '', False),
    ('not_approved', '', '  ', False), ('not_approved', '', 'Kurul uygun görmedi.', True)])
def test_new_status_validation_and_retained_values(db, monitoring, client, status, percentage, explanation, valid):
    d = monitoring
    row = m.rows(db, d['owned'])[0]
    response = client.post(f"/internal-control/monitoring/batches/{d['owned'].id}/entry", data={
        'id': row['target'].id, 'expected_revision': 0,
        **values(status=status, percentage=percentage, explanation=explanation)}, follow_redirects=False)
    assert response.status_code == (303 if valid else 422)
    if valid:
        entry = m.rows(db, d['owned'])[0]['entry']
        assert entry.status == status and entry.percentage == (int(percentage) if percentage else None)
    else:
        assert f'value="{status}" selected' in response.text
        assert '1–99 arasında yüzde girin' in response.text if status == 'partial' else 'Açıklama alanına gerekçeyi yazın' in response.text
        assert m.state(db, d['owned'].id)[1] == 0


def test_official_draft_permission_and_status_phrases(db, monitoring, client):
    from app.main import app, require_user
    d = monitoring
    url = f"/internal-control/monitoring/periods/{d['period'].id}/official-report.xlsx"
    assert client.get(url + '?draft=true').status_code == 403
    app.dependency_overrides[require_user] = lambda: d['strategy']
    assert client.get(url).status_code == 409
    fill(db, d, d['owned'], status='partial', percentage='45')
    fill(db, d, d['body'], status='not_approved', percentage='', explanation='Kurul uygun görmedi.')
    response = client.get(url + '?draft=true')
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book['Özet']['A1'].value == 'TASLAK'
    assert any('(KISMEN TAMAMLANDI (DEVAM EDİYOR))' in str(row[-1]) for row in book['Resmî İzleme'].values)
    assert any('Eylem Uygun Görülmemiştir.' in str(row[-1]) for row in book['Resmî İzleme'].values)
    doc = Document(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'docx', draft=True)))
    assert 'TASLAK' in doc.paragraphs[0].text and 'TASLAK' in doc.sections[0].header.paragraphs[0].text
    assert client.get(url + '?draft=yanlis').status_code == 422


@pytest.mark.parametrize('official_plan', ['', 'warning'], indirect=True)
def test_warning_exclusion_cover_and_responsive_screens(db, official_plan, client):
    from playwright.sync_api import sync_playwright, expect
    from app.main import app, require_user
    d = official_plan
    metrics = m.official_metrics(db, d['strategy'], d['period'])
    has_warning = metrics['eligible_total'] == 29
    assert metrics['plan_total'] == 30
    assert metrics['plan_rate'] == ('31,03% (9/29)' if has_warning else '30,00% (9/30)')
    warning = 'Bitiş tarihi okunamayan eylemler: BİS13.2.1' if has_warning else ''
    assert metrics['warning'] == warning
    for fmt in ('docx', 'xlsx') if has_warning else ():
        content = export.official_export(db, d['strategy'], d['period'].id, fmt)
        if fmt == 'docx':
            assert warning in [p.text for p in Document(BytesIO(content)).paragraphs]
        else:
            assert (warning, None, None, None, None, None, None) in list(load_workbook(BytesIO(content))['Özet'].values)
    app.dependency_overrides[require_user] = lambda: d['strategy']
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        options = dict(headless=True, args=['--no-sandbox'])
        if os.getenv('CIVENSA_CHROMIUM_PATH'):
            options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
        browser = p.chromium.launch(**options)
        try:
            page = browser.new_page()
            def serve(route):
                response = client.get(route.request.url)
                headers = dict(response.headers)
                headers.pop('set-cookie', None)
                route.fulfill(status=response.status_code, headers=headers, body=response.content)
            page.route('http://testserver/**', serve)
            for width in (1366, 390):
                page.set_viewport_size({'width': width, 'height': 900 if width == 1366 else 844})
                page.goto(f"http://testserver/internal-control/monitoring?period_id={d['period'].id}")
                if has_warning:
                    expect(page.get_by_role('alert')).to_contain_text(warning)
                else:
                    expect(page.get_by_text('30,00% (9/30)', exact=True)).to_be_visible()
                expect(page.get_by_text('36,00% (9/25)', exact=True)).to_be_visible()
                expect(page.get_by_role('link', name='Resmî format · Onaylı Word indir')).to_be_visible()
                assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                page.screenshot(path=str(ARTIFACTS / f"strateji-{'uyari-' if has_warning else ''}{width}.png"), full_page=True)
        finally:
            browser.close()


@pytest.mark.parametrize('new_status', ['partial', 'not_approved'])
def test_migration_roundtrip_and_refuses_new_history(new_status):
    base = create_engine(os.environ['CIVENSA_TEST_POSTGRES_URL'])
    schema = 'civ084_migration_' + uuid4().hex
    with base.begin() as connection:
        connection.execute(CreateSchema(schema))
    url = make_url(os.environ['CIVENSA_TEST_POSTGRES_URL']).update_query_dict({'options': '-csearch_path=' + schema}).render_as_string(hide_password=False)
    env = {**os.environ, 'DATABASE_URL': url}
    engine = create_engine(url)
    def migrate(direction, target, success=True):
        result = subprocess.run([sys.executable, '-m', 'alembic', direction, target], cwd=ROOT,
                                env=env, capture_output=True, text=True)
        assert (result.returncode == 0) == success, result.stderr
        return result
    try:
        migrate('upgrade', 'c078a1b2c3d4')
        with engine.begin() as connection:
            connection.execute(text("INSERT INTO organizations(id,name) VALUES (1,'Sentetik')"))
            connection.execute(text("INSERT INTO departments(id,organization_id,name) VALUES (1,1,'Test Müdürlüğü')"))
            connection.execute(text("INSERT INTO users(id,email,full_name,role,password_hash,is_active,department_id,position,can_submit) VALUES (1,'strategy@civ084.invalid','Strateji','strategy','unused',true,1,'director',false),(2,'unit@civ084.invalid','Birim','department','unused',true,1,'staff',true)"))
        with Session(engine, expire_on_commit=False) as db:
            actor, unit = db.get(User, 1), db.get(User, 2)
            batch, _ = ic.create_import(db, actor, workbook(), 'migration-sentetik.xlsx')
            plan = ic.commit_import(db, actor, batch.id)
            db.commit()
            ic.activate_plan(db, actor, plan.id)
            db.commit()
            period = m.open_period(db, actor, plan.id, 2027, 1, m.today(db))
            db.commit()
            owned = db.scalar(select(Batch).where(Batch.period_id == period.id))
            target = m.rows(db, owned)[0]['target']
            m.save_entry(db, unit, owned.id, target.id, 0, values())
            db.commit()
            before = db.execute(text('SELECT * FROM ic_monitoring_entries ORDER BY id')).all()
            db.rollback()
            migrate('upgrade', 'head')
            assert db.execute(text('SELECT * FROM ic_monitoring_entries ORDER BY id')).all() == before
            db.rollback()
            migrate('downgrade', 'c078a1b2c3d4')
            migrate('upgrade', 'head')
            m.save_entry(db, unit, owned.id, target.id, 1, values(status=new_status,
                percentage='42' if new_status == 'partial' else '', explanation='Gerekçeli karar'))
            db.commit()
            result = migrate('downgrade', 'c078a1b2c3d4', success=False)
            assert 'Veriyi koruyan geçiş hazırlanmadan sürüm düşürülemez' in result.stderr
            assert db.scalar(select(Entry.status).order_by(Entry.id.desc()).limit(1)) == new_status
    finally:
        engine.dispose()
        with base.begin() as connection:
            connection.execute(DropSchema(schema, cascade=True))
        base.dispose()


def test_new_status_form_mobile_desktop(db, monitoring, client):
    from playwright.sync_api import sync_playwright, expect
    d = monitoring
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        options = dict(headless=True, args=['--no-sandbox'])
        if os.getenv('CIVENSA_CHROMIUM_PATH'):
            options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
        browser = p.chromium.launch(**options)
        try:
            page = browser.new_page()
            def serve(route):
                request = route.request
                response = client.request(request.method, request.url, content=request.post_data_buffer,
                    headers={'content-type': request.headers.get('content-type', ''), 'accept': 'text/html'})
                headers = dict(response.headers)
                headers.pop('set-cookie', None)
                route.fulfill(status=response.status_code, headers=headers, body=response.content)
            page.route('http://testserver/**', serve)
            for width in (1366, 390):
                page.set_viewport_size({'width': width, 'height': 900 if width == 1366 else 844})
                page.goto(f"http://testserver/internal-control/monitoring/batches/{d['owned'].id}")
                form = page.locator('form[action$="/entry"]').first
                expect(form.locator('option[value="partial"]')).to_have_text('Kısmen tamamlandı (devam ediyor)')
                expect(form.locator('option[value="not_approved"]')).to_have_text('Eylem uygun görülmemiştir')
                form.locator('[name="status"]').select_option('partial')
                form.locator('[name="percentage"]').fill('100')
                form.locator('[name="explanation"]').fill('Korunan kısmi tamamlanma açıklaması')
                form.locator('[name="reason"]').fill('Takvim güncellendi')
                form.get_by_role('button', name='Eylemi kaydet').click()
                form = page.locator('form[action$="/entry"]').first
                expect(form.locator('[name="status"]')).to_have_value('partial')
                expect(form.locator('[name="percentage"]')).to_have_value('100')
                expect(form.locator('[name="explanation"]')).to_have_value('Korunan kısmi tamamlanma açıklaması')
                expect(page.get_by_text('1–99 arasında yüzde girin', exact=False).first).to_be_visible()
                assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                page.screenshot(path=str(ARTIFACTS / f'kismi-giris-hata-{width}.png'), full_page=True)
        finally:
            browser.close()


def test_existing_export_payloads_unchanged(db, monitoring, monkeypatch):
    """ZIP container timestamps aside, every existing report member stays identical."""
    from zipfile import ZipFile
    from xml.etree import ElementTree
    def payload(archive, name):
        data = archive.read(name)
        if name == 'docProps/core.xml':
            root = ElementTree.fromstring(data)
            for tag in ('created', 'modified'):
                stamp = root.find('{http://purl.org/dc/terms/}' + tag)
                if stamp is not None:
                    stamp.text = '[creation time]'
            data = ElementTree.tostring(root)
        return data
    d = monitoring
    for batch in d['batches']:
        approve(db, d, batch)
    for fmt in ('docx', 'xlsx'):
        current = export.period_export(db, d['strategy'], d['period'].id, fmt)
        with monkeypatch.context() as patch:
            patch.setattr(m, 'STATUSES', {k: v for k, v in m.STATUSES.items() if k not in ('partial', 'not_approved')})
            original = export.period_export(db, d['strategy'], d['period'].id, fmt)
        with ZipFile(BytesIO(current)) as actual, ZipFile(BytesIO(original)) as previous:
            assert actual.namelist() == previous.namelist()
            assert all(payload(actual, name) == payload(previous, name) for name in actual.namelist())


def test_new_status_database_constraints(db, monitoring):
    from sqlalchemy.exc import IntegrityError
    d = monitoring
    batch = d['owned']
    target = m.rows(db, batch)[0]['target']
    for status, percentage, explanation in [('partial', None, ''), ('partial', 0, ''),
                                            ('partial', 100, ''), ('not_approved', None, '  ')]:
        with pytest.raises(IntegrityError):
            with db.begin_nested():
                event = m.event(db, d['user'], batch, 0, 'save', 'draft', 'draft')
                db.execute(Entry.__table__.insert().values(target_id=target.id, event_id=event.id,
                    status=status, percentage=percentage, explanation=explanation, reason='Gerekçe', evidence=''))
    assert m.state(db, batch.id)[1] == 0


def test_new_statuses_survive_approval_and_official_summary(db, monitoring):
    d = monitoring
    approve(db, d, d['owned'], status='partial', percentage='45')
    approve(db, d, d['coop'])
    approve(db, d, d['body'], status='not_approved', percentage='', explanation='Kurul uygun görmedi.')
    metrics = m.official_metrics(db, d['strategy'], d['period'])
    assert metrics['components'][0][1:6] == [3, 0, 1, 2, 0]
    book = load_workbook(BytesIO(export.official_export(db, d['strategy'], d['period'].id, 'xlsx')))
    explanations = [row[-1] for row in list(book['Resmî İzleme'].values)[1:]]
    assert sum(e.startswith('(KISMEN TAMAMLANDI (DEVAM EDİYOR))') for e in explanations) == 2
    assert any(e.startswith('Eylem Uygun Görülmemiştir.') and 'Kurul uygun görmedi.' in e for e in explanations)


@pytest.mark.parametrize('future_status', ['completed', 'ongoing', 'not_approved'])
def test_component_period_counts_preserve_all_completed_total_and_rates(db, official_plan, future_status):
    d = official_plan
    records = m.action_records(db, d['period'])
    # A completed action with an unreadable date also stays outside period counts.
    records[0]['due'] = None
    for record in records:
        if record['due'] is False:
            record['status'] = future_status
    metrics = m.official_metrics(db, d['strategy'], d['period'], records=records)
    completed_total = 14 if future_status == 'completed' else 9
    assert [sum(row[column] for row in metrics['components']) for column in range(1, 6)] == [30, 8, 0, 16, completed_total]
    assert metrics['period_rate'] == '33,33% (8/24)'
    assert metrics['plan_rate'] == ('44,83% (13/29)' if future_status == 'completed' else '27,59% (8/29)')
