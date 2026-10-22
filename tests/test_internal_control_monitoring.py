"""CIV-079 PostgreSQL acceptance: isolated synthetic import data only."""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
from uuid import uuid4

from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError, IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import internal_control_service as ic
from app import internal_control_monitoring_service as m
from app import internal_control_monitoring_export as export
from app.internal_control_monitoring_models import (ICMonitoringPeriod as Period,
    ICMonitoringBatch as Batch, ICMonitoringTarget as Target, ICMonitoringEntry as Entry,
    ICMonitoringEvent as Event, ICMonitoringNote as Note)
from app.models import Department, User
from app.internal_control_models import ICAction
from test_internal_control import workbook, action_row, saved

pytestmark = pytest.mark.postgres
ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / 'tests/artifacts/civ079'


@pytest.fixture
def monitoring(db, sample):
    other = Department(organization_id=sample['department'].organization_id, name='Park ve Bahçeler Müdürlüğü')
    db.add(other)
    db.flush()
    collaborator = User(email='coop@civ079.invalid', full_name='Katkı Personeli', role='department',
        department_id=other.id, password_hash='unused', can_submit=True)
    staff = User(email='ic-staff@civ079.invalid', full_name='Strateji Personeli', role='strategy',
        department_id=sample['department'].id, password_hash='unused', can_submit=True)
    db.add_all([collaborator, staff])
    db.commit()
    due = (m.today(db) - timedelta(days=1)).strftime('%d.%m.%Y')
    data = workbook([
        action_row(**{'8': other.name, '11': due}),
        action_row(**{'2': 'KOS1.2', '5': 'KOS1.2.1', '6': 'Ortak etik eğitim', '7': 'Tüm Müdürlükler'}),
        action_row(**{'2': 'KOS1.3', '5': 'KOS1.3.1', '6': 'Kurul değerlendirmesi', '7': 'İç Kontrol Kurulu'}),
    ])
    batch, report = ic.create_import(db, sample['strategy'], data, 'civ079-sentetik.xlsx')
    ic.map_unit(db, sample['strategy'], batch, 'İç Kontrol Kurulu', body_name='İç Kontrol Kurulu')
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    ic.activate_plan(db, sample['strategy'], plan.id)
    db.commit()
    period = m.open_period(db, sample['strategy'], plan.id, 2027, 1, m.today(db) + timedelta(days=1000))
    db.commit()
    batches = db.scalars(select(Batch).where(Batch.period_id == period.id).order_by(Batch.id)).all()
    owned = next(b for b in batches if b.department_id == sample['department'].id)
    coop = next(b for b in batches if b.department_id == other.id)
    body = next(b for b in batches if b.body_name)
    return {**sample, **dict(plan=plan, period=period, batches=batches, owned=owned, coop=coop, body=body,
        collaborator=collaborator, staff=staff)}


def values(**changes):
    return dict(status='ongoing', percentage='0', explanation='Eğitim hazırlığı sürüyor.',
        reason='Takvim güncellendi.', evidence='EBYS 2026/79 · 22.10.2026', **changes) if not changes else {**values(), **changes}


def fill(db, data, batch, actor=None, **changes):
    actor = actor or (data['staff'] if batch.body_name else
        data['user'] if batch.department_id == data['department'].id else data['collaborator'])
    for row in m.rows(db, batch):
        m.save_entry(db, actor, batch.id, row['target'].id, m.state(db, batch.id)[1], values(**changes))
    db.commit()
    return actor


def approve(db, data, batch, **changes):
    actor = fill(db, data, batch, **changes)
    m.transition(db, actor, batch.id, 'submit', m.state(db, batch.id)[1])
    db.commit()
    m.transition(db, data['strategy'], batch.id, 'approve', m.state(db, batch.id)[1])
    db.commit()


def test_periods_single_open_close_and_plan_untouched(db, monitoring):
    d = monitoring
    before = [(a.id, a.text, a.due_raw) for a in db.scalars(select(ICAction))]
    with pytest.raises(HTTPException) as exc:
        m.open_period(db, d['strategy'], d['plan'].id, 2027, 2, m.today(db))
    assert exc.value.status_code == 409
    with pytest.raises(HTTPException):
        m.close_period(db, d['strategy'], d['period'].id, 1, ' ')
    m.close_period(db, d['strategy'], d['period'].id, 1, 'İlk dönem veri girişi tamamlandı.')
    db.commit()
    second = m.open_period(db, d['strategy'], d['plan'].id, 2027, 2, m.today(db))
    db.commit()
    assert second.half == 2
    assert before == [(a.id, a.text, a.due_raw) for a in db.scalars(select(ICAction))]
    with pytest.raises(HTTPException):
        m.save_entry(db, d['user'], d['owned'].id, m.rows(db, d['owned'])[0]['target'].id, 0, values())
    with pytest.raises(IntegrityError):
        with db.begin_nested():
            db.add(Period(plan_id=d['plan'].id, year=2028, half=1, deadline=m.today(db)))
            db.flush()


@pytest.mark.parametrize('status,percentage,reason,valid', [
    ('ongoing', '', 'Var', False), ('ongoing', '0', 'Var', True),
    ('not_started', '', 'Var', True), ('not_started', '0', 'Var', True),
    ('completed', '100', '', True), ('delayed', '25', '', False),
    ('ongoing', '50', '', False), ('not_started', '', '', False),
    ('ongoing', '-1', 'Var', False), ('ongoing', '101', 'Var', False),
    ('ongoing', '1,5', 'Var', False), ('unknown', '0', 'Var', False),
])
def test_entry_validation_blank_zero_reason(db, monitoring, status, percentage, reason, valid):
    d = monitoring
    row = m.rows(db, d['owned'])[0]
    data = values(status=status, percentage=percentage, reason=reason)
    if not valid:
        with pytest.raises(HTTPException) as exc:
            m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 0, data)
        assert exc.value.status_code == 422 and exc.value.pp_errors
        assert m.state(db, d['owned'].id) == ('draft', 0)
    else:
        entry = m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 0, data)
        db.commit()
        assert entry.percentage == (int(percentage) if percentage else None)


def test_continuous_future_actions_need_reason_only_when_delayed(db, monitoring):
    d = monitoring
    row = m.rows(db, d['owned'])[1]
    entry = m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 0, values(reason=''))
    assert entry.reason == ''
    with pytest.raises(HTTPException):
        m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 1, values(status='delayed', reason=''))


def test_ownership_isolation_notes_and_named_missing(db, monitoring, client):
    d = monitoring
    assert len(m.rows(db, d['owned'])) == 2 and len(m.rows(db, d['coop'])) == 1
    assert len(m.rows(db, d['body'])) == 1
    with pytest.raises(HTTPException) as exc:
        m.get_batch(db, d['user'], d['coop'].id)
    assert exc.value.status_code == 403
    with pytest.raises(HTTPException):
        m.save_entry(db, d['strategy'], d['owned'].id, m.rows(db, d['owned'])[0]['target'].id, 0, values())
    with pytest.raises(HTTPException):
        m.save_entry(db, d['user'], d['body'].id, m.rows(db, d['body'])[0]['target'].id, 0, values())
    fill(db, d, d['body'])
    action = m.rows(db, d['owned'])[0]['action']
    m.add_note(db, d['collaborator'], d['period'].id, action.id, 'Eğitim için salon sağlandı.')
    db.commit()
    assert len(m.visible_notes(db, d['user'], d['period'])) == 1
    assert len(m.visible_notes(db, d['strategy'], d['period'])) == 1
    outsider_dep = Department(organization_id=d['department'].organization_id, name='Fen İşleri Müdürlüğü')
    db.add(outsider_dep)
    db.flush()
    outsider = User(email='outsider@civ079.invalid', full_name='Dış birim', role='department', department_id=outsider_dep.id, password_hash='unused')
    db.add(outsider)
    db.commit()
    assert m.visible_notes(db, outsider, d['period']) == []
    with pytest.raises(HTTPException):
        m.add_note(db, d['user'], d['period'].id, action.id, 'Yetkisiz katkı')
    rows, counts = m.overview(db, d['strategy'], d['period'])
    assert {r['batch'].label for r in rows if r['missing']} == {'Test Müdürlüğü', 'Park ve Bahçeler Müdürlüğü'}
    assert sum(t['missing'] for t in counts.values()) == 3
    assert client.get(f"/internal-control/monitoring/batches/{d['coop'].id}").status_code == 403
    assert client.get('/internal-control/monitoring').status_code == 403
    assert client.get(f"/internal-control/monitoring/periods/{d['period'].id}/report.xlsx?draft=true").status_code == 403
    page = client.get(f"/internal-control/our-actions?period_id={d['period'].id}")
    assert page.status_code == 200 and 'Eğitim için salon sağlandı.' in page.text
    assert 'Kurul değerlendirmesi' not in page.text


def test_workflow_permissions_reasons_reopen_history_and_sod(db, monitoring):
    d = monitoring
    b = d['owned']
    with pytest.raises(HTTPException):
        m.transition(db, d['user'], b.id, 'submit', 0)
    fill(db, d, b)
    d['user'].can_submit = False
    db.commit()
    with pytest.raises(HTTPException) as exc:
        m.transition(db, d['user'], b.id, 'submit', 2)
    assert exc.value.status_code == 403
    strategy_department = Department(organization_id=d['department'].organization_id, name='Strateji Geliştirme Müdürlüğü')
    db.add(strategy_department)
    db.flush()
    d['strategy'].department_id = strategy_department.id
    db.commit()
    d['user'].position = 'director'
    db.commit()
    m.transition(db, d['user'], b.id, 'submit', 2)
    db.commit()
    with pytest.raises(HTTPException):
        m.transition(db, d['staff'], b.id, 'approve', 3)
    with pytest.raises(HTTPException) as exc:
        m.transition(db, d['strategy'], b.id, 'return', 3, '')
    assert exc.value.status_code == 422
    m.transition(db, d['strategy'], b.id, 'return', 3, 'Kanıtı tamamlayın.')
    db.commit()
    assert m.state(db, b.id) == ('returned', 4)
    m.transition(db, d['user'], b.id, 'submit', 4)
    db.commit()
    m.transition(db, d['strategy'], b.id, 'approve', 5)
    db.commit()
    with pytest.raises(HTTPException):
        m.save_entry(db, d['user'], b.id, m.rows(db, b)[0]['target'].id, 6, values())
    with pytest.raises(HTTPException):
        m.transition(db, d['strategy'], b.id, 'reopen', 6, '')
    m.transition(db, d['strategy'], b.id, 'reopen', 6, 'EBYS numarası düzeltilecek.')
    db.commit()
    assert m.state(db, b.id) == ('draft', 7)
    assert db.scalar(select(Event).where(Event.batch_id == b.id, Event.action == 'approve')).snapshot_json
    body_actor = fill(db, d, d['body'], actor=d['strategy'])
    m.transition(db, body_actor, d['body'].id, 'submit', 1)
    db.commit()
    with pytest.raises(HTTPException) as exc:
        m.transition(db, body_actor, d['body'].id, 'approve', 2)
    assert exc.value.status_code == 403 and 'Görevler ayrılığı' in exc.value.detail


def test_approved_history_immutable_at_orm_and_database(db, monitoring):
    d = monitoring
    approve(db, d, d['owned'])
    entry = m.rows(db, d['owned'])[0]['entry']
    with pytest.raises(Exception):
        with db.begin_nested():
            entry.explanation = 'Değiştir'
            db.flush()
    for table in ('ic_monitoring_entries', 'ic_monitoring_events', 'ic_monitoring_targets',
                  'ic_monitoring_batches', 'ic_monitoring_period_events'):
        for statement in (f'DELETE FROM {table}', f'UPDATE {table} SET id=id', f'TRUNCATE {table} CASCADE'):
            with pytest.raises(DBAPIError):
                with db.begin_nested():
                    db.execute(text(statement))
    with pytest.raises(DBAPIError):
        with db.begin_nested():
            db.execute(text("UPDATE ic_monitoring_periods SET closed_at=now()"))
    with pytest.raises(DBAPIError):
        with db.begin_nested():
            db.execute(Entry.__table__.insert().values(target_id=entry.target_id, event_id=entry.event_id,
                status='completed', percentage=100, explanation='', reason='', evidence=''))
    assert m.rows(db, d['owned'])[0]['entry'].explanation == 'Eğitim hazırlığı sürüyor.'


def test_concurrent_approve_return_accepts_one_and_stale_save_rejected(db, monitoring):
    d = monitoring
    fill(db, d, d['owned'])
    m.transition(db, d['user'], d['owned'].id, 'submit', 2)
    db.commit()
    batch_id, actor_id = d['owned'].id, d['strategy'].id
    def decide(action):
        with Session(db.get_bind(), expire_on_commit=False) as session:
            try:
                m.transition(session, session.get(User, actor_id), batch_id, action, 3, 'Eşzamanlı karar')
                session.commit()
                return 'accepted'
            except HTTPException as exc:
                session.rollback()
                return exc.status_code
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(decide, ['approve', 'return']))
    assert sorted(results, key=str) == [409, 'accepted']
    assert db.scalars(select(Event).where(Event.batch_id == batch_id, Event.revision == 4)).one()
    with pytest.raises(HTTPException) as exc:
        m.transition(db, d['strategy'], batch_id, 'approve', 3)
    assert exc.value.status_code == 409


def test_exports_draft_counts_percent_turkish_annual_latest_and_samples(db, monitoring):
    d = monitoring
    with pytest.raises(HTTPException) as exc:
        export.period_export(db, d['strategy'], d['period'].id, 'xlsx')
    assert exc.value.status_code == 409
    draft_word = export.period_export(db, d['strategy'], d['period'].id, 'docx', draft=True)
    doc = Document(BytesIO(draft_word))
    assert 'TASLAK' in doc.paragraphs[0].text and 'TASLAK' in doc.sections[0].header.paragraphs[0].text
    draft_book = load_workbook(BytesIO(export.period_export(db, d['strategy'], d['period'].id, 'xlsx', draft=True)))
    assert draft_book['Özet']['A1'].value == 'TASLAK'
    for batch in d['batches']:
        approve(db, d, batch)
    docx = export.period_export(db, d['strategy'], d['period'].id, 'docx')
    xlsx = export.period_export(db, d['strategy'], d['period'].id, 'xlsx')
    doc = Document(BytesIO(docx))
    text_content = '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert 'Kontrol Ortamı' in text_content and 'KOS.1.1.1' in text_content
    assert 'EBYS 2026/79' in text_content and '%0' in text_content
    book = load_workbook(BytesIO(xlsx))
    assert book['İzleme'].max_row == 5
    assert book['İzleme']['H2'].value == 0 and book['İzleme']['H2'].number_format == '0%'
    assert any(row == ('Kontrol Ortamı', 'Devam ediyor', 4) for row in book['Özet'].values)
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    (ARTIFACTS / 'sentetik-izleme-raporu.docx').write_bytes(docx)
    (ARTIFACTS / 'sentetik-izleme-raporu.xlsx').write_bytes(xlsx)
    m.close_period(db, d['strategy'], d['period'].id, 1, 'İlk yarıyıl onaylandı.')
    db.commit()
    second = m.open_period(db, d['strategy'], d['plan'].id, 2027, 2, m.today(db) + timedelta(days=1000))
    db.commit()
    second_owned = db.scalar(select(Batch).where(Batch.period_id == second.id, Batch.department_id == d['department'].id))
    approve(db, d, second_owned, status='completed', percentage='100', reason='')
    annual = load_workbook(BytesIO(export.annual_export(db, d['strategy'], d['plan'].id, 2027)))
    statuses = [r[6] for r in list(annual['İzleme'].values)[1:]]
    assert statuses.count('Tamamlandı') == 2 and statuses.count('Devam ediyor') == 2
    m.transition(db, d['strategy'], second_owned.id, 'reopen', m.state(db, second_owned.id)[1], 'Düzeltme')
    db.commit()
    fill(db, d, second_owned, status='delayed', percentage='75')
    annual = load_workbook(BytesIO(export.annual_export(db, d['strategy'], d['plan'].id, 2027)))
    assert [r[6] for r in list(annual['İzleme'].values)[1:]].count('Tamamlandı') == 2


def test_error_retention_uses_pp_form_contract(db, monitoring, client):
    d = monitoring
    row = m.rows(db, d['owned'])[0]
    response = client.post(f"/internal-control/monitoring/batches/{d['owned'].id}/entry", data={
        'id': row['target'].id, 'expected_revision': 0, **values(percentage='105', explanation='<Korunan açıklama>')},
        headers={'accept': 'text/html'})
    assert response.status_code == 422
    assert '105' in response.text and '&lt;Korunan açıklama&gt;' in response.text
    assert '0–100 girin' in response.text and m.state(db, d['owned'].id)[1] == 0


def test_migration_upgrade_empty_downgrade_and_data_refusal():
    base = create_engine(os.environ['CIVENSA_TEST_POSTGRES_URL'])
    schema = 'civ079_migration_' + uuid4().hex
    with base.begin() as conn:
        conn.execute(CreateSchema(schema))
    url = make_url(os.environ['CIVENSA_TEST_POSTGRES_URL']).update_query_dict({'options': '-csearch_path=' + schema}).render_as_string(hide_password=False)
    env = {**os.environ, 'DATABASE_URL': url}
    engine = create_engine(url)
    def migrate(direction, target, success=True):
        result = subprocess.run([sys.executable, '-m', 'alembic', direction, target], cwd=ROOT,
            env=env, capture_output=True, text=True)
        assert (result.returncode == 0) == success, result.stderr
        return result
    try:
        migrate('upgrade', 'head')
        migrate('downgrade', 'c074a1b2c3d4')
        migrate('upgrade', 'head')
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO organizations(id,name) VALUES (1,'Sentetik')"))
            conn.execute(text("INSERT INTO departments(id,organization_id,name) VALUES (1,1,'Test Müdürlüğü')"))
            conn.execute(text("INSERT INTO users(id,email,full_name,role,password_hash,is_active,department_id,position) VALUES (1,'migration@civ079.invalid','Strateji','strategy','unused',true,1,'director')"))
        with Session(engine, expire_on_commit=False) as session:
            actor = session.get(User, 1)
            batch, report = ic.create_import(session, actor, workbook(), 'sentetik.xlsx')
            plan = ic.commit_import(session, actor, batch.id)
            session.commit()
            ic.activate_plan(session, actor, plan.id)
            session.commit()
            period = m.open_period(session, actor, plan.id, 2027, 1, m.today(session))
            session.commit()
            owned = session.scalar(select(Batch).where(Batch.period_id == period.id))
            staff = User(id=2, email='department@civ079.invalid', full_name='Müdürlük', role='department', department_id=1, can_submit=True, password_hash='unused')
            session.add(staff)
            session.commit()
            row = m.rows(session, owned)[0]
            entry = m.save_entry(session, staff, owned.id, row['target'].id, 0, values())
            session.commit()
            m.transition(session, staff, owned.id, 'submit', 1)
            session.commit()
            m.transition(session, actor, owned.id, 'approve', 2)
            session.commit()
            with pytest.raises(DBAPIError):
                with session.begin_nested():
                    session.execute(text('UPDATE ic_monitoring_entries SET percentage=55'))
        result = migrate('downgrade', 'c074a1b2c3d4', success=False)
        assert 'İzleme verisi var' in result.stderr
        with engine.connect() as conn:
            assert conn.execute(text('SELECT count(*) FROM ic_monitoring_entries')).scalar() == 1
    finally:
        engine.dispose()
        with base.begin() as conn:
            conn.execute(DropSchema(schema, cascade=True))
        base.dispose()


def test_chromium_turkish_entry_error_retention_and_overview_desktop_mobile(db, monitoring, client):
    from playwright.sync_api import sync_playwright, expect
    from app.main import app, require_user
    d = monitoring
    current = {'user': d['user']}
    app.dependency_overrides[require_user] = lambda: current['user']
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        options = {'headless': True, 'args': ['--no-sandbox']}
        if os.getenv('CIVENSA_CHROMIUM_PATH'):
            options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
        browser = p.chromium.launch(**options)
        try:
            with browser.new_context(viewport={'width': 1366, 'height': 900}) as context:
                def serve(route):
                    request = route.request
                    response = client.request(request.method, request.url, content=request.post_data_buffer,
                        headers={'content-type': request.headers.get('content-type', ''),
                                 'accept': request.headers.get('accept', '')}, follow_redirects=True)
                    headers = dict(response.headers)
                    headers.pop('set-cookie', None)
                    route.fulfill(status=response.status_code, headers=headers, body=response.content)
                context.route('http://testserver/**', serve)
                page = context.new_page()
                page.goto(f"http://testserver/internal-control/our-actions?period_id={d['period'].id}")
                expect(page.get_by_role('heading', name='İç Kontrol Eylemlerimiz', exact=True)).to_be_visible()
                expect(page.locator('meta[name="pp-period-id"]')).to_have_count(0)
                form = page.locator('form[action$="/entry"]').first
                form.locator('[name="status"]').select_option('ongoing')
                form.locator('[name="percentage"]').fill('105')
                form.locator('[name="explanation"]').fill('Korunan Türkçe açıklama: eğitim hazırlıkları.')
                form.locator('[name="reason"]').fill('Takvim güncellendi.')
                form.locator('[name="evidence"]').fill('EBYS 2026/79')
                form.get_by_role('button', name='Eylemi kaydet').click()
                expect(page.locator('[name="percentage"]').first).to_have_value('105')
                expect(page.locator('[name="explanation"]').first).to_have_value('Korunan Türkçe açıklama: eğitim hazırlıkları.')
                expect(page.get_by_text('0–100 girin', exact=False).first).to_be_visible()
                page.screenshot(path=str(ARTIFACTS / 'giris-hata-1366.png'), full_page=True)
                page.set_viewport_size({'width': 390, 'height': 844})
                page.screenshot(path=str(ARTIFACTS / 'giris-hata-390.png'), full_page=True)
                page.locator('form[action$="/entry"]').first.locator('[name="percentage"]').fill('0')
                page.locator('form[action$="/entry"]').first.get_by_role('button', name='Eylemi kaydet').click()
                expect(page.locator('[name="percentage"]').first).to_have_value('0')
                for width in (1366, 390):
                    page.set_viewport_size({'width': width, 'height': 900 if width == 1366 else 844})
                    page.goto(f"http://testserver/internal-control/our-actions?period_id={d['period'].id}")
                    expect(page.get_by_role('heading', name='Sorumlu olduğumuz eylemler', exact=True)).to_be_visible()
                    assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                    page.screenshot(path=str(ARTIFACTS / f'mudurluk-giris-{width}.png'), full_page=True)
                current['user'] = d['strategy']
                for width in (1366, 390):
                    page.set_viewport_size({'width': width, 'height': 900 if width == 1366 else 844})
                    page.goto(f"http://testserver/internal-control/monitoring?period_id={d['period'].id}")
                    expect(page.get_by_role('heading', name='Sorumlu birimlerin durumu')).to_be_visible()
                    expect(page.get_by_role('link', name='Park ve Bahçeler Müdürlüğü', exact=True)).to_be_visible()
                    expect(page.get_by_role('link', name='TASLAK DOCX indir')).to_be_visible()
                    assert page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')
                    page.screenshot(path=str(ARTIFACTS / f'strateji-izleme-{width}.png'), full_page=True)
        finally:
            browser.close()


def test_blank_and_zero_stay_distinct_in_reports(db, monitoring):
    d = monitoring
    batch = d['owned']
    first, second = m.rows(db, batch)
    m.save_entry(db, d['user'], batch.id, first['target'].id, 0, values(status='not_started', percentage=''))
    m.save_entry(db, d['user'], batch.id, second['target'].id, 1, values(status='not_started', percentage='0'))
    m.transition(db, d['user'], batch.id, 'submit', 2)
    db.commit()
    m.transition(db, d['strategy'], batch.id, 'approve', 3)
    db.commit()
    data = export.period_export(db, d['strategy'], d['period'].id, 'xlsx', draft=True)
    book = load_workbook(BytesIO(data))
    selected = [r for r in list(book['İzleme'].values)[1:] if r[5] == batch.label]
    assert [r[7] for r in selected] == [None, 0]
    doc = Document(BytesIO(export.period_export(db, d['strategy'], d['period'].id, 'docx', draft=True)))
    selected = [row for table in doc.tables[1:] for row in table.rows[1:] if row.cells[4].text == batch.label]
    assert [row.cells[6].text for row in selected] == ['', '%0']


def test_deadline_readonly_and_stale_edits(db, monitoring, monkeypatch):
    d = monitoring
    row = m.rows(db, d['owned'])[0]
    m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 0, values())
    db.commit()
    with pytest.raises(HTTPException) as exc:
        m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 0, values(percentage='75'))
    assert exc.value.status_code == 409
    monkeypatch.setattr(m, 'today', lambda db: d['period'].deadline + timedelta(days=1))
    with pytest.raises(HTTPException) as exc:
        m.save_entry(db, d['user'], d['owned'].id, row['target'].id, 1, values())
    assert exc.value.status_code == 409 and 'süresi doldu' in exc.value.detail
    with pytest.raises(HTTPException):
        m.transition(db, d['user'], d['owned'].id, 'submit', 1)
    action_id = row['action'].id
    with pytest.raises(HTTPException):
        m.add_note(db, d['collaborator'], d['period'].id, action_id, 'Geç not')


def test_two_simultaneous_opens_accept_one(db, monitoring):
    d = monitoring
    m.close_period(db, d['strategy'], d['period'].id, 1, 'Yeni yarıyıl için kapatıldı.')
    db.commit()
    plan_id, actor_id = d['plan'].id, d['strategy'].id
    deadline = m.today(db)
    def open_half(half):
        with Session(db.get_bind(), expire_on_commit=False) as session:
            try:
                m.open_period(session, session.get(User, actor_id), plan_id, 2028, half, deadline)
                session.commit()
                return 'accepted'
            except HTTPException as exc:
                session.rollback()
                return exc.status_code
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(open_half, [1, 2]))
    assert sorted(results, key=str) == [409, 'accepted']


def test_authenticated_monitoring_writes_require_csrf(db, monitoring, client):
    import re
    from app.main import app, require_user
    from app.security import hash_password
    d = monitoring
    app.dependency_overrides.pop(require_user)
    d['user'].password_hash = hash_password('civ079-local-test-only')
    db.commit()
    assert client.post('/login', data={'email': d['user'].email,
        'password': 'civ079-local-test-only'}, follow_redirects=False).status_code == 303
    path = f"/internal-control/monitoring/batches/{d['owned'].id}"
    page = client.get(path)
    assert page.status_code == 200
    token = re.search('name="csrf_token" value="([^"]+)"', page.text).group(1)
    row = m.rows(db, d['owned'])[0]
    data = {'id': row['target'].id, 'expected_revision': 0, **values()}
    assert client.post(path + '/entry', data=data).status_code == 403
    assert m.state(db, d['owned'].id)[1] == 0
    data['csrf_token'] = token
    assert client.post(path + '/entry', data=data, follow_redirects=False).status_code == 303
    assert m.state(db, d['owned'].id)[1] == 1


def test_new_active_plan_without_monitoring_keeps_readonly_catalogue(db, monitoring, client):
    d = monitoring
    data = workbook([action_row(**{'6': 'Yeni planın eylemi'})], metadata=['Yeni Uyum Planı', 2027, 2029, 'v2', '22.10.2026'])
    batch, report = ic.create_import(db, d['strategy'], data, 'yeni-plan.xlsx')
    plan = ic.commit_import(db, d['strategy'], batch.id, supersedes_id=d['plan'].id)
    db.commit()
    ic.activate_plan(db, d['strategy'], plan.id)
    db.commit()
    page = client.get('/internal-control/our-actions')
    assert page.status_code == 200 and 'Yeni planın eylemi' in page.text
    assert 'Geçmiş izleme dönemlerini gör' in page.text
    history = client.get(f"/internal-control/our-actions?period_id={d['period'].id}")
    assert history.status_code == 200 and 'Eğitim düzenle' in history.text


def test_overview_also_names_directorates_without_responsible_actions(db, monitoring, client):
    from app.main import app, require_user
    d = monitoring
    department = Department(organization_id=d['department'].organization_id, name='Yeni Hizmet Müdürlüğü')
    db.add(department)
    db.commit()
    app.dependency_overrides[require_user] = lambda: d['strategy']
    response = client.get(f"/internal-control/monitoring?period_id={d['period'].id}")
    assert response.status_code == 200 and 'Yeni Hizmet Müdürlüğü' in response.text
    assert 'Sorumlu eylem yok; gönderim beklenmiyor.' in response.text
    response = client.get(f"/internal-control/monitoring/periods/{d['period'].id}/report.xlsx?draft=gecersiz")
    assert response.status_code == 422 and 'Taslak seçimi geçersiz' in response.text
