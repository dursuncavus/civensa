"""CIV-026: lossless, submit-ready Word programme imports."""
import io
import json
from copy import deepcopy
from decimal import Decimal

import pytest
from docx import Document
from sqlalchemy import select, func

from app.data_importer import parse_performance_docx
from app.import_service import apply_import_batch, revert_import_batch
from app.models import Activity, ActivityCostValue, EconomicValue, ImportBatch, Indicator, ReportCard
from app.program_service import validate_department
from test_word_import import word_document, import_context, import_client

pytestmark = pytest.mark.postgres


def source_document(*, merged=False, missing=False, personnel='100', outside='0'):
    doc = Document(io.BytesIO(word_document(values=('-', '', '0'))))
    header = doc.tables[0].rows[6]
    header._tr.addnext(deepcopy(header._tr))
    costs = doc.tables[1]
    # Finance agrees with the compact economic table, including explicit zero outside funding.
    for row, amount in [(1, '100'), (2, '100'), (3, '0'), (4, personnel)]:
        costs.cell(row, 1).text = amount
    for label in ['2026 Bütçe', '2026 Haziran Sonu Harcama', '2026 Yıl Sonu Harcama Tahmini']:
        cells = costs.add_column(100000).cells
        cells[0].text = label
        for cell in cells[1:5]: cell.text = '-'
    econ = doc.add_table(rows=0, cols=6)
    for texts in [
        ('Ekonomik Ad', '2026 Bütçe', '2026 Haziran Sonu Harcama', '2027 Bütçe', '2028 Tahmin', '2029 Tahmin'),
        ('Personel Giderleri', '80', '80', '80', '0', '0'),
        ('Sosyal Güvenlik Kurumuna Devlet Primi Giderler', '20', '20', '20', '0', '0'),
        ('Mal ve Hizmet Alım Giderleri', '100', '100', '100', '0', '0'),
        ('BÜTÇE İÇİ TOPLAM KAYNAK', '200', '200', '200', '0', '0'),
        ('BÜTÇE DIŞI TOPLAM KAYNAK', outside, outside, outside, '0', '0'),
        ('FAALİYET MALİYETİ TOPLAMI', '200', '200', '200', '0', '0'),
    ]:
        for cell, text in zip(econ.add_row().cells, texts): cell.text = text
    if merged: econ.cell(3, 3).merge(econ.cell(3, 4)).text = '100'
    if missing: econ.cell(3, 3).text = ''
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def import_parsed(db, context, parsed):
    batch = ImportBatch(period_id=context.period.id, filename='civ026.docx', status='preview',
                        payload_json=json.dumps(parsed), created_by=context.user.id)
    db.add(batch); db.commit()
    result = apply_import_batch(db, batch); db.commit()
    return batch, result


def test_repeated_headers_never_become_indicators():
    parsed = parse_performance_docx(source_document())
    block = parsed['departments'][0]['blocks'][0]
    assert [i['name'] for i in block['indicators']] == ['Park sayısı']
    assert block['metadata']['calculation_method'] == 'Parkları saymak'


def test_dash_activity_is_not_applicable_but_indicator_needs_check(db, import_context):
    parsed = parse_performance_docx(source_document())
    assert parsed['departments'][0]['blocks'][0]['indicators'][0]['values'][0]['state'] == 'needs_review'
    costs = parsed['departments'][0]['activities'][0]['costs']
    assert all(c['state'] == 'not_applicable' for c in costs if c['year'] == 2026)
    batch, _ = import_parsed(db, import_context, parsed)
    historical = db.scalars(select(ActivityCostValue).where(ActivityCostValue.year == 2026)).all()
    assert len(historical) == 9
    assert all(v.amount is None and v.raw_text == '-' and v.state == 'not_applicable' for v in historical)
    errors, warnings = validate_department(db, batch.period_id, db.scalar(select(ReportCard.department_id)))
    assert not any('2026' in e and ('finansman' in e or 'tutarı' in e) for e in errors)
    assert any('2026' in w and 'kontrol' in w for w in warnings)


def test_economic_values_columns_provenance_and_safe_revert(db, import_context):
    parsed = parse_performance_docx(source_document())
    table = parsed['departments'][0]['economic_tables'][0]
    assert {(v['year'], v['kind']) for v in table['values']} == {
        (2026, 'budget'), (2026, 'midyear_spend'), (2027, 'budget'), (2028, 'forecast'), (2029, 'forecast')}
    assert not parsed['errors']
    batch, result = import_parsed(db, import_context, parsed)
    assert result['economic_values'] == 30
    row = db.scalar(select(EconomicValue).where(EconomicValue.year == 2027,
                    EconomicValue.category == 'goods_services', EconomicValue.funding_type == 'inside'))
    assert (row.amount, row.raw_text, row.state, row.source_table_number) == (Decimal('100'), '100', 'verified', 3)
    assert json.loads(row.source_context_json)['card_ids'] == [row.source_card_id]
    errors, _ = validate_department(db, batch.period_id, row.department_id)
    assert not any('ekonomik' in e for e in errors), errors
    revert_import_batch(db, batch); db.commit()
    assert db.scalar(select(func.count()).select_from(EconomicValue)) == 0
    assert db.scalar(select(func.count()).select_from(ReportCard)) == 0


@pytest.mark.parametrize('option', ['merged', 'missing'])
def test_ambiguous_economic_cells_stay_null_and_listed(db, import_context, option):
    parsed = parse_performance_docx(source_document(**{option: True}))
    table = parsed['departments'][0]['economic_tables'][0]
    value = next(v for v in table['values'] if v['year'] == 2027 and v['category'] == 'goods_services' and v['funding_type'] == 'inside')
    assert value['amount'] is None and value['state'] == 'needs_review'
    assert any(r['table_number'] == 3 and 'Mal ve Hizmet' in r['record'] for r in parsed['review_items'])
    batch, _ = import_parsed(db, import_context, parsed)
    row = db.scalar(select(EconomicValue).where(EconomicValue.year == 2027,
                    EconomicValue.category == 'goods_services', EconomicValue.funding_type == 'inside'))
    assert row.amount is None and row.state == 'needs_review'
    errors, _ = validate_department(db, batch.period_id, row.department_id)
    assert any('ekonomik kaynak' in e for e in errors)


def test_personnel_conflict_keeps_both_source_assertions(db, import_context):
    parsed = parse_performance_docx(source_document(personnel='999'))
    personnel = parsed['departments'][0]['personnel']
    assert next(v['amount'] for v in personnel if v['year'] == 2027) == '999'
    values = parsed['departments'][0]['economic_tables'][0]['values']
    assert all(v['state'] == 'needs_review' for v in values if v['year'] == 2027 and v['category'] in {'personnel','sgk'})
    assert any('ortak personel tablosundan farklı' in r['message'] for r in parsed['review_items'])
    import_parsed(db, import_context, parsed)
    amounts = db.scalars(select(EconomicValue.amount).where(EconomicValue.year==2027,
                 EconomicValue.category.in_(['personnel','sgk']), EconomicValue.funding_type=='inside')).all()
    assert sorted(amounts) == [Decimal('20'), Decimal('80')]


def test_nonzero_outside_total_never_fabricates_category_split():
    parsed = parse_performance_docx(source_document(outside='5'))
    values = parsed['departments'][0]['economic_tables'][0]['values']
    assert all(v['state']=='needs_review' and v['amount'] is None for v in values if v['year']==2027)


def test_preview_exposes_all_review_cells(import_client, import_context):
    response = import_client.post(f'/api/periods/{import_context.period.id}/data-import-preview',
                                 files={'file': ('civ026.docx', source_document(merged=True))})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['errors'] == [] and preview['can_apply']
    assert preview['review_items'] and any('Mal ve Hizmet' in r['record'] for r in preview['review_items'])
    detail = import_client.get(f"/api/import-batches/{preview['batch_id']}").json()
    assert detail['review_items'] == preview['review_items']
    assert preview['summary']['needs_review_count'] == len(preview['review_items'])


@pytest.mark.parametrize('kind,label', [('actual','Gerçekleşme'), ('target','Hedef'),
    ('forecast','Tahmin'), ('budget','Bütçe'), ('year_end_estimate','Yıl Sonu Tahmini'), ('value','Değer')])
def test_internal_dimensions_are_turkish_in_validation_and_visible_html(db, sample, client, kind, label):
    from app.models import IndicatorValue
    from html.parser import HTMLParser
    indicator = sample['indicator']
    db.add(IndicatorValue(indicator_id=indicator.id, year=2025, value_kind=kind,
                          value=None, raw_text='belirsiz', state='needs_review'))
    cost = db.scalar(select(ActivityCostValue).where(ActivityCostValue.year==2028, ActivityCostValue.funding_type=='outside'))
    cost.amount = None; cost.state='needs_review'
    db.commit()
    errors, warnings = validate_department(db, sample['period'].id, sample['department'].id)
    messages = '\n'.join(errors+warnings)
    assert f'2025 {label}' in messages
    assert '2028 Tahmin / Bütçe Dışı' in messages
    internal = ('actual','target','forecast','budget','year_end_estimate','inside','outside',' value')
    assert not any(word in messages for word in internal)
    response = client.get(f"/program/{sample['department'].id}?period_id={sample['period'].id}")
    assert response.status_code == 200
    class VisibleText(HTMLParser):
        def __init__(self):
            super().__init__(); self.parts = []; self.hidden = False
        def handle_starttag(self, tag, attrs):
            if tag in {'script', 'style'}: self.hidden = True
        def handle_endtag(self, tag):
            if tag in {'script', 'style'}: self.hidden = False
        def handle_data(self, data):
            if not self.hidden: self.parts.append(data)
    parser = VisibleText(); parser.feed(response.text)
    visible = ' '.join(parser.parts)
    assert not any(word in visible for word in internal)


def test_dash_does_not_equal_blank_or_zero_and_future_dash_blocks(db, import_context):
    parsed = parse_performance_docx(source_document())
    values = parsed['departments'][0]['blocks'][0]['indicators'][0]['values']
    assert [(v['value'],v['state']) for v in values] == [(None,'needs_review'),(None,'missing'),('0','verified')]
    costs = parsed['departments'][0]['activities'][0]['costs']
    next(v for v in costs if v['year']==2028 and v['funding_type']=='inside').update(amount=None, raw_text='-', state='not_applicable')
    batch, _ = import_parsed(db, import_context, parsed)
    errors, _ = validate_department(db, batch.period_id, db.scalar(select(ReportCard.department_id)))
    assert any('2028 Tahmin' in e and 'doğrulanmalı' in e for e in errors)


def test_numbered_shared_personnel_and_activity_descriptions():
    doc = Document(io.BytesIO(source_document()))
    doc.tables[1].cell(4, 0).text = '0006.Personel ve SGK Giderleri Toplamı'
    doc.tables[1].cell(6, 0).text = '0001.Faaliyet Adı: Park bakımı'
    output = io.BytesIO(); doc.save(output)
    parsed = parse_performance_docx(output.getvalue())
    assert next(v['amount'] for v in parsed['departments'][0]['personnel'] if v['year']==2027) == '100'
    assert parsed['departments'][0]['activities'][0]['explanation'] == 'Parkların bakımı yapılır.'


def test_economic_review_can_be_resolved_explicitly(db, sample, client):
    row = db.scalar(select(EconomicValue).where(EconomicValue.year==2027, EconomicValue.category=='goods_services', EconomicValue.funding_type=='inside'))
    row.amount = None; row.raw_text = 'hatalı'; row.state = 'needs_review'; row.issue = 'Kaynağı kontrol edin.'
    db.commit()
    response = client.post(f"/program/{sample['department'].id}/verify-money/economic/{row.id}",
                          data={'amount':'123,45', 'program_group':'service'}, follow_redirects=False)
    assert response.status_code == 303, response.text
    db.refresh(row)
    assert (row.amount,row.raw_text,row.state,row.issue)==(Decimal('123.45'),'hatalı','verified','')
    assert not validate_department(db,sample['period'].id,sample['department'].id)[0]


def test_economic_migration_preserves_data_and_refuses_lossy_downgrade(db, sample):
    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from alembic.script import ScriptDirectory
    from alembic.config import Config
    from sqlalchemy import text
    root = Path(__file__).resolve().parents[1]
    path = root / 'migrations/versions/b0262027a001_economic_import_provenance.py'
    spec = importlib.util.spec_from_file_location('civ026_migration',path)
    migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
    before = db.execute(text('SELECT id, amount FROM ' + db.bind.get_execution_options()['schema_translate_map'][None] + '.economic_values ORDER BY id')).all()
    connection = db.connection()
    schema = db.bind.get_execution_options()['schema_translate_map'][None]
    connection.execute(text(f'SET LOCAL search_path TO "{schema}"'))
    with Operations.context(MigrationContext.configure(connection)):
        migration.downgrade()
        migration.upgrade()
        after = connection.execute(text('SELECT id, amount FROM economic_values ORDER BY id')).all()
        assert before == after
        connection.execute(text("UPDATE economic_values SET state='needs_review', raw_text='-' WHERE id=(SELECT min(id) FROM economic_values)"))
        with pytest.raises(RuntimeError, match='korunmadan'):
            migration.downgrade()
        assert connection.execute(text("SELECT count(*) FROM economic_values WHERE state='needs_review' AND raw_text='-' ")).scalar() == 1
    assert ScriptDirectory.from_config(Config(str(root/'alembic.ini'))).get_heads() == [migration.revision]


def test_unresolved_economic_programme_scope_is_listed_before_apply(db, import_context):
    from app.import_service import build_department_plan
    doc = Document(io.BytesIO(source_document()))
    earlier = deepcopy(doc.tables[0]._tbl)
    doc.tables[0]._tbl.addprevious(earlier)
    doc.tables[0].cell(2, 2).text = 'Örnek Belediyesi / Fen İşleri Müdürlüğü'
    output = io.BytesIO(); doc.save(output)
    parsed = parse_performance_docx(output.getvalue())
    before = deepcopy(parsed)
    plan = build_department_plan(db, parsed)
    assert parsed == before  # The stored source evidence is immutable during planning.
    assert plan['summary']['needs_review_count'] == len(plan['parsed']['review_items'])
    assert any('program bloklarının müdürlük kapsamı' in r['message'] for r in plan['parsed']['review_items'])
    _, result = import_parsed(db, import_context, parsed)
    assert result['economic_values'] == 30
    assert all(v.state=='needs_review' for v in db.scalars(select(EconomicValue)))


def repeated_block_document(*, program=None, department=None, value=None, name=None):
    doc = Document(io.BytesIO(source_document()))
    original = doc.tables[0]
    original._tbl.addnext(deepcopy(original._tbl))
    repeated = doc.tables[1]
    for row, text in [(1, program), (2, department)]:
        if text is not None:
            for cell in repeated.rows[row].cells[2:]: cell.text = text
    if value is not None: repeated.cell(8, 4).text = value
    if name is not None: repeated.cell(8, 0).text = name
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


@pytest.mark.parametrize('options,causes', [
    ({}, ['duplicate_programme_block', 'duplicate_indicator']),
    ({'program': 'Başka Program'}, ['duplicate_indicator']),
    ({'value': '99'}, ['duplicate_programme_block']),
    ({'name': 'Başka gösterge'}, ['duplicate_programme_block']),
    ({'department': 'Başka Müdürlüğü'}, []),
])
def test_duplicate_source_blocks_are_reviewed_without_changing_values(options, causes):
    parsed = parse_performance_docx(repeated_block_document(**options))
    warnings = [r for r in parsed['review_items'] if r.get('cause', '').startswith('duplicate_')]
    assert [r['cause'] for r in warnings] == causes
    assert parsed['summary']['indicator_count'] == 2
    assert sum(len(d['blocks']) for d in parsed['departments']) == 2
    assert all(r['state'] == 'needs_review' and r['table_numbers'] == [1, 2] for r in warnings)
    assert all('kaynak' in r['message'].lower() and 'kontrol' in r['message'].lower() for r in warnings)


def test_duplicate_warning_survives_preview_and_apply(import_client, import_context, db):
    response = import_client.post(f'/api/periods/{import_context.period.id}/data-import-preview',
                                 files={'file': ('repeated.docx', repeated_block_document())})
    preview = response.json()
    assert response.status_code == 200 and preview['can_apply'], preview
    warnings = [r for r in preview['review_items'] if r.get('cause', '').startswith('duplicate_')]
    assert len(warnings) == 2
    detail = import_client.get(f"/api/import-batches/{preview['batch_id']}").json()
    assert detail['review_items'] == preview['review_items']
    batch = db.get(ImportBatch, preview['batch_id'])
    result = apply_import_batch(db, batch); db.commit()
    assert result['cards'] == result['indicators'] == 2
    assert db.scalar(select(func.count()).select_from(Indicator)) == 2


@pytest.mark.parametrize('row,raw,cause,fragments', [
    (5, '50.000', 'outside_funding_unsplit', ['Bütçe Dışı Toplam 50.000,00 TL', 'kategori', 'dağılım']),
    (4, '12.500.000', 'inside_total_mismatch', ['Bütçe İçi Toplam 12.500.000,00 TL', 'toplamı 200,00 TL', 'fark 12.499.800,00 TL']),
    (6, '12.500.000', 'funding_total_mismatch', ['Faaliyet Maliyeti Toplamı 12.500.000,00 TL', '200,00 TL', 'fark 12.499.800,00 TL']),
    (3, 'belirsiz', 'source_cell_unverified', ['Mal ve Hizmet Alım Giderleri', 'belirsiz', 'doğrulanmış']),
    (5, '', 'source_cell_unverified', ['Bütçe Dışı Toplam', 'boş', 'doğrulanmış']),
])
def test_economic_review_names_each_cause_and_source_column(row, raw, cause, fragments):
    doc = Document(io.BytesIO(source_document()))
    doc.tables[2].cell(row, 3).text = raw
    output = io.BytesIO(); doc.save(output)
    parsed = parse_performance_docx(output.getvalue())
    values = parsed['departments'][0]['economic_tables'][0]['values']
    value = next(v for v in values if v['year']==2027 and v['category']=='goods_services' and v['funding_type']=='inside')
    reasons = [c for c in value['review_causes'] if c['cause'] == cause]
    assert len(reasons) == 1
    assert all(fragment in reasons[0]['message'] for fragment in ['Tablo 3, 2027 Bütçe:', *fragments])
    assert 'Kaynak tabloyu kontrol edin.' in reasons[0]['message']
    assert 'Ekonomik satırlar, kaynak toplamları' not in value['issue']


def test_mixed_programme_group_keeps_all_economic_causes():
    doc = Document(io.BytesIO(source_document(outside='5')))
    doc.tables[0]._tbl.addnext(deepcopy(doc.tables[0]._tbl))
    for cell in doc.tables[1].rows[1].cells[2:]: cell.text = 'Yönetim ve Destek Programı'
    output = io.BytesIO(); doc.save(output)
    parsed = parse_performance_docx(output.getvalue())
    value = next(v for v in parsed['departments'][0]['economic_tables'][0]['values'] if v['year']==2027)
    reasons = {c['cause']: c['message'] for c in value['review_causes']}
    assert set(reasons) == {'outside_funding_unsplit', 'funding_total_mismatch', 'program_group_unallocated'}
    assert all('Tablo 4' in message for message in reasons.values())
    assert 'Program grubu dağılımı kaynakta ayrılmamış.' in reasons['program_group_unallocated']
    assert 'Yönetim ve Destek Programı' in reasons['program_group_unallocated']
    assert 'Çevre Programı' in reasons['program_group_unallocated']
