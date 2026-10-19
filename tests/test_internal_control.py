"""CIV-048 PostgreSQL acceptance, validation, immutable history and access."""
from datetime import date
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine, func, select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError, IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateSchema, DropSchema

from app import internal_control_service as ic
from app.internal_control_models import (ICPlan, ICCondition, ICAction, ICActionParty,
    ICUnitAlias, ICImport, ICImportAudit, ICStandard, STANDARD_COMPONENTS)
from app.internal_control_routes import catalogue
from app.models import Department, User

pytestmark = pytest.mark.postgres
ROOT = Path(__file__).resolve().parents[1]


def workbook(rows=None, mutate=None, metadata=None):
    book = load_workbook(BytesIO(ic.template_bytes()))
    info = metadata or ['Test Uyum Planı', 2027, 2029, 'v1', '18.10.2026']
    for idx, value in enumerate(info, 2):
        book['Bilgi'].cell(idx, 2, value)
    for row in rows if rows is not None else [action_row()]:
        book['Plan'].append(row)
    if mutate:
        mutate(book)
    return ic.export_book(book)


def action_row(**changes):
    row = ['KOS', 'KOS1', 'KOS 1.1', 'Etik değerler', 'Geliştirilecek', 'KOS.1.1.1',
           'Eğitim düzenle', 'Test Müdürlüğü', '', 'Eğitim raporu', '2027', 'Sürekli', '', 'Etik Değerler']
    for key, val in changes.items():
        row[int(key)] = val
    return row


def stage(db, sample, data=None):
    batch, report = ic.create_import(db, sample['strategy'], data or workbook(), 'plan.xlsx')
    db.commit()
    return batch, report


def saved(db, sample, data=None):
    batch, report = stage(db, sample, data)
    assert report['can_commit'], report['issues']
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    return plan


def test_template_round_trip_seed_and_raw_values(db, sample):
    book = load_workbook(BytesIO(ic.template_bytes()))
    assert book.sheetnames == ['Plan', 'Bilgi', 'Yönerge']
    assert list(next(book['Plan'].values)) == ic.HEADERS + [ic.TITLE_HEADER]
    rows = [action_row(**{'10': date(2027, 1, 15), '11': date(2029, 12, 31)})]
    plan = saved(db, sample, workbook(rows))
    action = db.scalar(select(ICAction).where(ICAction.plan_id == plan.id))
    assert action.due_date == date(2029, 12, 31)
    assert action.start_raw == '2027-01-15' and action.due_raw == '2029-12-31'
    assert action.action_code == 'KOS1.1.1' and action.action_code_raw == 'KOS.1.1.1'
    assert plan.approval_date == date(2026, 10, 18) and plan.status == 'draft'
    assert plan.source_filename == 'plan.xlsx' and len(plan.source_sha256) == 64
    standards = db.scalars(select(ICStandard)).all()
    assert {s.code: s.component_code for s in standards} == STANDARD_COMPONENTS
    assert db.get(ICStandard, 'KOS1').title == 'Etik Değerler'
    assert db.scalar(select(func.count()).select_from(ICImportAudit)) == 2


@pytest.mark.parametrize('raw, expected', [('KOS.2.1', 'KOS2.1'), ('KOS 2.1', 'KOS2.1'),
    ('KOS 01.01.02', 'KOS1.1.2'), ('bis.13.1', 'BİS13.1'), ('İS 17.1.1', 'İS17.1.1')])
def test_code_normalisation(raw, expected):
    assert ic.normalize_code(raw) == expected


@pytest.mark.parametrize('raw, expected', [('1-A2-B', ['A', 'B']), ('1-A\n2-B', ['A', 'B']),
    ('A;B\nC', ['A', 'B', 'C']), ('1 – A2 – B', ['A', 'B'])])
def test_numbered_parties(raw, expected):
    assert ic.split_units(raw) == expected


@pytest.mark.parametrize('change, fragment', [
    ({'1': 'KOS99', '2': 'KOS99.1', '5': 'KOS99.1.1'}, 'Standart kodu'),
    ({'2': 'KOS2.1'}, 'Genel şart kodu standartla'),
    ({'7': ''}, 'sorumlu birimi yok'),
    ({'7': 'Bilinmeyen Müdürlük'}, 'eşleştirilemedi'),
    ({'6': ''}, 'eylem metni boş'),
    ({'0': 'RDS'}, 'Bileşen ve standart'),
    ({'5': ''}, 'Eylem kodu eksik'),
    ({'3': ''}, 'Genel şart metni boş'),
])
def test_blocking_validation_rules(db, sample, change, fragment):
    _, report = stage(db, sample, workbook([action_row(**change)]))
    assert not report['can_commit']
    assert any(i['level'] == 'blocking' and fragment in i['message'] for i in report['issues'])
    for issue in report['issues']:
        assert set(issue) == {'level', 'sheet', 'row', 'column', 'value', 'message', 'todo'}


def test_duplicate_action_codes_normalized(db, sample):
    rows = [action_row(), action_row(**{'5': 'KOS 1.1.1'})]
    _, report = stage(db, sample, workbook(rows))
    assert [i['column'] for i in report['issues']] == ['F']
    assert 'yinelenmiş' in report['issues'][0]['message']


@pytest.mark.parametrize('mutate', [lambda b: setattr(b['Plan']['H1'], 'value', 'Sorumlu'),
    lambda b: b.remove(b['Plan']), lambda b: b.remove(b['Bilgi']),
    lambda b: setattr(b['Plan']['N1'], 'value', 'Eylem Kod No')])
def test_missing_or_renamed_headers_and_sheets(db, sample, mutate):
    _, report = stage(db, sample, workbook(mutate=mutate))
    assert not report['can_commit']
    assert len(report['issues']) == 1
    assert not any('boş' in i['message'] for i in report['issues'])


@pytest.mark.parametrize('data, filename', [(b'not an xlsx', 'bad.xlsx'), (b'PKbroken', 'bad.xlsx'),
                                           (b'not an xlsx', 'bad.xls')])
def test_invalid_xlsx_checklist(db, sample, data, filename):
    _, report = ic.create_import(db, sample['strategy'], data, filename)
    assert [i['sheet'] for i in report['issues']] == ['Dosya']
    assert not report['can_commit']


@pytest.mark.parametrize('change, fragment', [({'11': '2027 ikinci yarıyıl'}, 'Tarih kesin'),
    ({'10': 'gelecek ay'}, 'Tarih kesin'), ({'5': 'KOS2.1.1'}, 'önek')])
def test_warning_rules_allow_commit_and_keep_raw(db, sample, change, fragment):
    data = workbook([action_row(**change)])
    batch, report = stage(db, sample, data)
    assert report['can_commit']
    assert len(report['issues']) == 1 and fragment in report['issues'][0]['message']
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    a = db.scalar(select(ICAction).where(ICAction.plan_id == plan.id))
    if '11' in change:
        assert a.due_date is None and a.due_raw == change['11']


def test_conflicting_condition_warning_is_traceable(db, sample):
    rows = [action_row(), action_row(**{'3': 'Farklı metin', '4': 'Farklı durum', '5': 'KOS 1.1.2'})]
    batch, report = stage(db, sample, workbook(rows))
    assert report['can_commit'] and len(report['issues']) == 1
    assert 'ilk satır' in report['issues'][0]['message']
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    condition = db.scalar(select(ICCondition).where(ICCondition.plan_id == plan.id))
    assert condition.text == rows[0][3] and condition.current_state_text == rows[0][4]
    assert 'Farklı metin' in batch.source_json


def test_condition_without_action_and_continuous(db, sample):
    no_action = ['KOS', 'KOS1', 'KOS 1.2', 'İkinci şart', 'Mevcut durumda makul güvence sağlanmıştır. Yeni bir eylem belirlemeye ihtiyaç yoktur.'] + [''] * 9
    plan = saved(db, sample, workbook([action_row(), no_action]))
    assert db.scalar(select(func.count()).select_from(ICCondition).where(ICCondition.plan_id == plan.id)) == 2
    a = db.scalar(select(ICAction).where(ICAction.plan_id == plan.id))
    assert a.is_continuous and a.due_date is None and a.due_raw == 'Sürekli'


def test_acceptance_fixture_mapping_reuse_counts_and_scoping(db, sample, client):
    b = Department(organization_id=sample['department'].organization_id, name='Diğer Müdürlüğü')
    db.add(b)
    db.commit()
    rows = [action_row(), action_row(**{'2': 'KOS1.2', '5': 'KOS1.2.1', '7': '1-Test Müdürlüğü2-Diğer Müdürlüğü'}),
        action_row(**{'2': 'KOS1.3', '5': 'KOS1.3.1', '7': 'Üst Yönetim'}),
        action_row(**{'2': 'KOS1.4', '5': 'KOS1.4.1', '7': 'Tüm Müdürlükler'}),
        action_row(**{'2': 'KOS1.5', '5': 'KOS1.5.1', '7': 'Yeni Birim'}),
        action_row(**{'2': 'KOS1.6', '5': 'KOS1.6.1', '6': 'Yalnız diğer birimin gizli eylemi', '7': 'Diğer Müdürlüğü', '8': 'Test Müdürlüğü'}),
        action_row(**{'0': 'RDS', '1': 'RDS99', '2': 'RDS99.1', '5': 'RDS99.1.1', '7': 'Diğer Müdürlüğü'})]
    rows += [['KOS', 'KOS1', 'KOS1.7', 'Eylemsiz şart', 'Makul güvence sağlanmıştır.'] + [''] * 9]
    batch, report = stage(db, sample, workbook(rows))
    assert len(report['issues']) == 2, report['issues']
    assert {i['column'] for i in report['issues']} == {'B', 'H'}
    with pytest.raises(HTTPException, match='422'):
        ic.commit_import(db, sample['strategy'], batch.id)
    ic.map_unit(db, sample['strategy'], batch, 'Yeni Birim', sample['department'].id)
    db.commit()
    rows[6][1:3] = ['RDS5', 'RDS5.1']
    rows[6][5] = 'RDS5.1.1'
    fixed, report = stage(db, sample, workbook(rows))
    assert report['issues'] == []
    plan = ic.commit_import(db, sample['strategy'], fixed.id)
    ic.activate_plan(db, sample['strategy'], plan.id)
    db.commit()
    assert db.scalar(select(func.count()).select_from(ICUnitAlias)) == 1
    marker = db.scalar(select(ICActionParty).where(ICActionParty.all_departments.is_(True)))
    assert marker.department_id is None and marker.body_name is None
    own = ic.directorate_actions(db, sample['user'], 'responsible')
    coop = ic.directorate_actions(db, sample['user'], 'cooperating')
    assert {a.action_code for a in own} == {'KOS1.1.1', 'KOS1.2.1', 'KOS1.4.1', 'KOS1.5.1'}
    assert [a.action_code for a in coop] == ['KOS1.6.1']
    payload = catalogue(db, plan)
    assert payload['counts']['KOS']['conditions'] == 7 and payload['counts']['KOS']['actions'] == 6
    assert payload['counts']['RDS']['conditions'] == 1 and payload['counts']['RDS']['actions'] == 1
    assert len(payload['counts']['KOS']['departments'][sample['department'].name]) == 4
    assert len(payload['counts']['KOS']['departments'][b.name]) == 3
    filtered = catalogue(db, plan, department_id=b.id, q='gizli')
    assert len(filtered['groups']['KOS']['KOS1']['conditions']) == 1
    response = client.get('/internal-control/our-actions')
    assert response.status_code == 200 and 'İşbirliği yapacağımız' in response.text
    assert client.get(f'/internal-control/plans/{plan.id}').status_code == 403


def test_non_directorate_mapping_reused(db, sample):
    data = workbook([action_row(**{'7': 'Yerel Çalışma Ekibi'})])
    batch, report = stage(db, sample, data)
    assert not report['can_commit']
    ic.map_unit(db, sample['strategy'], batch, 'Yerel Çalışma Ekibi', body_name='Çalışma ekibi')
    db.commit()
    report = ic.dry_run(db, sample['strategy'], batch)
    assert report['can_commit']
    plan = ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    party = db.scalar(select(ICActionParty).join(ICAction).where(ICAction.plan_id == plan.id))
    assert party.body_name == 'Çalışma ekibi' and party.department_id is None
    assert stage(db, sample, data)[1]['can_commit']


def test_atomic_rollback_after_partial_write(db, sample, monkeypatch):
    batch, report = stage(db, sample)
    def failure(*args, **kwargs):
        raise RuntimeError('injected audit failure after plan/actions')
    monkeypatch.setattr(ic, 'audit', failure)
    with pytest.raises(RuntimeError):
        ic.commit_import(db, sample['strategy'], batch.id)
    db.commit()
    assert db.scalar(select(func.count()).select_from(ICPlan)) == 0
    assert db.scalar(select(func.count()).select_from(ICCondition)) == 0
    assert db.scalar(select(func.count()).select_from(ICAction)) == 0
    assert db.scalar(select(func.count()).select_from(ICActionParty)) == 0
    assert db.scalar(select(func.count()).select_from(ICImportAudit)) == 1


def test_single_active_plan_and_version_lineage(db, sample):
    first = saved(db, sample)
    ic.activate_plan(db, sample['strategy'], first.id)
    db.commit()
    old_action_id = db.scalar(select(ICAction.id).where(ICAction.plan_id == first.id))
    batch, _ = stage(db, sample, workbook(metadata=['Test Uyum Planı', 2027, 2029, 'v2', '']))
    second = ic.commit_import(db, sample['strategy'], batch.id, supersedes_id=first.id)
    db.commit()
    assert second.status == 'draft' and second.supersedes_id == first.id
    ic.activate_plan(db, sample['strategy'], second.id)
    db.commit()
    db.refresh(first)
    assert first.status == 'archived' and second.status == 'active'
    assert db.get(ICAction, old_action_id) is not None
    assert db.scalar(select(func.count()).select_from(ICPlan).where(ICPlan.status == 'active')) == 1
    with pytest.raises(HTTPException):
        ic.activate_plan(db, sample['strategy'], first.id)
    with pytest.raises(HTTPException):
        ic.commit_import(db, sample['strategy'], batch.id)


def test_directorate_isolation_no_import_and_no_inactive_accounts(db, sample, client):
    b = Department(organization_id=sample['department'].organization_id, name='Diğer Müdürlüğü')
    db.add(b)
    db.commit()
    plan = saved(db, sample, workbook([action_row(**{'7': b.name, '6': 'Başka birime özel eylem'})]))
    ic.activate_plan(db, sample['strategy'], plan.id)
    db.commit()
    assert 'Başka birime özel eylem' not in client.get('/internal-control/our-actions').text
    assert client.get('/internal-control/template.xlsx').status_code == 403
    assert client.post('/internal-control/imports', files={'file': ('plan.xlsx', workbook())}).status_code == 403
    for path in [f'/internal-control/imports/{plan.import_id}', f'/internal-control/imports/{plan.import_id}/validation.xlsx']:
        assert client.get(path).status_code == 403
    for suffix in ['commit', 'validate', 'mappings']:
        assert client.post(f'/internal-control/imports/{plan.import_id}/{suffix}').status_code == 403
    assert client.post(f'/internal-control/plans/{plan.id}/activate').status_code == 403
    for role in ['admin', 'viewer']:
        sample['user'].role = role
        db.commit()
        assert client.get('/internal-control').status_code == 403
    sample['user'].role = 'department'
    sample['user'].is_active = False
    db.commit()
    assert client.get('/internal-control/our-actions').status_code == 403


def test_strategy_ui_round_trip(client, db, sample):
    from app.main import app, require_user
    app.dependency_overrides[require_user] = lambda: sample['strategy']
    assert client.get('/internal-control').status_code == 200
    assert client.get('/internal-control/template.xlsx').status_code == 200
    response = client.post('/internal-control/imports', files={'file': ('plan.xlsx', workbook())})
    assert response.status_code == 200 and 'Engelleyici madde yok' in response.text
    batch = db.scalar(select(ICImport))
    assert client.get(f'/internal-control/imports/{batch.id}/validation.xlsx').status_code == 200
    response = client.post(f'/internal-control/imports/{batch.id}/commit')
    assert response.status_code == 200 and 'Etik Değerler' in response.text
    assert 'Kaydeden: Strateji' in response.text
    plan = db.scalar(select(ICPlan))
    assert client.get(f'/internal-control/plans/{plan.id}?component=&standard=&department_id=&q=').status_code == 200
    assert client.post(f'/internal-control/plans/{plan.id}/activate').status_code == 200
    assert 'Etkin' in client.get('/internal-control').text


def test_validation_export_never_executes_formula(db, sample):
    _, report = stage(db, sample, workbook([action_row(**{'7': '=HYPERLINK("bad")'})]))
    book = load_workbook(BytesIO(ic.validation_bytes(report)), data_only=False)
    assert book['Doğrulama']['E2'].value == '=HYPERLINK("bad")'
    assert book['Doğrulama']['E2'].data_type == 's'


def test_model_immutable_history(db, sample):
    plan = saved(db, sample)
    action = db.scalar(select(ICAction))
    action.text = 'overwrite'
    with pytest.raises(ValueError, match='değiştirilemez'):
        db.flush()
    db.rollback()
    plan.title = 'overwrite'
    with pytest.raises(ValueError, match='değiştirilemez'):
        db.flush()
    db.rollback()


def test_migration_upgrade_downgrade_preserves_data_and_sql_guards():
    url = os.environ['CIVENSA_TEST_POSTGRES_URL']
    base = create_engine(url)
    schema = 'ic_migration_' + uuid4().hex
    with base.begin() as conn:
        conn.execute(CreateSchema(schema))
    schema_url = make_url(url).update_query_dict({'options': '-csearch_path=' + schema}).render_as_string(hide_password=False)
    env = {**os.environ, 'DATABASE_URL': schema_url}
    engine = create_engine(schema_url)
    def migrate(direction, target):
        result = subprocess.run([sys.executable, '-m', 'alembic', direction, target], cwd=ROOT, env=env, text=True, capture_output=True)
        assert result.returncode == 0, result.stderr
    try:
        migrate('upgrade', 'c044a1b2c3d4')
        with engine.begin() as conn:
            conn.execute(text("INSERT INTO organizations(id,name) VALUES (1,'Test')"))
            conn.execute(text("INSERT INTO departments(id,organization_id,name) VALUES (1,1,'Test Müdürlüğü')"))
            conn.execute(text("INSERT INTO users(id,email,full_name,role,password_hash,is_active) VALUES (1,'ic@test.invalid','Test','strategy','unused',true)"))
        migrate('upgrade', 'head')
        with Session(engine, expire_on_commit=False) as db:
            actor = db.get(User, 1)
            batch, report = ic.create_import(db, actor, workbook(), 'plan.xlsx')
            assert report['can_commit']
            plan = ic.commit_import(db, actor, batch.id)
            db.commit()
            plan_id = plan.id
            for statement in ["UPDATE ic_actions SET text='changed'", 'DELETE FROM ic_conditions',
                              "UPDATE ic_plans SET title='changed'", 'DELETE FROM ic_import_audits',
                              "INSERT INTO ic_conditions(plan_id,standard_code,standard_title,condition_code,condition_code_raw,text,current_state_text,sort_order) VALUES (1,'KOS1','','KOS1.9','KOS1.9','x','',99)",
                              "INSERT INTO ic_action_parties(action_id,role,department_id,body_name,all_departments,raw_name,sort_order) SELECT id,'responsible',1,NULL,false,'x',99 FROM ic_actions"]:
                with pytest.raises(DBAPIError):
                    with db.begin_nested():
                        db.execute(text(statement))
            ic.activate_plan(db, actor, plan.id)
            db.commit()
            # Unique partial index rejects a second active plan even through SQL.
            second_batch, _ = ic.create_import(db, actor, workbook(), 'v2.xlsx')
            second = ic.commit_import(db, actor, second_batch.id)
            db.commit()
            with pytest.raises(IntegrityError):
                with db.begin_nested():
                    db.execute(text("UPDATE ic_plans SET status='active' WHERE id=:id"), {'id': second.id})
        migrate('downgrade', 'c044a1b2c3d4')
        with engine.connect() as conn:
            assert conn.execute(text('SELECT count(*) FROM ic_actions')).scalar() == 2
            assert conn.execute(text('SELECT count(*) FROM organizations')).scalar() == 1
        migrate('upgrade', 'head')
        with engine.connect() as conn:
            assert conn.execute(text('SELECT count(*) FROM ic_standards')).scalar() == 18
            assert conn.execute(text('SELECT status FROM ic_plans WHERE id=:id'), {'id': plan_id}).scalar() == 'active'
        heads = subprocess.run([sys.executable, '-m', 'alembic', 'heads'], cwd=ROOT, text=True, capture_output=True)
        assert heads.returncode == 0 and len(heads.stdout.strip().splitlines()) == 1
        from alembic.config import Config
        from alembic.script import ScriptDirectory
        scripts = ScriptDirectory.from_config(Config(str(ROOT / 'alembic.ini')))
        assert scripts.get_revision('c048a1b2c3d4').down_revision == 'c044a1b2c3d4'
    finally:
        engine.dispose()
        with base.begin() as conn:
            conn.execute(DropSchema(schema, cascade=True))
        base.dispose()


def test_concurrent_activation_serializes_first_and_successor(db, sample):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    first, second = saved(db, sample), saved(db, sample)
    actor_id = sample['strategy'].id
    barrier = Barrier(2)
    def activate(plan_id):
        with Session(db.get_bind(), expire_on_commit=False) as concurrent:
            barrier.wait(timeout=10)
            ic.activate_plan(concurrent, concurrent.get(User, actor_id), plan_id)
            concurrent.commit()
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [pool.submit(activate, pid) for pid in (first.id, second.id)]
        for future in futures:
            future.result(timeout=20)
    db.expire_all()
    assert db.scalar(select(func.count()).select_from(ICPlan).where(ICPlan.status == 'active')) == 1
    assert db.scalar(select(func.count()).select_from(ICPlan).where(ICPlan.status == 'archived')) == 1
    assert db.scalar(select(func.count()).select_from(ICImportAudit).where(ICImportAudit.result == 'activation')) == 2


def test_browser_upload_mapping_save_activate_and_filters(db, sample, tmp_path):
    from fastapi.testclient import TestClient
    from playwright.sync_api import sync_playwright, expect
    from app.main import app
    from app.db import get_db
    from app.security import hash_password
    sample['strategy'].password_hash = hash_password('ic-browser-test-only')
    db.commit()
    def database():
        yield db
    app.dependency_overrides[get_db] = database
    try:
        with sync_playwright() as p:
            options = {'headless': True}
            if os.getenv('CIVENSA_CHROMIUM_PATH'):
                options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
            browser = p.chromium.launch(**options)
            with TestClient(app) as client, browser.new_context(viewport={'width': 1366, 'height': 900}) as context:
                assert client.post('/login', data={'email': sample['strategy'].email,
                    'password': 'ic-browser-test-only'}, follow_redirects=False).status_code == 303
                def serve(route):
                    request = route.request
                    response = client.request(request.method, request.url, content=request.post_data_buffer,
                        headers={'content-type': request.headers.get('content-type', '')}, follow_redirects=True)
                    headers = dict(response.headers)
                    headers.pop('set-cookie', None)
                    route.fulfill(status=response.status_code, headers=headers, body=response.content)
                context.route('http://testserver/**', serve)
                page = context.new_page()
                page.goto('http://testserver/internal-control')
                page.locator('input[name="file"]').set_input_files({'name': 'plan.xlsx', 'mimeType':
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    'buffer': workbook([action_row(**{'7': 'Yeni Birim'})])})
                page.get_by_role('button', name='Yükle ve doğrula', exact=True).click()
                expect(page.get_by_role('heading', name='Eşleşmeyen birimler')).to_be_visible()
                expect(page.get_by_role('button', name='Uyarıları kabul et ve taslak kaydet')).to_be_disabled()
                page.locator('select[name="department_id"]').select_option(str(sample['department'].id))
                page.get_by_role('button', name='Eşleştir ve tekrar doğrula').click()
                expect(page.get_by_text('Engelleyici madde yok. Plan kaydedilebilir.')).to_be_visible()
                page.get_by_role('button', name='Uyarıları kabul et ve taslak kaydet').click()
                expect(page.get_by_role('heading', name='Test Uyum Planı', exact=True)).to_be_visible()
                page.get_by_role('button', name='Planı etkinleştir').click()
                expect(page.get_by_role('button', name='Planı etkinleştir')).to_have_count(0)
                page.locator('select[name="component"]').select_option('KOS')
                page.get_by_role('button', name='Uygula', exact=True).click()
                expect(page.get_by_text('KOS.1.1.1 · Eğitim düzenle', exact=True)).to_be_visible()
                page.screenshot(path=str(tmp_path / 'ic-plan.png'), full_page=True)
            browser.close()
    finally:
        app.dependency_overrides.clear()


def test_corrupt_worksheet_is_a_blocking_checklist(db, sample):
    from zipfile import ZipFile
    original = workbook()
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, 'w') as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                data = data[:len(data) // 2]
            target.writestr(member, data)
    batch, report = stage(db, sample, output.getvalue())
    assert not report['can_commit']
    assert any(i['sheet'] == 'Dosya' for i in report['issues'])


def test_long_codes_fail_validation_without_database_error(db, sample):
    _, report = stage(db, sample, workbook([action_row(**{'2': 'KOS1.' + '9' * 5000,
                                                        '5': 'KOS1.1.' + '9' * 5000})]))
    assert not report['can_commit']
    assert {i['column'] for i in report['issues'] if i['level'] == 'blocking'} == {'C', 'F'}


def test_reordered_headers_report_the_actual_source_column(db, sample):
    def reorder(book):
        for row in book['Plan'].iter_rows():
            row[0].value, row[7].value = row[7].value, row[0].value
    _, report = stage(db, sample, workbook([action_row(**{'7': 'Eşleşmeyen Birim'})], mutate=reorder))
    assert len(report['issues']) == 1
    assert report['issues'][0]['row'] == 2 and report['issues'][0]['column'] == 'A'
    assert report['issues'][0]['value'] == 'Eşleşmeyen Birim'
