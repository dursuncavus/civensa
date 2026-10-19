"""CIV-041: economic tables require a card and retain code/name categories."""
import io
from copy import deepcopy
from decimal import Decimal

import pytest
from docx import Document
from sqlalchemy import select

from app.data_importer import parse_performance_docx
from app.import_service import import_validation_errors
from app.models import EconomicValue, ImportBatch
from test_pp_import_submit_ready import source_document, import_parsed
from test_word_import import import_client, import_context

pytestmark = pytest.mark.postgres


def save_document(doc):
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def rejected_document(case):
    doc = Document(io.BytesIO(source_document()))
    if case == 'after_completed_card':
        doc.element.body.append(deepcopy(doc.tables[2]._tbl))
        number = 4
    elif case == 'before_card':
        doc.tables[0]._tbl.addprevious(deepcopy(doc.tables[2]._tbl))
        number = 1
    else:
        doc.tables[2].cell(0, 3).text = (
            '2028 Bütçe' if case == 'wrong_year' else '2027 Tahmin')
        number = 3
    return save_document(doc), number


@pytest.mark.parametrize('case', ['after_completed_card', 'before_card', 'wrong_year', 'forecast_only'])
def test_unattached_economic_table_is_recorded_and_blocks_validation(case):
    data, number = rejected_document(case)
    parsed = parse_performance_docx(data)
    headers = [cell.text for cell in Document(io.BytesIO(data)).tables[number-1].rows[0].cells]
    rejected = parsed['rejected_economic_tables']
    assert len(rejected) == 1
    assert rejected[0]['table_number'] == number
    assert rejected[0]['header_row'] == headers
    assert parsed['summary']['economic_table_count'] == (1 if number == 3 else 2)
    attached = parsed['departments'][0].get('economic_tables', [])
    assert len(attached) == (0 if number == 3 else 1)
    assert all(table['table_number'] != number for table in attached)
    assert parsed['summary']['economic_value_count'] == (0 if number == 3 else 30)
    errors = import_validation_errors(parsed)
    assert len(errors) == 1
    assert f'Kart tablosuna bağlı olmayan ekonomik tablo (tablo {number}):' in errors[0]
    assert ' | '.join(headers) in errors[0]
    assert 'Bu tabloyu ilgili müdürlüğün kartından hemen sonra yerleştirin veya kaynaktan çıkarın.' in errors[0]
    if number == 3:
        assert '2027 Bütçe' in errors[0]
    # Recorded source evidence also protects a stored preview without its error list.
    parsed['errors'] = []
    assert import_validation_errors(parsed) == errors


@pytest.mark.parametrize('case', ['after_completed_card', 'wrong_year'])
def test_rejected_economic_table_prevents_preview_apply(db, import_client, import_context, case):
    data, _ = rejected_document(case)
    response = import_client.post(
        f'/api/periods/{import_context.period.id}/data-import-preview',
        files={'file': ('economic.docx', data)})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['can_apply'] is False
    assert 'Kart tablosuna bağlı olmayan ekonomik tablo' in preview['errors'][0]
    detail = import_client.get(f"/api/import-batches/{preview['batch_id']}").json()
    assert detail['errors'] == preview['errors']
    assert import_client.post(f"/api/import-batches/{preview['batch_id']}/apply").status_code == 409
    assert db.get(ImportBatch, preview['batch_id']).status == 'preview'
    assert db.scalar(select(EconomicValue)) is None


def test_invalid_card_does_not_attach_economic_table_to_an_earlier_card():
    doc = Document(io.BytesIO(source_document()))
    invalid_card = deepcopy(doc.tables[0]._tbl)
    doc.tables[2]._tbl.addprevious(invalid_card)
    for cell in doc.tables[2].rows[2].cells[2:]:
        cell.text = ''
    parsed = parse_performance_docx(save_document(doc))
    assert not parsed['departments'][0].get('economic_tables')
    assert parsed['rejected_economic_tables'][0]['table_number'] == 4
    assert any('Tablo 3: İdare/Birim Adı' in error for error in import_validation_errors(parsed))
    assert any('ekonomik tablo (tablo 4)' in error for error in import_validation_errors(parsed))


def test_wrong_year_table_does_not_leave_a_pending_card_for_the_next_table():
    data, _ = rejected_document('wrong_year')
    doc = Document(io.BytesIO(data))
    doc.element.body.append(deepcopy(Document(io.BytesIO(source_document())).tables[2]._tbl))
    parsed = parse_performance_docx(save_document(doc))
    assert not parsed['departments'][0].get('economic_tables')
    assert [table['table_number'] for table in parsed['rejected_economic_tables']] == [3, 4]
    assert parsed['summary']['economic_value_count'] == 0


@pytest.mark.parametrize('layout', ['code_name', 'code_only', 'name_first', 'name_second'])
def test_economic_categories_and_real_values_are_imported(db, import_context, layout):
    doc = Document(io.BytesIO(source_document()))
    old = doc.tables[2]
    table = doc.add_table(rows=0, cols=7)
    headers = ['Ekonomik Kod', 'Ekonomik Ad', *[c.text for c in old.rows[0].cells[1:]]]
    for cell, text in zip(table.add_row().cells, headers):
        cell.text = text
    categories = [
        ('01', 'Personel Gideri', 'personnel', '80'),
        ('02', 'Sosyal Güvenlik Kurumuna Devlet Primi Giderleri', 'sgk', '20'),
        ('03', 'Mal ve Hizmet Alım Giderleri', 'goods_services', '100'),
        ('04', 'Faiz Giderleri', 'interest', '4'),
        ('05', 'Cari Transferler', 'current_transfers', '5'),
        ('06', 'Sermaye Giderleri', 'capital', '6'),
        ('07', 'Sermaye Transferleri', 'capital_transfers', '7'),
        ('08', 'Borç Verme', 'lending', '8'),
        ('09', 'Yedek Ödenekler', 'reserve', '9'),
    ]
    for code, name, _, amount in categories:
        labels = {'code_name': (code, name), 'code_only': (code, ''),
                  'name_first': (name, ''), 'name_second': ('', name)}[layout]
        for cell, text in zip(table.add_row().cells, (*labels, amount, amount, amount, '0', '0')):
            cell.text = text
    for name, amount in [('BÜTÇE İÇİ TOPLAM KAYNAK', '239'),
                         ('BÜTÇE DIŞI TOPLAM KAYNAK', '0'),
                         ('FAALİYET MALİYETİ TOPLAMI', '239')]:
        for cell, text in zip(table.add_row().cells, (name, '', amount, amount, amount, '0', '0')):
            cell.text = text
    old._tbl.addprevious(table._tbl)
    old._tbl.getparent().remove(old._tbl)
    parsed = parse_performance_docx(save_document(doc))
    assert parsed['errors'] == []
    economic = parsed['departments'][0]['economic_tables'][0]
    assert set(economic['categories']) == {category for _, _, category, _ in categories}
    assert len(economic['values']) == 90
    assert all(value['state'] == 'verified' for value in economic['values'])
    _, result = import_parsed(db, import_context, parsed)
    assert result['economic_values'] == 90
    rows = db.scalars(select(EconomicValue).where(
        EconomicValue.year == 2027, EconomicValue.funding_type == 'inside')).all()
    assert {row.category: row.amount for row in rows} == {
        category: Decimal(amount) for _, _, category, amount in categories}
    assert all(row.state == 'verified' and row.raw_text and row.source_table_number == 3 for row in rows)
