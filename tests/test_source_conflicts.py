"""CIV-011: explicit financial decisions with real login and CSRF on PostgreSQL."""
import io
import json
from copy import deepcopy
from decimal import Decimal
from importlib import import_module

import pytest
from docx import Document
from fastapi import HTTPException
from sqlalchemy import select, text

from app.data_importer import parse_performance_docx
from app.import_service import build_department_plan
from app.models import ActivityCostValue, AuditLog, DepartmentCostValue, ImportBatch
from app.program_service import save_activity_costs
from test_live_pilot_hardening import live_client, login
from test_word_import import word_document

pytestmark = pytest.mark.postgres


def duplicated_personnel(second='30,00'):
    doc = Document(io.BytesIO(word_document()))
    doc.tables[1].cell(4, 1).text = '20,00'
    duplicate = deepcopy(doc.tables[1]._tbl)
    doc.element.body.append(duplicate)
    doc.tables[2].cell(4, 1).text = second
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def source_total(db, sample, *, matching=False):
    activity = sample['activity']
    total = next(v for v in activity.cost_values if (v.year, v.value_kind, v.funding_type) == (2027, 'budget', 'total'))
    total.amount = total.source_amount = Decimal('140.00')
    total.source_raw_text = total.raw_text = '140,00 TL'
    total.state = 'needs_review'
    inside = next(v for v in activity.cost_values if (v.year, v.value_kind, v.funding_type) == (2027, 'budget', 'inside'))
    outside = next(v for v in activity.cost_values if (v.year, v.value_kind, v.funding_type) == (2027, 'budget', 'outside'))
    inside.amount = Decimal('100.00')
    outside.amount = Decimal('40.00' if matching else '33.50')
    db.commit()
    return total


@pytest.mark.parametrize('chosen', ['20,00', '30,00'])
def test_import_conflict_http_decision_apply_revert(db, sample, live_client, chosen):
    browser = live_client()
    token = login(browser, sample['strategy'].email)
    headers = {'X-CSRF-Token': token}
    sample['period'].template_status = 'verified'
    db.commit()
    response = browser.post(f"/api/periods/{sample['period'].id}/data-import-preview",
                            headers=headers, files={'file': ('kaynak.docx', duplicated_personnel())})
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview['can_apply'] is False
    assert len(preview['conflicts']) == 1
    conflict = preview['conflicts'][0]
    assert conflict['department'] == 'Park ve Bahçeler Müdürlüğü'
    assert all(part in conflict['message'] for part in ('2027', 'Bütçe', 'Personel ve SGK', '20,00', '30,00', 'Tablo 2, satır 5', 'Tablo 3, satır 5'))
    assert [c['locations'][0]['table_number'] for c in conflict['candidates']] == [2, 3]
    assert [c['locations'][0]['row_number'] for c in conflict['candidates']] == [5, 5]
    batch_id = preview['batch_id']
    assert browser.post(f'/api/import-batches/{batch_id}/apply', headers=headers).status_code == 409
    endpoint = f"/api/import-batches/{batch_id}/conflicts/{conflict['id']}/decision"
    assert browser.post(endpoint, json={'chosen_amount': '20,00', 'reason': 'Kaynak kontrol edildi.'}).status_code == 403
    for body in ({'chosen_amount': '20,00'}, {'chosen_amount': '20,00', 'reason': 'kısa'},
                 {'chosen_amount': '40,00', 'reason': 'Kaynak kontrol edildi.'}):
        assert browser.post(endpoint, headers=headers, json=body).status_code == 422
    decision_response = browser.post(endpoint, headers=headers, json={'chosen_amount': chosen, 'reason': 'Seçilen kaynak tutarı kontrol edildi.'})
    assert decision_response.status_code == 200, decision_response.text
    assert decision_response.json()['can_apply'] is True
    decision = decision_response.json()['decisions'][0]
    assert decision['actor'] == sample['strategy'].id and decision['time']
    assert decision['candidates'] == conflict['candidates']
    audit = db.scalar(select(AuditLog).where(AuditLog.action == 'import_source_conflict_decision'))
    assert audit.user_id == sample['strategy'].id and json.loads(audit.details) == decision
    assert browser.post(f'/api/import-batches/{batch_id}/apply', headers=headers).status_code == 200
    row = db.scalar(select(DepartmentCostValue).where(DepartmentCostValue.department_id != sample['department'].id,
                     DepartmentCostValue.year == 2027, DepartmentCostValue.value_kind == 'budget'))
    assert row.amount == Decimal(chosen.replace(',', '.')) and row.raw_text == chosen
    response = browser.post(f'/api/import-batches/{batch_id}/revert', headers=headers)
    assert response.status_code == 200, response.text
    detail = browser.get(f'/api/import-batches/{batch_id}').json()
    assert detail['status'] == 'reverted' and detail['decisions'] == [decision]
    db.refresh(audit)
    assert json.loads(audit.details) == decision


def test_equal_duplicates_keep_locations_without_conflict(db):
    parsed = parse_performance_docx(duplicated_personnel('20,00'))
    assert all('table_number' in value and 'row_number' in value for dept in parsed['departments']
               for activity in dept['activities'] for value in activity['costs'])
    plan = build_department_plan(db, parsed)
    assert plan['conflicts'] == []
    row = next(r for r in plan['parsed']['departments'][0]['personnel'] if r['year'] == 2027)
    assert Decimal(row['amount']) == Decimal('20.00')
    assert [(loc['table_number'], loc['row_number']) for loc in row['candidates'][0]['locations']] == [(2, 5), (3, 5)]


def test_total_disagreement_http_unchanged_and_explicit_decision(db, sample, live_client):
    total = source_total(db, sample)
    browser = live_client()
    token = login(browser, sample['user'].email)
    headers = {'X-CSRF-Token': token}
    endpoint = f"/api/activities/{sample['activity'].id}"
    body = {'cost_2027_budget_inside': '100,00', 'cost_2027_budget_outside': '33,50'}
    before = (total.amount, total.state, total.raw_text, total.source_amount, total.source_raw_text)
    assert browser.put(endpoint, json=body).status_code == 403
    response = browser.put(endpoint, headers=headers, json=body)
    assert response.status_code == 422, response.text
    message = response.json()['detail']
    assert all(part in message for part in ('Ortak faaliyet', '2027', 'Bütçe', '140,00', '133,50', 'Hesaplanan toplamı kullan', 'gerekçe'))
    db.refresh(total)
    assert (total.amount, total.state, total.raw_text, total.source_amount, total.source_raw_text) == before
    page = browser.get(f"/department/{sample['department'].id}")
    assert page.status_code == 200 and 'Kaynak: 140,00' in page.text and '133,50' in page.text
    assert 'confirm-source-total' in page.text
    # The old money-verification route must not bypass a reasoned total decision.
    verify = browser.post(f"/program/{sample['department'].id}/verify-money/activity/{total.id}",
                          headers=headers, data={'amount': '133,50'})
    assert verify.status_code == 422 and '140,00' in verify.json()['detail']
    body['cost_2027_budget_use_computed'] = True
    assert browser.put(endpoint, headers=headers, json=body).status_code == 422
    body['cost_2027_budget_reason'] = 'Finansman dağılımı kontrol edildi.'
    response = browser.put(endpoint, headers=headers, json=body)
    assert response.status_code == 200, response.text
    db.refresh(total)
    assert total.amount == Decimal('133.50') and total.state == 'verified'
    assert total.source_amount == Decimal('140.00') and total.source_raw_text == total.raw_text == '140,00 TL'
    audit = db.scalar(select(AuditLog).where(AuditLog.action == 'activity_source_total_decision'))
    decision = json.loads(audit.details)
    assert decision == json.loads(total.source_resolution_json)
    assert decision['actor'] == sample['user'].id and decision['activity'] == sample['activity'].id
    assert decision['source_total'] == '140.00' and decision['new_total'] == '133.50' and decision['time']
    # Source provenance is available through the existing structured snapshot.
    from app.program_service import department_snapshot
    costs = department_snapshot(db, sample['period'].id, sample['department'].id)['activities'][0]['costs']
    assert next(c for c in costs if c['id'] == total.id)['source_amount'] == '140.00'


def test_changed_split_refused_and_all_columns_atomic(db, sample):
    total = source_total(db, sample)
    before = {v.id: (v.amount, v.state) for v in sample['activity'].cost_values}
    with pytest.raises(HTTPException) as error:
        save_activity_costs(db, sample['activity'], 2027, {'cost_2027_budget_inside': '101,00', 'cost_2027_budget_outside': '33,50'}, actor=sample['user'])
    assert error.value.status_code == 422 and '134,50' in error.value.detail
    assert {v.id: (v.amount, v.state) for v in sample['activity'].cost_values} == before
    with pytest.raises(HTTPException):
        save_activity_costs(db, sample['activity'], 2027, {'cost_2027_budget_inside': '100,00', 'cost_2027_budget_outside': '33,50',
            'cost_2027_budget_use_computed': True, 'cost_2027_budget_reason': 'Finansman doğrulandı.', 'cost_2028_forecast_inside': 'geçersiz'}, actor=sample['user'])
    assert total.source_resolution_json == '{}'
    assert db.scalar(select(AuditLog).where(AuditLog.action == 'activity_source_total_decision')) is None


def test_matching_split_http_saves_normally(db, sample, live_client):
    total = source_total(db, sample, matching=True)
    browser = live_client()
    token = login(browser, sample['user'].email)
    response = browser.put(f"/api/activities/{sample['activity'].id}", headers={'X-CSRF-Token': token},
                           json={'cost_2027_budget_inside': '100,00', 'cost_2027_budget_outside': '40,00'})
    assert response.status_code == 200, response.text
    db.refresh(total)
    assert total.amount == total.source_amount == Decimal('140.00') and total.state == 'verified'
    assert total.source_raw_text == '140,00 TL'
    assert db.scalar(select(AuditLog).where(AuditLog.action == 'activity_source_total_decision')) is None


def test_migration_upgrade_and_history_blocks_downgrade(db, sample):
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    migration = import_module('migrations.versions.b8d011e6f420_preserve_source_financial_conflicts')
    schema = db.bind.get_execution_options()['schema_translate_map'][None]
    db.execute(text(f'SET LOCAL search_path TO "{schema}"'))
    for column in ('source_resolution_json', 'source_raw_text', 'source_amount'):
        db.execute(text(f'ALTER TABLE activity_cost_values DROP COLUMN {column}'))
    db.execute(text('ALTER TABLE import_batches DROP COLUMN decisions_json'))
    db.execute(text("UPDATE activity_cost_values SET amount=140, raw_text='140,00' WHERE year=2027 AND funding_type='total'"))
    with Operations.context(MigrationContext.configure(db.connection())):
        migration.upgrade()
    row = db.execute(text("SELECT amount, source_amount, source_raw_text, state FROM activity_cost_values WHERE year=2027 AND funding_type='total'")).one()
    assert row == (Decimal('140.00'), Decimal('140.00'), '140,00', 'needs_review')
    with Operations.context(MigrationContext.configure(db.connection())):
        migration.downgrade()
        migration.upgrade()
    batch = ImportBatch(period_id=sample['period'].id, filename='kaynak.docx', decisions_json='[{"reason":"Karar geçmişi"}]')
    db.add(batch); db.flush()
    with Operations.context(MigrationContext.configure(db.connection())):
        with pytest.raises(RuntimeError, match='karar geçmişi'):
            migration.downgrade()
    assert db.execute(text('SELECT decisions_json FROM import_batches WHERE id=:id'), {'id': batch.id}).scalar() == batch.decisions_json
