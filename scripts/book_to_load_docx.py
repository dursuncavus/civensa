#!/usr/bin/env python3
"""Convert print IDML tables to the first-load Word format, without InDesign.

Usage: python scripts/book_to_load_docx.py BOOK.idml name_map.json OUT.docx
The adjacent OUT.docx.report.json records mapping and source table/page provenance.
Unplaced stories are reported, rather than treated as printed book data.
"""
import argparse
from collections import Counter
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def norm(text):
    text = (text or '').replace('İ', 'i').replace('I', 'ı').lower()
    return re.sub(r'[^0-9a-zçğıöşü]+', '', text)


def table_type(table):
    first = norm(table['rows'][0][0]) if table['rows'] else ''
    for prefix, kind in [('bütçeyılı', 'card'), ('altprogramkapsamında', 'membership'),
                         ('faaliyetler', 'cost'), ('ekonomikad', 'economic')]:
        if first.startswith(prefix):
            return kind
    return 'other'


def extract_tables(path, *, excluded=None):
    """Read the ZIP directly; retain the legacy span grid and print spread order."""
    with ZipFile(path) as archive:
        def xml(src):
            return ET.fromstring(archive.read(src))

        design = xml('designmap.xml')
        from scripts.idml_printed import story_locations
        stories, story_pages, spread_pages = story_locations(archive, design)
        tables = []
        for story in stories:
            sid = story.get('Self')
            spread_index, page = story_pages.get(sid, (9999, '?'))
            for node in story.iter('Table'):
                ncols = int(node.attrib['ColumnCount'])
                nrows = sum(int(node.get(k, 0)) for k in ('BodyRowCount', 'HeaderRowCount', 'FooterRowCount'))
                rows = [[None] * ncols for _ in range(nrows)]
                for cell in node.findall('Cell'):
                    col, row = map(int, cell.attrib['Name'].split(':'))
                    rs, cs = int(cell.get('RowSpan', 1)), int(cell.get('ColumnSpan', 1))
                    parts = [n.text or '' if n.tag == 'Content' else ' '
                             for n in cell.iter() if n.tag in {'Content', 'Br'}]
                    rows[row][col] = re.sub(r'\s+', ' ', ''.join(parts)).strip()
                    for r in range(row, min(row + rs, nrows)):
                        for c in range(col, min(col + cs, ncols)):
                            if (r, c) != (row, col):
                                rows[r][c] = '→' if r == row else '↓'
                table = dict(id=node.get('Self'), story=sid, spread=spread_index, page=page, rows=rows)
                if sid in story_pages:
                    tables.append(table)
                elif excluded is not None:
                    excluded.append(dict(table, spread_pages=spread_pages.get(sid, []),
                                         reason='pasteboard' if sid in spread_pages else 'unplaced'))
        return tables


def generate_docx(idml, name_map, output):
    mapping = json.loads(Path(name_map).read_text(encoding='utf-8'))
    names = {norm(k): v for k, v in mapping['map'].items()}
    names.update({norm(v): v for v in mapping['map'].values()})
    excluded = []
    extracted = extract_tables(idml, excluded=excluded)
    kinds = {'card': 0, 'membership': 1, 'cost': 2, 'economic': 3}
    tables = [t for t in extracted if table_type(t) in kinds and t['page'] != '?']
    tables.sort(key=lambda t: (int(t['page']) if t['page'].isdigit() else 10**6,
                               t['spread'], kinds[table_type(t)], t['id']))
    card_pages = Counter(t['page'] for t in tables if table_type(t) == 'card')
    ambiguous = sorted(page for page, count in card_pages.items() if count > 1)
    if ambiguous:
        raise ValueError('Aynı sayfada birden fazla kart bulundu; faaliyet üyeliği sırası belirsiz. '
                         'Kitap kart yerleşimini kontrol edin: ' + ', '.join(ambiguous))
    report = dict(departments=[], unmapped=[], tables=[],
                  unplaced=[{'id': t['id'], 'type': table_type(t)} for t in excluded
                            if t['reason'] == 'unplaced' and table_type(t) in kinds],
                  non_printed=[dict(id=t['id'], story=t['story'], spread_pages=t['spread_pages'],
                                    header=t['rows'][0][:4] if t['rows'] else [], reason=t['reason'])
                               for t in excluded])
    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(9)
    current = None

    def system_name(raw):
        base = (raw or '').split('/')[-1].strip()
        if norm(base) not in names:
            report['unmapped'].append(raw)
        return names.get(norm(base), base)

    for source in tables:
        kind = table_type(source)
        rows = [list(row) for row in source['rows']]
        indicator_rows = set()
        if kind == 'card':
            for row in rows:
                if norm(row[0]).startswith(('idarebirim', 'sorumluidare')):
                    row[1] = system_name(row[1])
            dept = next(row[1] for row in rows if norm(row[0]).startswith('idarebirim'))
            if dept != current:
                doc.add_heading(dept, level=2)
                current = dept
                report['departments'].append(dept)
            start = next(i for i, row in enumerate(rows) if (row[0] or '').startswith('Performans Göstergesi'))
            end = next((i for i, row in enumerate(rows) if norm(row[0]).startswith('göstergeyeilişkin')), len(rows))
            indicator_rows = set(range(start, end))
        provenance = dict(table_number=len(report['tables']) + 1, page=source['page'],
                          id=source['id'], type=kind, department=current)
        report['tables'].append(provenance)
        table = doc.add_table(rows=len(rows), cols=max(map(len, rows)))
        ncols = len(table.columns)
        cells = table._cells
        table.style = 'Table Grid'
        caption = OxmlElement('w:tblCaption')
        caption.set(qn('w:val'), json.dumps(provenance, ensure_ascii=False))
        table._tbl.tblPr.append(caption)
        for ri, row in enumerate(rows):
            col = 0
            while col < len(row):
                span = 1
                while col + span < len(row) and row[col + span] == '→':
                    span += 1
                cell = cells[ri * ncols + col]
                if span > 1 and ri not in indicator_rows:
                    cell = cell.merge(cells[ri * ncols + col + span - 1])
                cell.text = '' if row[col] in (None, '→', '↓') else row[col]
                col += span
        doc.add_paragraph('')
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    report['unmapped'] = sorted(set(report['unmapped']))
    report['counts'] = dict(Counter(t['type'] for t in report['tables']))
    Path(str(output) + '.report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('idml', type=Path)
    parser.add_argument('name_map', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    report = generate_docx(args.idml, args.name_map, args.output)
    print(json.dumps({'counts': report['counts'], 'unmapped': report['unmapped'],
                      'unplaced': report['unplaced'], 'non_printed': report['non_printed']}, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
