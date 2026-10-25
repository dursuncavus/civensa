#!/usr/bin/env python3
"""CIV-054/CIV-093: capture the PP guide or complete UI map on disposable data.

Run from the checkout: python scripts/guide_screenshots.py --out docs/screenshots/guide-pp-v1
UI map: python scripts/guide_screenshots.py --ui-map --out docs/screenshots/ui-map-v1
Only CIVENSA_TEST_POSTGRES_URL is an input; application database URLs are ignored.
"""
import argparse
from contextlib import closing, contextmanager, ExitStack
from copy import deepcopy
from dataclasses import dataclass
import hashlib
from http.server import BaseHTTPRequestHandler, HTTPServer
import logging
import os
from pathlib import Path
import secrets
import shutil
import signal
import subprocess
import sys
from tempfile import TemporaryDirectory
import threading
from urllib.parse import parse_qs, urlsplit

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'tests/fixtures/first_load/source/pp2027_print_d3.idml'
MAX_PNG_BYTES = 1_500_000
UI_READONLY_POSTS = {'/api/budget-reallocations/choices', '/api/budget-reallocations/preview'}


@dataclass(frozen=True)
class Shot:
    filename: str
    path: str
    visible: str
    account: str = 'strategy'
    action: str = 'top'
    full: bool = False


SHOTS = (
    Shot('01-genel-bakis-ust.png', '/dashboard', 'Genel Bakış · üst bölüm'),
    Shot('02-ilerleme-panosu.png', '/dashboard', 'Müdürlük İlerleme Panosu', action='progress'),
    Shot('03-mudurluk-strateji-gorunumu.png', '/department/{submitted}', 'Gönderilmiş paketin Strateji inceleme görünümü'),
    Shot('04-inceleme-kuyrugu.png', '/review', 'İnceleme Kuyruğu · Onayla / İade Et', full=True),
    Shot('05-iade-penceresi.png', '/review', 'İade penceresi · boş gerekçe · ardından Vazgeç', action='return'),
    Shot('06-faaliyet-eslestirme.png', '/activity-mapping', 'Faaliyet Eşleştirme · üst bölüm'),
    Shot('07-kurum-yapisi.png', '/structure', 'Kurum Yapısı · üst bölüm'),
    Shot('08-kullanicilar.png', '/users', 'Kullanıcılar · üst bölüm'),
    Shot('09-butce-aktarma.png', '/budget-reallocations', 'Bütçe Aktarma', full=True),
    Shot('10-sistem-durumu.png', '/operations', 'Sistem Durumu · üst bölüm'),
    Shot('11-denetim-gecmisi.png', '/audit', 'Denetim Geçmişi · üst bölüm'),
    Shot('12-moduller.png', '/modules', 'Modüller', full=True),
    Shot('13-giris.png', '/login', 'Oturum kapalı giriş ekranı · alanlar boş', 'anonymous'),
    Shot('14-veri-girisi-ust.png', '/department/{own}', 'Müdürlük veri girişi · üst bölüm', 'department'),
    Shot('15-faaliyet-maliyetleri.png', '/department/{own}', 'Faaliyet Maliyetleri sekmesi · başlık ve maliyet tablosu', 'department', 'costs'),
    Shot('16-gonder-bolumu.png', '/department/{own}', 'Müdürlük paketi · Stratejiye Gönder (tıklanmaz)', 'department', 'submit'),
    Shot('17-gonderim-yetkileri.png', '/submit-authorizations', 'Gönderim Yetkileri · üst bölüm', 'department'),
    Shot('18-parola-degistir.png', '/submit-authorizations', 'Parolamı Değiştir penceresi · alanlar boş · ardından Vazgeç', 'department', 'password'),
)


class GuideError(Exception):
    """Only fixed, credential-free messages may be passed here."""


def check(condition, message):
    if not condition:
        raise GuideError(message)


def setup(engine):
    from sqlalchemy import select
    from sqlalchemy.orm import Session
    from app.models import Department, ReportingPeriod, User
    from app.personal_accounts import prepare_user_creation
    from app.program_service import validate_department
    from app.released_permissions import audit_user_change, commit_user_change
    from app.security import hash_password
    from scripts.first_load_dryrun import inspect_import, parse_book

    print('2027 kitabı ayrıştırılıyor (kaynak tablo 195 hariç)…', flush=True)
    parsed, _ = parse_book(BOOK, exclude_source_tables=[195])
    print('Kitap verisi yükleniyor; hesaplar ve tek gönderim hazırlanıyor…', flush=True)
    with Session(engine, expire_on_commit=False) as db:
        inspection = inspect_import(db, deepcopy(parsed), BOOK.name)
        check(inspection['applied'] is not None and not inspection['apply_error'],
              '2027 kitap verisi gerçek içe aktarma servisinden yüklenemedi.')
        departments = db.scalars(select(Department).order_by(Department.name)).all()
        ready_names = {row['name'] for row in inspection['results'] if row['ready']}
        ready = [d for d in departments if d.name in ready_names]
        check(len(ready) >= 2, 'Kurulum için en az iki doğrulaması temiz müdürlük gerekli.')
        own = next((d for d in departments if d.name == 'Strateji Geliştirme Müdürlüğü'
                    and any(row['name'] == d.name and (row['stored']['indicators'] or
                            row['stored']['activities']) for row in inspection['results'])), None)
        fallback = own is None
        own = own or ready[0]
        submitted = next(d for d in ready if d.id != own.id)
        # The unique active-director constraint also covers strategy accounts.
        anchor = next(d for d in departments if d.id not in {own.id, submitted.id})
        period = db.scalar(select(ReportingPeriod).where(ReportingPeriod.year == 2027))
        check(period.template_status == 'verified', '2027 şablonu doğrulanmış değil.')
        password = secrets.token_urlsafe(32)
        password_hash = hash_password(password)
        strategy = db.scalar(select(User).where(User.email == 'dryrun@example.invalid'))
        strategy.email = 'guide.strategy@example.invalid'
        strategy.full_name = 'Strateji Müdürü (test)'
        strategy.password_hash = password_hash
        strategy.department_id = anchor.id
        db.commit()
        # Bootstrap the imported strategy actor through the real appointment checks.
        candidate = User(full_name=strategy.full_name, email=strategy.email, role='strategy',
                         department_id=anchor.id, position='director', is_active=True)
        _, reason = prepare_user_creation(db, strategy, candidate)
        strategy.position = 'director'
        audit_user_change(db, strategy, strategy, 'position', 'staff', 'director', reason)
        commit_user_change(db)

        department = User(email='guide.department@example.invalid',
                          full_name='Strateji Geliştirme Müdürlüğü (test)', role='department',
                          department_id=own.id, position='director', is_active=True,
                          password_hash=password_hash)
        sender = User(email='guide.setup@example.invalid', full_name='Gönderim kurulumu (test)',
                      role='department', department_id=submitted.id, position='director',
                      is_active=True, password_hash=password_hash)
        for target in (department, sender):
            prepare_user_creation(db, strategy, target)
            db.add(target)
            commit_user_change(db)

        errors, _ = validate_department(db, period.id, submitted.id)
        check(not errors, 'Seçilen müdürlüğün PP doğrulaması temiz değil.')
        # This is the production submit implementation, also called by the API.
        from app.main import activate_period, submit_department
        activated = activate_period(period.id, db=db, user=strategy)
        check(activated.get('ok'), '2027 dönemi gerçek aktivasyon servisinden etkinleştirilemedi.')
        result = submit_department(submitted.id, db=db, user=sender)
        check(isinstance(result, dict) and result.get('ok'), 'Gerçek gönderim servisi paketi gönderemedi.')
        sender.is_active = False  # Setup-only actor; never opens a browser session.
        db.commit()
        return dict(own=own.id, submitted=submitted.id, own_name=own.name,
                    submitted_name=submitted.name, anchor_name=anchor.name,
                    fallback=fallback, ready=len(ready), password=password,
                    departments=len(departments),
                    indicators=sum(len(row['stored']['indicators']) for row in inspection['results']),
                    activities=sum(len(row['stored']['activities']) for row in inspection['results']),
                    accounts={'strategy': strategy.email, 'department': department.email})


def persisted_state(engine):
    """Prove that capturing/cancelling dialogs did not save a decision or password."""
    from sqlalchemy import select
    from sqlalchemy.orm import Session
    from app.models import AuditLog, Submission, SubmissionEvent, User
    with Session(engine) as db:
        return (
            list(db.execute(select(Submission.id, Submission.status, Submission.revision,
                                   Submission.return_note, Submission.reviewed_at).order_by(Submission.id))),
            list(db.execute(select(SubmissionEvent.id, SubmissionEvent.event_type).order_by(SubmissionEvent.id))),
            list(db.execute(select(User.id, User.password_hash).order_by(User.id))),
            list(db.execute(select(AuditLog.id).order_by(AuditLog.id))),
        )


@contextmanager
def local_http(engine, data):
    """Reuse the UI tests' loopback TestClient bridge, with pre-browser logins."""
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import Session
    from app.db import get_db
    from app.main import app

    def database():
        with Session(engine) as db:
            yield db

    app.dependency_overrides[get_db] = database
    try:
        with ExitStack() as stack:
            # Do not enter TestClient itself: lifespan would migrate again.
            clients = {name: stack.enter_context(closing(TestClient(app, follow_redirects=False)))
                       for name in ('strategy', 'department', 'anonymous')}
            for name, email in data['accounts'].items():
                response = clients[name].post('/login', data={'email': email, 'password': data['password']})
                check(response.status_code == 303 and response.headers.get('location') == '/modules',
                      'Yerel test hesabı ile giriş yapılamadı.')

            class Handler(BaseHTTPRequestHandler):
                def do_GET(self):
                    # No arbitrary hosts, absolute request targets, or write methods.
                    if not self.path.startswith('/') or self.path.startswith('//'):
                        self.send_error(400)
                        return
                    client = clients.get(self.headers.get('X-Guide-Account', 'anonymous'))
                    if client is None:
                        self.send_error(403)
                        return
                    try:
                        key = (self.headers.get('X-Guide-Account', 'anonymous'),
                               self.headers.get('X-Guide-State', ''))
                        saved = data.get('_responses', {}).get(key)
                        response = saved if saved is not None and urlsplit(self.path).path == urlsplit(str(saved.request.url)).path else client.get(self.path, headers={'Accept': self.headers.get('Accept', '*/*')})
                    except Exception:
                        self.send_error(500)
                        return
                    self.reply(response)

                def do_POST(self):
                    if not data.get('ui_map') or self.path not in UI_READONLY_POSTS:
                        self.send_error(405)
                        return
                    client = clients.get(self.headers.get('X-Guide-Account', 'anonymous'))
                    length = int(self.headers.get('Content-Length', '0'))
                    if client is None or not 0 < length <= 100_000:
                        self.send_error(400)
                        return
                    response = client.post(self.path, content=self.rfile.read(length),
                        headers={'Content-Type': 'application/json',
                                 'X-CSRF-Token': self.headers.get('X-CSRF-Token', '')})
                    self.reply(response)

                def reply(self, response):
                    self.send_response(response.status_code)
                    for key, value in response.headers.items():
                        if key.lower() not in {'content-length', 'transfer-encoding', 'connection', 'set-cookie'}:
                            self.send_header(key, value)
                    self.send_header('Content-Length', str(len(response.content)))
                    self.end_headers()
                    self.wfile.write(response.content)

                def log_message(self, *args):
                    pass

            server = HTTPServer(('127.0.0.1', 0), Handler)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                data['_clients'] = clients
                yield f'http://127.0.0.1:{server.server_port}'
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)
    finally:
        app.dependency_overrides.pop(get_db, None)


def capture(base, folder, data, results):
    from PIL import Image
    from playwright.sync_api import sync_playwright, expect

    width = data.get('viewport_width', 1440)
    with sync_playwright() as pw:
        options = {'headless': True}
        if os.getenv('CIVENSA_CHROMIUM_PATH'):
            options['executable_path'] = os.environ['CIVENSA_CHROMIUM_PATH']
        with pw.chromium.launch(**options) as browser, ExitStack() as stack:
            contexts = {}
            blocked = []

            def guard(route):
                request = route.request
                if (urlsplit(request.url).netloc == urlsplit(base).netloc
                        and urlsplit(request.url).scheme == 'http'
                        and (request.method == 'GET' or data.get('ui_map') and request.method == 'POST'
                             and urlsplit(request.url).path in UI_READONLY_POSTS)):
                    route.continue_()
                else:
                    blocked.append(True)
                    route.abort()

            for name in ('strategy', 'department', 'anonymous'):
                context = stack.enter_context(browser.new_context(
                    viewport={'width': width, 'height': 900}, device_scale_factor=1,
                    color_scheme='light', locale='tr-TR', timezone_id='Europe/Istanbul',
                    reduced_motion='reduce', service_workers='block', accept_downloads=False,
                    extra_http_headers={'X-Guide-Account': name}))
                context.route('**/*', guard)
                contexts[name] = context

            data['browser'] = browser.version
            for shot in data.get('_shots', SHOTS):
                page = contexts[shot.account].new_page()
                errors = []
                page.on('pageerror', lambda error: errors.append(True))
                target = folder / shot.filename
                try:
                    path = shot.path.format(**data)
                    page.set_extra_http_headers({'X-Guide-Account': shot.account,
                        'X-Guide-State': data.get('_screen_meta', {}).get(shot.filename, {}).get('response', '')})
                    response = page.goto(base + path, wait_until='networkidle')
                    expected_url = base + path
                    displayed_period = page.locator('meta[name="pp-period-id"]')
                    if displayed_period.count() and 'period_id' not in parse_qs(urlsplit(path).query):
                        expected_url += ('&' if '?' in path else '?') + 'period_id=' + displayed_period.get_attribute('content')
                    check(response is not None and response.status == data.get('_screen_meta', {}).get(shot.filename, {}).get('status', 200) and page.url == expected_url,
                          'Sayfa beklenen URL ve HTTP durumuyla açılmadı.')
                    if data.get('ui_map'):
                        check('text/html' in response.headers.get('content-type', ''),
                              'Beklenen HTML yerine farklı bir yanıt türü döndü.')
                    page.evaluate('document.fonts.ready')
                    if shot.account != 'anonymous' and not data.get('ui_map'):
                        expect(page.locator('.userbox strong')).to_have_text(
                            'Strateji Müdürü (test)' if shot.account == 'strategy'
                            else 'Strateji Geliştirme Müdürlüğü (test)')
                    if data.get('ui_map'):
                        ui_expand(page, shot.action)
                    elif shot.action == 'progress':
                        page.get_by_role('heading', name='Müdürlük İlerleme Panosu', exact=True).evaluate(
                            "el => el.scrollIntoView({block: 'start', behavior: 'instant'})")
                    elif shot.action == 'costs':
                        page.get_by_role('button', name='Faaliyet Maliyetleri', exact=True).click()
                        heading = page.get_by_role('heading', name='Faaliyet Maliyetleri', exact=True)
                        expect(heading).to_be_visible()
                        heading.evaluate("el => el.scrollIntoView({block: 'start', behavior: 'instant'})")
                    elif shot.action == 'submit':
                        button = page.get_by_role('button', name='Stratejiye Gönder', exact=True)
                        button.evaluate("el => el.scrollIntoView({block: 'center', behavior: 'instant'})")
                        expect(button).to_be_visible()
                    elif shot.action == 'return':
                        row = page.get_by_role('row').filter(has=page.get_by_role(
                            'button', name='Onayla', exact=True))
                        expect(row).to_have_count(1)
                        row.get_by_role('button', name='İade Et', exact=True).click()
                        expect(page.locator('#reviewModal')).to_be_visible()
                    elif shot.action == 'password':
                        page.get_by_role('button', name='Parolamı Değiştir', exact=True).click()
                        expect(page.locator('#passwordModal')).to_be_visible()
                    else:
                        page.evaluate('window.scrollTo(0, 0)')
                    if shot.path == '/review' and not data.get('ui_map'):
                        expect(page.get_by_role('button', name='Onayla', exact=True)).to_have_count(1)
                    check(data['password'] not in page.content(), 'Parola sayfa içeriğine sızdı; görsel kaydedilmedi.')
                    check(not errors and not blocked, 'JavaScript hatası veya engellenmiş ağ/yazma isteği var.')
                    if data.get('ui_map'):
                        data['_screen_meta'][shot.filename]['dom'] = page.evaluate(UI_DOM)
                    if data.get('_verify_page'):
                        data['_verify_page'](page, shot, data['_screen_meta'][shot.filename])
                    page.screenshot(path=str(target), full_page=shot.full, animations='disabled', caret='hide')
                    if shot.action in {'return', 'password'}:
                        modal = page.locator('#reviewModal' if shot.action == 'return' else '#passwordModal')
                        modal.get_by_role('button', name='Vazgeç', exact=True).click()
                        expect(modal).to_be_hidden()
                    # Lossless optimization only; never scale or quantize guide images.
                    with Image.open(target) as png:
                        check(png.width == width and (shot.full or png.height == 900),
                              'Görsel boyutları istenen viewport ile eşleşmiyor.')
                        png.save(target, optimize=True)
                        if data.get('ui_map') and target.stat().st_size >= MAX_PNG_BYTES:
                            # Palette encoding keeps full dimensions on long index pages.
                            png.convert('RGB').quantize(colors=256).save(target, optimize=True)
                    check(target.stat().st_size < MAX_PNG_BYTES, 'PNG 1.500.000 bayt sınırının altında değil.')
                    check(not errors and not blocked, 'Yakalama sırasında JavaScript veya ağ/yazma hatası oluştu.')
                    results[shot.filename] = None
                except Exception as exc:
                    target.unlink(missing_ok=True)
                    results[shot.filename] = (str(exc) if isinstance(exc, GuideError) else
                                              f'Ekran adımı tamamlanamadı ({type(exc).__name__}); seçicileri ve yerel Chromium kurulumunu kontrol edin.')
                finally:
                    page.close()
                print(f'{shot.filename}: {"üretildi" if results[shot.filename] is None else "üretilemedi (INDEX.md)"}',
                      flush=True)
                if data.get('ui_map') and results[shot.filename]:
                    print(results[shot.filename], flush=True)
            check(not blocked, 'Yakalama sırasında yerel GET dışında bir istek engellendi.')


def write_index(folder, data, results):
    def cell(value):
        return str(value).replace('|', '&#124;').replace('\n', ' ')

    lines = ['# CIV-054 · Strateji PP kullanım kılavuzu · ekran görüntüleri', '',
             'Yenileme (depo kökünden):', '', '```sh',
             'python scripts/guide_screenshots.py --out docs/screenshots/guide-pp-v1', '```', '',
             '`CIVENSA_TEST_POSTGRES_URL` yalnızca yerel PostgreSQL 16+ ve CREATE DATABASE yetkili test rolünü göstermelidir. '
             'Uygulama bağlantı ayarları girdi olarak okunmaz. Geçici veritabanı başarıda, hatada ve SIGINT/SIGTERM sırasında silinir.', '',
             f'Kaynak: `{BOOK.relative_to(ROOT)}` · SHA-256: `{hashlib.sha256(BOOK.read_bytes()).hexdigest()}`.',
             '2027 dönemi aktif, şablon doğrulanmış; dry-run ile aynı gerçek servisler, yalnız kaynak tablo 195 hariç. '
             'Bu, mevcut prova verisidir; kitap eksikleri düzeltilmez, boş değerler doldurulmaz.', '',
             'Chromium · 1440×900 · deviceScaleFactor 1 · açık tema · tr-TR · Europe/Istanbul. '
             'FULL dışındaki görüntüler viewport çekimidir. PNG sınırı: <1.500.000 bayt.', '']
    if data:
        lines += [f"Yüklenen veri: {data['departments']} müdürlük, {data['indicators']} gösterge, {data['activities']} faaliyet.", '',
                  f"Başlangıçta doğrulaması temiz: {data['ready']} müdürlük. Tam bir paket gerçek gönderim servisiyle gönderildi: "
                  f"**{data['submitted_name']}** (ID {data['submitted']}); onay/iade yapılmadı.", '',
                  f"Veri giriş hesabının müdürlüğü: **{data['own_name']}** (ID {data['own']}); paket taslak bırakıldı.",
                  ('**Yedek müdürlük kullanıldı:** Strateji Geliştirme Müdürlüğü bulunamadı veya verisi yoktu; '
                   'ad sırasındaki ilk hazır müdürlük seçildi.' if data['fallback'] else
                   'Strateji Geliştirme Müdürlüğü bulundu ve kitap verisi içeriyor; yedek müdürlük kullanılmadı.'), '',
                  f"Strateji hesabının teknik müdürlük bağı: **{data['anchor_name']}**. Aynı müdürlükte iki aktif Müdürü "
                  'engelleyen uygulama kısıtı nedeniyle veri giriş hesabından farklıdır; rolü strategy, görevi director.', '',
                  'Hesaplar: **strategy** = Strateji Müdürü (test); **department** = Strateji Geliştirme Müdürlüğü (test), '
                  'görevi director; **anonymous** = oturum kapalı (13, müdürlük kullanıcısının giriş öncesi ekranı). '
                  'Ek **Gönderim kurulumu (test)** hesabı yalnız gönderim servisinde kullanılır ve tarayıcıdan önce pasifleştirilir. '
                  'Parola çalışma anında rastgele üretilir; yazdırılmaz veya dosyaya kaydedilmez.', '',
                  f"Chromium sürümü: {data.get('browser', 'başlatılamadı')}. Kaynak kod: "
                  f"`{subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()}`. "
                  'Zaman damgaları her yerel yüklemede yenilenir; görsellerin birebir bayt eşitliği beklenmez.', '']
    lines += ['Tarayıcı başlamadan önce oturumlar gerçek yerel giriş akışıyla açılır. '
              'Çekim sırasında yalnız loopback GET istekleri geçer; form/API yazmaları ve dış ağ engellenir. '
              'İade ve parola pencereleri Vazgeç ile kapatılır. Gönderimler, olaylar, parola hashleri ve denetim kayıtları '
              'çekim öncesi/sonrası karşılaştırılır.', '',
              '| Dosya | URL yolu | Görünen içerik | Hesap | Çekim | Sonuç |',
              '| --- | --- | --- | --- | --- | --- |']
    for shot in SHOTS:
        path = shot.path.format(**data) if data else shot.path
        error = results[shot.filename]
        file = f'`{shot.filename}`' if error else f'[{shot.filename}]({shot.filename})'
        lines.append('| ' + ' | '.join([file, f'`{path}`', cell(shot.visible), shot.account,
                                        'FULL' if shot.full else 'Viewport',
                                        cell('ÜRETİLEMEDİ: ' + error) if error else 'Üretildi']) + ' |')
    missing = [name for name, error in results.items() if error]
    lines += ['', 'Üretilemeyen: ' + (', '.join(missing) if missing else 'Yok; 18/18 PNG üretildi.'), '']
    (folder / 'INDEX.md').write_text('\n'.join(lines), encoding='utf-8')


# CIV-093 shares the disposable database, HTTP bridge and browser lifecycle above.
# These fixtures contain no book import, real institution or person data.
UI_PAGES = {
    "/login": ("Giriş", "Hesabın oturumunu açar."),
    "/modules": ("Modüller", "Kullanıcının erişebildiği modülleri seçmesini sağlar."),
    "/dashboard": (
        "Genel Bakış",
        "Dönemin veri girişini ve müdürlük paketlerini özetler.",
    ),
    "/department/{department_id}": (
        "Müdürlük veri girişi",
        "Gösterge hedeflerini ve faaliyet maliyetlerini gösterir ve düzenler.",
    ),
    "/source-comparison": (
        "PP Kaynak Karşılaştırması",
        "Kaynak belge ile dönem kayıtlarının farklarını gösterir.",
    ),
    "/program/{department_id}": (
        "Rapor Hazırlığı",
        "Performans kartlarını ve ekonomik tabloları düzenler.",
    ),
    "/program/{department_id}/economic-split": (
        "Ekonomik dağılım",
        "Kaynak ekonomik tutarları program gruplarına dağıtır.",
    ),
    "/report-builder": (
        "Nihai Rapor",
        "Genel bölümleri, kapsamı ve yayın hazırlığını yönetir.",
    ),
    "/report-builder/manifest": (
        "Tamamlanma listesi",
        "Nihai yayın için eksik veri ve onayları listeler.",
    ),
    "/report-builder/import/{batch_id}": (
        "Genel bölüm içe aktarma",
        "Word genel bölümlerinin aktarım önizlemesini gösterir.",
    ),
    "/activity-mapping": (
        "Faaliyet Eşleştirme",
        "Yerel faaliyetleri resmî faaliyet kataloğuna bağlar.",
    ),
    "/review": (
        "İnceleme Kuyruğu",
        "Gönderilen müdürlük paketlerini inceleme ve karar için listeler.",
    ),
    "/budget-reallocations": (
        "Bütçe Aktarma",
        "Kaydedilmiş bütçe aktarma işlemlerini listeler.",
    ),
    "/budget-reallocations/new": (
        "Yeni Bütçe Aktarması",
        "Aktarma tutarlarını, önizlemeyi ve hukuki dayanağı hazırlar.",
    ),
    "/budget-reallocations/{operation_id}": (
        "Bütçe aktarma ayrıntısı",
        "Aktarmanın dayanağını ve değişen tutarları gösterir.",
    ),
    "/activity-report": (
        "Faaliyet Raporu hazırlık alanı",
        "Seçilen yıldaki hedef ve gerçekleşme kayıtlarını gösterir.",
    ),
    "/activity-reports": (
        "Yayına bağlı Faaliyet Raporları",
        "PP yayınından oluşturulan faaliyet raporu revizyonlarını listeler.",
    ),
    "/activity-reports/revisions/{revision_id}": (
        "Gerçekleşme revizyonu",
        "Sabitlenmiş PP hedefleri için gerçekleşme ve kanıt girişini gösterir.",
    ),
    "/activity-reports/publications/{publication_id}/bindings": (
        "Gösterge kimlik eşlemesi",
        "Yayın göstergelerinin kalıcı kimlik bağlantılarını gösterir.",
    ),
    "/activity-sections": (
        "Faaliyet Raporu · EK-1",
        "İdare ve birim raporlarının bölüm ve onay durumlarını listeler.",
    ),
    "/activity-sections/reports/{document_id}": (
        "EK-1 bölüm editörü",
        "Seçilen rapor bölümünün metin, tablo ve durumunu gösterir.",
    ),
    "/activity-sections/reports/{document_id}/financial/imports/{import_id}": (
        "Malî içe aktarma önizlemesi",
        "Excel malî verilerinin farklarını ve satır hatalarını gösterir.",
    ),
    "/activity-sections/reports/{document_id}/assurances/{code}": (
        "Güvence beyanı",
        "Sabit beyan metninin yer, tarih ve imzacı alanlarını düzenler.",
    ),
    "/internal-control": (
        "İç Kontrol planları",
        "Uyum eylem planlarını ve plan yüklemeyi gösterir.",
    ),
    "/internal-control/imports/{batch_id}": (
        "İç Kontrol yükleme kontrolü",
        "Plan dosyasının doğrulamasını ve birim eşlemelerini gösterir.",
    ),
    "/internal-control/plans/{plan_id}": (
        "İç Kontrol plan kataloğu",
        "Bileşen, standart, şart ve sorumlu birimlere göre eylemleri gösterir.",
    ),
    "/internal-control/our-actions": (
        "İç Kontrol Eylemlerimiz",
        "Müdürlüğün sorumlu ve katkı verdiği eylemleri gösterir.",
    ),
    "/internal-control/monitoring": (
        "İç Kontrol izleme",
        "Altı aylık dönemleri, paket durumlarını ve raporları yönetir.",
    ),
    "/internal-control/monitoring/batches/{batch_id}": (
        "İç Kontrol izleme paketi",
        "Eylemlerin gerçekleşme, kanıt ve onay geçmişini gösterir.",
    ),
    "/structure": (
        "Kurum Yapısı",
        "Döneme bağlı müdürlük yapısı ve yapı kararlarını yönetir.",
    ),
    "/users": ("Kullanıcılar", "Kişisel test hesaplarını ve görevlerini yönetir."),
    "/submit-authorizations": (
        "Gönderim Yetkileri",
        "Müdürlüğün personel gönderim yetkilerini yönetir.",
    ),
    "/change-password": ("Parola değiştir", "Mevcut oturumun parolasını değiştirir."),
    "/operations": (
        "Sistem Durumu",
        "Dönem, şablon, içe aktarma ve yayın kontrollerini gösterir.",
    ),
    "/audit": (
        "Denetim Geçmişi",
        "Kullanıcı işlemlerinin denetim kayıtlarını filtreler.",
    ),
}


def ui_setup(engine):
    """Small, wholly synthetic records on the migrated disposable schema."""
    from datetime import date
    from decimal import Decimal
    from io import BytesIO
    import json
    from openpyxl import load_workbook
    from sqlalchemy import select
    from sqlalchemy.orm import Session
    from app import models as m
    from app.security import hash_password
    from app import activity_sections_service as sections
    from app import activity_contract_service as contract
    from app import activity_financial_service as finance
    from app import internal_control_service as ic
    from app import internal_control_monitoring_service as monitoring
    from app.internal_control_monitoring_models import ICMonitoringBatch
    from app.serialization import json_dumps

    with Session(engine, expire_on_commit=False) as db:
        org = m.Organization(name="Örnek Belediyesi (sentetik)")
        period = m.ReportingPeriod(
            year=2027, status="active", template_status="verified"
        )
        db.add_all([org, period])
        db.flush()
        own = m.Department(organization_id=org.id, name="Örnek Hizmet Müdürlüğü")
        anchor = m.Department(
            organization_id=org.id, name="Strateji Geliştirme Müdürlüğü"
        )
        submitted = m.Department(organization_id=org.id, name="Örnek Destek Müdürlüğü")
        db.add_all([own, anchor, submitted])
        db.flush()
        password = secrets.token_urlsafe(32)
        hashed = hash_password(password)
        strategy = m.User(
            email="ui.strategy@example.invalid",
            full_name="Strateji Müdürü (sentetik)",
            role="strategy",
            department_id=anchor.id,
            position="director",
            password_hash=hashed,
        )
        department = m.User(
            email="ui.department@example.invalid",
            full_name="Birim Müdürü (sentetik)",
            role="department",
            department_id=own.id,
            position="director",
            password_hash=hashed,
        )
        editor = m.User(
            email="ui.setup@example.invalid",
            full_name="Kurulum Personeli (sentetik)",
            role="strategy",
            department_id=anchor.id,
            password_hash=hashed,
        )
        staff = m.User(
            email="ui.staff@example.invalid",
            full_name="Birim Personeli (sentetik)",
            role="department",
            department_id=own.id,
            password_hash=hashed,
            can_submit=True,
        )
        db.add_all([strategy, department, editor, staff])
        db.flush()
        program = m.Program(period_id=period.id, name="Yerel Hizmetler Programı")
        db.add(program)
        db.flush()
        sp = m.SubProgram(
            period_id=period.id, program_id=program.id, name="Erişilebilir Hizmetler"
        )
        db.add(sp)
        db.flush()
        activities = []
        for dep in (own, submitted):
            db.add(
                m.DepartmentScope(
                    period_id=period.id, department_id=dep.id, status="required"
                )
            )
            db.add(
                m.Submission(
                    period_id=period.id,
                    department_id=dep.id,
                    status="draft" if dep == own else "submitted",
                    revision=1,
                )
            )
            card = m.ReportCard(
                period_id=period.id,
                department_id=dep.id,
                subprogram_id=sp.id,
                source_title="2027 · Hizmet kalitesi",
                goal="Hizmete erişimi artırmak",
                rationale="Sentetik ihtiyaç değerlendirmesi",
                responsible=org.name,
                calculation_method="Yıllık hizmet kayıtlarının sayımı",
                data_source="Sentetik hizmet kayıtları",
            )
            db.add(card)
            db.flush()
            indicator = m.Indicator(
                period_id=period.id,
                department_id=dep.id,
                subprogram_id=sp.id,
                card_id=card.id,
                code="G-1",
                name="Tamamlanan hizmet sayısı",
                unit="Adet",
            )
            activity = m.Activity(
                period_id=period.id,
                department_id=dep.id,
                subprogram_id=sp.id,
                card_id=card.id,
                code="F-1",
                name="Hizmet noktalarının iyileştirilmesi",
                explanation="Sentetik yıllık çalışma",
            )
            db.add_all([indicator, activity])
            db.flush()
            activities.append(activity)
            for year, kind in (
                (2027, "target"),
                (2028, "forecast"),
                (2029, "forecast"),
            ):
                db.add(
                    m.IndicatorValue(
                        indicator_id=indicator.id,
                        year=year,
                        value_kind=kind,
                        value="120",
                    )
                )
            official = m.OfficialActivity(
                period_id=period.id,
                program_code="1",
                program_name=program.name,
                subprogram_code="1.1",
                subprogram_name=sp.name,
                code=f"1.1.{dep.id}",
                name="Hizmet sunumu",
            )
            db.add(official)
            db.flush()
            db.add(
                m.ActivityResponsibility(
                    official_activity_id=official.id, department_id=dep.id
                )
            )
            if dep == submitted:
                db.add(
                    m.ActivityCatalogueLink(
                        activity_id=activity.id, official_activity_id=official.id
                    )
                )
            for year, kind in (
                (2027, "budget"),
                (2028, "forecast"),
                (2029, "forecast"),
            ):
                for fund, amount in (
                    ("inside", "12500.23"),
                    ("outside", "0"),
                    ("total", "12500.23"),
                ):
                    db.add(
                        m.ActivityCostValue(
                            activity_id=activity.id,
                            year=year,
                            value_kind=kind,
                            funding_type=fund,
                            amount=amount,
                        )
                    )
                db.add(
                    m.DepartmentCostValue(
                        period_id=period.id,
                        department_id=dep.id,
                        year=year,
                        value_kind=kind,
                        category="personnel_sgk",
                        amount="2000",
                    )
                )
                for cat, amount in (
                    ("goods_services", "12500.23"),
                    ("personnel", "1500"),
                    ("sgk", "500"),
                ):
                    for fund in ("inside", "outside"):
                        db.add(
                            m.EconomicValue(
                                period_id=period.id,
                                department_id=dep.id,
                                year=year,
                                value_kind=kind,
                                category=cat,
                                program_group=(
                                    "unallocated"
                                    if dep == own
                                    and year == 2029
                                    and cat == "goods_services"
                                    else "service"
                                ),
                                funding_type=fund,
                                amount=amount if fund == "inside" else "0",
                            )
                        )
        db.add(
            m.ReportSection(
                period_id=period.id,
                title="Genel Bilgiler",
                kind="text",
                position=1,
                content_json=json.dumps(
                    {"text": "Bu ekran haritasındaki bütün kayıtlar sentetiktir."}
                ),
            )
        )
        db.commit()
        # Frozen publication from the same canonical serializer used by publish.
        payload, definitions = contract.freeze_program_payload(db, period.id, editor)
        raw = json_dumps(
            payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        pub = m.Publication(
            period_id=period.id,
            version=1,
            snapshot_json=raw,
            sha256=hashlib.sha256(raw.encode()).hexdigest(),
            created_by=editor.id,
        )
        db.add(pub)
        db.flush()
        contract.attach_publication_definitions(db, pub, definitions, editor)
        revision = contract.create_report_revision(db, editor, pub.id, org.id)
        db.commit()
        for item in db.scalars(
            select(m.ReportIndicator).where(
                m.ReportIndicator.report_revision_id == revision.id
            )
        ):
            contract.save_actual(
                db,
                editor,
                item.id,
                0,
                "110",
                "verified",
                [
                    {
                        "title": "Sentetik hizmet defteri",
                        "locator": "Örnek sayfa 1",
                        "source_revision": "2027/1",
                    }
                ],
                "Yıllık hedefin bir bölümü sonraki döneme kaldı.",
                {"corrective_action": "Takvim gözden geçirilecek."},
            )
        docs = sections.create_year(db, strategy, org.id, period.id)
        db.commit()
        idare = next(d for d in docs if d.scope == "idare")
        unit = next(d for d in docs if d.department_id == own.id)
        for doc, actor in ((idare, strategy), (unit, department)):
            sections.save_section(
                db,
                actor,
                doc.id,
                "I.A",
                [
                    {
                        "type": "paragraph",
                        "text": "Hizmetleri erişilebilir, ölçülebilir ve şeffaf biçimde sunmak.",
                    }
                ],
                "draft",
                "",
                doc.revision,
            )
            db.commit()
            if not doc.department_id:
                continue
            sections.save_resources(
                db,
                actor,
                doc.id,
                "buildings",
                [
                    {
                        "name": "Örnek Hizmet Binası",
                        "category": "Hizmet binası",
                        "address": "Örnek Mahalle",
                        "area": Decimal("240.5"),
                        "ownership": "mülk",
                    }
                ],
                doc.revision,
            )
            db.commit()
        finance.save_budget(
            db,
            strategy,
            idare.id,
            "expenditure",
            [
                dict(
                    department_id=own.id,
                    economic_code="03",
                    initial=12500.23,
                    added=100,
                    deducted=50,
                    actual=10000,
                    cancelled=0,
                    carried=0,
                    explanation="Sentetik bütçe gerçekleşmesi",
                )
            ],
            idare.revision,
        )
        db.commit()
        book = load_workbook(BytesIO(finance.template("revenue")))
        book.active.append(["01", 12500.23, 11000, 0, "Sentetik gelir"])
        output = BytesIO()
        book.save(output)
        preview = finance.preview_import(
            db, strategy, idare.id, "revenue", output.getvalue(), idare.revision
        )
        db.commit()
        # A pending Word batch uses a complete parser-shaped synthetic payload.
        parsed = {
            "departments": [],
            "warnings": [],
            "general_sections": [
                {
                    "title": "Örnek genel bölüm",
                    "kind": "text",
                    "position": 1,
                    "content": {"text": "Sentetik içe aktarma metni"},
                    "source": "sentetik.docx",
                }
            ],
        }
        batch = m.ImportBatch(
            period_id=period.id,
            filename="sentetik.docx",
            status="preview",
            payload_json=json.dumps(parsed),
            decisions_json="[]",
            created_by=strategy.id,
        )
        db.add(batch)
        db.commit()
        # Plan import and monitoring go through real production services.
        book = load_workbook(BytesIO(ic.template_bytes()))
        for row, value in enumerate(
            ["Sentetik Uyum Eylem Planı", 2027, 2029, "v1", "18.10.2026"], 2
        ):
            book["Bilgi"].cell(row, 2, value)
        book["Plan"].append(
            [
                "KOS",
                "KOS1",
                "KOS 1.1",
                "Etik değerler",
                "Geliştirilecek",
                "KOS.1.1.1",
                "Etik eğitim düzenlemek",
                own.name,
                "",
                "Eğitim raporu",
                "2027",
                "Sürekli",
                "",
                "Etik Değerler",
            ]
        )
        ic_batch, report = ic.create_import(
            db, strategy, ic.export_book(book), "sentetik-plan.xlsx"
        )
        check(report["can_commit"], "Sentetik İç Kontrol planı doğrulanamadı.")
        plan = ic.commit_import(db, strategy, ic_batch.id)
        db.commit()
        ic.activate_plan(db, strategy, plan.id)
        db.commit()
        mp = monitoring.open_period(db, strategy, plan.id, 2027, 1, date(2029, 12, 31))
        db.commit()
        mb = db.scalar(
            select(ICMonitoringBatch).where(
                ICMonitoringBatch.period_id == mp.id,
                ICMonitoringBatch.department_id == own.id,
            )
        )
        for row in monitoring.rows(db, mb):
            monitoring.save_entry(
                db,
                department,
                mb.id,
                row["target"].id,
                monitoring.state(db, mb.id)[1],
                dict(
                    status="ongoing",
                    percentage="50",
                    explanation="Eğitimin ilk oturumu tamamlandı.",
                    reason="İkinci oturum gelecek dönemde yapılacak.",
                    evidence="Sentetik eğitim tutanağı 2027/1",
                ),
            )
        db.commit()
        return dict(
            own=own.id,
            submitted=submitted.id,
            period=period.id,
            publication=pub.id,
            revision=revision.id,
            document=idare.id,
            unit=unit.id,
            import_id=preview["id"],
            pp_import=batch.id,
            ic_import=ic_batch.id,
            plan=plan.id,
            monitoring=mp.id,
            ic_batch=mb.id,
            password=password,
            accounts={"strategy": strategy.email, "department": department.email},
        )


def ui_routes():
    """Include schema-hidden routes and FastAPI's lazily included routers."""
    from app.main import app

    def walk(routes):
        for route in routes:
            if hasattr(route, "methods") and "GET" in route.methods:
                yield route.path
            elif hasattr(route, "original_router"):
                yield from walk(route.original_router.routes)

    return sorted(set(walk(app.routes)))


def ui_module(path):
    if path.startswith("/internal-control"):
        return "İç Kontrol"
    if path.startswith("/activity-"):
        return (
            "Faaliyet Raporu"
            if not path.startswith("/activity-mapping")
            else "Performans Programı"
        )
    if path.startswith(
        (
            "/program",
            "/report-builder",
            "/department",
            "/source-comparison",
            "/dashboard",
            "/review",
            "/budget",
            "/publications",
        )
    ):
        return "Performans Programı"
    return "Yönetim"


def ui_path(pattern, data, account):
    own = data["own"]
    params = dict(
        department_id=own,
        publication_id=data["publication"],
        revision_id=data["revision"],
        document_id=data["document"] if account == "strategy" else data["unit"],
        import_id=data["import_id"],
        operation_id=data.get("operation", 0),
        plan_id=data["plan"],
        code="EK-1" if account == "strategy" else "EK-2",
    )
    params["batch_id"] = (
        data["ic_batch"]
        if "/monitoring/" in pattern
        else (
            data["ic_import"]
            if pattern.startswith("/internal-control")
            else data["pp_import"]
        )
    )
    path = pattern.format(**params)
    if pattern in {
        "/dashboard",
        "/department/{department_id}",
        "/source-comparison",
        "/program/{department_id}",
        "/program/{department_id}/economic-split",
        "/report-builder",
        "/report-builder/manifest",
        "/report-builder/import/{batch_id}",
        "/review",
        "/activity-mapping",
        "/structure",
        "/operations",
        "/audit",
        "/users",
        "/submit-authorizations",
    }:
        path += f'?period_id={data["period"]}'
    if pattern == "/activity-report":
        path += "?year=2027"
    return path


def ui_state_digest(engine):
    """Hash all model tables, not just the original guide's mutation boundary."""
    from sqlalchemy import select
    from app.db import Base

    digest = hashlib.sha256()
    with engine.connect() as connection:
        for table in sorted(Base.metadata.tables.values(), key=lambda t: t.name):
            rows = [repr(tuple(row)) for row in connection.execute(select(table))]
            digest.update((table.name + "\n" + "\n".join(sorted(rows))).encode())
    return digest.hexdigest()


UI_DOM = r"""() => {
  const text = el => (el?.innerText || el?.textContent || '').replace(/\s+/g,' ').trim();
  const visible = el => !!el.getClientRects().length && getComputedStyle(el).visibility !== 'hidden'
    && !el.closest('[hidden]') && el.type !== 'hidden';
  const label = el => el.getAttribute('aria-label') || (['A','BUTTON','SUMMARY'].includes(el.tagName) ? text(el) : '') || (el.labels?.length ? [...el.labels].map(text).join(' / ') : '')
    || el.getAttribute('placeholder') || el.name || el.dataset.field || el.id || text(el);
  const context = el => {
    const row = el.closest('tr');
    const group = el.closest('article,section,details,.modal-card,.card');
    return (row ? text(row.querySelector('td')) : '') || text(group?.querySelector('h2,h3,summary'));
  };
  const controls = [...document.querySelectorAll('a,button,form,input,select,textarea,summary')].filter(visible).map(el => {
    const form = el.form || el.closest('form');
    const row = el.closest('[data-id],[data-activity],[data-user],[data-operation],[data-new]');
    const modal = el.closest('.modal');
    const item = {tag:el.tagName.toLowerCase(), label:label(el), id:el.id,
      classes:[...el.classList], name:el.name || el.dataset.field || '', type:el.type || '',
      href:el.getAttribute('href') || '', action:form ? (form.getAttribute('action') || location.pathname) : '',
      method:(form?.getAttribute('method') || 'get').toUpperCase(), formId:form?.id || '',
      disabled:el.matches(':disabled'), readonly:el.readOnly || false, context:context(el),
      dataset:{...el.dataset}, row:row ? {...row.dataset} : {}, modal:modal?.id || '',
      navigation:!!el.closest('aside'), required:el.required || false};
    if (el.tagName === 'FORM') item.label = [...el.querySelectorAll('button,input[type=submit]')].filter(visible).map(label).join(' / ') || context(el) || 'Form';
    if (['INPUT','TEXTAREA','SELECT'].includes(el.tagName) && el.type !== 'password' && el.type !== 'file') item.value=el.value;
    if (el.tagName==='SELECT') item.options=[...el.options].map(o=>({label:text(o),value:o.value,selected:o.selected}));
    if (el.type==='checkbox') item.value=el.checked ? 'İşaretli' : 'İşaretli değil';
    return item;
  });
  const blocks = [...document.querySelectorAll('.brand,.userbox,h1,h2,h3,h4,p,table,.workflow-strip,.badge,.empty-state,[role=alert]')]
    .filter(visible).filter(el => !el.closest('template')).map(el => ({tag:el.tagName.toLowerCase(),text:text(el)})).filter(x=>x.text);
  const covered = '.brand,.userbox,h1,h2,h3,h4,p,table,.workflow-strip,.badge,.empty-state,[role=alert],a,button,label,summary,script,style,template,select,textarea';
  const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT);
  while (walker.nextNode()) {
    const node=walker.currentNode, el=node.parentElement;
    const value=node.textContent.replace(/\s+/g,' ').trim();
    if (value && el && visible(el) && !el.closest(covered)) blocks.push({tag:el.tagName.toLowerCase(),text:value});
  }
  return {title:document.title,controls,blocks};
}"""


def ui_expand(page, action):
    """Only reveal existing UI controls; never submit a form from Chromium."""
    if action.startswith("click:"):
        page.locator(action.removeprefix("click:")).first.click()
        page.wait_for_load_state("networkidle")
    elif action.startswith("activity:"):
        page.locator("[data-tab=activities]").click()
        page.locator(
            ".edit-activity" if action == "activity:edit" else "#addActivity"
        ).first.click()
    elif action.startswith("budget:"):
        from playwright.sync_api import expect

        _, step, own, other = action.split(":")
        for i, did in enumerate((own, other)):
            row = page.locator(".leg").nth(i)
            row.locator(".department").select_option(did)
            row.locator(".direction").select_option(
                "decrease" if i == 0 else "increase"
            )
            row.locator(".leg-amount").fill("1.000,00")
        page.locator("#preview-button").click()
        expect(page.locator("#preview-result")).to_contain_text("1.000,00")
        if step == "reference":
            page.locator("#continue-reference").click()
            page.locator("#reference-date").fill("2026-10-01")
            page.locator("#reference-number").fill("SENTETIK-093")
            page.locator("#reason").fill("Sentetik örnek aktarma gerekçesi.")
        page.wait_for_load_state("networkidle")
    elif action.startswith("structure:"):
        step = int(action.split(":")[1])
        wizard = page.locator("#structure-wizard")
        for current in range(1, step):
            if current == 2:
                wizard.locator("[name=effective_date]").fill("2027-01-01")
            elif current == 3:
                wizard.locator("[name=council_decision_no]").fill("SENTETIK-093")
                wizard.locator("[name=council_decision_date]").fill("2026-10-01")
            elif current == 4:
                wizard.locator("[name=from_department_ids]").first.check()
                wizard.locator("[name=new_name]").fill("Örnek Yeni Hizmet Müdürlüğü")
            wizard.locator("#structure-next").click()
        wizard.scroll_into_view_if_needed()
    # Disclosure content is part of this screen; include it in the full-page PNG.
    page.locator("details").evaluate_all("nodes => nodes.forEach(el => el.open = true)")
    page.evaluate("window.scrollTo(0,0)")


# Explicit meanings for JS-only controls, read from the existing static scripts.
UI_JS = {
    "changePasswordButton": ("#passwordModal", "Parola değiştirme penceresini açar."),
    "cancelPasswordChange": (
        "#passwordModal",
        "Parola penceresini kaydetmeden kapatır.",
    ),
    "savePasswordChange": (
        "POST /api/me/change-password",
        "Mevcut ve yeni parolayı doğrulayarak parolayı değiştirir.",
    ),
    "notificationButton": (
        "GET /api/notifications → #notificationPanel",
        "Bildirim listesini açar.",
    ),
    "closeNotifications": ("#notificationPanel", "Bildirim panelini kapatır."),
    "newPeriodBtn": ("#newPeriodForm", "Yeni raporlama dönemi formunu açar."),
    "cancelPeriod": ("#newPeriodForm", "Dönem oluşturma formunu kapatır."),
    "createPeriod": ("POST /api/periods", "Seçilen yıl için taslak dönem oluşturur."),
    "applySections": (
        "POST /api/import-batches/{batchId}/apply → /report-builder/manifest",
        "Önizlenen Word bölümlerini uygular ve tamamlanma listesine geçer.",
    ),
    "revertSections": (
        "POST /api/import-batches/{batchId}/revert → /report-builder/manifest",
        "Uygun içe aktarmayı geri alır.",
    ),
    "structure-next": (
        "Bir sonraki sihirbaz adımı",
        "Geçerli adımın alanlarını doğrulayıp sonraki adıma geçer.",
    ),
    "structure-previous": (
        "Bir önceki sihirbaz adımı",
        "Kurum yapısı sihirbazında önceki adıma döner.",
    ),
    "previewBulkUsers": (
        "POST /api/users/bulk-preview",
        "CSV dosyasındaki kişisel hesapları oluşturmadan doğrular.",
    ),
    "applyBulkUsers": (
        "POST /api/users/bulk-import",
        "Önizlenen kişisel hesapları oluşturur.",
    ),
    "downloadBulkCredentials": (
        "Tarayıcı belleğinde üretilen CSV indirmesi",
        "Oluşturulan hesapların geçici parola dosyasını bir kez indirir.",
    ),
    "retireSelectedUsers": (
        "POST /api/users/bulk-deactivate",
        "Seçilen hesapları gerekçeyle pasifleştirir.",
    ),
    "saveAllMappings": (
        "POST /api/activity-mappings/bulk",
        "Seçilen alt program bağlantılarını kaydeder.",
    ),
    "previewMappingCsv": (
        "POST /api/activity-mappings/import-preview",
        "CSV eşleştirme dosyasını kontrol eder.",
    ),
    "applyMappingCsv": (
        "POST /api/activity-mappings/import-apply",
        "Doğrulanmış CSV eşleştirmelerini uygular.",
    ),
    "newUserBtn": ("#newUserForm", "Yeni kişisel hesap formunu açar."),
    "cancelUser": ("#newUserForm", "Yeni hesap formunu kapatır."),
    "saveUser": ("POST /api/users", "Girilen bilgilerle kişisel kullanıcı oluşturur."),
    "cancelReturn": ("#reviewModal", "İade gerekçesi penceresini kapatır."),
    "confirmReturn": (
        "POST /api/submissions/{id}/decision",
        "Gerekçeyi kaydedip seçilen paketi müdürlüğe iade eder.",
    ),
    "submitDepartment": (
        "POST /api/departments/{department}/submit",
        "Kontrolleri geçen müdürlük paketini Strateji incelemesine gönderir.",
    ),
    "savePersonnel": (
        "POST /api/departments/{department}/personnel-costs",
        "Personel ve SGK maliyetlerini kaydeder.",
    ),
    "addActivity": ("Faaliyet tablosu", "Yeni faaliyet satırını açar."),
    "add-leg": ("#legs", "Aktarmaya yeni müdürlük satırı ekler."),
    "preview-button": (
        "POST /api/budget-reallocations/choices + /api/budget-reallocations/preview",
        "Uygun ekonomik/faaliyet kayıtlarını ve aktarma önizlemesini hesaplar.",
    ),
    "refresh-preview": (
        "POST /api/budget-reallocations/preview",
        "Kayıt seçimlerine göre önizlemeyi yeniden hesaplar.",
    ),
    "continue-reference": ("#reference-section", "Aktarmanın dayanak formunu açar."),
    "apply-button": (
        "POST /api/budget-reallocations/apply → /budget-reallocations/{id}",
        "Dayanağı ve önizlemedeki tutarları tek işlemde kaydeder.",
    ),
    "reverse-button": (
        "POST /api/budget-reallocations/{operation}/reverse",
        "Önceki tutarlara dönüşü yeni bir telafi işlemi olarak kaydeder.",
    ),
    "selectStrongSuggestions": (
        "Eşleştirme tablosu",
        "Güçlü faaliyet eşleştirme önerilerini seçer.",
    ),
    "saveMappings": (
        "POST /api/activity-mappings/bulk",
        "Seçilen faaliyetlerin alt program bağlantılarını kaydeder.",
    ),
    "saveMappingsBottom": (
        "POST /api/activity-mappings/bulk",
        "Seçilen faaliyetlerin alt program bağlantılarını kaydeder.",
    ),
    "previewMappingImport": (
        "POST /api/activity-mappings/import-preview",
        "CSV eşleştirme dosyasını doğrular.",
    ),
    "applyMappingImport": (
        "POST /api/activity-mappings/import-apply",
        "Doğrulanmış CSV eşleştirmelerini uygular.",
    ),
    "clearNotifications": ("Bildirim listesi", "Bildirim görünümünü kapatır."),
}
UI_CLASSES = {
    "cancel-new": ("Aynı faaliyet tablosu", "Yeni faaliyet satırını kaydetmeden kaldırır."),
    "save-new": ("POST /api/activities", "Yeni faaliyet ve maliyetlerini kaydeder."),
    "tab": ("#{tab}", "Seçilen veri giriş sekmesini gösterir."),
    "save-indicator": (
        "POST /api/indicators/{id}/save",
        "Gösterge alanlarını ve yıllık değerlerini kaydeder, kontrolleri yeniler.",
    ),
    "issue-jump": (
        "Aynı gösterge kartındaki sorunlu alan",
        "Kontrol mesajının ilgili alanına odaklanır.",
    ),
    "edit-activity": ("Faaliyet satırı", "Faaliyet alanlarını düzenlemeye açar."),
    "save-activity": (
        "PUT /api/activities/{activity}",
        "Faaliyet bilgileri ve maliyetlerini kaydeder.",
    ),
    "cancel-activity": (
        "Faaliyet satırı",
        "Kaydedilmemiş satır değişikliklerini iptal eder.",
    ),
    "delete-activity": (
        "DELETE /api/activities/{activity}",
        "Onay sonrası taslak faaliyeti siler.",
    ),
    "approve-submission": (
        "POST /api/submissions/{id}/decision",
        "Gönderilen müdürlük paketini onaylar.",
    ),
    "return-submission": (
        "#reviewModal",
        "Seçilen paket için iade gerekçesi penceresini açar.",
    ),
    "toggle-submit-authorization": (
        "POST /api/users/{user}/submit-authorization",
        "Personelin paket gönderim yetkisini açar veya kaldırır.",
    ),
    "save-user-position": (
        "POST /api/users/{user}/position",
        "Kullanıcının personel veya müdür görevini kaydeder.",
    ),
    "toggle-user-active": (
        "POST /api/users/{user}/active",
        "Hesabın giriş erişimini etkinleştirir veya kapatır.",
    ),
    "show-password-reset": (
        "Hesabın geçici parola alanı",
        "Geçici parola yenileme alanını açar veya kapatır.",
    ),
    "save-password-reset": (
        "POST /api/users/{user}/reset-password",
        "Hesabın geçici parolasını yeniler.",
    ),
    "create-department-user": (
        "#newUserForm",
        "Seçilen müdürlük için hesap formunu açar.",
    ),
    "import-item": (
        "GET /api/import-batches/{import}",
        "Word içe aktarma kaydının ayrıntılarını aynı ekranda gösterir.",
    ),
    "snapshot-item": (
        "GET /api/template-snapshots/{snapshot}",
        "Şablon analizinin ayrıntılarını aynı ekranda gösterir.",
    ),
    "choose-suggestion": (
        "Eşleştirme seçimi",
        "Bu faaliyet için önerilen alt programı seçer.",
    ),
    "revert-mapping-batch": (
        "POST /api/activity-mappings/batches/{batch}/revert",
        "Seçilen eşleştirme aktarımını geri alır.",
    ),
    "remove-leg": (
        "Aktarma satırı",
        "Müdürlük satırını taslak aktarma formundan kaldırır.",
    ),
    "ar-add-row": ("Aynı kaynak/malî tablo", "Yeni boş veri satırı ekler."),
    "ar-remove-row": (
        "Aynı kaynak/malî tablo",
        "Satırı formdan kaldırır; kalıcı işlem kaydedince uygulanır.",
    ),
    "confirm-source-total": (
        "PUT /api/activities/{activity}",
        "Gerekçeyle hesaplanan faaliyet toplamını kabul eder.",
    ),
}


def ui_control(control):
    """DOM destinations first; add Turkish behavior from the route/JS contract."""
    c = control
    if c["formId"] == "firstPasswordForm":
        return (
            "POST /api/me/change-password → /modules",
            "Geçici parolayı doğrular; kişisel parolayı belirleyip Modüller ekranına geçer.",
        )
    if "review-action" in c["classes"]:
        return (
            ("#reviewModal", "Seçilen paket için iade gerekçesi penceresini açar.")
            if c["dataset"].get("decision") == "return"
            else (
                "POST /api/submissions/" + c["dataset"]["id"] + "/decision",
                "Paketi onaylar; veri girişini kilitler.",
            )
        )
    if "edit-activity" in c["classes"] and c["label"] == "Kaydet":
        return (
            "PUT /api/activities/" + c["row"]["activity"],
            "Düzenlenen faaliyet bilgilerini ve maliyetlerini kaydeder.",
        )
    if "delete-activity" in c["classes"] and c["dataset"].get("mode") == "cancel":
        return (
            "Aynı sayfayı yeniden yükle",
            "Faaliyet satırındaki kaydedilmemiş değişikliklerden vazgeçer.",
        )
    if c["tag"] == "a":
        return c["href"], "Bağlantıdaki ekranı veya dosyayı açar: " + c["label"] + "."
    if c["tag"] == "summary":
        return (
            "Aynı sayfadaki ayrıntı bölümü",
            "İçerik ve form ayrıntılarını açar veya daraltır.",
        )
    if c["tag"] in {"input", "select", "textarea"}:
        field_actions = {
            "templateUpload": "POST /api/periods/{period_id}/template-upload",
            "dataImportUpload": "POST /api/periods/{period_id}/data-import-preview",
            "bulkUserFile": "POST /api/users/bulk-preview → /api/users/bulk-import",
            "mappingCsvFile": "POST /api/activity-mappings/import-preview → /api/activity-mappings/import-apply",
            "currentPassword": "POST /api/me/change-password",
            "newPassword": "POST /api/me/change-password",
            "confirmPassword": "POST /api/me/change-password",
            "returnNote": "POST /api/submissions/{id}/decision",
            "retireUsersReason": "POST /api/users/bulk-deactivate",
            "mappingDepartmentFilter": "Aynı sayfadaki faaliyet tablosunu müdürlüğe göre filtrele",
            "mappingSearch": "Aynı sayfadaki faaliyet tablosunda ada göre ara",
        }
        target = (
            c["method"] + " " + c["action"]
            if c["action"]
            else field_actions.get(c["id"], "Aynı ekranın JavaScript işlem verisi")
        )
        if "contract-bound" in c["classes"]:
            target = (
                "POST /api/activities/" + c["dataset"]["activityId"] + "/contract-bound"
            )
        elif c["row"].get("new") == "1":
            target = "POST /api/activities"
        elif c["row"].get("activity"):
            target = "PUT /api/activities/" + c["row"]["activity"]
        elif c["row"].get("id") and not c["action"]:
            target = "POST /api/indicators/" + c["row"]["id"] + "/save"
        elif "personnel-cost" in c["classes"]:
            target = "POST /api/departments/{department_id}/personnel-costs"
        meaning = (
            "Listeden seçim yapar."
            if c["tag"] == "select"
            else (
                "Dosya seçer; ilgili işlemle yüklenir."
                if c["type"] == "file"
                else (
                    "Seçeneği işaretler veya kaldırır."
                    if c["type"] == "checkbox"
                    else "İşlem için " + (c["label"] or c["name"]) + " değerini girer."
                )
            )
        )
        if c["readonly"] or c["disabled"]:
            meaning = "Mevcut değeri gösterir; bu görünümde değiştirilemez."
        return target, meaning
    if c["action"] and (c["tag"] == "form" or c["type"] == "submit"):
        target = c["method"] + " " + c["action"]
        action = c["action"]
        meanings = [
            ("/logout", "Oturumu kapatıp giriş ekranına döner."),
            ("/login", "E-posta ve parola ile oturum açar; Modüller ekranına geçer."),
            ("/reopen", "Gerekçeyle onaylı raporu/paketi yeniden düzenlemeye açar."),
            ("/freeze", "Kontrolleri geçen raporun değişmez sürümünü oluşturur."),
            (
                "/transition",
                "Seçilen gönderim, onay veya iade kararını gerekçesiyle kaydeder.",
            ),
            ("/apply", "Önizlemedeki aktarımı doğrulayıp uygular."),
            ("/activate", "Seçilen dönemi veya planı etkinleştirir."),
            ("/validate", "Yüklenen dosyanın doğrulamasını yeniler."),
            ("/mappings", "Ham birim adını seçilen müdürlük veya kurulla eşler."),
            ("/commit", "Doğrulanmış yüklemeyi plan olarak kaydeder."),
            ("/close", "İzleme dönemini gerekçeyle kapatır."),
            ("/open", "Plan için altı aylık izleme dönemi açar."),
            ("/notes", "Eyleme katkı notunu kaydeder."),
            ("/entry", "Eylemin gerçekleşme, ilerleme ve kanıt bilgilerini kaydeder."),
            ("/remove-draft", "Seçilen taslak kaydı gerekçeyle kaldırır."),
            ("/imports", "Seçilen plan dosyasını yükleyip kontrol ekranına geçer."),
        ]
        if c["method"] == "GET":
            meaning = "Alanlarla listeyi veya rapor görünümünü filtreler."
        else:
            meaning = next(
                (v for k, v in meanings if action.endswith(k)),
                (c["label"] or "Form")
                + " işlemini form alanlarıyla gönderir; doğrulama sonrası kaydeder veya hata gösterir.",
            )
        return target, meaning
    value = UI_JS.get(c["id"]) or next(
        (UI_CLASSES[k] for k in c["classes"] if k in UI_CLASSES), None
    )
    if value:
        params = {**c["row"], **c["dataset"]}
        target = value[0]
        for key, val in params.items():
            target = target.replace("{" + key + "}", val)
        return target, value[1]
    if "addEvidence" in c["dataset"]:
        return "Aynı gerçekleşme formu", "Yeni kanıt kaynak satırı ekler."
    if "removeEvidence" in c["dataset"]:
        return (
            "Aynı gerçekleşme formu",
            "Kanıt satırını formdan kaldırır; kayıt gönderildiğinde uygulanır.",
        )
    if "cancelRemoval" in c["dataset"]:
        return "Aynı taslak kaldırma kutusu", "Kaldırma kutusunu kaydetmeden kapatır."
    if c["dataset"].get("arAddRow") is not None:
        return "Aynı malî tablo", "Tabloya yeni satır ekler."
    if c["dataset"].get("arRemoveRow") is not None:
        return "Aynı malî tablo", "Satırı kaydetme öncesi formdan kaldırır."
    return "BELİRSİZ", "JavaScript işlem eşlemesi eksik."


def ui_add(
    data,
    pattern,
    account,
    state="Dolu",
    path=None,
    action="top",
    status=200,
    response="",
):
    title, purpose = UI_PAGES[pattern]
    if status >= 400:
        purpose = "İşlenemeyen istek veya eksik kayıt için hata mesajını ve düzeltme adımını gösterir."
    number = len(data["_screen_meta"]) + 1
    slug = pattern.strip("/").split("/")[0]
    filename = f"{number:03d}-{slug}-{account}.png"
    path = path or ui_path(pattern, data, account)
    shot = Shot(filename, path, title, account, action, True)
    data["_shots"].append(shot)
    data["_screen_meta"][filename] = dict(
        pattern=pattern,
        path=path,
        title=title,
        purpose=purpose,
        account=account,
        state=state,
        status=status,
        response=response,
    )
    return filename


def ui_probe(data):
    """Every registered GET route gets an explicit inventory disposition."""
    data["_inventory"] = []
    for pattern in ui_routes():
        if pattern == "/":
            data["_inventory"].append(
                dict(
                    pattern=pattern,
                    reason="Oturum yoksa /login; normal oturumda /modules; geçici parola gerekiyorsa /change-password yönlendirmesi (303), ayrı HTML ekranı yok.",
                )
            )
            continue
        if pattern not in UI_PAGES:
            reason = (
                "FastAPI otomatik API belgesi / OAuth yardımcı sayfası; ürün iş ekranı değil."
                if pattern in {"/docs", "/docs/oauth2-redirect", "/redoc"}
                else (
                    "OpenAPI JSON şeması; HTML ekranı değil."
                    if pattern == "/openapi.json"
                    else None
                )
            )
            reason = reason or (
                "API / JSON yanıtı; HTML ekranı değil."
                if pattern.startswith("/api/") or pattern == "/health"
                else (
                    "JavaScript kaynağı; HTML ekranı değil."
                    if pattern.endswith(".js")
                    else (
                        "Dosya indirme / dışa aktarma; HTML ekranı değil."
                        if any(
                            x in pattern
                            for x in (
                                ".csv",
                                ".xlsx",
                                ".json",
                                ".docx",
                                ".pdf",
                                ".{format}",
                            )
                        )
                        else None
                    )
                )
            )
            data["_inventory"].append(
                dict(
                    pattern=pattern,
                    reason=reason or "Bilinmeyen GET rotası; ekran tanımı eklenmeli.",
                )
            )
            check(
                reason is not None,
                "Yeni GET rotası sınıflandırılmamış; UI_PAGES envanterini güncelleyin.",
            )
            continue
        item = dict(pattern=pattern, roles={})
        data["_inventory"].append(item)
        if pattern == "/login":
            item["roles"]["anonymous"] = ui_add(
                data, pattern, "anonymous", state="Oturum kapalı"
            )
            continue
        if pattern == "/budget-reallocations/{operation_id}":
            item["reason"] = (
                "İlk aşamada işlem yok; ikinci aşamada gerçek aktarma servisiyle doldurulup yakalanır."
            )
            continue
        for account in ("strategy", "department"):
            path = ui_path(pattern, data, account)
            response = data["_clients"][account].get(
                path, headers={"Accept": "text/html"}
            )
            if response.status_code == 200 and "text/html" in response.headers.get(
                "content-type", ""
            ):
                item["roles"][account] = ui_add(
                    data,
                    pattern,
                    account,
                    state=(
                        "Boş liste"
                        if pattern == "/budget-reallocations"
                        else "Dolu / taslak"
                    ),
                )
            else:
                reason = f"HTTP {response.status_code}"
                if response.is_redirect:
                    reason += (
                        " → "
                        + response.headers.get("location", "")
                        + "; ayrı ekran yok."
                    )
                elif response.status_code == 403:
                    reason += "; bu rolün bu ekrana erişimi yok."
                else:
                    reason += "; beklenen normal HTML ekranı açılamadı."
                item["roles"][account] = reason
                # Unexpected errors are captured as real failures, not concealed.
                if response.status_code not in (302, 303, 403):
                    ui_add(
                        data,
                        pattern,
                        account,
                        state="Açılamadı (envanter)",
                        status=response.status_code,
                    )
                    data["_unexpected"].append(
                        pattern + " / " + account + " / " + reason
                    )
    # Distinct section types share a GET path but have different controls.
    for account in ("strategy", "department"):
        for code in (
            "SUNUS",
            "I.C.1",
            "I.C.2",
            "I.C.3",
            "I.C.4",
            "III.A.1",
            "III.A.2",
            "III.A.3",
            "III.B.2.i",
            "EKLER",
        ):
            if account == "department" and code == "SUNUS":
                continue
            base = ui_path("/activity-sections/reports/{document_id}", data, account)
            ui_add(
                data,
                "/activity-sections/reports/{document_id}",
                account,
                state="Bölüm " + code,
                path=base + "?code=" + code,
            )
        ui_add(
            data,
            "/department/{department_id}",
            account,
            state="Faaliyet Maliyetleri sekmesi",
            action='click:[data-tab="activities"]',
        )
    ui_add(
        data,
        "/dashboard",
        "strategy",
        state="Yeni dönem formu",
        action="click:#newPeriodBtn",
    )
    for step in range(2, 7):
        ui_add(
            data,
            "/structure",
            "strategy",
            state=f"Yapı sihirbazı · adım {step}/6",
            action=f"structure:{step}",
        )
    for state in ("edit", "new"):
        ui_add(
            data,
            "/department/{department_id}",
            "department",
            state="Faaliyet satırı · " + state,
            action="activity:" + state,
        )
    ui_add(
        data,
        "/users",
        "strategy",
        state="Geçici parola yenileme alanı",
        action="click:.show-password-reset",
    )
    ui_add(
        data,
        "/users",
        "strategy",
        state="Yeni kullanıcı formu",
        action="click:#newUserBtn",
    )
    ui_add(
        data,
        "/review",
        "strategy",
        state="İade gerekçesi penceresi",
        action="click:[data-decision=return]",
    )
    ui_add(
        data,
        "/modules",
        "department",
        state="Parola penceresi",
        action="click:#changePasswordButton",
    )
    ui_add(
        data,
        "/modules",
        "strategy",
        state="Bildirim paneli",
        action="click:#notificationButton",
    )
    ui_add(
        data,
        "/activity-report",
        "department",
        state="Boş yıl listesi",
        path="/activity-report?year=2026",
    )
    ui_add(
        data,
        "/activity-sections/reports/{document_id}",
        "strategy",
        state="Hata sayfası (404)",
        path="/activity-sections/reports/999999",
        status=404,
    )


def ui_validation(data):
    """Render a real failed POST before the read-only browser capture."""
    import re

    client = data["_clients"]["strategy"]
    html = client.get("/modules").text
    token = re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)
    response = client.post(
        "/report-builder/section",
        data={
            "csrf_token": token,
            "period_id": data["period"],
            "title": "",
            "kind": "text",
            "content": "Sentetik doğrulama denemesi",
            "position": "1",
            "placement": "before",
        },
    )
    check(
        response.status_code == 422
        and "text/html" in response.headers.get("content-type", ""),
        "Gerçek form doğrulama yanıtı üretilemedi.",
    )
    data["_responses"] = {("strategy", "validation"): response}
    ui_add(
        data,
        "/report-builder",
        "strategy",
        state="Başlığı boş form · doğrulama hatası",
        path="/report-builder/section?period_id=" + str(data["period"]),
        status=422,
        response="validation",
    )


def ui_approve(engine, data):
    from sqlalchemy import select
    from sqlalchemy.orm import Session
    from app.models import User, Submission
    from app import budget_service as budget
    from app import internal_control_monitoring_service as monitoring

    with Session(engine) as db:
        strategy = db.scalar(
            select(User).where(User.email == data["accounts"]["strategy"])
        )
        department = db.scalar(
            select(User).where(User.email == data["accounts"]["department"])
        )
        monitoring.transition(
            db,
            department,
            data["ic_batch"],
            "submit",
            monitoring.state(db, data["ic_batch"])[1],
        )
        db.commit()
        monitoring.transition(
            db,
            strategy,
            data["ic_batch"],
            "approve",
            monitoring.state(db, data["ic_batch"])[1],
        )
        db.commit()
        for sub in db.scalars(select(Submission)):
            sub.status = "approved"
        db.commit()
        request = dict(
            period_id=data["period"],
            year=2027,
            value_kind="budget",
            change_total=False,
            change_outside=False,
            legs=[
                dict(department_id=data["own"], direction="decrease", amount="1000.00"),
                dict(
                    department_id=data["submitted"],
                    direction="increase",
                    amount="1000.00",
                ),
            ],
        )
        preview = budget.preview(db, strategy, request)
        operation = budget.apply(
            db,
            strategy,
            dict(
                token=preview["token"],
                phase="before_adoption",
                reference_type="mayor_instruction",
                reference_date="2026-10-01",
                reference_number="SENTETIK-093",
                reason="Sentetik örnek: hizmet noktaları arasında kaynak düzenlemesi.",
            ),
        )
        db.commit()
        data["operation"] = operation.id
    data["_shots"] = []
    for pattern in (
        "/department/{department_id}",
        "/program/{department_id}",
        "/internal-control/monitoring/batches/{batch_id}",
        "/internal-control/our-actions",
    ):
        for account in ("strategy", "department"):
            if pattern == "/internal-control/our-actions" and account == "strategy":
                continue
            ui_add(data, pattern, account, state="Onaylı · salt okunur")
    for step in ("preview", "reference"):
        ui_add(
            data,
            "/budget-reallocations/new",
            "strategy",
            state="Aktarma sihirbazı · " + step,
            action=f'budget:{step}:{data["own"]}:{data["submitted"]}',
        )
    for pattern in (
        "/budget-reallocations",
        "/budget-reallocations/{operation_id}",
        "/review",
        "/internal-control/monitoring",
    ):
        filename = ui_add(data, pattern, "strategy", state="Onay sonrası / dolu")
        if pattern == "/budget-reallocations/{operation_id}":
            item = next(x for x in data["_inventory"] if x["pattern"] == pattern)
            item.pop("reason", None)
            item["roles"] = {
                "strategy": filename,
                "department": "HTTP 403; bütçe aktarma yetkisi yok.",
            }


def ui_write(folder, data, results):
    import json
    import re
    from urllib.parse import urlsplit

    def cell(value):
        return (
            str(value)
            .replace("|", "&#124;")
            .replace("\n", " ")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    screens = data["_screen_meta"]
    for filename in screens:
        results.setdefault(filename, "Yakalama başlamadı.")
    captured = [
        (f, s) for f, s in screens.items() if results.get(f) is None and "dom" in s
    ]
    lines = [
        "# CIV-093 · Arayüz ekran haritası",
        "",
        "Bu harita mevcut arayüzü belgeler. Bütün kurum, hesap, gösterge, tutar, kanıt ve plan kayıtları sentetiktir; "
        "canlı veritabanı ve kitap dosyası kullanılmaz. Strateji hesabı `strategy/director`, müdürlük hesabı "
        "`department/director` rolündedir. Personel ve kurulum hesapları tarayıcı açmaz. Rastgele parolalar kaydedilmez.",
        "",
        "## Tek komutla yeniden üretim",
        "",
        "Önkoşullar: Python 3.12, PostgreSQL 16+ üzerinde CREATE DATABASE yetkili yerel test rolü; "
        "`CIVENSA_TEST_POSTGRES_URL` bu sunucuyu göstermelidir. Uygulamanın DATABASE_URL / MIGRATION_DATABASE_URL ayarları "
        "yok sayılır. Temiz checkout içinde depo kökünden:",
        "",
        "```sh",
        "python -m pip install -r requirements-dev.txt && python -m playwright install --with-deps chromium && python scripts/guide_screenshots.py --ui-map --out docs/screenshots/ui-map-v1",
        "```",
        "",
        "Sistem Chromium kurulumu varsa `CIVENSA_CHROMIUM_PATH` ile seçilebilir. Geçici veritabanı gerçek Alembic "
        "geçişleriyle kurulur, bitişte veya hatada silinir. Betik CIV-054 ile aynı PostgreSQL/HTTP/Playwright yakalama "
        "akışını kullanır; `--ui-map` verilmezse eski kılavuz modu çalışır.",
        "",
        f"{data.get('viewport_width', 1440)} px genişlik, tam sayfa PNG, deviceScaleFactor 1, açık tema, tr-TR, Europe/Istanbul; her PNG <1.500.000 bayt. "
        "Uzun ekranlarda boyut sınırını sağlamak için 256 renkli PNG kodlaması kullanılabilir; genişlik ve içerik korunur. Ayrıntı kutuları açık çekilir. Sekme ve pencereler ayrıca çekilir. DOM tablosu aşağı kaydırılarak ulaşılabilen "
        "öğeleri de kapsar; kapalı sekme/pencerelerin gizli öğeleri o durumun ayrı çekiminde bulunur. "
        "Yatay taşan tabloların tüm sütunları DOM dökümünde yer alır.",
        "",
        "Tarayıcı yalnız loopback GET ve iki salt okunur bütçe hesaplama POST rotasını kullanır (`choices`, `preview`); diğer yazmalar engellenir. Doğrulama örneği, tarayıcıdan önce gerçek oturum/CSRF ile yapılan "
        "başarısız POST yanıtının aynısıdır; ekran sahte hata metniyle değiştirilmez. Her yakalama aşamasından önce/sonra "
        "bütün model tablolarının özeti karşılaştırılır. İkinci aşamada sentetik PP paketleri onaylı fixture durumuna "
        "alınır, İç Kontrol paketi gerçek gönderim/onay servislerinden geçer ve gerçek bütçe aktarma servisi çalışır. "
        "Sabit PP yayını gerçek yayın serileştiricisiyle kurulur; resmî yayın önkoşullarının tamamlandığı iddia edilmez.",
        "",
        f"Chromium: {data.get('browser','başlatılamadı')}. Rota envanteri: {len(data.get('_inventory',[]))} kayıtlı GET; "
        f"HTML ekran türü: {len(UI_PAGES)}; üretilen çekim: {len(captured)}/{len(screens)}.",
        "",
        "Her E numarası DOM sırasındaki tek öğedir. Formlar ve düğmeleri ayrı satırlardadır; yinelenen etiketlerde bağlam "
        "hangi satır/karta ait olduğunu belirtir. URL ve eylemler DOM href/action değerlerinden, JavaScript düğmelerinin "
        "anlamları mevcut statik betiklerden alınır. **Pasif** düğme hedefi belgeler ama bu durumda çalıştırılamaz. "
        "Parola/CSRF/oturum içerikleri döküme alınmaz.",
        "",
        "[Modüllere göre gezinme diyagramı](navigation.md) · [Makinece okunabilir DOM ve rota envanteri](screens.json)",
        "",
        "## Ekran dizini",
        "",
        "| No | Ekran | Rol | Durum | Görsel |",
        "| --- | --- | --- | --- | --- |",
    ]
    for i, (filename, screen) in enumerate(screens.items(), 1):
        lines.append(
            f'| {i:03d} | [{cell(screen["title"])}](#ekran-{i:03d}) | {screen["account"]} | {cell(screen["state"])} | '
            + (
                f"[{filename}]({filename})"
                if results.get(filename) is None
                else "ÜRETİLEMEDİ"
            )
            + " |"
        )
    lines += [
        "",
        "## GET rota kapsamı ve yakalanamayanlar",
        "",
        "Bu tablo uygulamanın kayıtlı GET rotalarından üretilir; "
        "HTML dışı rotalar da dışlama gerekçeleriyle yer alır. 403 ve yönlendirmeler yetki kapsamını gösterir. "
        "İki rolün erişebildiği normal HTML ekranları için açıklamasız rota bırakılmaz.",
        "",
        "| URL kalıbı | Yakalama / gerekçe |",
        "| --- | --- |",
    ]
    for item in data.get("_inventory", []):
        parts = []
        for role, value in item.get("roles", {}).items():
            parts.append(
                role
                + ": "
                + (
                    f"[{value}]({value})"
                    if value in screens and results.get(value) is None
                    else value
                )
            )
        if item.get("reason"):
            parts.append(item["reason"])
        lines.append("| `" + item["pattern"] + "` | " + cell("; ".join(parts)) + " |")
    lines += [
        "",
        "HTML rotası için beklenmeyen durumlar: "
        + ("; ".join(data["_unexpected"]) or "Yok."),
        "",
        "Yakalama hataları: "
        + ("; ".join(f"{f}: {e}" for f, e in results.items() if e) or "Yok."),
        "",
        "## Ekran ayrıntıları",
        "",
    ]
    unknown = []
    for i, (filename, s) in enumerate(screens.items(), 1):
        lines += [
            f'<a id="ekran-{i:03d}"></a>',
            "",
            f'### {i:03d} · {s["title"]} · {s["state"]}',
            "",
            f'URL kalıbı: `{s["pattern"]}` · Çekilen URL: `{s["path"]}` · Rol: **{s["account"]}** · HTTP {s["status"]}.',
            "",
            s["purpose"],
            "",
        ]
        if results.get(filename) or "dom" not in s:
            lines += [
                "ÜRETİLEMEDİ: " + str(results.get(filename) or "Yakalama başlamadı."),
                "",
            ]
            continue
        lines += [
            f'![{s["title"]} · {s["account"]} · {s["state"]}]({filename})',
            "",
            "**Görünen metinler ve tablolar:**",
            "",
        ]
        for block in s["dom"]["blocks"]:
            lines.append("- `" + block["tag"] + "`: " + cell(block["text"]))
        lines += [
            "",
            "**Bağlantı, düğme, form ve alanlar:**",
            "",
            "| Öğe | Tür / etiket / bağlam | Hedef URL veya eylem | Ne yapar / mevcut değer |",
            "| --- | --- | --- | --- |",
        ]
        for n, c in enumerate(s["dom"]["controls"], 1):
            target, meaning = ui_control(c)
            c["target"] = target
            c["meaning"] = meaning
            if target == "BELİRSİZ":
                unknown.append((filename, c["id"], c["classes"], c["label"]))
            detail = meaning
            if c["disabled"]:
                detail += " **Pasif.**"
            if c.get("required"):
                detail += " Zorunlu."
            if "value" in c and c["value"] != "":
                detail += " Değer: " + str(c["value"]) + "."
            if c.get("options"):
                detail += (
                    " Seçenekler: " + "; ".join(o["label"] for o in c["options"]) + "."
                )
            label = c["tag"] + " · " + (c["label"] or "(etiketsiz)")
            if c["id"]:
                label += " (#" + c["id"] + ")"
            if c["context"]:
                label += " · " + c["context"]
            lines.append(
                "| E"
                + f"{n:03d}"
                + " | "
                + cell(label)
                + " | `"
                + cell(target)
                + "` | "
                + cell(detail)
                + " |"
            )
        lines.append("")
    (folder / "INDEX.md").write_text("\n".join(lines), encoding="utf-8")
    # No response bodies, cookie jars, tokens or passwords are serialized.
    manifest = dict(
        version=1,
        viewport_width=data.get('viewport_width', 1440),
        browser=data.get("browser"),
        inventory=data.get("_inventory", []),
        screens=[
            dict(
                filename=f,
                **{k: v for k, v in s.items() if k != "response"},
                error=results.get(f),
            )
            for f, s in screens.items()
        ],
    )
    (folder / "screens.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    ui_navigation(folder, data)
    if unknown:
        print(
            "Eşleştirilmemiş JS öğeleri: "
            + repr(sorted(set((b, tuple(c), d) for a, b, c, d in unknown))),
            flush=True,
        )
    return len(unknown)


def ui_navigation(folder, data):
    import re
    from urllib.parse import urlsplit

    nodes = {p: f"n{i}" for i, p in enumerate(UI_PAGES, 1)}

    def match(href):
        path = urlsplit(href).path
        return next(
            (
                p
                for p in UI_PAGES
                if re.fullmatch(re.sub(r"\{[^}]+\}", r"[^/]+", p), path)
            ),
            None,
        )

    nav = {"strategy": set(), "department": set()}
    edges = set()
    for screen in data["_screen_meta"].values():
        for c in screen.get("dom", {}).get("controls", []):
            if c["tag"] != "a":
                continue
            target = match(c["href"])
            if not target:
                continue
            source = screen["pattern"]
            if c["navigation"]:
                nav[screen["account"]].add(target)
            elif source != target:
                edges.add((source, target))
    lines = [
        "# Ekranlar arası gezinme",
        "",
        "Düğümler URL kalıplarını temsil eder; rol, sekme, bölüm ve onay durumları "
        "INDEX içindeki ayrı görsellere ayrılmıştır. Düz oklar DOM bağlantılarıdır. Ortak sol menü her ekranda tekrar "
        "çizilmez: rol menüsü düğümünden erişilir. Kesik oklar aynı ekran içindeki pencere/sekme veya POST sonrası "
        "dönüşü gösterir; düğme ve form eylemlerinin tam URL’leri INDEX tablolarındadır. Dosya indirmeleri ekran değildir.",
        "",
        "```mermaid",
        "flowchart LR",
        '  strategy_menu["Strateji ortak sol menüsü"]',
        '  department_menu["Müdürlük ortak sol menüsü"]',
    ]
    for i, module in enumerate(
        ("Performans Programı", "Faaliyet Raporu", "İç Kontrol", "Yönetim"), 1
    ):
        lines.append(f'  subgraph m{i}["{module}"]')
        for p, node in nodes.items():
            if ui_module(p) == module:
                lines.append(f'    {node}["{UI_PAGES[p][0]}<br/>{p}"]')
        lines.append("  end")
    for role, targets in nav.items():
        for target in sorted(targets):
            lines.append(f"  {role}_menu --> {nodes[target]}")
    for source, target in sorted(edges):
        lines.append(f"  {nodes[source]} --> {nodes[target]}")
    lines += [
        f'  {nodes["/login"]} -. "POST giriş" .-> {nodes["/modules"]}',
        f'  {nodes["/modules"]} -. "role göre ortak menü" .-> strategy_menu',
        f'  {nodes["/modules"]} -. "role göre ortak menü" .-> department_menu',
        f'  {nodes["/budget-reallocations/new"]} -. "önizleme + uygulama" .-> {nodes["/budget-reallocations/{operation_id}"]}',
        f'  {nodes["/internal-control"]} -. "plan yükle" .-> {nodes["/internal-control/imports/{batch_id}"]}',
        f'  {nodes["/internal-control/imports/{batch_id}"]} -. "planı kaydet" .-> {nodes["/internal-control/plans/{plan_id}"]}',
        f'  {nodes["/activity-reports"]} -. "revizyon oluştur" .-> {nodes["/activity-reports/revisions/{revision_id}"]}',
        "```",
        "",
        "Ortak menülerde **Parolamı Değiştir** aynı ekranın penceresini açar; **Çıkış** POST `/logout` "
        "sonrası `/login` ekranına döner. Bildirimler GET `/api/notifications` ile aynı ekranda panel açar; "
        "okunmamış bir bildirime geçiş POST `/api/notifications/{id}/read` sonrasında bildirimin bağlantısını açar. "
        "Müdürlük veri girişinin Gösterge / Faaliyet sekmeleri aynı URL içinde kalır.",
        "",
    ]
    (folder / "navigation.md").write_text("\n".join(lines), encoding="utf-8")


def ui_publish(folder, destination, screens):
    import json
    import re

    previous = destination / "screens.json"
    owned = set(screens)
    if previous.exists():
        for screen in json.loads(previous.read_text(encoding="utf-8")).get(
            "screens", []
        ):
            name = screen.get("filename", "")
            if re.fullmatch(r"\d{3}-[a-z-]+\.png", name):
                owned.add(name)
    destination.mkdir(parents=True, exist_ok=True)
    for name in owned:
        if not (folder / name).exists():
            (destination / name).unlink(missing_ok=True)
    for file in folder.iterdir():
        shutil.copyfile(file, destination / file.name)


def ui_main(args, url):
    data = None
    results = {}
    with TemporaryDirectory(prefix="civensa_ui_map_") as temporary:
        folder = Path(temporary)
        try:
            from scripts.first_load_dryrun import disposable_database

            with disposable_database(url) as engine:
                import app.db as app_db

                app_db.engine.dispose()
                app_db.engine = engine
                app_db.SessionLocal.configure(bind=engine)
                # A second viewport in the same process must not retain the
                # first (already dropped) disposable database in health().
                import app.main as app_main
                app_main.engine = engine
                data = ui_setup(engine)
                data.update(ui_map=True, _shots=[], _screen_meta={}, _unexpected=[],
                            viewport_width=getattr(args, 'viewport_width', 1440))
                with local_http(engine, data) as base:
                    ui_probe(data)
                    ui_validation(data)
                    for phase in ("draft", "approved"):
                        if phase == "approved":
                            ui_approve(engine, data)
                        for shot in data["_shots"]:
                            results[shot.filename] = "Yakalama tamamlanmadı."
                        before = ui_state_digest(engine)
                        capture(base, folder, data, results)
                        check(
                            before == ui_state_digest(engine),
                            "Yakalama sırasında kalıcı durum değişti; çıktılar geçersiz.",
                        )
                unknown = ui_write(folder, data, results)
                check(
                    not unknown,
                    "Bazı JavaScript düğmelerinin anlamı eksik; envanteri tamamlayın.",
                )
                check(
                    not data["_unexpected"],
                    "Bazı normal HTML ekranları açılamadı; envanteri kontrol edin.",
                )
                check(
                    all(error is None for error in results.values()),
                    "Bazı ekran görüntüleri üretilemedi; INDEX raporuna bakın.",
                )
        except (Exception, KeyboardInterrupt) as exc:
            if not isinstance(exc, (GuideError, KeyboardInterrupt)):
                # Only code locations; exception messages may contain DB parameters.
                import traceback
                print('Yakalama hata konumu: ' + ' → '.join(
                    f'{Path(frame.filename).name}:{frame.lineno} ({frame.name})'
                    for frame in traceback.extract_tb(exc.__traceback__)), file=sys.stderr)
            message = (
                str(exc)
                if isinstance(exc, GuideError)
                else f"UI haritası tamamlanamadı ({type(exc).__name__})."
            )
            print(message, file=sys.stderr)
            if isinstance(exc, GuideError) and "kalıcı durum değişti" in message:
                results = {name: message for name in data["_screen_meta"]}
                for file in folder.glob("*.png"):
                    file.unlink()
            # Diagnostic output remains reviewable, but never replaces a complete map.
            if data and data.get("_screen_meta"):
                ui_write(folder, data, results)
                ui_publish(folder, args.out, data["_screen_meta"])
            return 1
        ui_publish(folder, args.out, data["_screen_meta"])
    print(f'{len(results)}/{len(results)} ekran üretildi: {args.out / "INDEX.md"}')
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out", type=Path, default=ROOT / "docs/screenshots/guide-pp-v1"
    )
    parser.add_argument(
        "--ui-map",
        action="store_true",
        help="CIV-093: bütün HTML ekranları ve DOM öğeleri",
    )
    parser.add_argument('--viewport-width', type=int, choices=(1440, 390), default=1440,
                        help='UI haritası görünüm genişliği (varsayılan: 1440).')
    args = parser.parse_args()
    if args.ui_map and args.out == ROOT / "docs/screenshots/guide-pp-v1":
        args.out = ROOT / "docs/screenshots/ui-map-v1"
    sys.path.insert(0, str(ROOT))
    os.chdir(ROOT)
    # Overwrite, never inspect or use inherited application connection settings.
    os.environ["DATABASE_URL"] = "sqlite://"
    os.environ.pop("MIGRATION_DATABASE_URL", None)
    os.environ.update(
        CIVENSA_ENV="development",
        DEMO_SEED="0",
        COOKIE_SECURE="0",
        SESSION_SECRET=secrets.token_urlsafe(48),
    )
    logging.disable(logging.CRITICAL)
    from scripts.first_load_dryrun import disposable_database, local_test_url

    def interrupted(signum, frame):
        # Bypass per-shot Exception handlers so cancellation stops the whole run.
        raise KeyboardInterrupt

    signal.signal(signal.SIGTERM, interrupted)
    signal.signal(signal.SIGINT, interrupted)
    try:
        url = local_test_url()  # Reject unsafe URLs before any DB connection/output.
    except ValueError as exc:
        print(f"Ekran görüntüleri üretilemedi: {exc}", file=sys.stderr)
        return 1

    if args.ui_map:
        return ui_main(args, url)

    data = None
    results = {
        shot.filename: "Yerel kurulum veya tarayıcı tamamlanamadı." for shot in SHOTS
    }
    with TemporaryDirectory(prefix="civensa_guide_") as temporary:
        folder = Path(temporary)
        try:
            with disposable_database(url) as engine:
                import app.db as app_db

                # local_test_url imports app.db under a harmless SQLite URL. Bind
                # both normal request Sessions and health() to this disposable DB.
                app_db.engine.dispose()
                app_db.engine = engine
                app_db.SessionLocal.configure(bind=engine)
                data = setup(engine)
                with local_http(engine, data) as base:
                    before = persisted_state(engine)
                    check(sum(row[1] == 'submitted' for row in before[0]) == 1,
                          'İnceleme kuyruğunda tam bir gönderilmiş paket bulunmalı.')
                    capture(base, folder, data, results)
                    check(persisted_state(engine) == before, 'Çekim sırasında kalıcı durum değişti; görseller geçersiz.')
        except (Exception, KeyboardInterrupt) as exc:
            error = ('İşlem sinyal ile durduruldu.' if isinstance(exc, KeyboardInterrupt) else
                     str(exc) if isinstance(exc, GuideError) else
                     f'Yerel kurulum/çekim hatası ({type(exc).__name__}).')
            # A global safety/cleanup failure invalidates even already taken shots.
            results = {shot.filename: error for shot in SHOTS}
            for shot in SHOTS:
                (folder / shot.filename).unlink(missing_ok=True)
        if data:
            data.pop('password', None)
        write_index(folder, data, results)
        args.out.mkdir(parents=True, exist_ok=True)
        for shot in SHOTS:
            destination = args.out / shot.filename
            if results[shot.filename]:
                destination.unlink(missing_ok=True)  # Never leave stale success images.
            else:
                shutil.copyfile(folder / shot.filename, destination)
        shutil.copyfile(folder / 'INDEX.md', args.out / 'INDEX.md')
    produced = sum(error is None for error in results.values())
    print(f'{produced}/18 PNG üretildi. Ayrıntılar: {args.out / "INDEX.md"}')
    return 0 if produced == len(SHOTS) else 1


if __name__ == '__main__':
    raise SystemExit(main())
