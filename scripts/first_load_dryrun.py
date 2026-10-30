#!/usr/bin/env python3
"""Inspect a PP print IDML or Word book in disposable, local PostgreSQL.

Usage (from the repository root):
    python scripts/first_load_dryrun.py BOOK.docx --out docs/first-load/NAME.md

Set CIVENSA_TEST_POSTGRES_URL to a local PostgreSQL 16+ database whose user can
CREATE DATABASE. DATABASE_URL and MIGRATION_DATABASE_URL are never inputs.
The database created by this script is dropped even if import/validation fails.
"""
import argparse
from collections import Counter, defaultdict
from contextlib import contextmanager
from copy import deepcopy
from decimal import Decimal
import hashlib
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import re
import shlex
import sys
import unicodedata
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
NAME_MAP = ROOT / 'tests/fixtures/first_load/name_map.json'
YEAR = 2027
FUNDS = {'inside': 'Bütçe içi', 'outside': 'Bütçe dışı', 'total': 'Toplam'}


def label_key(text):
    text = str(text or '').lower().replace('ı', 'i')
    text = ''.join(c for c in unicodedata.normalize('NFKD', text)
                   if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', text)


def scan_names(parsed):
    """Independent of parser header skipping; retain source locations for review."""
    labels = {label_key(s) for s in (
        'Performans Göstergesi', 'Performans Göstergeleri', 'Ölçü Birimi',
        'Gösterge Adı', 'Gösterge Kodu', 'Sıra No', 'Bütçe Yılı', 'Program Adı',
        'Alt Program Adı', 'İdare Birim Adı', 'Gerçekleşme', 'Hedef', 'Tahmin',
        'Göstergeye İlişkin Açıklama', 'Göstergelere İlişkin Açıklamalar',
        'Hesaplama Yöntemi', 'Verinin Kaynağı', 'Sorumlu İdare',
        'Planlanan', 'Bütçe', 'Değer', 'Faaliyetler', 'Ekonomik Ad',
    )}
    findings = []
    for department in parsed['departments']:
        name = department['name']
        for block in department['blocks']:
            for indicator in block['indicators']:
                text = indicator.get('name', '').strip()
                key = label_key(text)
                if not text:
                    cause = 'empty_indicator'
                elif key in labels or re.fullmatch(
                    r'20\d{2}(?:yili)?(?:gerceklesme|hedef|tahmin|butce|planlanan|deger|'
                    r'haziransonuharcama|yilsonu(?:harcama)?tahmin)(?:i|si)?', key
                ):
                    cause = 'header_label_indicator'
                elif re.fullmatch(r'20\d{2}(?:yili)?(?:20\d{2})?', key):
                    cause = 'year_indicator'
                else:
                    continue
                findings.append({'cause': cause, 'department': name, 'name': text,
                                 'table': block.get('table_number')})
        codes = defaultdict(list)
        for activity in department['activities']:
            text = activity.get('name', '').strip()
            location = [{'table': c.get('table_number'), 'row': c.get('row_number')}
                        for c in activity.get('costs', [])[:1]]
            if not text:
                findings.append({'cause': 'empty_activity', 'department': name,
                                 'name': text, 'locations': location})
            code = activity.get('code') or re.match(r'^\s*(\d+)[.)\s]', text)
            if code:
                code = code.group(1) if isinstance(code, re.Match) else str(code)
                codes[code].append({'name': text, 'locations': location})
        for code, activities in codes.items():
            if len(activities) > 1:
                findings.append({'cause': 'duplicate_activity_code', 'department': name,
                                 'code': code, 'activities': activities})
    return findings


@contextmanager
def environment(**values):
    previous = {key: os.environ.get(key) for key in values}
    try:
        for key, value in values.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def local_test_url():
    from sqlalchemy.engine import make_url

    raw = os.getenv('CIVENSA_TEST_POSTGRES_URL')
    if not raw:
        raise ValueError('CIVENSA_TEST_POSTGRES_URL gerekli; yerel PostgreSQL test veritabanını belirtin.')
    # Importing app.db creates an engine, so mask application URLs even here.
    with environment(DATABASE_URL='sqlite://', MIGRATION_DATABASE_URL=None):
        from app.db import normalize_database_url
        try:
            url = make_url(normalize_database_url(raw))
        except Exception:
            raise ValueError('Test veritabanı URL biçimi geçersiz.') from None
    if (url.drivername != 'postgresql+psycopg'
            or url.host not in {'localhost', '127.0.0.1', '::1'}
            or not url.database or url.query):
        raise ValueError('Yalnızca localhost/127.0.0.1/::1 üzerinde, sorgu parametresiz '
                         'PostgreSQL test URL adresi kullanılabilir.')
    return url


@contextmanager
def disposable_database(url):
    from sqlalchemy import create_engine
    from sqlalchemy.pool import NullPool

    # Pin localhost to loopback; hostaddr also defeats inherited PGHOSTADDR.
    connect_args = {'hostaddr': '::1' if url.host == '::1' else '127.0.0.1',
                    'connect_timeout': 5}
    # libpq service/options must not redirect connections or alter session setup.
    cleared = {key: None for key in os.environ if key.startswith('PG')}
    name = 'civensa_first_load_' + uuid4().hex
    with environment(**cleared):
        admin = create_engine(url, isolation_level='AUTOCOMMIT', poolclass=NullPool,
                              connect_args=connect_args)
        created = False
        runtime = None
        try:
            with admin.connect() as connection:
                if int(connection.exec_driver_sql('SHOW server_version_num').scalar()) < 160000:
                    raise ValueError('Yerel PostgreSQL 16 veya üzeri gerekli.')
                connection.exec_driver_sql(f'CREATE DATABASE "{name}"')
                created = True
            temporary = url.set(database=name).update_query_dict(
                {'hostaddr': connect_args['hostaddr'], 'connect_timeout': '5'})
            rendered = temporary.render_as_string(hide_password=False)
            with environment(DATABASE_URL=rendered, MIGRATION_DATABASE_URL=rendered, DEMO_SEED='0'):
                from alembic import command
                from alembic.config import Config

                command.upgrade(Config(str(ROOT / 'alembic.ini')), 'head')
                runtime = create_engine(temporary, poolclass=NullPool)
                yield runtime
        finally:
            if runtime is not None:
                runtime.dispose()
            try:
                if created:
                    with admin.connect() as connection:
                        connection.exec_driver_sql(f'DROP DATABASE "{name}" WITH (FORCE)')
            finally:
                admin.dispose()


def inspect_import(db, parsed, filename):
    from sqlalchemy import select
    from sqlalchemy.exc import SQLAlchemyError
    from app.import_service import (DISPLAY_OVERRIDES,
                                    build_department_plan, department_key,
                                    import_validation_errors)
    from app.models import Department, ImportBatch, Organization, ReportingPeriod, User
    from app.program_service import validate_department, department_snapshot
    from app.first_load_service import apply_first_load_batch

    names = json.loads(NAME_MAP.read_text(encoding='utf-8'))
    system_names = set(names['map'].values())
    for source in parsed['departments']:
        override = DISPLAY_OVERRIDES.get(department_key(source['name']))
        if override:
            system_names.add(override)
    organization = Organization(name='İlk yükleme denemesi Belediyesi')
    strategy = User(email='dryrun@example.invalid', full_name='Strateji deneme kullanıcısı',
                    role='strategy', password_hash='unused')
    period = ReportingPeriod(year=YEAR, status='draft', template_status='verified')
    db.add_all([organization, strategy, period])
    db.flush()
    db.add_all(Department(organization_id=organization.id, name=name)
               for name in sorted(system_names))
    db.commit()
    plan = build_department_plan(db, parsed)
    errors = import_validation_errors(parsed)
    unresolved = [m for m in plan['mappings'] if m['action'] != 'match']
    batch = ImportBatch(period_id=period.id, filename=filename, status='preview',
                        payload_json=json.dumps(parsed, ensure_ascii=False), created_by=strategy.id)
    db.add(batch)
    db.commit()
    applied = None
    catalogue = None
    apply_error = None
    # Exercise the real gate even if preview reports blocking issues. Never choose
    # source candidates or remove a bad directorate to make the book applicable.
    try:
        applied, catalogue, _ = apply_first_load_batch(db, batch, strategy)
        db.commit()
    except ValueError as exc:
        db.rollback()
        applied = None
        apply_error = str(exc)
    except SQLAlchemyError as exc:
        db.rollback()
        applied = None
        # SQL/driver messages can contain credentials or complete source payloads.
        apply_error = f'Veritabanı uygulama hatası ({type(exc).__name__}); paket uygulanmadı.'
    results = []
    for department in db.scalars(select(Department).order_by(Department.name)):
        validation_errors, warnings = validate_department(db, period.id, department.id)
        results.append({'name': department.name, 'ready': applied is not None and not validation_errors,
                        'errors': validation_errors, 'warnings': warnings,
                        'stored': department_snapshot(db, period.id, department.id)})
    return {'plan': plan, 'errors': errors, 'unresolved': unresolved, 'applied': applied,
            'apply_error': apply_error, 'results': results, 'pending': names.get('pending_ali', []),
            'catalogue': catalogue}


def budget_sum(rows, expected):
    """Unknown/absent cells never turn into zero; retain a known subtotal too."""
    rows = list(rows)
    known = [Decimal(str(row['amount'])) for row in rows if row.get('amount') is not None]
    subtotal = sum(known, Decimal(0))
    complete = len(rows) == expected and len(known) == expected and expected > 0
    return subtotal if complete else None, subtotal, expected - len(known)


def cross_tables(parsed):
    checks = []
    for department in parsed['departments']:
        activities = department['activities']
        tables = department.get('economic_tables', [])
        for fund in FUNDS:
            activity = budget_sum((v for a in activities for v in a['costs']
                                   if (v['year'], v['kind'], v['funding_type']) == (YEAR, 'budget', fund)),
                                  len(activities))
            economic = budget_sum((v for t in tables for v in t.get('totals', {}).get(fund, [])
                                   if (v['year'], v['kind']) == (YEAR, 'budget')), len(tables))
            checks.append({'department': department['name'], 'fund': fund,
                           'activity': activity, 'economic': economic,
                           'status': 'Eksik/belirsiz' if None in (activity[0], economic[0]) else
                           ('Eşit' if activity[0] == economic[0] else 'FARK')})
    return checks


def cell(value):
    return str(value).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('|', '&#124;').replace('\n', '<br>')


def money(value):
    if value is None:
        return 'Bilinmiyor'
    return f'{value:,.2f}'.replace(',', '_').replace('.', ',').replace('_', '.')


def amount_text(result):
    amount, subtotal, unknown = result
    return money(amount) if amount is not None else f'Bilinmiyor (bilinen: {money(subtotal)}; eksik: {unknown})'


def state_counts(values):
    counts = Counter(v.get('state', 'missing') for v in values)
    return '; '.join(f'{key}: {value}' for key, value in sorted(counts.items())) or '—'


def source_pages(parsed, department_name, error=None):
    from app.import_service import department_key
    sources = parsed.get('source_tables', {})
    departments = [d for d in parsed['departments']
                   if department_key(d['name']) == department_key(department_name)]
    tables = set()
    if error:
        for department in departments:
            for activity in department['activities']:
                if error.startswith(activity['name'] + ':'):
                    tables.update(v.get('table_number') for v in activity['costs'])
            for block in department['blocks']:
                if any(error.startswith(i['name'] + ':') for i in block['indicators']):
                    tables.add(block['table_number'])
            if not tables and 'ekonomik' in error:
                tables.update(t['table_number'] for t in department.get('economic_tables', []))
                if 'toplam' in error or 'ortak' in error:
                    tables.update(v.get('table_number') for a in department['activities'] for v in a['costs'])
    pages = {str(sources[str(t)]['page']) for t in tables if str(t) in sources}
    if not pages:
        pages = {str(t['page']) for t in sources.values()
                 if department_key(t.get('department', '')) == department_key(department_name)}
    return sorted(pages, key=lambda p: int(p) if p.isdigit() else 10**6)


def reference_text(cells):
    groups = defaultdict(set)
    for value in cells:
        groups[(value['page'], value['table'], value['column'])].add(value['row'])
    return '; '.join(f's.{page} / T{table} / sütun {col} / satır ' + ','.join(map(str, sorted(rows)))
                     for (page, table, col), rows in groups.items())


def proof_text(cells, stage):
    return '; '.join(f"s.{v['page']} T{v['table']} r{v['row']}c{v['column']} = «{v.get(stage) or 'BOŞ'}»"
                     for v in cells)


def institution_reference(value):
    return f"s.{value['page']} / {value['table']} / r{value['row']}c{value['column']}"


def institution_evidence(audit):
    if not audit.get('evidence'):
        return []
    lines = ['### Büyük farklar — katkı ve kapsam kanıtı', '',
             'Aşağıdaki hesaplar kaynak hücrelerine kadar izlenir. CHECK satırındaki hesaplanan tutar '
             'tanısal aday toplamdır; farklı faaliyet kapsamlarını eşit kabul etmez ve basılı tutarın '
             'düzeltmesi değildir. UNKNOWN, iki basılı kaynak arasındaki farkın hangi taraftan '
             'geldiğinin belirlenemediğini gösterir. Kart listesi kaynak bağlamıdır; aynı ekonomik '
             'hücre kart sayısı ile çarpılmaz. Tüm kaynak hücreleri fiziksel tablo / satır / sütun '
             'kimliğine göre bir kez sayılır.', '']
    for i, f in enumerate(audit.get('evidence', []), 1):
        lines += [f"#### Kanıt {i} — s.{f['page']} / {f['table']} / r{f['row']}c{f['column']}", '',
                  f"Sınıf: **{f['classification'] or '—'}**; basılı: **{money(f['printed'])}**; "
                  f"hesaplanan: **{money(f['computed'])}**; Fark: **{money(f['difference'])}**.", '',
                  cell(f['coverage_basis']), '',
                  'Basılı toplamın doğrudan kapsadığı satır / bileşenler:', '',
                  '| Kaynak hücresi | Basılı satır | Basılı değer |', '| --- | --- | ---: |']
        for ref in f['printed_components']:
            lines.append('| ' + ' | '.join(cell(v) for v in (
                institution_reference(ref), ref['label'], ref['raw'] if ref['raw'] is not None else 'BOŞ')) + ' |')
        if f.get('printed_leaves'):
            lines += ['', 'Bu alt programların kapsadığı faaliyet satırları (alt toplamlarla birlikte tekrar toplanmaz):', '',
                      '| Kaynak hücresi | Faaliyet | Basılı değer |', '| --- | --- | ---: |']
            for ref in f['printed_leaves']:
                lines.append('| ' + ' | '.join(cell(v) for v in (
                    institution_reference(ref), ref['label'], ref['raw'] if ref['raw'] is not None else 'BOŞ')) + ' |')
        if f.get('coverage'):
            c = f['coverage']
            lines += ['', f"Kapsam karşılaştırması: **{'Eşleşiyor' if c['comparable'] else 'Eşleşmiyor / tekil değil'}**. "
                      f"Yalnız kurumda: {len(c['summary_only'])}; yalnız müdürlüklerde: {len(c['directorate_only'])}. "
                      'Bu sayılar tutar dağılımı veya yeni program eşlemesi üretmez.', '',
                      'Kurum tablosundaki sonraki Kurumsal Genel Yönetim Giderleri / Diğer Destek, Denetim ve '
                      'Hizmet Birimleri satırları, program başlığının bütün müdürlük kartlarıyla aynı kapsam '
                      'olduğunu kanıtlamaz; bu satırlardan programa tutar taşınmadı.', '']
        if f['check'] == 'Sütun toplamı':
            lines += ['', 'Hesaplanan tutar yukarıda gösterilen doğrudan basılı satırların toplamıdır. '
                      'Alt programlar ve faaliyetler ayrıca eklenmedi; sonraki kurumsal giderler dışarıda bırakıldı.', '']
            continue
        lines += ['', 'Hesaplamaya giren kaynak hücreleri:', '',
                  '| Müdürlük | Bağlı kartlar (sayfa / tablo; program / alt program) | Kaynak hücresi | Kaynak satır | Ham değer | Katkı |',
                  '| --- | --- | --- | --- | --- | ---: |']
        for ref in f['contributors']:
            cards = '; '.join(f"s.{c['page']} / {c['table']} r{c['programme_row']}c2 / r{c['subprogramme_row']}c2; "
                              f"{c['programme']} / {c['subprogramme']}"
                              for c in ref.get('cards', [])) or '—'
            lines.append('| ' + ' | '.join(cell(v) for v in (
                ref.get('department', '—'), cards, institution_reference(ref), ref['label'] or '—',
                ref['raw'] if ref['raw'] is not None else 'BOŞ', money(ref['amount']))) + ' |')
        lines += ['', f"Tekil kaynak hücresi: {len(f['contributors'])}; "
                  f"katkılar toplamı: {money(sum((r['amount'] for r in f['contributors'] if r['amount'] is not None), Decimal(0)))}. "
                  'Bu bilinen ara toplamdır; eksik hücre varsa genel hesap Bilinmiyor kalır.', '']
    return lines


def institution_report(audit):
    from app.numbers import format_number

    findings = audit['findings']
    failures = sum(f['status'] == 'FAIL' for f in findings)
    classes = Counter(f['classification'] for f in findings if f['status'] == 'FAIL')
    if any(k not in {'BOOK', 'CHECK', 'UNKNOWN'} for k in classes):
        raise ValueError('Kurum bulgusunun BOOK / CHECK / UNKNOWN sınıfı eksik; denetim çıktısını kontrol edin.')
    lines = ['## Kurum özet tabloları', '',
             f"Kurum kontrolü: **{len(findings) - failures} PASS / {failures} FAIL**. "
             'Bu bulgular müdürlük SOURCE/SYSTEM mesajlarına ve kitap kontrol maddesi sayımına eklenmez.', '',
             f"FAIL sınıfları: **BOOK {classes['BOOK']} / CHECK {classes['CHECK']} / UNKNOWN {classes['UNKNOWN']}**.", '',
             'BOOK = aynı basılı tablonun kendi kapsamı içinde aritmetik yanlış; aşağıda toplanan hücreler '
             'gösterilir. CHECK = kapsam / program eşlemesi / tekil dağılım veya okuma sınırı; kitap hatası '
             'sayılmaz. UNKNOWN = farklı basılı kaynaklar uyuşmuyor veya hücre okunamıyor; kaynaktan hangi '
             'tarafın yanlış olduğu belirlenemedi. Sayısal fark tek başına BOOK kanıtı değildir.', '',
             'Tablolar başlık metniyle bulunur; sayfa numarası yalnız kaynak kanıtıdır. '
             'Fark = basılı − hesaplanan. Boş, çizgi veya okunamayan tutar sıfır sayılmaz; '
             'eksik toplam ve hesaplanamayan karşılaştırma FAIL olur. '
             'Yedek Ödenekler yalnız kurum düzeyinde ayrı gösterilir. '
             'Ekonomik çapraz toplamlar ayrıştırılmış kaynak tutarlarını kullanır; '
             'kaynak inceleme durumunu veya yükleme engelini değiştirmez.', '',
             'Okunan tablolar: ' + ('; '.join(f"s.{t['page']} / {t['id']} — {t['name']}" for t in audit['tables']) or 'Yok') + '.', '',
             'Sınırlar: IDML sayfası hikâyenin ilk basılı sayfasıdır; çok sayfalı tabloda hücrenin kesin '
             'devam sayfası çıkarılamaz. Program ayrıntısı, başlığında “Ekonomik Kod” yazmasına rağmen '
             'program / alt program / faaliyet hiyerarşisidir; ekonomik kategori dökümü içermez. '
             'Program çapraz toplamında kartlara bağlı üyelik tutarları kullanılır; ekonomik kategoriler '
             'program gruplarına tahminen dağıtılmaz. Yalnız genel “Program/Programı” soneki eşdeğer sayılır; '
             'farklı program adları eşlenmez ve ayrı FAIL olarak gösterilir. Döner Sermaye satırı satır toplamında denetlenir; '
             'kitapta bütçe içi + döner sermaye genel toplamı verilmediğinde böyle bir toplam üretilmez.', '',
             'Yedek Ödenekler — yalnız kurum düzeyi; müdürlük farkına katılmaz. Aşağıdaki tutarlar basılı '
             'hücrelerdir; doğrulukları satır / sütun kontrollerinde ayrıca gösterilir.', '',
             '| Kitap sayfası / tablo / satır | Yıl | Hizmet | Yönetim ve Destek | Program Dışı | Toplam |',
             '| --- | --- | ---: | ---: | ---: | ---: |']
    for reserve in audit.get('reserves', []):
        values = [f"s.{reserve['page']} / {reserve['table']} / {reserve['row']}", reserve['year'], *reserve['values']]
        lines.append('| ' + ' | '.join(cell(v) for v in values) + ' |')
    lines += ['', '| Durum | Sınıf | Kontrol | Kitap sayfası | Tablo | Satır | Sütun | Basılı değer | Hesaplanan değer | Fark | Açıklama |',
              '| --- | --- | --- | --- | --- | --- | --- | ---: | ---: | ---: | --- |']
    for f in findings:
        def display(value):
            if value is None:
                return 'Bilinmiyor'
            return format_number(value, money=not f['check'].startswith('Yüzde'))
        raw = display(f['printed']) if f['printed'] is not None else f"«{f['raw'] or 'BOŞ'}»"
        note = f['note']
        if f['status'] == 'FAIL' and f['classification'] == 'BOOK':
            note += ' Kanıt: ' + '; '.join(f"{institution_reference(r)} = «{r['raw']}»" for r in f['contributors'])
        values = [f['status'], f['classification'] or '—', f['check'], f['page'], f"{f['table']} — {f['name']}",
                  f"{f['row']} — {f['label']}", f['column'], raw,
                  display(f['computed']), display(f['difference']), note]
        lines.append('| ' + ' | '.join(cell(v) for v in values) + ' |')
    return lines + [''] + institution_evidence(audit) + ['']


def report(parsed, inspection, path, digest, output_path=None):
    from scripts.first_load_audit import audit_results, checklist, GROUPS, BookAudit, money
    audit = audit_results(parsed, inspection['results'])
    institution = audit.get('institution', dict(findings=[], tables=[]))
    institution_clean = all(f['status'] == 'PASS' for f in institution['findings'])
    presentation = checklist(parsed, inspection['results'], audit)
    findings = audit['findings']
    ready = sum(r['ready'] for r in inspection['results'])
    excluded = parsed.get('excluded_source_tables', [])
    flags = ''.join(f" --exclude-source-table {e['table_number']}" for e in excluded)
    lines = ['# 2027 PP ilk yükleme — kaynak kontrol listesi', '']
    lines.extend([f"Kitap temiz mi: **{'Evet' if presentation['clean'] and institution_clean else 'Hayır'}**.", '',
                  f"Kitap kontrol maddesi: **{len(presentation['items'])}**.", '',
                  f"Yalnız program grubu dağılımı bekleyen müdürlük: **{len(presentation['only_allocation'])}**.", '',
                  f'Gönderime hazır: **{ready}/{len(inspection["results"])}** (yalnızca bu prova).', ''])
    for item in excluded:
        lines.extend([f"**Provada çıkarıldı — nihai kitapta düzeltilmeli: T{item['table_number']}, "
                      f"kitap s.{item['page']}; başlık: {cell(' | '.join(item['header']))}.**", '',
                      f"Yapılacak: s.{item['page']}, tablo {item['table_number']}, satır 1, yıl sütunları: "
                      'dönemin Bütçe yılı/sütunlarını ve ilgili kart bağlantısını düzeltin. '
                      'Nihai kitabı tablo çıkarmadan yeniden deneyin.', ''])
    lines.extend([
                  f'Kaynak: `{path.name}`; SHA256: `{digest}`.', '',
                  f"Mesaj sayımı: **{presentation['book_messages']} kitap mesajı + "
                  f"{presentation['allocation_messages']} yalnız program grubu mesajı = "
                  f"{presentation['book_messages'] + presentation['allocation_messages']}**. "
                  f"Kitap mesajlarından {presentation['mixed_messages']} adedi ayrıca dağılım gerektirir; "
                  'kitap nedenleri listede, dağılım hücreleri ayrı bölümde gösterilir ve mesaj ikinci kez sayılmaz.', '',
                  f"Doğrulama: {sum(f['side'] == 'SOURCE' for f in findings)} SOURCE mesajı, "
                  f"{sum(f['side'] == 'SYSTEM' for f in findings)} SYSTEM mesajı; "
                  f"{len(audit['transfers'])} saklama aktarım sorunu; "
                  f"{sum(len(r['warnings']) for r in inspection['results'])} engelleyici olmayan uyarı.", '',
                  'Aynı kök nedenden üreyen mesajlar müdürlük listesinde tek kontrol maddesinde toplanır. '
                  'Hazır = gerçek paket uygulandı ve validate_department hata listesi boş. '
                  'Çıkarılan tablo ve ek üyelik sorunu doğrulama mesajı sayısına eklenmez. '
                  'Kitap maddeleri dağılım işlerini içermez. Fark = kitaptaki/kontrol edilen toplam − hesaplanan tutar. '
                  'Koşullu 0,00 önerisi veri düzeltmesi değildir; kaynak yokluğunu Strateji teyit etmelidir.', '',
                  '## Çalıştırma', '', '```sh',
                  '# CIVENSA_TEST_POSTGRES_URL: yerel, disposable PostgreSQL 16+; CREATE DATABASE yetkisi.',
                  f'python scripts/first_load_dryrun.py {shlex.quote(path.as_posix())}{flags} --out {output_path or "docs/first-load/dryrun-d3.md"}',
                  '```', '',
                  'Yalnızca localhost test URL adresi okunur; '
                  'her prova ayrı veritabanı oluşturur ve sonunda siler. Basılmayan IDML hikâyeleri ortak '
                  'geometri kontrolüyle hem Word yüklemesinden hem bütün denetimlerden çıkarılır ve aşağıda listelenir. '
                  'İsteğe bağlı numaralı prova dışlaması yalnız geçici Word üzerinde yapılır; '
                  'basılı kurum özetlerinin denetimini kaldırmaz.', '',
                  'Tek kurum, doğrulanmış dönem şablonu, ad eşlemesi ve denetimli/idempotent faaliyet kataloğu '
                  'gerçek servislerle hazırlanır. Boş, çizgi ve sıfır ayrıdır. Kaynak sayıları değiştirilmedi; '
                  'doğrulama kuralları gevşetilmedi. Sayfalar IDML hikâyesinin ilk basılı sayfasıdır. '
                  'T numarası dışlamadan önce üretilen Word tablosunu; r/satır ve c/sütun 1 tabanlı hücreyi gösterir.', ''])
    applied = inspection.get('applied')
    if applied:
        catalogue = inspection['catalogue']
        lines.extend([f"Uygulama: {applied['departments']} müdürlük, {applied['cards']} kart, "
                      f"{applied['indicators']} gösterge, {applied['activities']} faaliyet, "
                      f"{applied['economic_values']} ekonomik değer; alt programa bağlanamayan faaliyet: "
                      f"{applied['unlinked_activities']}. Katalog: {catalogue['entries']} faaliyet, "
                      f"{catalogue['responsibilities']} sorumluluk, {catalogue['links']} yerel bağlantı.", ''])
    lines.extend([f"İçe aktarma hatası: {len(inspection['errors'])}; kaynak çelişkisi: "
                  f"{len(inspection['plan']['conflicts'])}; çözümlenemeyen müdürlük: {len(inspection['unresolved'])}.", ''])
    if inspection.get('apply_error'):
        lines.extend(['Uygulama engeli: ' + cell(inspection['apply_error']), ''])
    for error in inspection['errors']:
        lines.append('- ' + cell(error))
    lines.extend(institution_report(institution))
    lines.extend(['## Yapılan SYSTEM düzeltmeleri', '',
                  '- Para ayrıştırıcısı baştaki ₺ işaretini tanır; ham metin korunur. Veteriner s.107 ekonomik '
                  'toplamı «₺10.424.938,95» artık 10424938.95 olarak okunur; rastgele ek metin kabul edilmez.',
                  '- Kodlu faaliyet açıklamaları sıra yerine faaliyet kodu ve tekil başlık ile eşlenir. '
                  'Halkla İlişkiler s.192’de 0006 başlığındaki “Faaliyet Adı:” eksikliği sonraki dört açıklamayı '
                  'kaydırmaz. Kültür s.167’deki yanlış 0008/0009 sağlık başlıkları 0015/0016’ya taşınmaz; '
                  'kaynakta uygun açıklama olmayan 0015–0017 aşağıda SOURCE olarak kalır.',
                  '- CIV-041 ile s.243 ekonomik tablosunun son müdürlüğe yanlış bağlanması önlendi; '
                  'ret listesi ile üyelik/sayfa kökeni birlikte korunur.', '',
                  '## Bütün hata gruplarının sınıflandırması', '',
                  '| Grup | Sınıf | Doğrulama mesajı |', '| --- | --- | ---: |'])
    for item in audit['counts']:
        lines.append(f"| {GROUPS[item['group']]} | {item['side']} | {item['count']} |")
    lines.extend(['', 'SYSTEM = üretilen Word, ayrıştırıcı, eşleme/ekim, saklama veya doğrulama kaynaklı hata. '
                  'SOURCE = kitap hücresi/dağılımının kendisi. Tanınmayan hata veya kitap–Word/saklama farkı '
                  'otomatik SOURCE sayılmaz; SYSTEM olarak raporlanır. Program Grubu Belirlenmeli kitap '
                  'düzeltme listesine girmez; kategori × program grubu dağılımı yükleme sonrası ayrı Strateji işidir. '
                  'Paylaştırma oranı veya varsayılan grup uydurulmadı.', '',
                  '## Her gruptan kitap → Word → saklanan değer kanıtı', ''])
    for group in GROUPS:
        example = next((f for f in findings if f['group'] == group), None)
        if not example:
            continue
        lines.extend([f"### {GROUPS[group]} — {example['side']}", '', cell(example['department']), '',
                      '- Kitap hücreleri: ' + cell(proof_text(example['cells'], 'book')),
                      '- Üretilen Word hücreleri: ' + cell(proof_text(example['cells'], 'word')),
                      '- Saklanan değerler (ondalık depolama gösterimi): ' + cell(json.dumps(example['stored'], ensure_ascii=False, default=str)),
                      '- Neden: ' + cell(example['details']), ''])
    # Show successful system repairs too, including the total preserved inside
    # EconomicValue.source_context_json rather than pretending it is a category.
    trace = BookAudit(parsed)
    lines.extend(['### SYSTEM düzeltmesinin saklama kanıtı', ''])
    for result in inspection['results']:
        for value in result['stored']['economics']:
            context = json.loads(value.get('source_context_json') or '{}')
            matched = next((v for v in context.get('totals', {}).get('inside', [])
                            if (v.get('raw_text') or '').startswith('₺')), None)
            if matched:
                table = value['source_table_number']
                cells = trace.economic_cells({'table_number': table}, (matched['year'], matched['kind']), ['BÜTÇE İÇİ TOPLAM KAYNAK'])
                lines.extend(['- Kitap: ' + cell(proof_text(cells, 'book')),
                              '- Word: ' + cell(proof_text(cells, 'word')),
                              '- Saklanan EconomicValue.source_context_json.totals.inside: ' + cell(json.dumps(matched, ensure_ascii=False)), ''])
                break
        if result['name'] == 'Halkla İlişkiler Müdürlüğü':
            source = next(d for d in parsed['departments'] if d['name'] == result['name'])
            activity = next(a for a in source['activities'] if a['name'].startswith('0006.'))
            loc = activity.get('explanation_source')
            if loc:
                cells = [trace.cell(loc['table_number'], loc['title_row']-1, 0), trace.cell(loc['table_number'], loc['row_number']-1, 0)]
                lines.extend(['- Kitap: ' + cell(proof_text(cells, 'book')),
                              '- Word: ' + cell(proof_text(cells, 'word')),
                              '- Saklanan Activity.explanation: ' + cell(trace.activity_stored(result['stored'], activity)[0]['explanation']), ''])
    lines.extend(['## Kitapta düzeltilmez — yükleme sonrası Strateji girer (program grubu dağılımı)', '',
                  'Bu bölüm kitap kontrol maddesi sayısına dahil değildir. Her kaynak kategori hücresi için '
                  'Hizmet + Yönetim ve Destek tutarları kaynak tutarını korumalıdır. Bütçe dışı 0,00 değerleri '
                  'kitabın Bütçe Dışı Toplam satırından gelir; kategori hücresindeki bütçe içi tutarla karıştırılmamalı. '
                  'Mevcut uygulamada Strateji için iki gruba atomik bölme akışı bulunmuyor; bu bölüm yapılacak '
                  'yükleme sonrası işi tanımlar (PR açıklamasında route/validator incelemesi vardır).', ''])
    if not presentation['allocations']:
        lines.extend(['Dağılım bekleyen hücre yok.', ''])
    for name, allocation in presentation['allocations'].items():
        lines.extend([f'### {name}', '',
                      f"{allocation['messages']} doğrulama mesajında dağılım gereksinimi var; "
                      'başka kitap nedeni de taşıyan mesajlar kitap sayımında bir kez sayılmıştır.', '',
                      '| Kart | Program / alt program | Grup |', '| --- | --- | --- |'])
        for card in allocation['cards']:
            lines.append(f"| {cell(card['title'])} | {cell(card['program'])} / {cell(card['subprogram'])} | {card['group']} |")
        lines.extend(['', '| Yıl / sütun | Kategori | Finansman | Dağıtılacak tutar | Kaynak hücresi |',
                      '| --- | --- | --- | ---: | --- |'])
        from app.program_service import ECONOMIC_CATEGORIES
        from app.value_labels import value_kind_label, funding_label
        for value in sorted(allocation['cells'], key=lambda v: (v['year'], v['kind'], v['category'], v['funding_type'])):
            lines.append(f"| {value['year']} {value_kind_label(value['kind'])} | {ECONOMIC_CATEGORIES[value['category']]} | "
                         f"{funding_label(value['funding_type'])} | {money(value['amount'])} | {cell(reference_text(value['cells']))} |")
        lines.append('')
    lines.extend(['## Pazartesi — müdürlük bazında yalnız kitap düzeltme listesi', ''])
    for result in inspection['results']:
        selected = [i for i in presentation['items'] if i['department'] == result['name']]
        lines.extend([f"### {result['name']}", '',
                      f"Gönderime hazır: **{'Evet' if result['ready'] else 'Hayır'}**; "
                      f"{sum(len(i['messages']) for i in selected)} kitap mesajı / {len(selected)} kitap kontrol maddesi.", ''])
        if result['name'] in presentation['only_allocation']:
            lines.extend(['Kitap açısından temiz; program grubu dağılımı bekliyor.', ''])
        for item in selected:
            lines.extend(['- [ ] ' + cell(item['title']), '',
                          '  **Hücreler:** ' + cell(reference_text(item['cells'])), '',
                          '  Yapılacak: ' + cell(item['action']), ''])
            if item['note']:
                lines.extend(['  ' + cell(item['note']), ''])
        if not selected and result['name'] not in presentation['only_allocation']:
            lines.append('Kitap düzeltmesi yok.')
        lines.append('')
    systems = [f for f in findings if f['side'] == 'SYSTEM']
    if systems or audit['transfers']:
        lines.extend(['## Çözülmemiş SYSTEM bulguları — kaynak listesine taşınmadı', ''])
        lines.extend('- ' + cell(f['department'] + ': ' + f['error'] + ' ' + f['details']) for f in systems)
        lines.extend('- ' + cell(json.dumps(t, ensure_ascii=False)) for t in audit['transfers'])
    conversion = inspection.get('conversion')
    if conversion:
        lines.extend(['## Basılmayan (pasteboard) tablolar — yüklenmedi', '',
                      'Sayfayla kısmen bile kesişen frame korunur. Threaded frame ve basılı hikâyeye '
                      'anchored frame aynı basılı hikâye kapsamındadır. Bu tablolar bütün denetimlerin dışındadır.', '',
                      f"{sum(t['reason'] == 'pasteboard' for t in conversion.get('non_printed', []))} pasteboard tablo; "
                      f"ayrıca {sum(t['reason'] == 'unplaced' for t in conversion.get('non_printed', []))} yerleşimsiz tablo (önceki okuyucu da yüklemiyordu).", '',
                      '| Story | Spread sayfaları | Tablo | İlk başlık hücreleri |',
                      '| --- | --- | --- | --- |'])
        for table in conversion.get('non_printed', []):
            lines.append('| ' + ' | '.join(cell(v) for v in (
                table['story'], '–'.join(table['spread_pages']) or 'Yerleşimsiz',
                table['id'], ' / '.join(v or 'BOŞ' for v in table['header']))) + ' |')
        if not conversion.get('non_printed'):
            lines.append('Yok.')
        lines.append('')
        lines.extend(['## Dönüşüm ve varsayımlar', '',
                      'Eşlenmeyen ad (ham metin korundu, yerleşik müdürlük eşlemesi uygulandı): '
                      + cell('; '.join(conversion['unmapped']) or 'Yok') + '.', '',
                      'Basılı sayfaya bağlı olmayan tablolar (dışlama seçeneğinden bağımsız; yüklemeye alınmaz): '
                      + cell(json.dumps(conversion['unplaced'], ensure_ascii=False)) + '.', ''])
    lines.extend(['Teyit bekleyen ad eşlemeleri: ' + ('; '.join(inspection['pending']) or 'Yok (D-94, 19 Ekim 2026 proje sahibi teyidi)') + '. '
                  'Program/alt program kodları kitapta yoktur ve boş bırakılır. Tekil müdürlük içi faaliyet '
                  'kodu üyeliği belirler. Resmî katalog 0006 Özel Kalem gibi maliyet karşılığı bulunmayan '
                  'üyelikleri korur; yeni maliyet üretilmez. Nihai kitap dışlamasız yeniden denenmelidir.', ''])
    return '\n'.join(lines)


def parse_book(path, name_map=NAME_MAP, exclude_source_tables=()):
    """Keep generated Word and mapping report temporary during a rehearsal."""
    from app.data_importer import parse_performance_docx
    with TemporaryDirectory(prefix='civensa_first_load_word_') as directory:
        word = Path(directory) / 'book.docx'
        if path.suffix.lower() == '.idml':
            from scripts.book_to_load_docx import generate_docx
            conversion = generate_docx(path, name_map, word)
        else:
            word.write_bytes(path.read_bytes())
            conversion = None
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(word)
        numbers = set(exclude_source_tables)
        if any(n < 1 or n > len(doc.tables) for n in numbers):
            raise ValueError('Provada çıkarılacak tablo numarası bulunamadı; oluşturulan Word tablo numaralarını kontrol edin.')
        excluded = []
        for number, table in enumerate(list(doc.tables), 1):
            if number not in numbers:
                continue
            caption = table._tbl.tblPr.find(qn('w:tblCaption'))
            source = json.loads(caption.get(qn('w:val'))) if caption is not None else {}
            excluded.append(dict(table_number=number, page=source.get('page', '?'),
                                 header=[c.text for c in table.rows[0].cells]))
            table._tbl.getparent().remove(table._tbl)
        doc.save(word)
        parsed = parse_performance_docx(word.read_bytes())
        parsed['excluded_source_tables'] = excluded
        book_tables = {}
        if path.suffix.lower() == '.idml':
            from scripts.book_to_load_docx import extract_tables
            from scripts.institution_summary import table_kind
            extracted = extract_tables(path)
            book_tables = {t['id']: t['rows'] for t in extracted}
            parsed['institution_tables'] = [t for t in extracted if table_kind(t)]
        else:
            from scripts.institution_summary import table_kind
            candidates = [dict(id=f'Word T{i}', page='?', rows=[[c.text for c in r.cells] for r in t.rows])
                          for i, t in enumerate(doc.tables, 1)]
            parsed['institution_tables'] = [t for t in candidates if table_kind(t)]
        parsed['rehearsal_tables'] = {}
        for index, table in enumerate(doc.tables, 1):
            source = parsed.get('source_tables', {}).get(str(index), {})
            parsed['rehearsal_tables'][str(index)] = dict(
                page=source.get('page', '?'), original_number=source.get('table_number', index),
                book=book_tables.get(source.get('id')), word=[[c.text for c in row.cells] for row in table.rows])
        return parsed, conversion


def comparison_tables(path, conversion):
    """Match printed tables by department/type/occurrence, never export-local IDs."""
    from scripts.book_to_load_docx import extract_tables
    from scripts.institution_summary import table_kind
    provenance = {t['id']: t for t in conversion['tables']}
    counts = Counter()
    result = {}
    for table in sorted(extract_tables(path), key=lambda t: (int(t['page']), t['id'])):
        source = provenance.get(table['id'])
        if source:
            base = ('department', source['department'], source['type'])
            title = {'card': 'Kart', 'membership': 'Kart faaliyet üyeliği',
                     'cost': 'Faaliyet maliyetleri', 'economic': 'Ekonomik tablo'}[source['type']]
            label = f"T{source['table_number']} / {title}"
        else:
            kind = table_kind(table)
            base = ('institution', kind) if kind else ('other', table['page'], label_key(table['rows'][0][0]))
            label = table['rows'][0][0] or 'Başlıksız tablo'
        counts[base] += 1
        result[(*base, counts[base])] = dict(table, department=source['department'] if source else 'Kurum / diğer',
                                          label=label)
    return result


def paired_rows(old, new):
    """Align insertions by row labels; keep positional label corrections visible."""
    from difflib import SequenceMatcher
    if len(old) == len(new):
        return [(i, i) for i in range(len(old))]
    pairs = []
    labels = lambda rows: [label_key(row[0]) for row in rows]
    for tag, a, b, c, d in SequenceMatcher(None, labels(old), labels(new), autojunk=False).get_opcodes():
        length = max(b - a, d - c)
        pairs.extend((a + i if a + i < b else None, c + i if c + i < d else None) for i in range(length))
    return pairs


def amount_changes(old_tables, new_tables):
    """Compare every printed numeric/sentinel cell, including institution totals.

    Missing rows/cells use YOK, a different state from a printed blank. Raw text
    is preserved beside formatted numeric values, including ambiguous numbers.
    """
    from app.numbers import source_cell, format_number
    def numeric(raw):
        return raw is not None and bool(re.fullmatch(r'[+\-−\d.,%₺()\s]+(?:TL|Mn|İçerik)?', str(raw).strip()))
    def display(raw, exists):
        if not exists:
            return 'YOK (satır/hücre yok)'
        if raw is None or raw == '':
            return 'BOŞ'
        return str(raw)
    changes = []
    for key in sorted(old_tables.keys() | new_tables.keys()):
        old, new = old_tables.get(key), new_tables.get(key)
        before, after = old['rows'] if old else [], new['rows'] if new else []
        for ri, rj in paired_rows(before, after):
            a, b = before[ri] if ri is not None else [], after[rj] if rj is not None else []
            for column in range(1, max(len(a), len(b))):
                ae, be = column < len(a), column < len(b)
                av, bv = a[column] if ae else None, b[column] if be else None
                if ae == be and av == bv:
                    continue
                # Merged-cell markers carry structure, never a value.
                if av in ('→', '↓') or bv in ('→', '↓'):
                    continue
                empty = lambda v: v is None or v in ('', '-', '–', '—')
                if not (numeric(av) or numeric(bv) or (empty(av) and empty(bv))):
                    continue
                an = source_cell(av)['number'] if ae else None
                bn = source_cell(bv)['number'] if be else None
                delta = (format_number(Decimal(bn) - Decimal(an),
                                       money=key[0] == 'institution' or (key[0] == 'department' and key[2] != 'card'))
                         if an is not None and bn is not None else 'Hesaplanamaz')
                label = (b or a)[0] or 'BOŞ satır etiketi'
                if a and b and a[0] != b[0]:
                    label = f'{a[0] or "BOŞ"} → {b[0] or "BOŞ"}'
                table = new or old
                header = next((row[column] for row in table['rows'] if len(row) > column
                               and ('Performans Göstergesi' in (row[0] or '') or row is table['rows'][0])), None)
                # Card metadata uses merged columns; indicator columns are named
                # by the indicator header rather than the budget-year row.
                if key[0] == 'department' and key[2] == 'card':
                    header = next((row[column] for row in table['rows'] if len(row) > column
                                   and (row[0] or '').startswith('Performans Göstergesi')), header)
                changes.append(dict(department=table['department'],
                    page=f"{old['page'] if old else 'YOK'} → {new['page'] if new else 'YOK'}",
                    table=f"{old['label'] if old else 'YOK'} → {new['label'] if new else 'YOK'}",
                    row=f"r{ri + 1 if ri is not None else 'YOK'} → r{rj + 1 if rj is not None else 'YOK'} / {label}",
                    column=f"c{column + 1} / {header or 'BOŞ başlık'}", old=display(av, ae), new=display(bv, be), difference=delta))
    return changes


def difference_report(old_path, new_path, old_parsed, new_parsed, old_inspection, new_inspection,
                      old_conversion, new_conversion, output_path):
    """Owner-facing evidence, reproduced together with the approved dry-run."""
    from scripts.first_load_audit import audit_results, checklist
    import shlex
    audits = [audit_results(p, i['results']) for p, i in (
        (old_parsed, old_inspection), (new_parsed, new_inspection))]
    checks = [checklist(p, i['results'], a) for p, i, a in zip(
        (old_parsed, new_parsed), (old_inspection, new_inspection), audits)]
    old_items, new_items = checks[0]['items'], checks[1]['items']
    changes = amount_changes(comparison_tables(old_path, old_conversion), comparison_tables(new_path, new_conversion))
    lines = ['# Onaylı 2027 PP — 16 Ekim kaynağına göre fark', '',
             'Muhatap: proje sahibi ve Strateji. 2027 ilk yüklemesinin tek kaynağı onaylı IDML’dir (D-316). '
             'Bu rapor kaynak değerlerini ve yükleyici/doğrulama kurallarını değiştirmez.', '',
             f"Eski: `{old_path.name}`; SHA256 `{hashlib.sha256(old_path.read_bytes()).hexdigest()}`.", '',
             f"Yeni: `{new_path.name}`; SHA256 `{hashlib.sha256(new_path.read_bytes()).hexdigest()}`.", '',
             'Karşılaştırma, iki kaynağı da aynı basılı sayfa geometrisi filtresiyle ve ayrı geçici yerel '
             'PostgreSQL veritabanlarında gerçek servislerle çalıştırır. '
             'Eski rapor `docs/first-load/dryrun-printed.md` ve `dryrun-d3.md` değiştirilmez. '
             'Onaylı arşivdeki 240 tablonun 229’u basılı, 7’si pasteboard ve 4’ü yerleşimsizdir; '
             'yerleşimsiz dört tablo önceki okuyucu tarafından da yüklenmiyordu.', '',
             '## Yeniden üretme', '', '```sh',
             '# CIVENSA_TEST_POSTGRES_URL: yerel, disposable PostgreSQL 16+; CREATE DATABASE yetkisi.',
             f'python scripts/first_load_dryrun.py {shlex.quote(new_path.as_posix())} '
             f'--out docs/first-load/dryrun-onayli.md --compare-source {shlex.quote(old_path.as_posix())} '
             f'--diff-out {output_path.as_posix()}', '```', '',
             '## Müdürlük sayımları', '',
             'Her hücre eski → yeni. Ekonomik değer sayısı kategori kayıtlarını sayar; toplam satırları ayrıca denetlenir. '
             'YOK müdürlüğün kaynakta bulunmadığı anlamına gelir.', '',
             '| Müdürlük | Kart | Gösterge | Faaliyet | Ekonomik değer | Kitap maddesi | Hazır |',
             '| --- | ---: | ---: | ---: | ---: | ---: | --- |']
    departments = [{d['name']: d for d in p['departments']} for p in (old_parsed, new_parsed)]
    def counts(d):
        return (len(d['blocks']), sum(len(b['indicators']) for b in d['blocks']), len(d['activities']),
                sum(len(t['values']) for t in d['economic_tables'])) if d else ('YOK',) * 4
    totals = [[0] * 4, [0] * 4]
    for name in sorted(departments[0].keys() | departments[1].keys()):
        values = [counts(d.get(name)) for d in departments]
        for index in range(2):
            totals[index] = [x + (y if isinstance(y, int) else 0) for x, y in zip(totals[index], values[index])]
        item_counts = [sum(i['department'] == name for i in items) for items in (old_items, new_items)]
        ready = [next(('Evet' if r['ready'] else 'Hayır' for r in inspection['results']
                       if r['name'] == name), 'YOK') for inspection in (old_inspection, new_inspection)]
        lines.append('| ' + ' | '.join([cell(name)] + [f'{a} → {b}' for a, b in zip(*values)]
                                      + [f'{item_counts[0]} → {item_counts[1]}', ' → '.join(ready)]) + ' |')
    lines.extend(['| **Toplam** | ' + ' | '.join(f'**{a} → {b}**' for a, b in zip(*totals))
                  + f" | **{len(old_items)} → {len(new_items)}** | "
                  f"**{sum(r['ready'] for r in old_inspection['results'])} → {sum(r['ready'] for r in new_inspection['results'])}** |", '',
                  f"Müdürlük: {len(departments[0])} → {len(departments[1])}; "
                  f"yeni: {', '.join(sorted(departments[1].keys() - departments[0].keys())) or 'Yok'}; "
                  f"eksilen: {', '.join(sorted(departments[0].keys() - departments[1].keys())) or 'Yok'}.", '',
                  '## Değişen bütün tutar / gösterge hücreleri', '',
                  f'**{len(changes)} değişen hücre**. Kapsam: bütün basılı IDML tabloları; kurum özetleri ve diğer '
                  'sayısal tablolar dahildir. T numarası üretilen Word tablosudur. Basılmayan tablolar kapsam dışıdır. '
                  'Fark = yeni − eski; BOŞ, çizgi, sıfır ve YOK ayrı tutulur. Ham değerler aynen korunur; '
                  'yalnız fark Türkçe sayı biçimiyle yazılır. Ham biçim değişip sayısal değer aynı kalırsa fark 0’dır. '
                  'Belirsiz/okunamayan veya eksik bir taraf varsa fark hesaplanmaz.', '',
                  'Tablolar müdürlük + tür + basılı sıra ile; kurum tabloları başlık türüyle; diğer tablolar '
                  'sayfa + ilk başlık + sıra ile eşlenir. IDML kimlikleri dışa aktarmada yenilendiğinden eşleme anahtarı değildir. '
                  'Aynı boyutlu tablolarda satır konumu korunur; eklenen satırlar etiket dizisiyle hizalanır. '
                  'Sütunlar fiziksel sıradadır; satır/sütun konumu ve başlık birlikte kanıttır.', '',
                  '| Müdürlük / kapsam | Sayfa eski → yeni | Tablo eski → yeni | Satır / etiket | Sütun / etiket | Eski ham değer | Yeni ham değer | Fark |',
                  '| --- | --- | --- | --- | --- | ---: | ---: | ---: |'])
    for change in changes:
        lines.append('| ' + ' | '.join(cell(change[k]) for k in (
            'department', 'page', 'table', 'row', 'column', 'old', 'new', 'difference')) + ' |')
    lines.extend(['', '## Kitap kontrolü ve gönderime hazırlık', '',
                  f"Kitap kontrol maddesi: **{len(old_items)} → {len(new_items)}**; "
                  f"gönderime hazır müdürlük: **{sum(r['ready'] for r in old_inspection['results'])}/{len(old_inspection['results'])} → "
                  f"{sum(r['ready'] for r in new_inspection['results'])}/{len(new_inspection['results'])}**.", '',
                  f"Mesajlar: kitap {checks[0]['book_messages']} → {checks[1]['book_messages']}; "
                  f"yalnız program grubu {checks[0]['allocation_messages']} → {checks[1]['allocation_messages']}; "
                  f"SYSTEM {sum(f['side'] == 'SYSTEM' for f in audits[0]['findings'])} → "
                  f"{sum(f['side'] == 'SYSTEM' for f in audits[1]['findings'])}.", ''])
    def item_map(items):
        result = {}
        for item in items:
            key = item['department'], item['title']
            if key in result:
                raise ValueError('Fark raporunda kontrol maddesi kimliği belirsiz: ' + str(key))
            result[key] = item
        return result
    before, after = item_map(old_items), item_map(new_items)
    resolved, remaining, added = before.keys() - after.keys(), before.keys() & after.keys(), after.keys() - before.keys()
    lines.extend([f'Eski {len(old_items)} madde: **{len(resolved)} çözüldü, {len(remaining)} açık**. '
                  f'**{len(added)} yeni madde** var. Kimlik müdürlük + kontrol başlığı/yılıdır; '
                  'durum aynı kontrol alanının geçip geçmediğini gösterir, gerekçe ve tutar değişmiş olabilir. '
                  'Örneğin Sağlık toplam etiketi düzelmiştir; 2026 Haziran ve 2029 ekonomik kontrolleri '
                  'yeni aritmetik farklar nedeniyle açık kalır. Kurum özet bulguları bu sayıya dahil değildir.', '',
                  '| Durum | Müdürlük | Kontrol maddesi | Eski sayfa / tablo / hücre | Yeni sayfa / tablo / hücre |',
                  '| --- | --- | --- | --- | --- |'])
    def refs(item):
        return reference_text(item['cells']) if item else 'YOK'
    for status, keys in (('Çözüldü', resolved), ('Açık', remaining), ('Yeni', added)):
        for key in sorted(keys):
            lines.append('| ' + ' | '.join(cell(v) for v in (
                status, *key, refs(before.get(key)), refs(after.get(key)))) + ' |')
    lines.extend(['', '## Okuma sınırları, kaynak sorunları ve teyit', '',
                  'Onaylı PDF bu çalışmada ayrıca açılıp doğrulanmadı; PDF ile eşleşme bilgisi görevde verilen '
                  'D-320 teyididir. IDML SHA256 değeri doğrulandı. Çok sayfalı story için ilk basılı sayfa gösterilir; '
                  'InDesign dizgi motoru kullanılmadığından hücrenin kesin devam sayfası doğrulanamaz.', '',
                  'Özel Kalem’de 0006 emekli ikramiyesi maliyet satırı onaylı kaynakta vardır: faaliyet sayısının '
                  '290’dan 291’e çıkması bu satırdır; maliyet uydurulmadı. Sağlık ekonomik tablosunun eski BOŞ '
                  'toplam etiketi onaylı kaynakta “BÜTÇE İÇİ TOPLAM KAYNAK” olmuştur; aynı tutarın artık tanınması '
                  'sayısal düzeltmeyle karıştırılmamalıdır.', '',
                  'İşletme ve İştirakler’de basılı kart üyelik kodları 0011/0012/0013/0014/0015, maliyet kodları '
                  '0001/0002/0033/0004/0005’tir. Tekil kod eşlemesi beş faaliyeti bağlayamaz: 5 üyelik kontrol '
                  'maddesi ve 10 SYSTEM bağlantı mesajı vardır. Bu, kaynak kod uyuşmazlığıdır; parser kusuru '
                  'olduğu kanıtlanmadı. Kodlar düzeltilmedi, ad benzerliğiyle bağlanmadı, tablo dışlanmadı.', '',
                  'Diğer çözülmemiş SYSTEM bulguları onaylı dry-run’da ayrı listelenir. Bu görevde pasteboard '
                  'dışında kanıtlanmış parser kusuru düzeltilmedi. Belirsiz sayı biçimi, eksik kaynak ve '
                  'doğrulama sorunları raporda korunur. Canlı veritabanına bağlanılmadı.', ''])
    for p, a, c in zip((old_path, new_path), audits, (old_conversion, new_conversion)):
        findings = a.get('institution', {}).get('findings', [])
        states = Counter(f['status'] for f in findings)
        classes = Counter(f.get('classification') for f in findings if f['status'] == 'FAIL')
        lines.extend([f"`{p.name}`: kurum {states['PASS']} PASS / {states['FAIL']} FAIL; "
                      f"sınıflar {dict(classes)}. Geometriyle dışlanan: "
                      f"{sum(t['reason'] == 'pasteboard' for t in c['non_printed'])} pasteboard + "
                      f"{sum(t['reason'] == 'unplaced' for t in c['non_printed'])} yerleşimsiz tablo.", ''])
    return '\n'.join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('book', type=Path)
    parser.add_argument('--exclude-source-table', type=int, action='append', default=[],
                        help='Yalnızca provada çıkarılacak Word tablo numarası; tekrarlanabilir')
    parser.add_argument('--out', type=Path, help='Markdown çıktı dosyası; verilmezse stdout')
    parser.add_argument('--compare-source', type=Path, help='Eski IDML kaynağını aynı kurallarla karşılaştır')
    parser.add_argument('--diff-out', type=Path, help='Türkçe kaynak fark raporu; --compare-source ile birlikte')
    args = parser.parse_args()
    if bool(args.compare_source) != bool(args.diff_out):
        parser.error('--compare-source ve --diff-out birlikte gerekli')
    if args.compare_source and (args.exclude_source_table or args.book.suffix.lower() != '.idml'
                                or args.compare_source.suffix.lower() != '.idml'):
        parser.error('Kaynak karşılaştırması dışlamasız iki IDML gerektirir')
    sys.path.insert(0, str(ROOT))
    try:
        url = local_test_url()  # Refuse unsafe targets before any connection.
        data = args.book.read_bytes()
        from sqlalchemy.orm import Session

        parsed, conversion = parse_book(args.book, exclude_source_tables=args.exclude_source_table)
        with disposable_database(url) as engine:
            with Session(engine, expire_on_commit=False) as db:
                inspection = inspect_import(db, deepcopy(parsed), args.book.name)
        inspection['conversion'] = conversion
        content = report(parsed, inspection, args.book, hashlib.sha256(data).hexdigest(), args.out)
        if args.out:
            args.out.parent.mkdir(parents=True, exist_ok=True)
            args.out.write_text(content, encoding='utf-8')
            print(f'Rapor yazıldı: {args.out}')
        else:
            print(content)
        if args.compare_source:
            old_parsed, old_conversion = parse_book(args.compare_source)
            with disposable_database(url) as engine:
                with Session(engine, expire_on_commit=False) as db:
                    old_inspection = inspect_import(db, deepcopy(old_parsed), args.compare_source.name)
            comparison = difference_report(args.compare_source, args.book, old_parsed, parsed,
                                           old_inspection, inspection, old_conversion, conversion, args.diff_out)
            args.diff_out.parent.mkdir(parents=True, exist_ok=True)
            args.diff_out.write_text(comparison, encoding='utf-8')
            print(f'Fark raporu yazıldı: {args.diff_out}')
    except ValueError as exc:
        print(f'Kuru deneme çalıştırılamadı: {exc}', file=sys.stderr)
        return 1
    except Exception as exc:
        # DB exceptions can include connection details; never print them/tracebacks.
        print(f'Kuru deneme çalıştırılamadı ({type(exc).__name__}); yerel PostgreSQL, '
              'dosya ve CREATE DATABASE yetkisini kontrol edin.', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
