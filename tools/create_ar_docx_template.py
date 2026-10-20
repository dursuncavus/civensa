"""Create the initial CIV-046 template. Subsequent design edits belong in Word."""
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


def create(path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(10)
    section.different_first_page_header_footer = True
    compat = doc.settings.element.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat')
        doc.settings.element.append(compat)
    for node in list(compat):
        if node.tag == qn('w:useFELayout') or node.get(qn('w:name')) == 'compatibilityMode':
            compat.remove(node)
    mode = OxmlElement('w:compatSetting')
    mode.set(qn('w:name'), 'compatibilityMode')
    mode.set(qn('w:uri'), 'http://schemas.microsoft.com/office/word')
    mode.set(qn('w:val'), '15')
    compat.append(mode)
    styles = doc.styles
    defaults = styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    defaults.find(qn('w:lang')).set(qn('w:val'), 'tr-TR')
    fonts = defaults.find(qn('w:rFonts'))
    for key in ('asciiTheme', 'hAnsiTheme'):
        fonts.attrib.pop(qn('w:' + key), None)
    for key in ('ascii', 'hAnsi'):
        fonts.set(qn('w:' + key), 'Arial')
    defaults.find(qn('w:sz')).set(qn('w:val'), '20')
    styles['Normal'].font.name = 'Arial'
    styles['Normal'].font.size = Pt(10)
    body = styles.add_style('Civensa Gövde', WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = styles['Normal']
    body.font.name = 'Arial'
    body.font.size = Pt(10)
    body.font.color.rgb = RGBColor.from_string('243746')
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.widow_control = True
    for depth, size in enumerate([18, 14, 12, 11], 1):
        style = styles.add_style(f'Civensa Başlık {depth}', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = body
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('164E63')
        style.paragraph_format.space_before = Pt(15 if depth == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), str(depth - 1))
        style.element.get_or_add_pPr().append(outline)
        style.quick_style = True
    for level in range(1, 4):
        toc = styles.add_style(f'TOC {level}', WD_STYLE_TYPE.PARAGRAPH)
        toc.base_style = body
        toc.paragraph_format.left_indent = Mm((level - 1) * 5)
        toc.paragraph_format.space_after = Pt(4)
        toc.paragraph_format.line_spacing = 1
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), '9')
        toc.element.get_or_add_pPr().append(outline)
    bullet = styles.add_style('Civensa Madde', WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = body
    # Reuse the template's real Word bullet definition, without literal bullet text.
    numpr = OxmlElement('w:numPr')
    numid = OxmlElement('w:numId')
    numid.set(qn('w:val'), '1')
    numpr.append(numid)
    bullet.element.get_or_add_pPr().append(numpr)
    bullet.paragraph_format.left_indent = Mm(5)
    bullet.paragraph_format.first_line_indent = Mm(-3)
    table = styles.add_style('Civensa Tablo', WD_STYLE_TYPE.TABLE)
    table.base_style = styles['Table Grid']
    table.font.name = body.font.name
    table.font.size = Pt(9)
    caption = styles.add_style('Civensa Tablo Başlık', WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = body
    caption.font.bold = True
    caption.font.size = Pt(9)
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_after = Pt(4)
    note = styles.add_style('Civensa Not', WD_STYLE_TYPE.PARAGRAPH)
    note.base_style = body
    note.font.color.rgb = RGBColor.from_string('52616B')
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(8)
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'F0F4F6')
    note.element.get_or_add_pPr().append(shade)
    cover = styles.add_style('Civensa Kapak Başlık', WD_STYLE_TYPE.PARAGRAPH)
    cover.base_style = body
    cover.font.size = Pt(24)
    cover.font.bold = True
    cover.font.color.rgb = RGBColor.from_string('164E63')
    cover.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(28)
    cover.paragraph_format.space_after = Pt(14)
    for style in styles:
        if style.name.startswith('Civensa ') or style.name in {'TOC 1', 'TOC 2', 'TOC 3'}:
            style.font.name = 'Arial'
            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), 'tr-TR')
            style.element.get_or_add_rPr().append(lang)
    section.first_page_header.paragraphs[0].style = body
    section.first_page_footer.paragraphs[0].style = body
    doc.core_properties.title = 'Civensa Faaliyet Raporu Word Şablonu'
    doc.core_properties.author = 'Civensa'
    doc.add_paragraph('Bu belgenin Civensa stillerini Word’de düzenleyip aynı adla kaydedin.', body)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == '__main__':
    create(Path(__file__).resolve().parents[1] / 'app/templates_docx/faaliyet_raporu.docx')
